"""Copy the decks of one dataset run into deliverables/ with a summary table.

Usage:
    python examples/collect_deliverables.py \
        --report slide-workspace/final-2026-09-28k/dataset_report.json \
        --out deliverables
"""

from __future__ import annotations

import argparse
import io
import json
import re
import shutil
import tempfile
import zipfile
from pathlib import Path, PureWindowsPath

from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
NAMES = {
    "VK Tech шаблон": "vk-tech",
    "VK_WorkSpace_Клиентская_конференция_Шаблон_03": "vk-workspace",
    "Шаблон презентации VK Education": "vk-education",
}


def compact_copy(source: Path, target: Path, *, jpeg_over: int = 500_000) -> None:
    """Publishable copy: unused template layouts dropped, large photos as JPEG.

    The service keeps every template layout so the deck can be edited further
    in the brand; the published copy keeps only the layouts its slides use and
    stores big opaque PNG pictures as JPEG. Slides stay native objects.
    """
    prs = Presentation(source)
    used = {id(slide.slide_layout.part) for slide in prs.slides}
    for master in prs.slide_masters:
        for layout in list(master.slide_layouts):
            if id(layout.part) not in used:
                master.slide_layouts.remove(layout)
    with tempfile.TemporaryDirectory() as folder:
        pruned = Path(folder) / "pruned.pptx"
        prs.save(pruned)
        with zipfile.ZipFile(pruned) as package:
            members = {name: package.read(name) for name in package.namelist()}
    renamed: dict[str, str] = {}
    for name in list(members):
        if not (name.startswith("ppt/media/") and name.endswith(".png")):
            continue
        if len(members[name]) <= jpeg_over:
            continue
        with Image.open(io.BytesIO(members[name])) as image:
            if image.mode in {"RGBA", "LA", "P"} and image.convert("RGBA").getextrema()[3][0] < 255:
                continue  # keep transparency
            buffer = io.BytesIO()
            image.convert("RGB").save(buffer, format="JPEG", quality=88, optimize=True)
        new_name = name[:-4] + ".jpeg"
        members[new_name] = buffer.getvalue()
        del members[name]
        renamed[name.rsplit("/", 1)[1]] = new_name.rsplit("/", 1)[1]
    if renamed:
        pattern = re.compile("|".join(re.escape(old) for old in renamed))
        for name in list(members):
            if name.endswith(".rels"):
                text = members[name].decode("utf-8")
                members[name] = pattern.sub(lambda m: renamed[m.group(0)], text).encode("utf-8")
        types = members["[Content_Types].xml"].decode("utf-8")
        if 'Extension="jpeg"' not in types:
            types = types.replace(
                "<Default ",
                '<Default Extension="jpeg" ContentType="image/jpeg"/><Default ',
                1,
            )
        members["[Content_Types].xml"] = types.encode("utf-8")
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as package:
        for name, data in members.items():
            package.writestr(name, data)
    Presentation(target)  # the copy must still open


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument("--out", type=Path, default=ROOT / "deliverables")
    args = parser.parse_args()
    report = json.loads(args.report.read_text(encoding="utf-8"))
    runs = report.get("runs") or []
    if len(runs) != len(NAMES):
        raise ValueError(f"Expected {len(NAMES)} completed template runs, found {len(runs)}")
    for run in runs:
        batch = run.get("batch") or {}
        variants = batch.get("variants") or []
        if (
            run.get("error")
            or batch.get("status") != "completed"
            or batch.get("diversity_status") != "passed"
            or len(variants) != 3
            or not run.get("source_unchanged", False)
        ):
            raise ValueError(f"Template run is incomplete: {run.get('template')}")
        for variant in variants:
            if (
                variant.get("status") != "completed"
                or variant.get("qa", {}).get("status") != "passed"
                or variant.get("exports", {}).get("status") != "passed"
                or variant.get("images", {}).get("placed", 0) < 2
            ):
                raise ValueError(
                    f"Variant is incomplete: {run.get('template')} / "
                    f"{variant.get('variant', {}).get('id')}"
                )
            directory = Path(variant["presentation_dir"])
            if not (directory / "output.pptx").is_file() or not (
                directory / "exports" / "output.pdf"
            ).is_file():
                raise ValueError(f"Missing source deliverable: {directory}")
    if (
        report.get("source_matches_git_head_at_start") is not True
        or report.get("source_unchanged_during_run") is not True
        or report.get("inputs_unchanged_during_run") is not True
    ):
        raise ValueError("The dataset run has no stable committed source and inputs")
    args.out.mkdir(parents=True, exist_ok=True)
    rows = []
    for run in runs:
        stem = PureWindowsPath(run["template"]).stem
        name = NAMES.get(stem, stem)
        for variant in run["batch"]["variants"]:
            source = Path(variant["presentation_dir"])
            target = f"{name}-{variant['variant']['id']}"
            compact_copy(source / "output.pptx", args.out / f"{target}.pptx")
            shutil.copy2(source / "exports" / "output.pdf", args.out / f"{target}.pdf")
            qa = variant["qa"]
            editorial = sum(
                1 for issue in qa.get("issues", []) if issue.get("check_type") == "contextual"
            )
            images = variant["images"]
            rows.append(
                f"| {stem} | {variant['variant']['id']} | [PPTX]({target}.pptx) · [PDF]({target}.pdf) "
                f"| {qa['score']} | {images['placed']} из {images['available']} | {editorial} |"
            )
    timings = " · ".join(
        f"{PureWindowsPath(run['template']).stem}: {run['elapsed_seconds']:.0f} с"
        for run in report["runs"]
    )
    readme = [
        "# Девять презентаций: три шаблона × три варианта",
        "",
        "Краткий бриф `examples/brief_feature.md` и две картинки из",
        "`examples/acceptance_images`; структуру и текст пишет локальная Qwen3.5-9B,",
        "вёрстка, аудит и экспорт — сервис. Конфиг: `examples/brief_demo.json`.",
        f"Время трёх вариантов на шаблон: {timings}.",
        "",
        "| Шаблон | Вариант | Файлы | QA | Картинки | Подсказки редактору |",
        "|---|---|---|---:|---|---:|",
        *rows,
        "",
        "Копии для публикации: из PPTX удалены неиспользуемые макеты шаблона, крупные",
        "фотографии сохранены как JPEG. Слайды — те же нативные объекты; полную версию",
        "со всеми макетами шаблона создаёт сервис.",
        "",
        "QA — балл структурных и визуальных проверок (см. AUDIT.md). Подсказки",
        "редактору — замечания текстовой модели после одной автоматической",
        "переписи; они не входят в балл и требуют проверки человеком.",
        "",
    ]
    (args.out / "README.md").write_text("\n".join(readme), encoding="utf-8")
    print("\n".join(rows))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
