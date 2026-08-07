"""Create three intentionally different PPTX templates and Russian demo content."""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

STYLES = {
    "tech-dark": {
        "background": "101828",
        "text": "F9FAFB",
        "muted": "98A2B3",
        "accent": "7F56D9",
        "accent2": "2E90FA",
        "font": "Aptos Display",
    },
    "editorial": {
        "background": "F7F1E8",
        "text": "2B2118",
        "muted": "786A5B",
        "accent": "B54708",
        "accent2": "4E5D43",
        "font": "Georgia",
    },
    "clean-blue": {
        "background": "FFFFFF",
        "text": "16324F",
        "muted": "5D7285",
        "accent": "1367D1",
        "accent2": "2AA198",
        "font": "Arial",
    },
}


def _font(shape, style, size, bold=False, color=None):
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = style["font"]
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(color or style["text"])


def _decorate(slide, prs, style, variant):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(style["background"])
    if variant == "tech-dark":
        stripe = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), prs.slide_height
        )
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = RGBColor.from_string(style["accent"])
        stripe.line.fill.background()
    elif variant == "editorial":
        stripe = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(0.45), Inches(0.06), Inches(6.4)
        )
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = RGBColor.from_string(style["accent"])
        stripe.line.fill.background()
    else:
        stripe = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.23), prs.slide_width, Inches(0.27)
        )
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = RGBColor.from_string(style["accent"])
        stripe.line.fill.background()
    stripe.name = f"Brand decoration {variant}"


def create_template(path: Path, variant: str) -> None:
    style = STYLES[variant]
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    cover = prs.slides.add_slide(prs.slide_layouts[0])
    _decorate(cover, prs, style, variant)
    cover.shapes.title.text = "Название продукта"
    cover.placeholders[1].text = "Короткий дескриптор и дата"
    _font(cover.shapes.title, style, 38, True)
    _font(cover.placeholders[1], style, 18, color=style["muted"])

    content = prs.slides.add_slide(prs.slide_layouts[1])
    _decorate(content, prs, style, variant)
    content.shapes.title.text = "Ключевой тезис"
    body = content.placeholders[1]
    body.text = "Первый аргумент\nВторой аргумент\nТретий аргумент"
    _font(content.shapes.title, style, 30, True)
    _font(body, style, 18, color=style["text"])
    for index in range(3):
        card = content.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(7.1 + index * 1.75),
            Inches(2.15),
            Inches(1.45),
            Inches(1.45),
        )
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor.from_string(
            style["accent"] if index == 0 else style["accent2"]
        )
        card.line.fill.background()

    comparison = prs.slides.add_slide(prs.slide_layouts[3])
    _decorate(comparison, prs, style, variant)
    comparison.shapes.title.text = "Сравнение вариантов"
    _font(comparison.shapes.title, style, 30, True)
    for placeholder, text in zip(
        list(comparison.placeholders)[1:],
        ["Вариант A\nПреимущества", "Вариант B\nПреимущества"],
    ):
        if placeholder.has_text_frame:
            placeholder.text = text
            _font(placeholder, style, 17)

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)


CONTENT = """# Pulse — аналитика клиентского опыта

## Проблема
- Команды тратят до 12 часов в неделю на сбор обратной связи.
- Данные распределены между пятью каналами.
- Критические сигналы доходят до продукта с задержкой.

## Решение
- Pulse объединяет отзывы, обращения и продуктовые события.
- Модель группирует темы и показывает причины изменений.
- Команда получает приоритетный список действий.

## Эффект пилота
- Время подготовки отчёта сократилось с 12 часов до 40 минут.
- Доля обработанных сигналов выросла с 54% до 91%.
- Среднее время реакции сократилось на 37%.

## Архитектура
- Коннекторы принимают данные из CRM, поддержки и аналитики.
- Пайплайн очищает, классифицирует и связывает сигналы.
- Дашборд и API возвращают результаты бизнес-командам.

## Следующие шаги
- Подключить два новых источника данных.
- Провести пилот в трёх продуктовых командах.
- Измерить влияние на удержание пользователей.
"""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("output", nargs="?", default="demo-assets")
    args = parser.parse_args()
    output = Path(args.output)
    for name in STYLES:
        create_template(output / f"{name}.pptx", name)
    (output / "content.md").write_text(CONTENT, encoding="utf-8")
    print(f"Created demo assets in {output.resolve()}")


if __name__ == "__main__":
    main()
