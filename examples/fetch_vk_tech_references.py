"""Download public VK Tech decks and convert them to editable PPTX fixtures."""

from __future__ import annotations

import argparse
import hashlib
import json
import time
import urllib.parse
import urllib.request
from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from slide_agent.composer import _set_picture_background
from slide_agent.pdf_importer import pdf_to_pptx
from slide_agent.utils import write_json

BLUE_REFERENCE = {
    "id": "vk-tech-blue-corporate",
    "title": "VK Tech — ведущий российский разработчик корпоративного ПО",
    "source_page": "https://ppt-online.org/1691195",
    "slide_base_url": "https://cf5.ppt-online.org/files5/slide/9/9IE3CY4t6WUzD2mlVpiZGvedALBoukOgPn5FSQ/slide-{index}.jpg",
    "selected_slide_indices": [0, 3, 4, 8, 12, 14],
}

REFERENCES = [
    {
        "id": "vk-tech-private-cloud-2025",
        "title": "ПАК как опорная точка цифрового суверенитета",
        "speaker": "Станислав Погоржельский, VK Cloud / VK Tech",
        "source_page": "https://spb.dcforum.ru/archive/programm?day=1",
        "download_url": "https://spb.dcforum.ru/sites/default/files/spb/15.00-15.20_pogorzhelskiy_2025_17_06_ivent_cod_v_spb_v2_short-1.pdf",
    },
    {
        "id": "vk-tech-training-2026",
        "title": "Как сделать линейку тренингов на основе матрицы компетенций",
        "speaker": "Павел Лапаев и Владислав Грищенко, VK Tech",
        "source_page": "https://edu-forum.pro/",
        "yandex_public_url": "https://disk.360.yandex.ru/i/aKyEeARywFax5g",
    },
]


def _download_url(reference: dict[str, str]) -> str:
    if reference.get("download_url"):
        return reference["download_url"]
    public_url = reference["yandex_public_url"]
    api = (
        "https://cloud-api.yandex.net/v1/disk/public/resources/download?public_key="
        + urllib.parse.quote(public_url, safe="")
    )
    with urllib.request.urlopen(api, timeout=30) as response:
        return json.loads(response.read().decode("utf-8"))["href"]


def _download_bytes(url: str, timeout: int = 45, attempts: int = 4) -> bytes:
    request = urllib.request.Request(
        url,
        headers={"User-Agent": "Mozilla/5.0 BrandDeck reference validator"},
    )
    last_error: Exception | None = None
    for attempt in range(attempts):
        try:
            with urllib.request.urlopen(request, timeout=timeout) as response:
                return response.read()
        except Exception as exc:  # noqa: BLE001 - retry transient public-host errors
            last_error = exc
            if attempt + 1 < attempts:
                time.sleep(0.5 * (attempt + 1))
    raise RuntimeError(f"Unable to download public reference: {url}") from last_error


def _transparent_wordmark(source: bytes) -> BytesIO:
    with Image.open(BytesIO(source)) as image:
        logo = image.convert("RGBA").crop((12, 516, 146, 568))
    pixels = logo.load()
    background = pixels[0, 0][:3]
    stack = [
        *((x, 0) for x in range(logo.width)),
        *((x, logo.height - 1) for x in range(logo.width)),
        *((0, y) for y in range(logo.height)),
        *((logo.width - 1, y) for y in range(logo.height)),
    ]
    cleared: set[tuple[int, int]] = set()
    while stack:
        x, y = stack.pop()
        if (x, y) in cleared or not (0 <= x < logo.width and 0 <= y < logo.height):
            continue
        red, green, blue, _ = pixels[x, y]
        distance = max(
            abs(red - background[0]),
            abs(green - background[1]),
            abs(blue - background[2]),
        )
        if distance > 105:
            continue
        cleared.add((x, y))
        pixels[x, y] = (red, green, blue, 0)
        stack.extend(((x - 1, y), (x + 1, y), (x, y - 1), (x, y + 1)))
    for y in range(logo.height):
        for x in range(logo.width):
            red, green, blue, alpha = pixels[x, y]
            if alpha and x > 48 and max(red, green, blue) < 180:
                pixels[x, y] = (70, 128, 194, alpha)
    payload = BytesIO()
    logo.save(payload, format="PNG")
    payload.seek(0)
    return payload


def _prepare_blue_reference(output: Path, max_pages: int | None) -> dict[str, object]:
    selected = BLUE_REFERENCE["selected_slide_indices"]
    if max_pages:
        selected = selected[:max_pages]
    downloads: list[tuple[int, bytes]] = []
    source_dir = output / "vk-tech-blue-reference"
    source_dir.mkdir(parents=True, exist_ok=True)
    for index in selected:
        url = str(BLUE_REFERENCE["slide_base_url"]).format(index=index)
        target = source_dir / f"slide-{index + 1:02d}.jpg"
        legacy_cache = output / "blue-reference" / target.name
        if target.exists():
            payload = target.read_bytes()
        elif legacy_cache.exists():
            payload = legacy_cache.read_bytes()
            target.write_bytes(payload)
        else:
            payload = _download_bytes(url)
            target.write_bytes(payload)
        downloads.append((index, payload))

    if not downloads:
        raise RuntimeError("The blue VK Tech reference did not return any slides")
    content_logo_source = next(
        (payload for index, payload in downloads if index == 4), downloads[-1][1]
    )
    wordmark = _transparent_wordmark(content_logo_source)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for position, (_, payload) in enumerate(downloads):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _set_picture_background(slide, payload)
        wordmark.seek(0)
        top = 0.36 if position < 2 else 6.82
        mark = slide.shapes.add_picture(
            wordmark, Inches(0.48), Inches(top), width=Inches(1.28)
        )
        mark.name = "VK Tech wordmark"
    pptx_path = output / f"{BLUE_REFERENCE['id']}.pptx"
    prs.save(pptx_path)
    pptx_payload = pptx_path.read_bytes()
    return {
        **BLUE_REFERENCE,
        "fixture_kind": "public_slide_image_reference",
        "slide_count": len(downloads),
        "pptx_sha256": hashlib.sha256(pptx_payload).hexdigest(),
        "pptx_size_bytes": len(pptx_payload),
        "slides": [
            {
                "source_index": index,
                "sha256": hashlib.sha256(payload).hexdigest(),
                "size_bytes": len(payload),
            }
            for index, payload in downloads
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("output", nargs="?", default="external-fixtures/vk-tech")
    parser.add_argument(
        "--max-pages", type=int, help="Limit pages per source for quick checks"
    )
    parser.add_argument(
        "--only",
        choices=["all", BLUE_REFERENCE["id"], *(item["id"] for item in REFERENCES)],
        default="all",
        help="Prepare one reference instead of the full validation set",
    )
    args = parser.parse_args()
    output = Path(args.output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    manifest = []
    if args.only in {"all", BLUE_REFERENCE["id"]}:
        manifest.append(_prepare_blue_reference(output, args.max_pages))
    for reference in REFERENCES:
        if args.only not in {"all", reference["id"]}:
            continue
        pdf_path = output / f"{reference['id']}.pdf"
        pptx_path = output / f"{reference['id']}.pptx"
        payload = _download_bytes(_download_url(reference), timeout=90)
        pdf_path.write_bytes(payload)
        conversion = pdf_to_pptx(pdf_path, pptx_path, max_pages=args.max_pages)
        pptx_payload = pptx_path.read_bytes()
        manifest.append(
            {
                **reference,
                "pdf_sha256": hashlib.sha256(payload).hexdigest(),
                "pdf_size_bytes": len(payload),
                "pptx_sha256": hashlib.sha256(pptx_payload).hexdigest(),
                "pptx_size_bytes": len(pptx_payload),
                "conversion": conversion,
            }
        )
        print(f"Prepared {reference['id']}: {conversion['slide_count']} slides")
    write_json(output / "manifest.json", {"references": manifest})


if __name__ == "__main__":
    main()
