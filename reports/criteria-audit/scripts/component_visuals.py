"""Controlled composer probe, NOT an inference or end-to-end generation test."""
import json
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from slide_agent.composer import compose_presentation
from slide_agent.planner import assign_patterns
from slide_agent.qa import inspect_presentation

ROOT=Path(__file__).resolve().parents[3]
OUT=ROOT/'reports/criteria-audit/artifacts/component-visuals/vk-education'
OUT.mkdir(parents=True,exist_ok=True)
meta=json.loads((ROOT/'reports/criteria-audit/artifacts/vk-education/balanced/manifest.json').read_text())
template=Path(meta['template_dir'])
design=json.loads((template/'design_system.json').read_text())
catalog=json.loads((template/'pattern_catalog.json').read_text())
def slide(title,visual=None,role='data'):
    return dict(title=title,role=role,subtitle='',body='',bullets=[],visual=visual,speaker_notes='Synthetic controlled component fixture; no model called.')
slides=[
    slide('Проверка нативных объектов',role='cover'),
    slide('Два подразделения: 60 и 90 участников',{'type':'bar_chart','categories':['Продажи','Поддержка'],'series':[{'name':'Сотрудники','values':[60,90]}],'unit':'человек'}),
    slide('Доля участников по подразделениям',{'type':'pie_chart','categories':['Продажи','Поддержка'],'values':[60,90]}),
    slide('Пилот рассчитан на шесть недель',{'type':'table','headers':['Этап','Недели','Ответственный'],'rows':[['Подготовка','1–2','Аналитик'],['Запуск','3–4','Инженер'],['Оценка','5–6','Менеджер']]}),
    slide('Проверяем путь обращения',{'type':'process','items':['Задать вопрос','Найти регламент','Показать источник','Передать специалисту']}),
    slide('Качество улучшаем после проверки',{'type':'cycle','items':['Собрать','Проверить','Запустить','Измерить']}),
    slide('У каждого участника есть роль',{'type':'hierarchy','root':'Команда пилота','items':[{'label':'Аналитик','detail':'Готовит материалы'},{'label':'Инженер','detail':'Настраивает сервис'},{'label':'Менеджер','detail':'Измеряет результаты'}]}),
    slide('Измеряем три результата',{'type':'icon_grid','items':['Время поиска ответа','Доля обращений без специалиста','Оценка сотрудника']}),
    slide('Решение принимаем по итогам пилота',{'type':'table','headers':['Метрика','Статус'],'rows':[['Время поиска','Будет измерено'],['Качество ответа','Будет проверено']]}),
    slide('Объекты доступны для редактирования',role='closing'),
]
plan=assign_patterns({'title':'Контроль компонентов','language':'ru','slides':slides},catalog)
(OUT/'controlled-plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2))
output=OUT/'component-visuals.pptx'
result=compose_presentation(template_dir=template,plan=plan,output_path=output,design_system=design)
(OUT/'compose.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
(OUT/'qa.json').write_text(json.dumps(inspect_presentation(output,design_system=design,expected_slide_count=len(slides)),ensure_ascii=False,indent=2))
prs=Presentation(output)
chart=next(sh.chart for sh in prs.slides[1].shapes if sh.has_chart)
data=CategoryChartData();data.categories=['Продажи','Поддержка'];data.add_series('Сотрудники',[61,89]);chart.replace_data(data)
table=next(sh.table for sh in prs.slides[3].shapes if sh.has_table);table.cell(1,0).text='Изменено аудитором'
text=next(sh for sh in prs.slides[0].shapes if sh.has_text_frame and sh.text.strip());text.text='Проверка правки текста'
edited=ROOT/'reports/criteria-audit/artifacts/component-edited/vk-education';edited.mkdir(parents=True,exist_ok=True)
edited_file=edited/'component-edited.pptx';prs.save(edited_file)
check=Presentation(edited_file)
result={'method':'python-pptx object edit, save, reopen; no GUI editing claimed','chart_values':list(next(s.chart for s in check.slides[1].shapes if s.has_chart).series[0].values),'table_cell':next(s.table for s in check.slides[3].shapes if s.has_table).cell(1,0).text,'cover_text':[s.text for s in check.slides[0].shapes if s.has_text_frame],'source_untouched':str(output)}
(OUT/'edit-roundtrip.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
print(json.dumps(result,ensure_ascii=False))
