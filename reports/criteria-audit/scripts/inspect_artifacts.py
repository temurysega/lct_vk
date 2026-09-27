"""Independent file inventory plus clearly identified application QA heuristics."""
from pathlib import Path
from collections import Counter
import json,hashlib,re,zipfile,sys
import fitz
from PIL import Image,ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from slide_agent.render_audit import inspect_rendered_fill,inspect_rendered_contrast
from slide_agent.visual_diversity import inspect_visual_diversity

ROOT=Path(__file__).resolve().parents[3]
OUT=ROOT/'reports/criteria-audit'
def walk(shapes):
 for s in shapes:
  yield s
  if s.shape_type==MSO_SHAPE_TYPE.GROUP:yield from walk(s.shapes)

def main():
 rows=[];diversity=[]
 for base in [OUT/'artifacts',OUT/'with-renderer/artifacts',OUT/'additional-unseen/artifacts']:
  if not base.exists():continue
  for source in sorted(base.glob('*/*/*.pptx')):
   if source.name.startswith('~$'):continue
   prs=Presentation(source);native=[]
   for i,s in enumerate(prs.slides,1):
    shapes=list(walk(s.shapes));texts=[x.text for x in shapes if getattr(x,'has_text_frame',False) and x.text.strip()]
    tables=[x.table for x in shapes if x.has_table];charts=[x.chart for x in shapes if x.has_chart]
    sizes=[r.font.size.pt for x in shapes if x.has_text_frame for p in x.text_frame.paragraphs for r in p.runs if r.font.size and r.text.strip()]
    native.append({'slide':i,'text_objects':len(texts),'tables':len(tables),'charts':len(charts),'groups':sum(x.shape_type==MSO_SHAPE_TYPE.GROUP for x in shapes),'pictures':sum(x.shape_type==MSO_SHAPE_TYPE.PICTURE for x in shapes),'min_explicit_font_pt':min(sizes) if sizes else None,'text':texts,'table_cells':[[[c.text for c in row.cells] for row in table.rows] for table in tables],'chart_values':[[list(ser.values) for ser in chart.series] for chart in charts]})
   with zipfile.ZipFile(source) as z:
    integrity=z.testzip();names=z.namelist()
   record={'file':str(source.relative_to(OUT)),'sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'zip_error':integrity,'slide_count':len(prs.slides),'smartart_parts':sum(n.startswith('ppt/diagrams/data') for n in names),'chart_parts':sum(bool(re.match(r'ppt/charts/chart\d+\.xml$',n)) for n in names),'embedded_workbooks':sum(n.startswith('ppt/embeddings/') for n in names),'slides':native}
   (source.parent/'object-inventory.json').write_text(json.dumps(record,ensure_ascii=False,indent=2));rows.append({k:v for k,v in record.items() if k!='slides'})
   for pdf in source.parent.glob('*.pdf'):
    renderer='PowerPoint' if pdf.name=='powerpoint.pdf' else 'LibreOffice'
    prefix='powerpoint' if renderer=='PowerPoint' else 'libreoffice'
    pages=[]
    with fitz.open(pdf) as doc:
     for i,page in enumerate(doc,1):
      png=source.parent/f'{prefix}-slide-{i:02d}.png';page.get_pixmap(matrix=fitz.Matrix(1200/page.rect.width,1200/page.rect.width),alpha=False).save(png)
      im=Image.open(png).convert('RGB');im.thumbnail((480,270));pages.append(im.copy())
     (source.parent/f'{prefix}-text.txt').write_text('\n'.join(f'=== SLIDE {i} ===\n{page.get_text()}' for i,page in enumerate(doc,1)))
    sheet=Image.new('RGB',(960,300*((len(pages)+1)//2)),'#ddd');draw=ImageDraw.Draw(sheet)
    for j,im in enumerate(pages):
     x=(j%2)*480;y=(j//2)*300;sheet.paste(im,(x,y+24));draw.text((x+8,y+5),f'{source.parent.parent.name}/{source.parent.name} | {renderer} | {j+1}',fill='black')
    sheet.save(source.parent/f'{prefix}-contact-sheet.jpg',quality=90)
    plan_path=source.parent/'deck_plan.final.json'
    if not plan_path.exists():plan_path=source.parent/'controlled-plan.json'
    if plan_path.exists():
     plan=json.loads(plan_path.read_text());roles=[s.get('role','content') for s in plan['slides']]
     qa={'renderer':renderer,'stage':'External acceptance checks on saved artifact; no repairs','fill':inspect_rendered_fill(pdf,source,roles),'contrast':inspect_rendered_contrast(pdf,source)}
     (source.parent/f'{prefix}-render-audit.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2))
  for template in base.iterdir():
   variants=[];roles=[]
   for variant in ['balanced','columns','focus']:
    folder=template/variant
    if not (folder/'manifest.json').exists():continue
    manifest=json.loads((folder/'manifest.json').read_text());plan=json.loads((folder/'deck_plan.final.json').read_text());roles=[s.get('role','content') for s in plan['slides']]
    pngs=sorted(folder.glob('libreoffice-slide-*.png'))
    variants.append({'variant':{'id':variant},'exports':{'previews':[{'slide':i,'path':str(p)} for i,p in enumerate(pngs,1)]}})
   if len(variants)==3 and any(v['exports']['previews'] for v in variants):
    result=inspect_visual_diversity(variants,roles);result['template']=str(template.relative_to(OUT));diversity.append(result)
 (OUT/'evidence/object-summary.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2))
 (OUT/'evidence/rendered-diversity.json').write_text(json.dumps(diversity,ensure_ascii=False,indent=2))
 print(json.dumps({'decks':len(rows),'diversity':[{k:d[k] for k in ['template','status']} for d in diversity]},ensure_ascii=False))

if __name__=='__main__':main()
