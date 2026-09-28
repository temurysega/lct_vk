"""Deterministic audit annotations and the contract for selective layout repairs."""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from typing import Any

from pptx import Presentation

from .qa import _inspected_shapes

REPAIRABLE = {
    "text_overflow_risk",
    "text_overlap",
    "out_of_bounds",
    "native_text_too_small",
    "placeholder_text",
    # Another layout removes the colliding or emptied exemplar block.
    "template_overlap",
    "emptied_template_block",
    "empty_photo_frame",
    "decor_overlap",
    # Another layout gives the text a larger box or keeps it off the edge.
    "text_too_small",
    "margin_intrusion",
}


_MIXED_SCRIPT = re.compile(r"\b(?=\w*[А-Яа-яЁё])(?=\w*[A-Za-z])\w+\b")
_LATIN_WORD = re.compile(r"(?<![\w@./:-])[a-z]{4,}(?![\w@./:-])")
_LATIN_TITLE = re.compile(r"(?<![\w@./:-])[A-Z][a-z]{3,}(?![\w@./:-])")
CONTENT_ROLES = {"content", "comparison", "data", "two_column", "image"}


def _language_issues(
    text: str, number: int, shape: str, plan: dict[str, Any]
) -> list[dict[str, Any]]:
    """Typos a model makes when it drifts between languages.

    A word mixing Cyrillic and Latin letters ("Сolution") is always a defect.
    In a Russian deck a lowercase Latin word or a line written wholly in English
    is flagged too; product names inside Russian text ("VK WorkSpace", "API")
    pass.
    """
    issues = []
    mixed = sorted(set(_MIXED_SCRIPT.findall(text)))
    if mixed:
        issues.append(
            {
                "severity": "warning",
                "code": "mixed_script_word",
                "slide": number,
                "shape": shape,
                "message": "Слово смешивает кириллицу и латиницу: " + ", ".join(mixed),
            }
        )
    if plan.get("language") == "ru":
        foreign = set(_LATIN_WORD.findall(text))
        for line in text.splitlines():
            # An English heading such as "Risks and Mitigation" has no Cyrillic.
            if not re.search(r"[А-Яа-яЁё]", line):
                foreign.update(_LATIN_TITLE.findall(line))
        foreign = sorted(foreign)
        if foreign:
            issues.append(
                {
                    "severity": "warning",
                    "code": "language_mix",
                    "slide": number,
                    "shape": shape,
                    "message": "Английские слова в русской колоде: "
                    + ", ".join(foreign),
                }
            )
    return issues


def language_problems(text: str, language: str) -> list[str]:
    """The same language checks for model text before it is laid out."""
    return [
        issue["message"]
        for issue in _language_issues(text, 0, "", {"language": language})
    ]


def enrich_audit(
    path: Path, report: dict[str, Any], plan: dict[str, Any]
) -> dict[str, Any]:
    prs = Presentation(path)
    issues = report["issues"]
    seen: dict[str, int] = {}
    for number, slide in enumerate(prs.slides, 1):
        texts = []
        for shape, _ in _inspected_shapes(slide.shapes):
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    texts.append(text)
                if str(shape.name).startswith("BrandDeck"):
                    issues.extend(_language_issues(text, number, shape.name, plan))
                if re.search(
                    r"lorem ipsum|\bTODO\b|\bXXX\b|вставьте текст", text, re.IGNORECASE
                ):
                    issues.append(
                        {
                            "severity": "warning",
                            "code": "placeholder_text",
                            "slide": number,
                            "shape": shape.name,
                            "message": "На слайде остался текст-заглушка",
                        }
                    )
            if getattr(shape, "has_table", False) and (
                len(shape.table.rows) > 7 or len(shape.table.columns) > 5
            ):
                issues.append(
                    {
                        "severity": "warning",
                        "code": "table_density",
                        "slide": number,
                        "shape": shape.name,
                        "message": "Таблица превышает 7 строк или 5 столбцов",
                    }
                )
            if getattr(shape, "has_chart", False) and len(shape.chart.series) > 5:
                issues.append(
                    {
                        "severity": "warning",
                        "code": "chart_density",
                        "slide": number,
                        "shape": shape.name,
                        "message": "В диаграмме больше 5 рядов",
                    }
                )
        fingerprint = "\n".join(texts)
        if fingerprint and fingerprint in seen:
            issues.append(
                {
                    "severity": "warning",
                    "code": "duplicate_slide",
                    "slide": number,
                    "message": f"Текст повторяет слайд {seen[fingerprint]}",
                }
            )
        seen[fingerprint] = number
        spec = plan["slides"][number - 1]
        bullets = spec.get("bullets", [])
        if spec.get("role") in CONTENT_ROLES and not (
            bullets or spec.get("body") or spec.get("visual")
        ):
            issues.append(
                {
                    "severity": "warning",
                    "code": "title_only",
                    "slide": number,
                    "message": "На слайде только заголовок, без содержания",
                }
            )
        if len(bullets) > 6 or any(len(b.split()) > 15 for b in bullets):
            issues.append(
                {
                    "severity": "warning",
                    "code": "bullet_density",
                    "slide": number,
                    "message": "Больше 6 тезисов или тезис длиннее 15 слов; требуется редактура",
                }
            )
    for index, issue in enumerate(issues):
        number = issue.get("slide")
        bounds = []
        if isinstance(number, int) and 1 <= number <= len(prs.slides):
            names = str(issue.get("shape", "")).split(" / ")
            for shape, _ in _inspected_shapes(prs.slides[number - 1].shapes):
                if shape.name in names:
                    bounds.append(
                        {
                            key: round(value / 914400, 4)
                            for key, value in zip(
                                ("x", "y", "w", "h"),
                                (shape.left, shape.top, shape.width, shape.height),
                            )
                        }
                    )
        identity = json.dumps(
            [number, issue.get("code"), issue.get("shape"), index], sort_keys=True
        )
        issue["id"] = hashlib.sha256(identity.encode()).hexdigest()[:16]
        issue["check_type"] = "deterministic"
        issue["bounds"] = bounds
        issue["repairable"] = bool(
            isinstance(number, int) and issue["code"] in REPAIRABLE
        )
    errors = sum(i["severity"] == "error" for i in issues)
    warnings = sum(i["severity"] == "warning" for i in issues)
    report["status"] = "failed" if errors else "warning" if warnings else "passed"
    report["score"] = min(report["score"], max(0, 100 - 18 * errors - 4 * warnings))
    report["schema_version"] = "2.0"
    return report
