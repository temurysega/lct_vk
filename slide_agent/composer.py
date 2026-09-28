from __future__ import annotations

import copy
import io
import math
import textwrap
from collections import Counter
from contextvars import ContextVar
from pathlib import Path
from typing import Any

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData, ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from .decor_audit import layout_art_boxes
from .diagrams import (
    DIAGRAM_TYPES,
    ITEM_LIMITS,
    contrast,
    diagram_style,
    on_fill,
    render_diagram,
    render_illustration,
)
from .images import crop_to_fill, crop_to_fill_salient
from .layout_geometry import (
    area,
    box,
    collides,
    gap,
    has_text,
    intersection,
    is_framing,
    is_generated,
    is_navigation_label,
    is_opaque,
    is_photo_frame,
    is_visible_box,
)
from .pictograms import add_icon, assign_icons
from .typography import (
    ESTIMATED_GLYPH_WIDTH_EM,
    effective_font_size,
    estimated_line_count,
    needs_cyrillic_font_fallback,
    readable_font,
)
from .utils import read_json, write_json

NATIVE_VISUALS = set(DIAGRAM_TYPES) | {"image", "table"}
# A table or card grid below this share of the slide is stretched over its
# zone: the slide would otherwise stay under a quarter filled.
FILL_SHARE = 0.27
# Statements whose wrapped text would cover less of the slide become cards.
SPARSE_FILL = 0.26
_CYRILLIC_THEME_FALLBACK: ContextVar[str | None] = ContextVar(
    "cyrillic_theme_fallback", default=None
)
# Average width of a rendered Cyrillic or Latin glyph at body sizes.
RENDERED_GLYPH_WIDTH_EM = 0.55
# Exemplar shapes that can be scaffolding; pictures are kept.
LEFTOVER_TYPES = {
    MSO_SHAPE_TYPE.AUTO_SHAPE,
    MSO_SHAPE_TYPE.GROUP,
    MSO_SHAPE_TYPE.FREEFORM,
    MSO_SHAPE_TYPE.TEXT_BOX,
    MSO_SHAPE_TYPE.LINE,
}

TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
SUBTITLE_TYPES = {PP_PLACEHOLDER.SUBTITLE}
BODY_TYPES = {
    PP_PLACEHOLDER.BODY,
    PP_PLACEHOLDER.OBJECT,
    PP_PLACEHOLDER.VERTICAL_BODY,
    PP_PLACEHOLDER.VERTICAL_OBJECT,
}


def _hex(value: str | None, fallback: str) -> str:
    candidate = str(value or "").lstrip("#").upper()
    if len(candidate) == 8:
        candidate = candidate[-6:]
    return candidate if len(candidate) == 6 else fallback


def _rgb(value: str | None, fallback: str = "000000") -> RGBColor:
    return RGBColor.from_string(_hex(value, fallback))


def _number(value: Any) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = "".join(
        char for char in str(value).replace(",", ".") if char in "-0123456789."
    )
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


def _theme_colors(design: dict[str, Any]) -> dict[str, str]:
    result = {
        str(item.get("role")): _hex(item.get("hex"), "000000")
        for item in design.get("colors", {}).get("theme", [])
    }
    observed = design.get("colors", {}).get("observed", [])
    fallbacks = [item.get("hex") for item in observed if item.get("hex")]
    result.setdefault("accent1", _hex(fallbacks[0] if fallbacks else None, "4472C4"))
    result.setdefault(
        "accent2", _hex(fallbacks[1] if len(fallbacks) > 1 else None, "70AD47")
    )
    result.setdefault(
        "accent3", _hex(fallbacks[2] if len(fallbacks) > 2 else None, "ED7D31")
    )
    result.setdefault("dk1", "1F2937")
    result.setdefault("lt1", "FFFFFF")
    result.setdefault("lt2", "F3F4F6")
    brand = design.get("brand", {})
    if brand:
        accent = _hex(brand.get("accent"), result["accent1"])
        # The detected brand accent is sometimes just the body ink (black or
        # white); a grey never replaces the theme's chromatic accent.
        if _chroma(accent) >= 40 or _chroma(result["accent1"]) < 40:
            result["accent1"] = accent
        result["dk1"] = _hex(brand.get("body"), result["dk1"])
        result["lt1"] = _hex(brand.get("heading"), result["lt1"])
    return result


# The template's type scale for the deck being composed (Appendix 1: sizes
# come from the template's typographic scale).
_TYPE_SCALE: ContextVar[tuple[float, ...]] = ContextVar("type_scale", default=())


def _snap_size(size: float, minimum: float) -> float:
    """Largest template size not above ``size``; ``size`` without a scale."""
    scale = _TYPE_SCALE.get()
    fitting = [value for value in scale if minimum <= value <= size + 0.05]
    if fitting:
        return max(fitting)
    larger = [value for value in scale if minimum <= value]
    return min(larger) if larger and min(larger) <= size * 1.25 else size


def _chroma(value: str) -> int:
    channels = [int(value[index : index + 2], 16) for index in (0, 2, 4)]
    return max(channels) - min(channels)


def _mix_color(first: str, second: str, ratio: float) -> str:
    ratio = min(1.0, max(0.0, ratio))
    left = _hex(first, "000000")
    right = _hex(second, "FFFFFF")
    channels = []
    for offset in (0, 2, 4):
        value = round(
            int(left[offset : offset + 2], 16) * (1 - ratio)
            + int(right[offset : offset + 2], 16) * ratio
        )
        channels.append(f"{value:02X}")
    return "".join(channels)


def _luminance(value: str) -> float:
    color = _hex(value, "000000")
    channels = [int(color[offset : offset + 2], 16) / 255 for offset in (0, 2, 4)]
    linear = [
        channel / 12.92 if channel <= 0.04045 else ((channel + 0.055) / 1.055) ** 2.4
        for channel in channels
    ]
    return 0.2126 * linear[0] + 0.7152 * linear[1] + 0.0722 * linear[2]


def _native_tokens(design: dict[str, Any], role: str = "content") -> dict[str, Any]:
    brand = design.get("brand", {})
    colors = _theme_colors(design)
    prefix = (
        "cover"
        if role in {"cover", "section"}
        else "closing"
        if role == "closing"
        else "content"
    )
    background = _hex(
        brand.get(f"{prefix}_background", brand.get("background")), "0B1220"
    )
    dark = _luminance(background) < 0.35
    heading = _hex(
        brand.get(f"{prefix}_heading", brand.get("heading")),
        colors.get("lt1") if dark else colors.get("dk1"),
    )
    body = _hex(brand.get(f"{prefix}_body", brand.get("body")), heading)
    accent = _hex(brand.get("accent"), colors.get("accent1"))
    surface = (
        _mix_color(background, "FFFFFF", 0.1)
        if dark
        else ("FFFFFF" if background != "FFFFFF" else "F4F8FD")
    )
    surface_alt = _mix_color(background, accent, 0.1 if dark else 0.07)
    typography = design.get("typography", {})
    return {
        "background": background,
        "heading": heading,
        "body": body,
        "accent": accent,
        "surface": surface,
        "surface_alt": surface_alt,
        "dark": dark,
        "accent_text": "FFFFFF" if _luminance(accent) < 0.52 else "10243A",
        "font": typography.get("primary_font", "Arial"),
        "heading_font": typography.get(
            "heading_font", typography.get("primary_font", "Arial")
        ),
        "title_size": float(typography.get("content_title_size_pt", 34)),
        "cover_size": float(typography.get("cover_title_size_pt", 46)),
        "body_size": max(16.0, float(typography.get("body_size_pt", 18))),
    }


def _all_layouts(prs: Presentation) -> list[Any]:
    return [layout for master in prs.slide_masters for layout in master.slide_layouts]


def _remove_slide_ids(prs: Presentation, slide_ids_to_remove: list[Any]) -> None:
    slide_ids = prs.slides._sldIdLst
    for slide_id in slide_ids_to_remove:
        rel_id = slide_id.rId
        slide_ids.remove(slide_id)
        try:
            prs.part.drop_rel(rel_id)
        except KeyError:
            pass


def _clone_slide(prs: Presentation, source: Any) -> Any:
    """Clone a source slide while preserving its shapes and media relationships."""
    destination = prs.slides.add_slide(source.slide_layout)
    destination_tree = destination.shapes._spTree
    for shape in list(destination.shapes):
        destination_tree.remove(shape.element)

    relationship_map: dict[str, str] = {}
    for rel in source.part.rels.values():
        if rel.reltype.endswith("/slideLayout") or rel.reltype.endswith("/notesSlide"):
            continue
        try:
            target = rel.target_ref if rel.is_external else rel.target_part
            relationship_map[rel.rId] = destination.part.relate_to(
                target, rel.reltype, is_external=rel.is_external
            )
        except (AttributeError, KeyError, ValueError):
            continue

    for shape in source.shapes:
        element = copy.deepcopy(shape.element)
        for child in element.iter():
            for attribute, value in list(child.attrib.items()):
                if value in relationship_map:
                    child.set(attribute, relationship_map[value])
        destination_tree.insert_element_before(element, "p:extLst")

    try:
        source_bg = source._element.cSld.bg
        if source_bg is not None:
            destination_bg = destination._element.cSld.bg
            if destination_bg is not None:
                destination._element.cSld.remove(destination_bg)
            destination._element.cSld.insert(0, copy.deepcopy(source_bg))
    except AttributeError:
        pass
    return destination


def _placeholder_type(shape: Any) -> Any:
    try:
        return shape.placeholder_format.type
    except (AttributeError, ValueError):
        return None


def _placeholder_groups(slide: Any) -> tuple[list[Any], list[Any], list[Any]]:
    titles: list[Any] = []
    subtitles: list[Any] = []
    bodies: list[Any] = []
    for shape in slide.placeholders:
        placeholder_type = _placeholder_type(shape)
        if placeholder_type in TITLE_TYPES:
            titles.append(shape)
        elif placeholder_type in SUBTITLE_TYPES:
            subtitles.append(shape)
        elif placeholder_type in BODY_TYPES and getattr(shape, "has_text_frame", False):
            bodies.append(shape)
    bodies.sort(key=lambda item: (item.left, item.top))
    return titles, subtitles, bodies


def _largest_font(shape: Any) -> float:
    sizes: list[float] = []
    if not getattr(shape, "has_text_frame", False):
        return 0.0
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size:
                sizes.append(run.font.size.pt)
    return max(sizes, default=0.0)


def _shape_text_style(shape: Any) -> dict[str, Any]:
    """Capture explicit source styling before replacing exemplar text."""
    if not getattr(shape, "has_text_frame", False):
        return {}
    style: dict[str, Any] = {}
    for paragraph in shape.text_frame.paragraphs:
        if paragraph.alignment is not None and "alignment" not in style:
            style["alignment"] = paragraph.alignment
        for run in paragraph.runs:
            font = run.font
            fill = font._element.find(qn("a:solidFill"))
            if fill is not None and "color_xml" not in style:
                style["color_xml"] = copy.deepcopy(fill)
            if font.name and "font_name" not in style:
                style["font_name"] = font.name
            if font.size and "font_size" not in style:
                style["font_size"] = font.size.pt
            if font.bold is not None and "bold" not in style:
                style["bold"] = bool(font.bold)
            try:
                rgb = font.color.rgb
            except (AttributeError, TypeError, ValueError):
                rgb = None
            if rgb is not None and "color" not in style:
                style["color"] = str(rgb)
            if "font_name" in style and "font_size" in style and "color" in style:
                return style
    return style


def _exemplar_text_shapes(
    shapes: Any, *, grouped: bool = False
) -> list[tuple[Any, bool]]:
    """Find editable exemplar text, including text nested in PowerPoint groups."""
    found: list[tuple[Any, bool]] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            transform = shape.element.grpSpPr.xfrm
            # Child boxes are in slide coordinates only when the group's
            # parent and child coordinate systems agree. A translated/scaled
            # group needs transformed geometry, so leave its text to the
            # normal cleanup and use a fresh slide text box instead.
            if (
                transform is not None
                and not (transform.rot or transform.flipH or transform.flipV)
                and all(
                    abs(first - second) <= Inches(0.02)
                    for first, second in (
                        (transform.off.x, transform.chOff.x),
                        (transform.off.y, transform.chOff.y),
                        (transform.ext.cx, transform.chExt.cx),
                        (transform.ext.cy, transform.chExt.cy),
                    )
                )
            ):
                found.extend(_exemplar_text_shapes(shape.shapes, grouped=True))
        elif getattr(shape, "has_text_frame", False) and shape.text.strip():
            found.append((shape, grouped))
    return found


def _infer_text_groups(
    slide: Any,
    titles: list[Any],
    subtitles: list[Any],
    bodies: list[Any],
    canvas_width: float,
    canvas_height: float,
) -> tuple[list[Any], list[Any], list[Any]]:
    if titles and bodies:
        return titles, subtitles, bodies
    occupied = {shape.element for shape in titles + subtitles + bodies}
    sources = [
        (shape, grouped)
        for shape, grouped in _exemplar_text_shapes(slide.shapes)
        if shape.element not in occupied
        and shape.width / 914400 >= 1.0
        and shape.height / 914400 >= 0.18
        # Some group transforms use coordinates outside the slide. Such a
        # child is not a reliable editable text slot in slide coordinates.
        and (
            not grouped
            or (
                shape.left >= 0
                and shape.top >= 0
                and shape.left + shape.width <= int((canvas_width + 0.02) * 914400)
                and shape.top + shape.height <= int((canvas_height + 0.02) * 914400)
            )
        )
    ]
    grouped_elements = {shape.element for shape, grouped in sources if grouped}
    candidates = [shape for shape, _ in sources]

    def substantive_group_text(shape: Any) -> bool:
        return shape.element not in grouped_elements or (
            shape.height >= Inches(0.3)
            and sum(char.isalpha() for char in shape.text) >= 15
        )

    content_candidates = [
        shape
        for shape in candidates
        if not is_navigation_label(shape.text, shape.width / 914400, canvas_width)
        and (
            shape.element not in grouped_elements
            or (
                sum(char.isalpha() for char in shape.text) >= 2
                and shape.top / 914400 < canvas_height * 0.9
            )
        )
    ]
    # A grouped page number alone is not a title or body slot. Keep the old
    # fallback for ordinary ungrouped source text, where it may be intentional.
    candidates = content_candidates or [
        shape for shape in candidates if shape.element not in grouped_elements
    ]
    if not titles and candidates:
        upper = [
            shape for shape in candidates if shape.top / 914400 < canvas_height * 0.28
        ]
        pool = upper or [
            shape for shape in candidates if shape.top / 914400 < canvas_height * 0.55
        ]
        pool = pool or candidates
        readable = [
            shape for shape in pool if sum(char.isalpha() for char in shape.text) >= 2
        ]
        if readable:
            pool = readable
        wide_titles = [
            shape
            for shape in pool
            if shape.width / 914400 >= canvas_width * 0.25
            and len(shape.text.strip()) >= 4
        ]
        if wide_titles:
            pool = wide_titles
        title = max(
            pool, key=lambda shape: (_largest_font(shape), shape.width, -shape.top)
        )
        titles = [title]
        candidates.remove(title)
    if not subtitles and candidates and titles:
        title_bottom = titles[0].top + titles[0].height
        nearby = [
            shape
            for shape in candidates
            if shape.top >= title_bottom
            and shape.top / 914400 < canvas_height * 0.85
            and _largest_font(shape) >= 11
            and substantive_group_text(shape)
        ]
        if nearby:
            subtitles = [min(nearby, key=lambda shape: shape.top)]
    subtitle_elements = {shape.element for shape in subtitles}
    if not bodies:
        bodies = [
            shape
            for shape in candidates
            if shape.element not in subtitle_elements
            and substantive_group_text(shape)
            and not (
                shape.top / 914400 > canvas_height * 0.9 and _largest_font(shape) <= 10
            )
        ]
        bodies.sort(key=lambda item: (item.left, item.top))
    return titles, subtitles, bodies


def _set_text_frame(
    shape: Any,
    paragraphs: list[str],
    *,
    font_name: str | None = None,
    font_size: float | None = None,
    color: str | None = None,
    bold_first: bool = False,
    bullets: bool = False,
    min_font_size: float = 9.0,
) -> None:
    existing_style = _shape_text_style(shape)
    if paragraphs and not str(shape.name).startswith("BrandDeck"):
        shape.name = f"BrandDeck Text {shape.shape_id}"
    font_name = font_name or existing_style.get("font_name")
    combined_text = "\n".join(str(paragraph) for paragraph in paragraphs)
    font_name = readable_font(font_name, combined_text)
    if font_name is None and any("\u0400" <= char <= "\u052f" for char in combined_text):
        font_name = _CYRILLIC_THEME_FALLBACK.get()
    font_size = (
        font_size or existing_style.get("font_size") or effective_font_size(shape)
    )
    color = color or existing_style.get("color")
    alignment = existing_style.get("alignment")
    inherited_bold = existing_style.get("bold")
    if font_size and paragraphs:
        frame = shape.text_frame
        width = max(
            0.01, (shape.width - frame.margin_left - frame.margin_right) / 914400
        )
        height = max(
            0.01, (shape.height - frame.margin_top - frame.margin_bottom) / 914400
        )
        fitted_size = float(font_size)
        while fitted_size > min_font_size:
            chars_per_line = max(
                1, int(width * 72 / (fitted_size * ESTIMATED_GLYPH_WIDTH_EM))
            )
            line_count = sum(
                estimated_line_count(str(text), chars_per_line)
                for paragraph_text in paragraphs
                for text in str(paragraph_text).split("\n")
            )
            paragraph_gap = max(0, len(paragraphs) - 1) * fitted_size * 0.12 / 72
            required_height = (
                line_count * fitted_size * 1.22 / 72 + paragraph_gap + 0.06
            )
            if required_height <= height:
                break
            fitted_size = max(min_font_size, fitted_size - 0.5)
        font_size = _snap_size(max(min_font_size, fitted_size), min_font_size)
    text_frame = shape.text_frame
    text_frame.clear()
    # The frame may inherit East Asian vertical writing from its layout.
    # Generated Cyrillic/Latin paragraphs always use horizontal writing.
    if paragraphs:
        text_frame._txBody.bodyPr.set("vert", "horz")
    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    text_frame.vertical_anchor = MSO_ANCHOR.TOP
    if not paragraphs:
        return
    # clear() keeps the exemplar's first paragraph properties (bullet, hanging
    # indent); new paragraphs repeat them instead of falling back to the
    # layout's list level, which drops the marker and shifts the indent.
    first_properties = text_frame.paragraphs[0]._p.pPr
    for index, text in enumerate(paragraphs):
        paragraph = (
            text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        )
        if index and first_properties is not None:
            if paragraph._p.pPr is not None:
                paragraph._p.remove(paragraph._p.pPr)
            paragraph._p.insert(0, copy.deepcopy(first_properties))
        paragraph.text = str(text)
        paragraph.level = 0
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(float(font_size or 18) * 0.15)
        paragraph.line_spacing = 1.15
        if not bullets:
            paragraph.alignment = alignment or PP_ALIGN.LEFT
        # A hard line break creates a second run in python-pptx. Style every
        # run or that line inherits the exemplar's original (often huge) size.
        for run in paragraph.runs:
            font = run.font
            if font_name:
                font.name = font_name
            if font_size:
                font.size = Pt(font_size)
            if color:
                font.color.rgb = _rgb(color)
            elif existing_style.get("color_xml") is not None:
                font._element.append(copy.deepcopy(existing_style["color_xml"]))
            if inherited_bold is not None:
                font.bold = inherited_bold
            if bold_first and index == 0:
                font.bold = True


def _add_textbox(
    slide: Any,
    text: str,
    box: tuple[float, float, float, float],
    *,
    font_name: str,
    font_size: float,
    color: str,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    min_font_size: float = 9.0,
    vertical_anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
) -> Any:
    x, y, w, h = box
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    _set_text_frame(
        shape,
        [text],
        font_name=font_name,
        font_size=font_size,
        color=color,
        bold_first=bold,
        min_font_size=min_font_size,
    )
    shape.text_frame.paragraphs[0].alignment = align
    shape.text_frame.margin_left = Inches(0.04)
    shape.text_frame.margin_right = Inches(0.04)
    shape.text_frame.margin_top = Inches(0.02)
    shape.text_frame.margin_bottom = Inches(0.02)
    shape.text_frame.vertical_anchor = vertical_anchor
    return shape


def _zones_intersect(
    first: tuple[float, float, float, float], second: tuple[float, float, float, float]
) -> bool:
    x1, y1, w1, h1 = first
    x2, y2, w2, h2 = second
    return min(x1 + w1, x2 + w2) > max(x1, x2) and min(y1 + h1, y2 + h2) > max(y1, y2)


def _zone_from_shape(shape: Any) -> tuple[float, float, float, float]:
    return (
        shape.left / 914400,
        shape.top / 914400,
        shape.width / 914400,
        shape.height / 914400,
    )


def _content_zone(
    slide: Any,
    design: dict[str, Any],
    title_shapes: list[Any],
) -> tuple[float, float, float, float]:
    canvas = design["canvas"]
    spacing = design["spacing"]
    left = float(spacing.get("typical_left_margin_inches", 0.6))
    right = float(spacing.get("typical_right_margin_inches", 0.6))
    title_bottom = 1.2
    if title_shapes:
        title_bottom = (
            max((item.top + item.height) / 914400 for item in title_shapes) + 0.15
        )
    return (
        left,
        title_bottom,
        max(2.0, float(canvas["width_inches"]) - left - right),
        max(1.0, float(canvas["height_inches"]) - title_bottom - 0.45),
    )


def _split_zone(
    zone: tuple[float, float, float, float], ratio: float = 0.44
) -> tuple[tuple[float, float, float, float], tuple[float, float, float, float]]:
    x, y, w, h = zone
    gap = 0.28
    left_w = max(1.0, w * ratio - gap / 2)
    return (x, y, left_w, h), (x + left_w + gap, y, max(1.0, w - left_w - gap), h)


def _distribute_paragraphs(
    shapes: list[Any], paragraphs: list[str], default_size: float
) -> list[list[str]]:
    """Partition consecutive paragraphs by available space, retaining their order.

    Minimize the worst estimated occupancy across slots. Empty assignments are
    allowed: a caption-sized placeholder need not receive body text. This is a
    layout estimate; the finished deck still needs structural and rendered QA.
    """

    def occupancy(shape: Any, texts: list[str]) -> float:
        if not texts:
            return 0.0
        frame = shape.text_frame
        width = max(
            0.01, (shape.width - frame.margin_left - frame.margin_right) / 914400
        )
        height = max(
            0.01, (shape.height - frame.margin_top - frame.margin_bottom) / 914400
        )
        size = float(effective_font_size(shape, default_size))
        chars = max(1, int(width * 72 / (size * ESTIMATED_GLYPH_WIDTH_EM)))
        lines = sum(
            estimated_line_count(line, chars)
            for text in texts
            for line in text.split("\n")
        )
        return (lines * size * 1.22 + max(0, len(texts) - 1) * size * 0.12) / (
            72 * height
        )

    if len(paragraphs) <= len(shapes) and all(
        occupancy(shape, [text]) <= 1.0 for shape, text in zip(shapes, paragraphs)
    ):
        return [
            [paragraphs[i]] if i < len(paragraphs) else [] for i in range(len(shapes))
        ]

    # state: number of assigned paragraphs -> (worst occupancy, partitions)
    states: dict[int, tuple[float, list[list[str]]]] = {0: (0.0, [])}
    for shape in shapes:
        following: dict[int, tuple[float, list[list[str]]]] = {}
        for assigned, (score, groups) in states.items():
            for end in range(assigned, len(paragraphs) + 1):
                group = paragraphs[assigned:end]
                candidate = max(score, occupancy(shape, group))
                if end not in following or candidate < following[end][0]:
                    following[end] = (candidate, groups + [group])
        states = following
    return states[len(paragraphs)][1]


def _whole_visual_zone(
    slide: Any,
    body_shapes: list[Any],
    zone: tuple[float, float, float, float],
    design: dict[str, Any],
) -> tuple[float, float, float, float]:
    """Give a diagram that carries all slide text the template's body area.

    Source text in exemplar bodies is cleared, and emptied body placeholders
    are removed so PowerPoint shows no prompt text under the diagram.
    """
    width = float(design["canvas"]["width_inches"])
    height = float(design["canvas"]["height_inches"])
    titles = [
        _zone_from_shape(shape)
        for shape in slide.placeholders
        if _placeholder_type(shape) in TITLE_TYPES
    ]
    usable = [
        shape
        for shape in body_shapes
        if shape.width >= Inches(1.0)
        and shape.height >= Inches(0.65)
        and shape.left >= 0
        # A column beside the title (title left, list right) starts at the
        # title's height; only a body under the title must start below it.
        and (
            shape.top >= Inches(zone[1] - 0.05)
            or not any(_zones_intersect(_zone_from_shape(shape), t) for t in titles)
        )
        and shape.left + shape.width <= Inches(width + 0.02)
        and shape.top + shape.height <= Inches(height + 0.02)
    ]
    result = zone
    if usable:
        left = min(shape.left for shape in usable) / 914400
        top = min(shape.top for shape in usable) / 914400
        right = max(shape.left + shape.width for shape in usable) / 914400
        bottom = max(shape.top + shape.height for shape in usable) / 914400
        # A narrow caption column would cramp a diagram: then use the whole
        # content area below the title instead of the template's text box.
        body_area = (right - left) * (bottom - top)
        if (
            right - left >= 2.8
            and bottom - top >= 1.2
            and body_area >= 0.45 * zone[2] * zone[3]
        ):
            # A short strip of exemplar labels keeps the template's content
            # width but extends down to the bottom of the content area.
            bottom = (
                max(bottom, zone[1] + zone[3])
                if bottom - top < zone[3] * 0.5
                else bottom
            )
            result = (left, top, right - left, bottom - top)
    # Source exemplars may place text into the footer. A new image or diagram
    # must stay inside the content area even when it inherits that tall slot.
    content_bottom = zone[1] + zone[3]
    if result[1] + result[3] > content_bottom:
        available = content_bottom - 0.08 - result[1]
        result = (
            (result[0], result[1], result[2], available)
            if available >= 0.8
            else zone
        )
    for shape in body_shapes:
        _set_text_frame(shape, [])
        if getattr(shape, "is_placeholder", False):
            shape.element.getparent().remove(shape.element)
    return result


def _add_body_text(
    slide: Any,
    slide_spec: dict[str, Any],
    body_shapes: list[Any],
    zone: tuple[float, float, float, float],
    design: dict[str, Any],
) -> tuple[float, float, float, float]:
    font = design["typography"]["primary_font"]
    size = float(design["typography"].get("body_size_pt", 18))
    colors = _theme_colors(design)
    body = slide_spec.get("body", "")
    bullets = slide_spec.get("bullets", [])
    visual = slide_spec.get("visual")
    if visual and visual.get("type") in NATIVE_VISUALS and not body and not bullets:
        return _whole_visual_zone(slide, body_shapes, zone, design)

    # Slide exemplars without native placeholders often contain many tiny
    # source-specific labels. Clear those labels and write new content into a
    # safe zone while retaining the exemplar's background and decoration.
    usable_shapes = [
        s
        for s in body_shapes
        if s.width >= Inches(1.0)
        and s.height >= Inches(0.65)
        and s.left >= 0
        and s.top >= 0
        and s.left + s.width <= Inches(design["canvas"]["width_inches"] + 0.02)
        and s.top + s.height <= Inches(design["canvas"]["height_inches"] + 0.02)
    ]
    if len(usable_shapes) >= 2:
        for shape in body_shapes:
            if shape not in usable_shapes:
                _set_text_frame(shape, [])
        body_shapes = usable_shapes
    if (
        body_shapes
        and len(usable_shapes) < 2
        and not any(getattr(shape, "is_placeholder", False) for shape in body_shapes)
    ):
        bullet_like = [
            shape
            for shape in body_shapes
            if shape.text.lstrip().startswith(("•", "-", "–", "—"))
        ]
        preferred = bullet_like if len(bullet_like) >= 2 else body_shapes
        if not bullet_like and len(preferred) > 2:
            preferred = sorted(
                preferred,
                key=lambda shape: shape.width * shape.height,
                reverse=True,
            )[: max(2, math.ceil(len(preferred) * 0.5))]
        style_source = max(
            preferred,
            key=lambda shape: shape.width * shape.height,
            default=body_shapes[0],
        )
        source_style = _shape_text_style(style_source)
        left = min(shape.left for shape in preferred) / 914400
        top = max(min(shape.top for shape in preferred) / 914400, zone[1])
        right = max(shape.left + shape.width for shape in preferred) / 914400
        bottom = max(shape.top + shape.height for shape in preferred) / 914400
        source_zone = (left, top, right - left, bottom - top)
        if source_zone[2] >= 2.8 and source_zone[3] >= 0.75:
            zone = source_zone
        elif zone[0] < left < zone[0] + zone[2] - 2.8:
            # Keep the exemplar's text indent: the column left of it usually
            # holds list markers or icons that stay on the slide.
            zone = (left, zone[1], zone[0] + zone[2] - left, zone[3])
        for shape in body_shapes:
            _set_text_frame(shape, [])
        if visual:
            text_zone, visual_zone = _split_zone(zone)
        else:
            text_zone, visual_zone = zone, zone
        paragraphs = ([body] if body else []) + list(bullets)
        if paragraphs:
            text_shape = _add_textbox(
                slide,
                "",
                text_zone,
                font_name=source_style.get("font_name", font),
                font_size=float(source_style.get("font_size", size)),
                color=source_style.get("color", colors.get("dk1", "1F2937")),
            )
            text_shape.name = "BrandDeck Body"
            _set_text_frame(
                text_shape,
                [f"• {item}" if bullets else item for item in paragraphs],
                font_name=source_style.get("font_name", font),
                font_size=float(source_style.get("font_size", size)),
                color=source_style.get("color", colors.get("dk1", "1F2937")),
            )
        return visual_zone

    if len(body_shapes) >= 2 and visual and visual.get("type") in NATIVE_VISUALS:
        # Pictures and diagrams need a coherent half of the body area; the
        # remaining small exemplar frames would make the text unreadable.
        paragraphs = ([body] if body else []) + list(bullets)
        style_source = max(body_shapes, key=lambda shape: shape.width * shape.height)
        source_style = _shape_text_style(style_source)
        area = _whole_visual_zone(slide, body_shapes, zone, design)
        text_zone, visual_zone = _split_zone(area)
        text_shape = _add_textbox(
            slide,
            "",
            text_zone,
            font_name=source_style.get("font_name", font),
            font_size=float(source_style.get("font_size", size)),
            color=source_style.get("color", colors.get("dk1", "1F2937")),
        )
        text_shape.name = "BrandDeck Body"
        _set_text_frame(
            text_shape,
            [f"• {item}" if bullets else item for item in paragraphs],
            font_name=source_style.get("font_name", font),
            font_size=max(size, float(source_style.get("font_size", size))),
            color=source_style.get("color", colors.get("dk1", "1F2937")),
        )
        return visual_zone

    if len(body_shapes) >= 2:
        text_shapes = sorted(
            body_shapes, key=lambda s: (round(s.top / 914400, 1), s.left)
        )
        visual_zone = zone
        if visual:
            # Reserve a separate slot; never draw a chart over assigned text.
            visual_shape = max(
                text_shapes, key=lambda shape: shape.width * shape.height
            )
            text_shapes.remove(visual_shape)
            visual_zone = _zone_from_shape(visual_shape)
            _set_text_frame(visual_shape, [])
        paragraphs = ([body] if body else []) + list(bullets)
        groups = _distribute_paragraphs(text_shapes, paragraphs, size)
        for shape, paragraphs in zip(text_shapes, groups):
            _set_text_frame(
                shape,
                paragraphs,
                bullets=bool(bullets),
                font_size=max(size, effective_font_size(shape, size)),
            )
        return visual_zone

    if body_shapes:
        body_zone = _zone_from_shape(body_shapes[0])
        if visual:
            text_zone, visual_zone = _split_zone(body_zone)
            body_shapes[0].left = Inches(text_zone[0])
            body_shapes[0].top = Inches(text_zone[1])
            body_shapes[0].width = Inches(text_zone[2])
            body_shapes[0].height = Inches(text_zone[3])
        else:
            visual_zone = body_zone
        paragraphs = ([body] if body else []) + list(bullets)
        _set_text_frame(
            body_shapes[0],
            paragraphs,
            bullets=bool(bullets),
            font_size=max(size, effective_font_size(body_shapes[0], size)),
        )
        return visual_zone

    if visual:
        text_zone, visual_zone = _split_zone(zone)
    else:
        text_zone, visual_zone = zone, zone
    paragraphs = ([body] if body else []) + list(bullets)
    if paragraphs:
        text_shape = _add_textbox(
            slide,
            "",
            text_zone,
            font_name=font,
            font_size=size,
            color=colors.get("dk1", "1F2937"),
        )
        text_shape.name = "BrandDeck Body"
        _set_text_frame(
            text_shape,
            [f"• {item}" if bullets else item for item in paragraphs],
            font_name=font,
            font_size=size,
            color=colors.get("dk1", "1F2937"),
        )
    return visual_zone


def _add_metric_cards(
    slide: Any,
    visual: dict[str, Any],
    zone: tuple[float, float, float, float],
    design: dict[str, Any],
) -> None:
    items = [item for item in visual.get("items", []) if isinstance(item, dict)][:6]
    if not items:
        return
    colors = _theme_colors(design)
    accents = [colors.get("accent1"), colors.get("accent2"), colors.get("accent3")]
    x, y, w, h = zone
    columns = 2 if len(items) <= 4 else 3
    rows = math.ceil(len(items) / columns)
    gap = 0.16
    card_w = (w - gap * (columns - 1)) / columns
    card_h = min(1.55, (h - gap * (rows - 1)) / rows)
    font = design["typography"]["primary_font"]
    for index, item in enumerate(items):
        col, row = index % columns, index // columns
        cx, cy = x + col * (card_w + gap), y + row * (card_h + gap)
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(cx),
            Inches(cy),
            Inches(card_w),
            Inches(card_h),
        )
        card.name = "BrandDeck Metric Card"
        card.fill.solid()
        card.fill.fore_color.rgb = _rgb(colors.get("lt2"), "F3F4F6")
        card.line.color.rgb = _rgb(accents[index % len(accents)], "4472C4")
        card.line.width = Pt(1.5)
        _add_textbox(
            slide,
            str(item.get("value", "")),
            (cx + 0.12, cy + 0.18, card_w - 0.24, card_h * 0.44),
            font_name=font,
            font_size=max(20, design["typography"].get("title_size_pt", 30) * 0.72),
            color=accents[index % len(accents)],
            bold=True,
        ).name = "BrandDeck Metric Value"
        _add_textbox(
            slide,
            str(item.get("label", "")),
            (cx + 0.12, cy + card_h * 0.58, card_w - 0.24, card_h * 0.28),
            font_name=font,
            font_size=max(10, design["typography"].get("body_size_pt", 16) * 0.8),
            color=colors.get("dk1", "1F2937"),
        ).name = "BrandDeck Metric Label"


def _add_bar_chart(
    slide: Any,
    visual: dict[str, Any],
    zone: tuple[float, float, float, float],
    design: dict[str, Any],
) -> None:
    categories = [str(item) for item in visual.get("categories", [])][:12]
    series = [item for item in visual.get("series", []) if isinstance(item, dict)][:5]
    if not categories or not series:
        return
    chart_data = CategoryChartData()
    chart_data.categories = categories
    for item in series:
        values = [_number(value) for value in item.get("values", [])[: len(categories)]]
        values.extend([0.0] * (len(categories) - len(values)))
        chart_data.add_series(str(item.get("name", "Series")), values)
    x, y, w, h = _chart_zone(slide, zone, design)
    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
        chart_data,
    )
    chart_shape.name = "BrandDeck Bar Chart"
    chart = chart_shape.chart
    chart.has_legend = len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    chart.value_axis.has_major_gridlines = True
    # Appendix 1: a chart needs its values and units, not only bar heights.
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.show_value = True
    plot.data_labels.font.size = Pt(_snap_size(11, 9))
    unit = str(visual.get("unit") or "").strip()
    if unit:
        chart.value_axis.has_title = True
        chart.value_axis.axis_title.text_frame.text = unit[:40]
    colors = _theme_colors(design)
    palette = [colors.get("accent1"), colors.get("accent2"), colors.get("accent3")]
    for index, chart_series in enumerate(chart.series):
        chart_series.format.fill.solid()
        chart_series.format.fill.fore_color.rgb = _rgb(
            palette[index % len(palette)], "4472C4"
        )


def _add_pie_chart(
    slide: Any,
    visual: dict[str, Any],
    zone: tuple[float, float, float, float],
    design: dict[str, Any],
) -> None:
    categories = [str(item) for item in visual.get("categories", [])][:10]
    values = [_number(value) for value in visual.get("values", [])[: len(categories)]]
    if not categories or len(values) != len(categories):
        return
    chart_data = ChartData()
    chart_data.categories = categories
    chart_data.add_series("Доли", values)
    x, y, w, h = _chart_zone(slide, zone, design)
    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.PIE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
        chart_data,
    )
    chart_shape.name = "BrandDeck Pie Chart"
    chart = chart_shape.chart
    # LibreOffice displays the single series name as an automatic chart title
    # unless autoTitleDeleted is set explicitly.
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(12)
    chart.plots[0].has_data_labels = True
    labels = chart.plots[0].data_labels
    labels.show_value = False
    labels.show_percentage = True
    labels.font.size = Pt(12)


def _chart_zone(
    slide: Any,
    zone: tuple[float, float, float, float],
    design: dict[str, Any],
) -> tuple[float, float, float, float]:
    """Use the content area when a chart-only slide inherited a narrow slot."""
    canvas = design["canvas"]
    canvas_area = float(canvas["width_inches"]) * float(canvas["height_inches"])
    if zone[2] >= float(canvas["width_inches"]) * 0.42 and area(zone) >= canvas_area * 0.24:
        return zone
    text_shapes = [shape for shape in slide.shapes if has_text(shape)]
    if any(
        shape.name == "BrandDeck Body"
        or (shape.top + shape.height) / 914400 > zone[1] + 0.05
        for shape in text_shapes
    ):
        return zone
    candidate = _content_zone(slide, design, text_shapes)
    candidate = _clear_layout_art(candidate, layout_art_boxes(slide, design))
    if area(candidate) > area(zone) * 1.4 and candidate[2] >= 4.0:
        return candidate
    return zone


def _add_table(
    slide: Any,
    visual: dict[str, Any],
    zone: tuple[float, float, float, float],
    design: dict[str, Any],
) -> None:
    headers = [str(value) for value in visual.get("headers", [])][:8]
    rows = [row for row in visual.get("rows", []) if isinstance(row, list)][:10]
    if not headers:
        return
    rows = [[str(value) for value in row[: len(headers)]] for row in rows]
    x, y, w, h = zone
    body_size = float(design["typography"].get("body_size_pt", 16))
    values = [headers] + rows
    column_w = w / len(headers) - 0.12

    def row_heights(size: float) -> list[float]:
        chars = max(1, int(column_w * 72 / (size * ESTIMATED_GLYPH_WIDTH_EM)))
        return [
            max(estimated_line_count(value, chars) for value in row) * size * 1.3 / 72
            + 0.16
            for row in values
        ]

    # Rows take the height of their text instead of stretching over the
    # zone (tall empty cells); the size steps down only when rows overflow.
    size = max(11.0, body_size * 0.85)
    while sum(row_heights(size)) > h and size > 9:
        size -= 0.5
    size = _snap_size(size, 9)
    heights = row_heights(size)
    canvas = design.get("canvas", {})
    canvas_area = float(canvas.get("width_inches", 0) or 0) * float(
        canvas.get("height_inches", 0) or 0
    )
    # Rows sized to a few words would leave the slide below a quarter filled
    # (Appendix 1); then, or after such a render, the rows share the zone.
    sparse = w * sum(heights) < FILL_SHARE * canvas_area
    if (visual.get("fill_zone") or sparse) and sum(heights) < h:
        extra = (h - sum(heights)) / len(heights)
        heights = [height + extra for height in heights]
    table_shape = slide.shapes.add_table(
        len(rows) + 1,
        len(headers),
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(min(h, sum(heights))),
    )
    table_shape.name = "BrandDeck Table"
    table = table_shape.table
    for row, height in zip(table.rows, heights):
        row.height = Inches(height)
    colors = _theme_colors(design)
    background = _hex(design.get("brand", {}).get("background"), "FFFFFF")
    body_fill = (
        _mix_color(background, "FFFFFF", 0.1)
        if _luminance(background) < 0.35
        else _hex(colors.get("lt2"), "F3F4F6")
    )
    header_fill, header_text = on_fill(colors.get("accent1", "4472C4"))
    body_text = max(("FFFFFF", "111111"), key=lambda color: contrast(color, body_fill))
    if contrast(body_text, body_fill) < 4.5:
        body_fill, body_text = background, max(
            ("FFFFFF", "111111"), key=lambda color: contrast(color, background)
        )
    font = design["typography"]["primary_font"]
    for row_index, row in enumerate(values):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = value
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(header_fill if row_index == 0 else body_fill)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.LEFT
                for run in paragraph.runs:
                    run.font.name = font
                    run.font.size = Pt(size)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = _rgb(header_text if row_index == 0 else body_text)


def _add_timeline(
    slide: Any,
    visual: dict[str, Any],
    zone: tuple[float, float, float, float],
    design: dict[str, Any],
) -> None:
    items = [item for item in visual.get("items", []) if isinstance(item, dict)][:6]
    if not items:
        return
    colors = _theme_colors(design)
    font = design["typography"]["primary_font"]
    x, y, w, h = zone
    center_y = y + h * 0.43
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(center_y), Inches(w), Inches(0.035)
    )
    line.name = "BrandDeck Timeline Line"
    line.fill.solid()
    line.fill.fore_color.rgb = _rgb(colors.get("accent1"), "4472C4")
    line.line.fill.background()
    step = w / max(1, len(items))
    for index, item in enumerate(items):
        cx = x + step * index + step / 2
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(cx - 0.11),
            Inches(center_y - 0.09),
            Inches(0.22),
            Inches(0.22),
        )
        dot.name = "BrandDeck Timeline Point"
        dot.fill.solid()
        dot.fill.fore_color.rgb = _rgb(colors.get("accent2"), "70AD47")
        dot.line.fill.background()
        text_y = center_y - 0.72 if index % 2 == 0 else center_y + 0.24
        _add_textbox(
            slide,
            str(item.get("label", "")),
            (cx - step * 0.43, text_y, step * 0.86, 0.28),
            font_name=font,
            font_size=max(10, design["typography"].get("body_size_pt", 16) * 0.75),
            color=colors.get("accent1", "4472C4"),
            bold=True,
            align=PP_ALIGN.CENTER,
        ).name = "BrandDeck Timeline Label"
        _add_textbox(
            slide,
            str(item.get("detail", "")),
            (cx - step * 0.43, text_y + 0.3, step * 0.86, 0.55),
            font_name=font,
            font_size=max(8, design["typography"].get("body_size_pt", 16) * 0.58),
            color=colors.get("dk1", "1F2937"),
            align=PP_ALIGN.CENTER,
        ).name = "BrandDeck Timeline Detail"


def _record(context: dict[str, Any] | None, record: dict[str, Any]) -> None:
    if context is not None:
        context.setdefault("records", []).append(
            {"slide": context.get("slide"), **record}
        )


def _visual_style(
    design: dict[str, Any], context: dict[str, Any] | None
) -> dict[str, Any]:
    context = context or {}
    style = diagram_style(
        design, context.get("background"), str(context.get("role", "content"))
    )
    canvas = design.get("canvas", {})
    style["canvas_area"] = float(canvas.get("width_inches", 0) or 0) * float(
        canvas.get("height_inches", 0) or 0
    )
    return style


def _asset_path(
    visual: dict[str, Any], context: dict[str, Any] | None
) -> tuple[Path | None, dict[str, Any]]:
    meta = ((context or {}).get("assets") or {}).get(str(visual.get("asset_id") or ""))
    if not isinstance(meta, dict):
        return None, {}
    path = Path(str(meta.get("path", "")))
    if path.suffix.lower() not in {".png", ".jpg", ".jpeg"} or not path.is_file():
        return None, meta
    return path, meta


def _set_picture_description(picture: Any, text: str) -> None:
    properties = picture._element.xpath("./p:nvPicPr/p:cNvPr")
    if properties and text:
        properties[0].set("descr", text[:500])


def _crop_picture(picture: Any, box_w: float, box_h: float) -> None:
    try:
        with Image.open(io.BytesIO(picture.image.blob)) as source:
            left, top, right, bottom = crop_to_fill_salient(
                source, (box_w, box_h)
            )
    except (AttributeError, OSError, ValueError):
        left, top, right, bottom = crop_to_fill(picture.image.size, (box_w, box_h))
    picture.crop_left, picture.crop_top = left, top
    picture.crop_right, picture.crop_bottom = right, bottom


def _add_image(
    slide: Any,
    visual: dict[str, Any],
    zone: tuple[float, float, float, float],
    design: dict[str, Any],
    context: dict[str, Any] | None,
    *,
    rounded: bool = False,
) -> dict[str, Any]:
    """Place an asset cropped to fill ``zone``; draw an illustration otherwise."""
    path, meta = _asset_path(visual, context)
    if path is None:
        record = render_illustration(
            slide,
            zone,
            _visual_style(design, context),
            text=str((context or {}).get("text", "")),
        )
        record["requested"] = visual.get("asset_id") or visual.get("request") or "image"
        return record
    x, y, w, h = zone
    caption = str(visual.get("caption") or "").strip()
    caption_h = 0.36 if caption and h > 1.6 else 0.0
    picture = slide.shapes.add_picture(
        str(path), Inches(x), Inches(y), Inches(w), Inches(h - caption_h)
    )
    picture.name = "BrandDeck Image"
    _crop_picture(picture, w, h - caption_h)
    _set_picture_description(picture, meta.get("label") or caption)
    if rounded:
        picture.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
    if caption_h:
        style = _visual_style(design, context)
        _add_textbox(
            slide,
            caption,
            (x, y + h - caption_h + 0.04, w, caption_h - 0.04),
            font_name=style["font"],
            font_size=12,
            color=style["text"],
            min_font_size=10,
        ).name = "BrandDeck Image Caption"
    return {
        "type": "image",
        "asset_id": meta.get("asset_id"),
        "source": meta.get("source"),
        "match": visual.get("match"),
        "placement": "zone",
    }


def _fill_picture_slot(
    slide: Any,
    slot: Any,
    visual: dict[str, Any] | None,
    design: dict[str, Any],
    context: dict[str, Any] | None,
) -> dict[str, Any]:
    """Reuse the template's own picture position, geometry and z-order."""
    box = _zone_from_shape(slot)
    path, meta = _asset_path(visual or {}, context)
    is_placeholder = bool(getattr(slot, "is_placeholder", False))
    if path is not None and is_placeholder:
        picture = slot.insert_picture(str(path))
        picture.name = "BrandDeck Image"
        _crop_picture(picture, box[2], box[3])
        _set_picture_description(picture, meta.get("label", ""))
        return {
            "type": "image",
            "asset_id": meta.get("asset_id"),
            "source": meta.get("source"),
            "match": (visual or {}).get("match"),
            "placement": "template_placeholder",
        }
    if path is not None:
        source_geometry = slot.element.find(f"{qn('p:spPr')}/{qn('a:prstGeom')}")
        x, y, w, h = box
        if source_geometry is None or source_geometry.get("prst") == "rect":
            # A rectangular bleed is invisible anyway; clip it to the canvas so
            # the new picture stays inside the slide.
            width = float(design["canvas"]["width_inches"])
            height = float(design["canvas"]["height_inches"])
            left, top = max(0.0, x), max(0.0, y)
            x, y, w, h = left, top, min(width, x + w) - left, min(height, y + h) - top
        picture = slide.shapes.add_picture(
            str(path), Inches(x), Inches(y), Inches(w), Inches(h)
        )
        picture.name = "BrandDeck Image"
        _crop_picture(picture, w, h)
        _set_picture_description(picture, meta.get("label", ""))
        target_geometry = picture._element.spPr.find(qn("a:prstGeom"))
        if source_geometry is not None and target_geometry is not None:
            target_geometry.addprevious(copy.deepcopy(source_geometry))
            picture._element.spPr.remove(target_geometry)
        slot.element.addprevious(picture._element)
        slot.element.getparent().remove(slot.element)
        return {
            "type": "image",
            "asset_id": meta.get("asset_id"),
            "source": meta.get("source"),
            "match": (visual or {}).get("match"),
            "placement": "template_slot",
        }
    slot.element.getparent().remove(slot.element)
    record = render_illustration(
        slide,
        box,
        _visual_style(design, context),
        text=str((context or {}).get("text", "")),
    )
    record["placement"] = "template_slot"
    return record


def _add_visual(
    slide: Any,
    visual: dict[str, Any] | None,
    zone: tuple[float, float, float, float],
    design: dict[str, Any],
    context: dict[str, Any] | None = None,
    *,
    rounded: bool = False,
) -> None:
    if not visual:
        return
    visual_type = visual.get("type")
    if visual_type == "image":
        _record(
            context, _add_image(slide, visual, zone, design, context, rounded=rounded)
        )
        return
    if visual_type in DIAGRAM_TYPES:
        try:
            style = _visual_style(design, context)
            if visual.get("fill_zone"):
                style = {**style, "fill_zone": True}
            record = render_diagram(
                slide,
                visual,
                zone,
                style,
                context=str((context or {}).get("title", "")),
            )
        except ValueError as exc:
            record = {"type": visual_type, "error": str(exc)[:200]}
        for key in ("origin", "estimated_fill"):
            if key in visual:
                record[key] = visual[key]
        _record(context, record)
        return
    if visual_type == "metric_cards":
        _add_metric_cards(slide, visual, zone, design)
    elif visual_type == "bar_chart":
        _add_bar_chart(slide, visual, zone, design)
    elif visual_type == "pie_chart":
        _add_pie_chart(slide, visual, zone, design)
    elif visual_type == "table":
        _add_table(slide, visual, zone, design)
    elif visual_type == "timeline":
        _add_timeline(slide, visual, zone, design)


def _picture_hash(shape: Any) -> str:
    try:
        return str(shape.image.sha1)
    except (AttributeError, ValueError):
        return ""


def _select_brand_source(
    source_slides: list[Any], repeated_image_hashes: set[str]
) -> Any | None:
    if not source_slides:
        return None
    start = max(1, len(source_slides) // 3)
    candidates = source_slides[start:-1] if len(source_slides) > 2 else source_slides
    candidates = candidates or source_slides

    def score(slide: Any) -> tuple[int, float]:
        count = 0
        area = 0.0
        for shape in slide.shapes:
            if (
                shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                and _picture_hash(shape) in repeated_image_hashes
            ):
                count += 1
                area += shape.width * shape.height
        return count, area

    return max(candidates, key=score, default=source_slides[0])


def _copy_brand_canvas(
    slide: Any,
    source: Any | None,
    repeated_image_hashes: set[str],
    design: dict[str, Any],
    tokens: dict[str, Any],
) -> list[Any]:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(tokens["background"])
    if source is None:
        return []
    canvas = design["canvas"]
    canvas_area = max(
        0.1, float(canvas["width_inches"]) * float(canvas["height_inches"])
    )
    background_candidates: list[tuple[float, Any]] = []
    brand_candidates: list[Any] = []
    for source_shape in source.shapes:
        if source_shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        image_hash = _picture_hash(source_shape)
        if not image_hash or image_hash not in repeated_image_hashes:
            continue
        area_ratio = (
            source_shape.width / 914400 * (source_shape.height / 914400) / canvas_area
        )
        width_ratio = source_shape.width / 914400 / float(canvas["width_inches"])
        height_ratio = source_shape.height / 914400 / float(canvas["height_inches"])
        if area_ratio >= 0.72 and width_ratio >= 0.84 and height_ratio >= 0.84:
            background_candidates.append((area_ratio, source_shape))
        elif area_ratio <= 0.08:
            brand_candidates.append(source_shape)

    if background_candidates:
        _, background_shape = max(background_candidates, key=lambda item: item[0])
        try:
            _set_picture_background(slide, background_shape.image.blob)
        except (AttributeError, ValueError):
            pass

    copied: list[Any] = []
    for source_shape in brand_candidates:
        try:
            picture = slide.shapes.add_picture(
                io.BytesIO(source_shape.image.blob),
                source_shape.left,
                source_shape.top,
                source_shape.width,
                source_shape.height,
            )
        except (AttributeError, ValueError):
            continue
        picture.name = "BrandDeck Brand Asset"
        for crop_name in ("crop_left", "crop_right", "crop_top", "crop_bottom"):
            try:
                setattr(picture, crop_name, getattr(source_shape, crop_name))
            except (AttributeError, ValueError):
                pass
        copied.append(picture)
    return copied


def _set_picture_background(slide: Any, blob: bytes) -> None:
    """Store a raster as the OOXML slide background, not as a selectable shape."""
    with Image.open(io.BytesIO(blob)) as source:
        converted = source.convert("RGB")
        payload = io.BytesIO()
        converted.save(payload, format="PNG")
    payload.seek(0)
    _, relationship_id = slide.part.get_or_add_image_part(payload)
    common_slide_data = slide.element.cSld
    existing = common_slide_data.find(qn("p:bg"))
    if existing is not None:
        common_slide_data.remove(existing)

    background = OxmlElement("p:bg")
    properties = OxmlElement("p:bgPr")
    blip_fill = OxmlElement("a:blipFill")
    blip_fill.set("dpi", "0")
    blip_fill.set("rotWithShape", "1")
    blip = OxmlElement("a:blip")
    blip.set(qn("r:embed"), relationship_id)
    stretch = OxmlElement("a:stretch")
    stretch.append(OxmlElement("a:fillRect"))
    blip_fill.append(blip)
    blip_fill.append(stretch)
    properties.append(blip_fill)
    properties.append(OxmlElement("a:effectLst"))
    background.append(properties)
    common_slide_data.insert(0, background)


def _native_right_limit(
    brand_shapes: list[Any], design: dict[str, Any], *, default_margin: float = 0.65
) -> float:
    canvas = design["canvas"]
    width = float(canvas["width_inches"])
    height = float(canvas["height_inches"])
    area = width * height
    right = width - default_margin
    for shape in brand_shapes:
        area_ratio = shape.width / 914400 * (shape.height / 914400) / area
        left = shape.left / 914400
        if 0.04 <= area_ratio <= 0.4 and left > width * 0.55:
            right = min(right, left - 0.28)
    return max(width * 0.62, right)


def _native_panel(
    slide: Any,
    box: tuple[float, float, float, float],
    tokens: dict[str, Any],
    *,
    alternate: bool = False,
) -> Any:
    x, y, w, h = box
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    panel.name = "BrandDeck Native Panel"
    panel.fill.solid()
    panel.fill.fore_color.rgb = _rgb(
        tokens["surface_alt"] if alternate else tokens["surface"]
    )
    panel.line.color.rgb = _rgb(tokens["accent"])
    panel.line.width = Pt(1.15 if not tokens["dark"] else 0.8)
    try:
        panel.adjustments[0] = 0.08
    except (IndexError, ValueError):
        pass
    return panel


def _native_card_content(
    slide: Any,
    item: str,
    number: int,
    box: tuple[float, float, float, float],
    tokens: dict[str, Any],
    *,
    alternate: bool = False,
    icon: str | None = None,
) -> None:
    x, y, w, h = box
    _native_panel(slide, box, tokens, alternate=alternate)
    accent = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x + 0.22),
        Inches(y + 0.2),
        Inches(0.46),
        Inches(0.46),
    )
    accent.name = "BrandDeck Native Accent"
    accent.fill.solid()
    accent.fill.fore_color.rgb = _rgb(tokens["accent"])
    accent.line.fill.background()
    try:
        accent.adjustments[0] = 0.18
    except (IndexError, ValueError):
        pass
    if icon:
        add_icon(slide.shapes, icon, (x + 0.3, y + 0.28, 0.3), tokens["accent_text"])
    else:
        _add_textbox(
            slide,
            f"{number:02d}",
            (x + 0.22, y + 0.29, 0.46, 0.22),
            font_name=tokens["heading_font"],
            font_size=11,
            color=tokens["accent_text"],
            bold=True,
            align=PP_ALIGN.CENTER,
            min_font_size=10,
        ).name = "BrandDeck Native Number"
    _add_textbox(
        slide,
        item,
        (x + 0.22, y + 0.88, max(0.6, w - 0.44), max(0.55, h - 1.08)),
        font_name=tokens["font"],
        font_size=tokens["body_size"],
        color=tokens["heading"],
        min_font_size=15,
        vertical_anchor=MSO_ANCHOR.TOP,
    ).name = "BrandDeck Native Card Text"


def _native_cards(
    slide: Any,
    items: list[str],
    zone: tuple[float, float, float, float],
    tokens: dict[str, Any],
    context: dict[str, Any] | None = None,
) -> None:
    x, y, w, h = zone
    items = items[:4]
    if not items:
        return
    # Pictograms replace numbers only when every card has a confident match.
    icons = assign_icons(
        items, context=str((context or {}).get("title", "")), require_all=True
    )
    if all(icons):
        _record(context, {"type": "card_icons", "icons": icons, "native": True})
    columns = len(items) if len(items) <= 3 else 2
    rows = math.ceil(len(items) / columns)
    gap = 0.2
    card_w = (w - gap * (columns - 1)) / columns
    available_h = min(h, 2.55 if rows == 1 else 4.15)
    card_h = (available_h - gap * (rows - 1)) / rows
    start_y = y + max(0.0, (h - available_h) * 0.38)
    for index, item in enumerate(items):
        col = index % columns
        row = index // columns
        _native_card_content(
            slide,
            item,
            index + 1,
            (
                x + col * (card_w + gap),
                start_y + row * (card_h + gap),
                card_w,
                card_h,
            ),
            tokens,
            alternate=index % 2 == 1,
            icon=icons[index],
        )


def _native_list(
    slide: Any,
    items: list[str],
    zone: tuple[float, float, float, float],
    tokens: dict[str, Any],
) -> None:
    x, y, w, h = zone
    items = items[:5]
    if not items:
        return
    gap = 0.13
    row_h = (h - gap * (len(items) - 1)) / len(items)
    for index, item in enumerate(items):
        row_y = y + index * (row_h + gap)
        _native_panel(
            slide,
            (x, row_y, w, row_h),
            tokens,
            alternate=index % 2 == 1,
        )
        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x + 0.22),
            Inches(row_y + row_h / 2 - 0.19),
            Inches(0.38),
            Inches(0.38),
        )
        badge.name = "BrandDeck Native Badge"
        badge.fill.solid()
        badge.fill.fore_color.rgb = _rgb(tokens["accent"])
        badge.line.fill.background()
        _add_textbox(
            slide,
            str(index + 1),
            (x + 0.22, row_y + row_h / 2 - 0.16, 0.38, 0.28),
            font_name=tokens["heading_font"],
            font_size=12,
            color=tokens["accent_text"],
            bold=True,
            align=PP_ALIGN.CENTER,
            min_font_size=11,
        ).name = "BrandDeck Native Badge Text"
        _add_textbox(
            slide,
            item,
            (x + 0.82, row_y + 0.13, w - 1.05, row_h - 0.26),
            font_name=tokens["font"],
            font_size=tokens["body_size"],
            color=tokens["heading"],
            min_font_size=15,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        ).name = "BrandDeck Native List Text"


def _native_split(
    slide: Any,
    items: list[str],
    visual: dict[str, Any] | None,
    zone: tuple[float, float, float, float],
    design: dict[str, Any],
    tokens: dict[str, Any],
    context: dict[str, Any] | None = None,
) -> None:
    if visual and not items:
        _add_visual(slide, visual, zone, design, context, rounded=True)
        return
    left, right = _split_zone(zone, 0.43)
    if items:
        _native_card_content(
            slide,
            "\n".join(f"• {item}" for item in items) if visual else items[0],
            1,
            left,
            tokens,
            alternate=True,
        )
    if visual:
        _add_visual(slide, visual, right, design, context, rounded=True)
    else:
        _native_list(slide, items[1:] or items[:1], right, tokens)


def _native_cover_motif(slide: Any, width: float, tokens: dict[str, Any]) -> None:
    """Add an editable accent motif that keeps the cover visually asymmetric."""
    specs = [
        (width - 3.8, 1.58, 2.45, 3.4, -8, 0.28),
        (width - 3.42, 1.38, 2.45, 3.4, 4, 0.48),
        (width - 3.0, 1.78, 2.12, 2.95, 12, 0.72),
    ]
    for index, (x, y, w, h, rotation, ratio) in enumerate(specs, 1):
        tile = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(w),
            Inches(h),
        )
        tile.name = f"BrandDeck Native Cover Motif {index}"
        tile.rotation = rotation
        tile.fill.solid()
        tile.fill.fore_color.rgb = _rgb(
            _mix_color(tokens["background"], tokens["accent"], ratio)
        )
        tile.line.color.rgb = _rgb(tokens["accent"])
        tile.line.width = Pt(1.1)
        try:
            tile.adjustments[0] = 0.12
        except (IndexError, ValueError):
            pass


def _short_navigation_label(value: str, fallback: str) -> str:
    words = [word.strip(".,:;!?()[]{}") for word in value.split() if word.strip()]
    label = " ".join(words[:2]).strip() or fallback
    return label if len(label) <= 17 else label[:16].rstrip() + "…"


def _native_navigation(
    slide: Any,
    design: dict[str, Any],
    tokens: dict[str, Any],
    titles: list[str],
    active_index: int,
) -> bool:
    rule = design.get("layout_rules", {}).get("header_navigation", {})
    if not rule.get("detected"):
        return False
    count = max(3, int(rule.get("count", 5)))
    labels = [
        _short_navigation_label(value, f"{index + 1:02d}")
        for index, value in enumerate(titles)
    ]
    while len(labels) < count:
        labels.append(f"{len(labels) + 1:02d}")
    window_start = min(max(0, active_index - count + 1), max(0, len(labels) - count))
    labels = labels[window_start : window_start + count]
    active_local = max(0, min(count - 1, active_index - window_start))

    left = float(rule.get("left_inches", 0.65))
    top = float(rule.get("top_inches", 0.28))
    right = float(rule.get("right_inches", design["canvas"]["width_inches"] - 0.65))
    height = max(0.18, float(rule.get("height_inches", 0.25)))
    gap = max(0.06, min(0.16, float(rule.get("gap_inches", 0.12))))
    pill_width = max(0.55, (right - left - gap * (count - 1)) / count)
    for index, label in enumerate(labels):
        x = left + index * (pill_width + gap)
        active = index == active_local
        fill_color = _hex(
            rule.get("active_fill" if active else "inactive_fill"),
            tokens["heading"] if active else tokens["surface"],
        )
        text_color = _hex(
            rule.get("active_text" if active else "inactive_text"),
            "FFFFFF" if active else tokens["body"],
        )
        pill = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(top),
            Inches(pill_width),
            Inches(height),
        )
        pill.name = f"BrandDeck Native Navigation {index + 1}"
        pill.fill.solid()
        pill.fill.fore_color.rgb = _rgb(fill_color)
        pill.line.color.rgb = _rgb(_hex(rule.get("border_color"), tokens["body"]))
        pill.line.width = Pt(0.45)
        try:
            pill.adjustments[0] = 0.5
        except (IndexError, ValueError):
            pass
        _add_textbox(
            slide,
            label,
            (x + 0.04, top + 0.012, pill_width - 0.08, height - 0.024),
            font_name=tokens["font"],
            font_size=7.2,
            color=text_color,
            bold=active,
            align=PP_ALIGN.CENTER,
            min_font_size=6.5,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        ).name = f"BrandDeck Native Navigation Label {index + 1}"
    return True


def _native_closing_panel(
    slide: Any,
    title: str,
    subtitle: str,
    design: dict[str, Any],
    tokens: dict[str, Any],
) -> bool:
    rule = design.get("layout_rules", {}).get("closing_panel", {})
    if not rule.get("detected"):
        return False
    width = float(design["canvas"]["width_inches"])
    height = float(design["canvas"]["height_inches"])
    panel_width = min(
        width * 0.62, max(width * 0.24, width * float(rule.get("width_ratio", 0.38)))
    )
    side = str(rule.get("side", "left"))
    panel_x = 0.0 if side == "left" else width - panel_width
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(panel_x),
        Inches(0),
        Inches(panel_width),
        Inches(height),
    )
    panel.name = "BrandDeck Native Closing Panel"
    panel.fill.solid()
    panel.fill.fore_color.rgb = _rgb(_hex(rule.get("background"), tokens["heading"]))
    panel.line.fill.background()
    try:
        panel.adjustments[0] = 0.08
    except (IndexError, ValueError):
        pass
    text_left = 0.72 if side == "left" else width - panel_width + 0.62
    text_width = panel_width - 1.15
    panel_heading = _hex(rule.get("heading"), "FFFFFF")
    marker = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(text_left),
        Inches(1.72),
        Inches(0.72),
        Inches(0.07),
    )
    marker.name = "BrandDeck Native Accent"
    marker.fill.solid()
    marker.fill.fore_color.rgb = _rgb(tokens["accent"])
    marker.line.fill.background()
    _add_textbox(
        slide,
        title,
        (text_left, 2.06, text_width, 1.8),
        font_name=tokens["heading_font"],
        font_size=min(40.0, tokens["cover_size"]),
        color=panel_heading,
        bold=True,
        min_font_size=28,
    ).name = "BrandDeck Native Closing Title"
    if subtitle:
        content_left = panel_width + 0.72 if side == "left" else 0.72
        content_width = width - panel_width - 1.35
        _add_textbox(
            slide,
            subtitle,
            (content_left, 2.35, content_width, 1.35),
            font_name=tokens["font"],
            font_size=max(18.0, tokens["body_size"]),
            color=tokens["heading"],
            bold=True,
            min_font_size=16,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        ).name = "BrandDeck Native Closing Subtitle"
    return True


def _native_slide(
    slide: Any,
    slide_spec: dict[str, Any],
    design: dict[str, Any],
    brand_source: Any | None,
    repeated_image_hashes: set[str],
    *,
    slide_number: int,
    slide_count: int,
    navigation_titles: list[str] | None = None,
    navigation_index: int = 0,
    context: dict[str, Any] | None = None,
) -> str:
    role = str(slide_spec.get("role", "content"))
    tokens = _native_tokens(design, role)
    if context is not None:
        context["background"] = tokens["background"]
    brand_shapes = _copy_brand_canvas(
        slide, brand_source, repeated_image_hashes, design, tokens
    )
    canvas = design["canvas"]
    width = float(canvas["width_inches"])
    right = _native_right_limit(brand_shapes, design)
    left = 0.62
    pattern_id = str(slide_spec.get("pattern_id", "native-cards"))
    title = str(slide_spec.get("title", ""))
    subtitle = str(slide_spec.get("subtitle", ""))

    if role in {"cover", "section"}:
        _native_cover_motif(slide, width, tokens)
        marker = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(left),
            Inches(1.46),
            Inches(0.72),
            Inches(0.07),
        )
        marker.name = "BrandDeck Native Accent"
        marker.fill.solid()
        marker.fill.fore_color.rgb = _rgb(tokens["accent"])
        marker.line.fill.background()
        _add_textbox(
            slide,
            title,
            (left, 1.78, min(7.2, max(5.0, right - left)), 2.35),
            font_name=tokens["heading_font"],
            font_size=tokens["cover_size"],
            color=tokens["heading"],
            bold=True,
            min_font_size=34,
        ).name = "BrandDeck Native Cover Title"
        if subtitle:
            _add_textbox(
                slide,
                subtitle,
                (left, 5.35, min(6.6, max(3.0, right - left)), 0.7),
                font_name=tokens["font"],
                font_size=max(16, tokens["body_size"]),
                color=tokens["accent"],
                min_font_size=15,
            ).name = "BrandDeck Native Cover Subtitle"
        mode = "native-cover"
    elif role == "closing":
        if _native_closing_panel(slide, title, subtitle, design, tokens):
            mode = "native-closing-split"
            notes = slide_spec.get("speaker_notes")
            if notes:
                try:
                    slide.notes_slide.notes_text_frame.text = notes
                except (AttributeError, NotImplementedError):
                    pass
            return mode
        marker = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(left),
            Inches(2.0),
            Inches(0.72),
            Inches(0.07),
        )
        marker.name = "BrandDeck Native Accent"
        marker.fill.solid()
        marker.fill.fore_color.rgb = _rgb(tokens["accent"])
        marker.line.fill.background()
        _add_textbox(
            slide,
            title,
            (left, 2.34, max(5.0, right - left), 1.35),
            font_name=tokens["heading_font"],
            font_size=tokens["cover_size"],
            color=tokens["heading"],
            bold=True,
            min_font_size=36,
        ).name = "BrandDeck Native Closing Title"
        if subtitle:
            _add_textbox(
                slide,
                subtitle,
                (left, 4.15, min(6.0, max(3.0, right - left)), 0.55),
                font_name=tokens["font"],
                font_size=tokens["body_size"],
                color=tokens["accent"],
                min_font_size=16,
            ).name = "BrandDeck Native Closing Subtitle"
        mode = "native-closing"
    else:
        has_navigation = _native_navigation(
            slide,
            design,
            tokens,
            navigation_titles or [],
            navigation_index,
        )
        title_top = 0.76 if has_navigation else 0.52
        marker_top = 1.56 if has_navigation else 1.38
        content_top = 1.88 if has_navigation else 1.78
        title_width = max(4.6, right - left)
        if (slide_spec.get("visual") or {}).get("type") == "image":
            title_backing = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(left - 0.12),
                Inches(title_top - 0.08),
                Inches(min(width - left, title_width + 0.24)),
                Inches(0.88),
            )
            title_backing.name = "BrandDeck Native Title Backing"
            title_backing.fill.solid()
            title_backing.fill.fore_color.rgb = _rgb(tokens["background"])
            title_backing.line.fill.background()
        _add_textbox(
            slide,
            title,
            (left, title_top, title_width, 0.72),
            font_name=tokens["heading_font"],
            font_size=tokens["title_size"],
            color=tokens["heading"],
            bold=True,
            min_font_size=28,
        ).name = "BrandDeck Native Title"
        marker = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(left),
            Inches(marker_top),
            Inches(0.72),
            Inches(0.06),
        )
        marker.name = "BrandDeck Native Accent"
        marker.fill.solid()
        marker.fill.fore_color.rgb = _rgb(tokens["accent"])
        marker.line.fill.background()
        items = ([str(slide_spec.get("body"))] if slide_spec.get("body") else []) + [
            str(item) for item in slide_spec.get("bullets", []) if str(item).strip()
        ]
        zone = (left, content_top, max(4.8, right - left), 6.4 - content_top)
        visual = slide_spec.get("visual")
        if pattern_id == "native-list" and not visual:
            _native_list(slide, items, zone, tokens)
            mode = "native-list"
        elif pattern_id == "native-split" or visual:
            _native_split(slide, items, visual, zone, design, tokens, context)
            mode = "native-split"
        else:
            _native_cards(slide, items, zone, tokens, context)
            mode = "native-cards"
        if (slide_spec.get("visual") or {}).get("type") != "image":
            _add_textbox(
                slide,
                f"{slide_number:02d} / {slide_count:02d}",
                (width - 1.25, 6.9, 0.72, 0.25),
                font_name=tokens["font"],
                font_size=9,
                color=tokens["accent"],
                align=PP_ALIGN.RIGHT,
                min_font_size=9,
            ).name = "BrandDeck Native Footer"

    notes = slide_spec.get("speaker_notes")
    if notes:
        try:
            slide.notes_slide.notes_text_frame.text = notes
        except (AttributeError, NotImplementedError):
            pass
    return mode


def _remove_unmatched_content_images(
    slide: Any,
    slide_spec: dict[str, Any],
    design: dict[str, Any],
    repeated_image_hashes: set[str],
    *,
    keep_slot: bool = False,
) -> Any | None:
    """Remove source content pictures; optionally keep the largest as a slot.

    Source pictures never survive into a generated slide. With ``keep_slot``
    the largest one stays until the caller replaces it, so the new image
    inherits the designer's position, crop frame and layer order.
    """
    canvas_area = max(
        0.1,
        float(design["canvas"]["width_inches"])
        * float(design["canvas"]["height_inches"]),
    )
    candidates = []
    for shape in list(slide.shapes):
        # Filled picture placeholders are reported as PLACEHOLDER by
        # python-pptx, although the underlying OOXML element is p:pic.  They
        # can carry source-specific QR codes and photos just like ordinary
        # pictures, so they must pass through the same cleanup.
        if shape.element.tag != qn("p:pic"):
            continue
        area_ratio = shape.width / 914400 * (shape.height / 914400) / canvas_area
        try:
            image_hash = str(shape.image.sha1)
        except (AttributeError, ValueError):
            image_hash = ""
        # Full-bleed imagery is treated as background. Every other non-recurring
        # picture is source content (logos, screenshots, icons, diagrams) and
        # must not leak into a text-only generated slide. Brand assets are
        # identified by recurring across a substantial share of the template.
        thin_strip = min(shape.width, shape.height) < Inches(0.04) and max(
            shape.width, shape.height
        ) > Inches(0.5)
        interior_fragment = (
            area_ratio < 0.02
            and shape.top >= Inches(0.7)
            and shape.top + shape.height
            <= Inches(float(design["canvas"]["height_inches"]) - 0.35)
        )
        if (
            (0.0005 <= area_ratio <= 0.8 and image_hash not in repeated_image_hashes)
            or interior_fragment
            or (keep_slot and thin_strip)
        ):
            candidates.append((area_ratio, shape))
    slot = None
    if keep_slot:
        usable = [item for item in candidates if item[0] >= 0.04]
        if usable:
            slot = max(usable, key=lambda item: item[0])[1]
    for _, shape in candidates:
        if shape is not slot:
            slide.shapes._spTree.remove(shape.element)
    return slot


def _constrain_title_width(slide: Any, title: Any) -> None:
    """Keep a title to the left of intersecting template artwork.

    Inspect inherited pictures too: a layout picture may cover slide text even
    though it is absent from slide.shapes. Full backgrounds do not start inside
    the title and therefore do not narrow it.
    """
    right = title.left + title.width
    for layer in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        for shape in layer.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            if (
                title.left < shape.left < right
                and shape.top < title.top + title.height
                and shape.top + shape.height > title.top
            ):
                right = min(right, shape.left - Inches(0.15))
    if right - title.left >= Inches(1.5):
        # Setting one dimension of an inherited placeholder can zero the others
        # in python-pptx; materialize the complete geometry first.
        left, top, height = title.left, title.top, title.height
        title.left, title.top = left, top
        title.width, title.height = right - left, height


def _compact_centered_cover(
    slide: Any,
    title_shape: Any,
    subtitle_shape: Any,
    title: str,
    subtitle: str,
    design: dict[str, Any],
) -> tuple[float, float, str] | None:
    """Fit a long centered cover stack above its background art.

    Some covers have short, centered placeholders over a full-bleed image.
    Their sample title and caption fit, but a brief-length replacement wraps
    into the illustration. Widen and lift both frames only when the original
    geometry identifies this particular cover arrangement and the text fits
    the wider frame at a readable size.
    """
    canvas = design["canvas"]
    width, height = float(canvas["width_inches"]), float(canvas["height_inches"])
    x, y, w, h = _zone_from_shape(title_shape)
    sx, sy, sw, _ = _zone_from_shape(subtitle_shape)
    center = x + w / 2
    original_title_size = float(effective_font_size(title_shape, 24) or 24)
    original_subtitle_size = float(effective_font_size(subtitle_shape, 14) or 14)
    preferred_title_size = min(original_title_size, 24 * width / 10)
    preferred_subtitle_size = min(original_subtitle_size, 14 * width / 10)
    if not (
        title
        and subtitle
        and len(title) * preferred_title_size * 0.52 / 72 > w - 0.2
        and len(subtitle) * preferred_subtitle_size * 0.52 / 72 > sw - 0.2
        and 0.35 * width <= w <= 0.60 * width
        and 0.20 * height <= y <= 0.34 * height
        and abs(center - width / 2) <= 0.06 * width
        and abs(sx + sw / 2 - center) <= 0.03 * width
        and abs(sw - w) <= 0.05 * width
        and -0.05 <= sy - y - h <= 0.25
    ):
        return None

    new_w = 0.76 * width
    new_x = center - new_w / 2
    inner_w = new_w - 0.2
    single_size = min(preferred_title_size, inner_w * 72 / (len(title) * 0.52))
    if single_size >= 20 * width / 10:
        title_text = title
        title_size = single_size
        new_y = y - min(0.2, 0.035 * height)
        title_h = max(0.50, min(0.65, 0.10 * height))
    else:
        words = title.split()
        if len(words) < 2:
            return None
        halves = [
            (" ".join(words[:index]), " ".join(words[index:]))
            for index in range(1, len(words))
        ]
        left, right = min(halves, key=lambda pair: abs(len(pair[0]) - len(pair[1])))
        title_text = f"{left}\n{right}"
        title_size = min(
            preferred_title_size,
            inner_w * 72 / (max(len(left), len(right)) * 0.52),
        )
        new_y = y - min(0.31, 0.055 * height)
        title_h = max(0.82, 2 * title_size * 1.22 / 72 + 0.06)
    subtitle_y = new_y + title_h + 0.11
    # Size and frame follow the same width estimate as the overflow check, so
    # QA does not reject the fitted stack. One smaller line keeps the stack
    # compact above the art; a caption that would be too small wraps once.
    readable = 11 * width / 10
    subtitle_size = min(
        preferred_subtitle_size,
        inner_w * 72 / (max(1, len(subtitle)) * ESTIMATED_GLYPH_WIDTH_EM),
    )
    if subtitle_size < readable:
        subtitle_size = preferred_subtitle_size

    def caption_lines(size: float) -> int:
        per_line = int(inner_w * 72 / (size * ESTIMATED_GLYPH_WIDTH_EM))
        return estimated_line_count(subtitle, max(1, per_line))

    while caption_lines(subtitle_size) > 2 and subtitle_size > readable:
        subtitle_size = max(readable, subtitle_size - 0.5)
    if subtitle_size < readable or caption_lines(subtitle_size) > 2:
        return None
    subtitle_h = max(
        0.42,
        min(0.50, subtitle_shape.height / 914400 + 0.1),
        caption_lines(subtitle_size) * subtitle_size * 1.22 / 72 + 0.2,
    )
    if new_x < 0.05 * width or new_x + new_w > 0.95 * width:
        return None
    if subtitle_y + subtitle_h >= height * 0.52:
        return None

    title_zone = (new_x, new_y, new_w, title_h)
    subtitle_zone = (new_x, subtitle_y, new_w, subtitle_h)
    for layer in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        for shape in layer.shapes:
            if shape.shape_type not in {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.GROUP}:
                continue
            artwork = _zone_from_shape(shape)
            # A full-slide picture is the intended cover background.
            if area(artwork) >= 0.8 * width * height:
                continue
            if intersection(artwork, title_zone) or intersection(artwork, subtitle_zone):
                return None

    if title_size < 18 * width / 10:
        return None
    _set_geometry(title_shape, *(Inches(value) for value in title_zone))
    _set_geometry(subtitle_shape, *(Inches(value) for value in subtitle_zone))
    return title_size, subtitle_size, title_text


def _zone_beside(
    zone: tuple[float, float, float, float], slot: tuple[float, float, float, float]
) -> tuple[float, float, float, float]:
    """Largest part of a fallback text zone left, right, above or below a slot."""
    x, y, w, h = zone
    sx, sy, sw, sh = slot
    if sx >= x + w or sx + sw <= x or sy >= y + h or sy + sh <= y:
        return zone
    gap = 0.2
    options = [
        (x, y, sx - gap - x, h),
        (sx + sw + gap, y, x + w - sx - sw - gap, h),
        (x, y, w, sy - gap - y),
        (x, sy + sh + gap, w, y + h - sy - sh - gap),
    ]
    usable = [option for option in options if option[2] >= 1.6 and option[3] >= 0.8]
    if not usable:
        return zone
    return max(usable, key=lambda option: option[2] * option[3])


def _expand_small_image_zone(
    slide: Any,
    zone: tuple[float, float, float, float],
    design: dict[str, Any],
) -> tuple[float, float, float, float]:
    """Use free space beside text when an exemplar offers a narrow photo zone."""
    x, y, w, h = zone
    canvas = design["canvas"]
    canvas_w = float(canvas["width_inches"])
    canvas_h = float(canvas["height_inches"])
    if w * h >= canvas_w * canvas_h * 0.16 or x + w >= canvas_w - 0.9:
        return zone
    candidate_w = canvas_w - 0.35 - x
    candidate_h = min(canvas_h - 0.45 - y, max(h, candidate_w / 1.65))
    # An exemplar may already use the full content height while leaving a
    # wide empty strip on the right. Widening alone still increases the
    # rendered content area without moving the neighboring text.
    candidate_w = max(w, candidate_w)
    candidate_h = max(h, candidate_h)
    if candidate_w <= w and candidate_h <= h:
        return zone
    candidate = (x, y, candidate_w, candidate_h)
    for shape in slide.shapes:
        if not has_text(shape):
            continue
        if intersection(box(shape), candidate) > 0.02:
            return zone
    return candidate


def _clear_layout_art(
    zone: tuple[float, float, float, float],
    art: list[tuple[float, float, float, float]],
    *,
    pad: float = 0.1,
    keep: float = 0.8,
) -> tuple[float, float, float, float]:
    """Pull a drawn visual off the layout's logos and ornaments.

    The zone gives up the side facing each piece of art; when that costs
    more than a fifth of it, the zone stays and the audit reports the
    overlap, so the slide gets another layout instead of a cramped visual.
    """
    original = zone
    for left, top, width, height in art:
        if intersection(zone, (left, top, width, height)) <= 0:
            continue
        x, y, w, h = zone
        options = [
            (x, y, w, top - pad - y),
            (x, top + height + pad, w, y + h - top - height - pad),
            (x, y, left - pad - x, h),
            (left + width + pad, y, x + w - left - width - pad, h),
        ]
        zone = max(options, key=area)
        if min(zone[2], zone[3]) <= 0 or area(zone) < keep * area(original):
            return original
    return zone


def _set_geometry(shape: Any, left: int, top: int, width: int, height: int) -> None:
    # Setting one dimension of an inherited placeholder can zero the others
    # in python-pptx; always write the complete geometry.
    shape.left, shape.top, shape.width, shape.height = left, top, width, height


def _contain_generated_text(slide: Any, width: int, height: int) -> None:
    """Keep generated text frames within the canvas after template fitting."""
    for shape in slide.shapes:
        if shape.name != "BrandDeck Body" or not shape.has_text_frame:
            continue
        if shape.left + shape.width > width:
            shape.left = max(0, width - shape.width)
        if shape.top + shape.height > height:
            shape.top = max(0, height - shape.height)
        shape.left = max(shape.left, 0)
        shape.top = max(shape.top, 0)


def _slot_text_changes(
    slot: Any, slide: Any, bodies: list[Any]
) -> list[tuple[Any, tuple[int, int, int, int]]] | None:
    """Plan how text makes room for a picture slot; ``None`` if it cannot.

    Text that merely runs into a side slot (a long title box, a wide body)
    is narrowed to end before it or to start after it. Text laid over the
    middle of the picture means it was a backdrop, not a slot. Nothing is
    changed here, so a rejected candidate leaves the slide untouched.
    """
    gap = Inches(0.2)
    minimum = Inches(1.6)
    candidates = list(bodies) + [
        shape
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
        and shape.text.strip()
        and shape.element is not slot.element
        and shape not in bodies
    ]
    slot_right = slot.left + slot.width
    slot_bottom = slot.top + slot.height
    changes: list[tuple[Any, tuple[int, int, int, int]]] = []
    for shape in candidates:
        left, top = shape.left, shape.top
        right, bottom = left + shape.width, top + shape.height
        overlap_w = min(right, slot_right) - max(left, slot.left)
        overlap_h = min(bottom, slot_bottom) - max(top, slot.top)
        if overlap_w <= 0 or overlap_h <= 0:
            continue
        smaller = max(1, min(shape.width * shape.height, slot.width * slot.height))
        if overlap_w * overlap_h <= 0.05 * smaller:
            continue
        if slot.left - gap - left >= minimum:
            changes.append((shape, (left, top, slot.left - gap - left, shape.height)))
        elif right - (slot_right + gap) >= minimum and slot.left <= left:
            changes.append(
                (shape, (slot_right + gap, top, right - slot_right - gap, shape.height))
            )
        else:
            return None
    return changes


def _apply_text_changes(changes: list[tuple[Any, tuple[int, int, int, int]]]) -> None:
    for shape, geometry in changes:
        _set_geometry(shape, *geometry)
        if shape.text.strip():
            # Titles already hold new text: refit it to the narrower box.
            _set_text_frame(shape, [shape.text])


def _empty_frames(slide: Any, design: dict[str, Any]) -> list[Any]:
    """Text-free exemplar shapes sized like a photo frame (not panels or bars)."""
    width = float(design["canvas"]["width_inches"])
    height = float(design["canvas"]["height_inches"])
    frames = []
    for shape in slide.shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE or getattr(
            shape, "is_placeholder", False
        ):
            continue
        if str(shape.name).startswith("BrandDeck") or (
            getattr(shape, "has_text_frame", False) and shape.text.strip()
        ):
            continue
        x, y, w, h = _zone_from_shape(shape)
        if x < -0.02 or y < -0.02 or x + w > width + 0.02 or y + h > height + 0.02:
            continue
        if 0.03 <= w * h / (width * height) <= 0.4 and 0.5 <= w / max(h, 0.01) <= 2:
            frames.append(shape)
    return sorted(frames, key=lambda shape: shape.width * shape.height, reverse=True)


def _choose_picture_slot(
    slide: Any, content_slot: Any | None, bodies: list[Any], design: dict[str, Any]
) -> Any | None:
    """First usable slot: the exemplar's photo, picture placeholders, frames."""
    candidates = [content_slot] if content_slot is not None else []
    candidates += sorted(
        _empty_picture_placeholders(slide),
        key=lambda shape: shape.width * shape.height,
        reverse=True,
    )
    candidates += _empty_frames(slide, design)
    chosen = None
    for candidate in candidates:
        changes = _slot_text_changes(candidate, slide, bodies)
        if changes is not None:
            _apply_text_changes(changes)
            chosen = candidate
            break
    if content_slot is not None and (
        chosen is None or chosen.element is not content_slot.element
    ):
        # An unused source photo must not survive into the generated slide.
        content_slot.element.getparent().remove(content_slot.element)
    return chosen


def _clear_group_text(group: Any, assigned_text: set[Any] | None = None) -> None:
    assigned_text = assigned_text or set()
    for shape in group.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            _clear_group_text(shape, assigned_text)
        elif (
            shape.element not in assigned_text
            and getattr(shape, "has_text_frame", False)
            and shape.text.strip()
        ):
            _set_text_frame(shape, [])


def _empty_picture_placeholders(slide: Any) -> list[Any]:
    return [
        shape
        for shape in slide.placeholders
        if _placeholder_type(shape) == PP_PLACEHOLDER.PICTURE
        and shape.element.tag == qn("p:sp")
    ]


def _text_block_area(texts: list[str], size_pt: float, width: float) -> float:
    """Area of the rectangle wrapped text occupies, as the rendered audit measures it."""
    glyph = size_pt * RENDERED_GLYPH_WIDTH_EM / 72
    per_line = max(1, int(width / glyph))
    lines, widest = 0, 0
    for text in texts:
        wrapped = textwrap.wrap(str(text), per_line, break_on_hyphens=False) or [""]
        lines += len(wrapped)
        widest = max(widest, max(len(line) for line in wrapped))
    return min(width, widest * glyph) * lines * size_pt * 1.25 / 72


def _sparse_statement_cards(
    slide_spec: dict[str, Any],
    titles: list[Any],
    bodies: list[Any],
    design: dict[str, Any],
) -> dict[str, Any] | None:
    """Card grid for a few short statements that would leave the slide empty.

    Appendix 1 treats a slide filled below a quarter as a defect, and text
    alone cannot fill a large zone: 150–250 characters at the template body
    size cover 10–20 % of a slide. When the template gives these statements
    no visible cards, they become a native card grid in the same zone. The
    estimate wraps the text in its box and measures the rectangle it takes,
    like the rendered audit; its threshold sits a little above a quarter,
    and the rendered audit, not this estimate, decides the final verdict.
    """
    bullets = [str(item) for item in slide_spec.get("bullets", []) if str(item).strip()]
    minimum, maximum = ITEM_LIMITS["icon_grid"]
    # A data slide without a chart or table carries statements too.
    if (
        slide_spec.get("role", "content") not in {"content", "data"}
        or slide_spec.get("visual")
        or slide_spec.get("body")
        or not minimum <= len(bullets) <= maximum
        or any(len(item.split()) > 22 for item in bullets)
        or (
            slide_spec.get("layout_hint") != "cards"
            and any(is_visible_box(shape) for shape in bodies)
        )
    ):
        return None
    if slide_spec.get("layout_hint") == "cards":
        # The rendered audit measured this slide below a quarter.
        return {"type": "icon_grid", "items": bullets, "origin": "rendered_fill"}
    body_size = float(design["typography"].get("body_size_pt", 18))
    if bodies:
        body_size = max(body_size, float(effective_font_size(bodies[0], body_size)))
    title_size = max((_largest_font(shape) for shape in titles), default=0.0)
    canvas_w = float(design["canvas"]["width_inches"])
    canvas = canvas_w * float(design["canvas"]["height_inches"])

    def inner(shapes: list[Any]) -> float:
        # Text frames keep 0.1 inch margins on each side.
        return max(0.5, box(shapes[0])[2] - 0.2) if shapes else canvas_w * 0.8

    estimate = (
        _text_block_area(bullets, body_size, inner(bodies))
        + _text_block_area(
            [shape.text for shape in titles], title_size or body_size * 1.6, inner(titles)
        )
    ) / canvas
    if estimate >= SPARSE_FILL:
        return None
    return {
        "type": "icon_grid",
        "items": bullets,
        "origin": "sparse_text",
        "estimated_fill": round(estimate, 3),
    }


def _remove_exemplar_leftovers(
    slide: Any,
    design: dict[str, Any],
    original: set[Any],
    had_text: set[Any],
    *,
    diagram_only: bool = False,
) -> list[dict[str, str]]:
    """Remove exemplar scaffolding the new content does not use.

    A cloned exemplar keeps its cards, labels, icons and photo frames after
    the sample text is erased. Left in place they turn into empty cards,
    icons without captions, frames without photos and shapes under the new
    text (Appendix 1: overlaps, empty blocks). A shape goes when its text
    was erased, when it is attached to such a shape (icon, dot, card), when
    it lies under new content, or when it is an empty photo or icon frame; a line
    goes only when it crosses new content. When a generated diagram carries
    all of the slide's text, exemplar decoration beside it (bars of a chart
    drawn with shapes, markers of a replaced list) belongs to the removed
    content and goes too. Bands, panels, anything hosting new content and
    pictures are kept.
    """
    canvas = (
        float(design["canvas"]["width_inches"]),
        float(design["canvas"]["height_inches"]),
    )
    shapes = list(slide.shapes)
    content = [
        (shape, box(shape))
        for shape in shapes
        if has_text(shape) or (is_generated(shape) and is_opaque(shape))
    ]
    erased = [
        box(shape)
        for shape in shapes
        if shape.element in had_text and not has_text(shape)
    ]
    diagrams = [
        box(shape)
        for shape in shapes
        if diagram_only and str(shape.name).startswith("BrandDeck Diagram")
    ]
    # The diagram's horizontal band across the whole slide.
    bands = [(0.0, y, canvas[0], h) for _, y, _, h in diagrams]
    removed: list[dict[str, str]] = []
    for shape in shapes:
        # A template text slot the plan left empty keeps its new name but is
        # still an erased sample.
        emptied_slot = is_generated(shape) and shape.element in had_text
        if (
            (is_generated(shape) and not emptied_slot)
            or shape.is_placeholder
            or shape.shape_type not in LEFTOVER_TYPES
            or has_text(shape)
        ):
            continue
        zone = box(shape)
        if is_framing(shape, zone, canvas, content, shape.element in had_text):
            continue
        reason = None
        if shape.shape_type == MSO_SHAPE_TYPE.LINE:
            # Dividers next to content are design; only crossing lines go.
            if any(
                collides(shape, zone, item, other, False) for item, other in content
            ):
                reason = "under_content"
            elif any(intersection(zone, band) > 0 for band in bands):
                reason = "beside_diagram"  # an axis of a shape-drawn chart
        elif shape.element in had_text:
            reason = "erased_sample"
        elif any(gap(zone, other) <= 0.3 for other in erased):
            reason = "attached_to_erased"
        elif any(
            collides(shape, zone, item, other, item.element in original)
            for item, other in content
        ):
            reason = "under_content"
        elif any(
            intersection(zone, band) >= 0.5 * max(area(zone), 1e-6) for band in bands
        ):
            reason = "beside_diagram"
        if reason:
            removed.append({"shape": str(shape.name), "reason": reason})
            shape.element.getparent().remove(shape.element)
    # Frames are judged last: the photo or icon they held may have gone above.
    remaining = list(slide.shapes)
    for shape in remaining:
        if is_generated(shape) or shape.is_placeholder or has_text(shape):
            continue
        zone = box(shape)
        occupants = [box(other) for other in remaining if other is not shape]
        if not is_framing(shape, zone, canvas, content, False) and is_photo_frame(
            shape, zone, content, occupants
        ):
            removed.append({"shape": str(shape.name), "reason": "empty_photo_frame"})
            shape.element.getparent().remove(shape.element)
    return removed


def _fill_slide(
    slide: Any,
    slide_spec: dict[str, Any],
    design: dict[str, Any],
    repeated_image_hashes: set[str],
    context: dict[str, Any] | None = None,
) -> None:
    original = {shape.element for shape in slide.shapes}
    had_text = {shape.element for shape in slide.shapes if has_text(shape)}
    # Example-slide charts and tables carry source-specific data. Keep the
    # surrounding style, but replace those data objects with the new plan.
    for shape in list(slide.shapes):
        if getattr(shape, "has_chart", False) or getattr(shape, "has_table", False):
            slide.shapes._spTree.remove(shape.element)
    visual = slide_spec.get("visual") or {}
    wants_picture = visual.get("type") == "image" or slide_spec.get("role") == "image"
    content_slot = _remove_unmatched_content_images(
        slide, slide_spec, design, repeated_image_hashes, keep_slot=wants_picture
    )
    titles, subtitles, bodies = _placeholder_groups(slide)
    titles, subtitles, bodies = _infer_text_groups(
        slide,
        titles,
        subtitles,
        bodies,
        float(design["canvas"]["width_inches"]),
        float(design["canvas"]["height_inches"]),
    )
    if (
        not subtitles
        and slide_spec.get("subtitle")
        and slide_spec.get("role") in {"cover", "closing", "section"}
        and titles
    ):
        # Some corporate covers encode the subtitle as BODY, not SUBTITLE.
        # A closing slide with statements keeps a large body for them: only a
        # caption-sized frame (or a spare body) becomes the subtitle.
        title_bottom = max(shape.top + shape.height for shape in titles)
        has_statements = bool(slide_spec.get("bullets") or slide_spec.get("body"))
        candidates = [
            shape
            for shape in bodies
            if shape.top >= title_bottom
            and (not has_statements or len(bodies) > 1 or shape.height <= Inches(1.2))
        ]
        if candidates:
            subtitle_shape = min(candidates, key=lambda shape: (shape.top, shape.left))
            subtitles = [subtitle_shape]
            bodies = [shape for shape in bodies if shape is not subtitle_shape]
    if not slide_spec.get("subtitle") and slide_spec.get("role") == "content":
        bodies = subtitles + bodies
        subtitles = []
    assigned_text = {shape.element for shape in titles + subtitles + bodies}
    for shape in slide.shapes:
        if (
            shape.element not in assigned_text
            and getattr(shape, "has_text_frame", False)
            and shape.text.strip()
        ):
            _set_text_frame(shape, [])
        elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            # Exemplar groups (chart legends, callouts) carry source text too.
            _clear_group_text(shape, assigned_text)
    font = design["typography"]["primary_font"]
    colors = _theme_colors(design)
    title = slide_spec.get("title", "")
    subtitle = slide_spec.get("subtitle", "")
    cover_sizes = None
    if slide_spec.get("role") == "cover" and titles and subtitles:
        cover_sizes = _compact_centered_cover(
            slide, titles[0], subtitles[0], title, subtitle, design
        )

    if titles:
        _constrain_title_width(slide, titles[0])
        if cover_sizes:
            _set_text_frame(
                titles[0],
                [cover_sizes[2]],
                font_size=cover_sizes[0],
                min_font_size=cover_sizes[0],
            )
        else:
            _set_text_frame(titles[0], [title])
    else:
        margin = float(design["spacing"].get("typical_left_margin_inches", 0.6))
        new_title = _add_textbox(
            slide,
            title,
            (margin, 0.45, design["canvas"]["width_inches"] - margin * 2, 0.8),
            font_name=font,
            font_size=float(design["typography"].get("title_size_pt", 30)),
            color=colors.get("dk1", "1F2937"),
            bold=True,
        )
        new_title.name = "BrandDeck Title"
        titles.append(new_title)

    if subtitles:
        if slide_spec.get("role") in {"cover", "closing"} and subtitle and not cover_sizes:
            # Corporate exemplars often provide a one-line caption frame that
            # is too short for an expanded brief. Give its real text enough
            # height before fitting instead of shrinking it to illegibility.
            subtitle_shape = subtitles[0]
            _set_geometry(
                subtitle_shape,
                subtitle_shape.left,
                subtitle_shape.top,
                subtitle_shape.width,
                max(subtitle_shape.height, Inches(0.55)),
            )
        if cover_sizes:
            _set_text_frame(
                subtitles[0],
                [subtitle],
                font_size=cover_sizes[1],
                min_font_size=cover_sizes[1],
            )
        else:
            _set_text_frame(subtitles[0], [subtitle])
    elif subtitle and slide_spec.get("role") in {"cover", "closing", "section"}:
        canvas = design["canvas"]
        box = (0.8, canvas["height_inches"] * 0.64, canvas["width_inches"] - 1.6, 0.65)
        align = PP_ALIGN.CENTER
        # A closing slide with next steps may reuse a content layout: its text
        # zones stay free, so the subtitle moves under the title or is dropped.
        occupied = (
            [_zone_from_shape(shape) for shape in bodies]
            if slide_spec.get("bullets") or slide_spec.get("body")
            else []
        )
        if any(_zones_intersect(box, zone) for zone in occupied):
            x, y, w, h = _zone_from_shape(titles[0])
            box, align = (x, y + h + 0.05, w, 0.5), PP_ALIGN.LEFT
        if not any(_zones_intersect(box, zone) for zone in occupied):
            _add_textbox(
                slide,
                subtitle,
                box,
                font_name=font,
                font_size=float(design["typography"].get("body_size_pt", 18)),
                color=colors.get("dk1", "1F2937"),
                align=align,
            ).name = "BrandDeck Subtitle"

    slot = (
        _choose_picture_slot(slide, content_slot, bodies, design)
        if wants_picture
        else None
    )
    zone = _content_zone(slide, design, titles + subtitles)
    if slot is not None:
        zone = _zone_beside(zone, _zone_from_shape(slot))
    else:
        cards = _sparse_statement_cards(slide_spec, titles, bodies, design)
        if cards:
            slide_spec = {**slide_spec, "bullets": [], "visual": cards}
    # The plan's visual claims remain identical across variants. Apply the
    # slide-level layout hint only to the temporary rendering specification,
    # including diagrams synthesized from sparse text during composition.
    render_visual = slide_spec.get("visual")
    layout_variant = slide_spec.get("layout_variant")
    if (
        isinstance(render_visual, dict)
        and render_visual.get("type") == "icon_grid"
        and layout_variant in {"balanced", "columns", "focus"}
    ):
        slide_spec = {
            **slide_spec,
            "visual": {**render_visual, "layout_variant": layout_variant},
        }
    # A template picture slot keeps the full body area for text.
    text_spec = {**slide_spec, "visual": None} if slot is not None else slide_spec
    visual_zone = _add_body_text(slide, text_spec, bodies, zone, design)
    if slot is not None:
        _record(
            context, _fill_picture_slot(slide, slot, visual or None, design, context)
        )
    else:
        if visual.get("type") == "image":
            visual_zone = _expand_small_image_zone(slide, visual_zone, design)
        elif slide_spec.get("visual") and visual_zone:
            visual_zone = _clear_layout_art(visual_zone, layout_art_boxes(slide, design))
        _add_visual(slide, slide_spec.get("visual"), visual_zone, design, context)
    if slide_spec.get("role") not in {"cover", "closing", "section"}:
        # A picture slot left empty looks broken; draw a native illustration.
        for placeholder in _empty_picture_placeholders(slide):
            changes = _slot_text_changes(placeholder, slide, [])
            if changes is None:
                continue  # text covers it; an empty placeholder does not render
            _apply_text_changes(changes)
            _record(
                context, _fill_picture_slot(slide, placeholder, None, design, context)
            )
    final_visual = slide_spec.get("visual") or {}
    removed = _remove_exemplar_leftovers(
        slide,
        design,
        original,
        had_text,
        diagram_only=slot is None
        and final_visual.get("type") in DIAGRAM_TYPES
        and not slide_spec.get("bullets")
        and not slide_spec.get("body"),
    )
    if removed and context is not None:
        context.setdefault("cleanup", []).append(
            {"slide": context.get("slide"), "removed": removed}
        )

    notes = slide_spec.get("speaker_notes")
    if notes:
        try:
            slide.notes_slide.notes_text_frame.text = notes
        except (AttributeError, NotImplementedError):
            pass


def outline_markdown(plan: dict[str, Any]) -> str:
    lines = [f"# {plan.get('title', 'Presentation')}", ""]
    for index, slide in enumerate(plan.get("slides", []), 1):
        lines.extend(
            [
                f"## {index}. {slide.get('title', '')}",
                "",
                f"Pattern: `{slide.get('pattern_id', 'unassigned')}` · Role: `{slide.get('role', 'content')}`",
                "",
            ]
        )
        if slide.get("subtitle"):
            lines.append(slide["subtitle"])
            lines.append("")
        if slide.get("body"):
            lines.append(slide["body"])
            lines.append("")
        for bullet in slide.get("bullets", []):
            lines.append(f"- {bullet}")
        if slide.get("visual"):
            lines.extend(["", f"Visual: `{slide['visual'].get('type')}`"])
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def _visual_summary(records: list[dict[str, Any]]) -> dict[str, Any]:
    diagrams = [
        r for r in records if r.get("type") in DIAGRAM_TYPES and not r.get("error")
    ]
    icons = sum(len(r.get("icons", [])) for r in records)
    return {
        "diagrams": len(diagrams),
        "pictograms": icons
        + sum(1 for r in records if r.get("type") == "illustration"),
        "images": sum(1 for r in records if r.get("type") == "image"),
        "illustrations": sum(1 for r in records if r.get("type") == "illustration"),
        "by_type": dict(Counter(str(r.get("type")) for r in records)),
        "slides": records,
        "note": "Diagrams are native grouped shapes (SmartArt-style), not OOXML SmartArt parts.",
    }


def compose_presentation(
    *,
    template_dir: Path,
    plan: dict[str, Any],
    output_path: Path,
    design_system: dict[str, Any],
) -> dict[str, Any]:
    # Preserve the template's design tokens, but do not force Cyrillic text
    # through a Japanese/Chinese font chosen for the source presentation.
    generated_design = design_system
    fallback = None
    if plan.get("language") == "ru":
        typography = design_system.get("typography", {})
        if any(
            needs_cyrillic_font_fallback(typography.get(key))
            for key in ("primary_font", "heading_font")
        ):
            generated_design = copy.deepcopy(design_system)
            for key in ("primary_font", "heading_font"):
                original_font = typography.get(key)
                generated_design["typography"][key] = readable_font(
                    original_font, "Русский текст"
                )
            fallback = "Arial"
    scale = tuple(
        sorted(
            float(value)
            for value in generated_design.get("typography", {}).get("observed_sizes_pt", [])
            if 8 <= float(value) <= 96
        )
    )
    token = _TYPE_SCALE.set(scale)
    font_token = _CYRILLIC_THEME_FALLBACK.set(fallback)
    try:
        return _compose(template_dir, plan, output_path, generated_design)
    finally:
        _CYRILLIC_THEME_FALLBACK.reset(font_token)
        _TYPE_SCALE.reset(token)


def _compose(
    template_dir: Path,
    plan: dict[str, Any],
    output_path: Path,
    design_system: dict[str, Any],
) -> dict[str, Any]:
    original = template_dir / "original.pptx"
    prs = Presentation(original)
    layouts = _all_layouts(prs)
    source_slides = list(prs.slides)
    image_hash_counts: Counter[str] = Counter()
    for source_slide in source_slides:
        for shape in source_slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            try:
                image_hash_counts[str(shape.image.sha1)] += 1
            except (AttributeError, ValueError):
                continue
    recurring_threshold = max(3, math.ceil(len(source_slides) * 0.4))
    repeated_image_hashes = {
        image_hash
        for image_hash, count in image_hash_counts.items()
        if count >= recurring_threshold
    }
    pdf_shape_count = sum(
        str(shape.name).lower().startswith("pdf ")
        for source_slide in source_slides
        for shape in source_slide.shapes
    )
    total_shape_count = sum(len(source_slide.shapes) for source_slide in source_slides)
    source_model = design_system.get("source_model", {})
    native_grid_source = (
        source_model.get("composition_mode") == "native_grid"
        or bool(source_model.get("fragmented"))
        or (
            pdf_shape_count >= 10
            and pdf_shape_count / max(1, total_shape_count) >= 0.25
        )
    )
    brand_source = _select_brand_source(source_slides, repeated_image_hashes)
    cover_brand_source = source_slides[0] if source_slides else brand_source
    closing_brand_source = source_slides[-1] if source_slides else brand_source
    blank_layout = min(layouts, key=lambda item: len(item.placeholders))
    original_slide_ids = list(prs.slides._sldIdLst)
    catalog = read_json(template_dir / "pattern_catalog.json")
    patterns = {item["id"]: item for item in catalog.get("patterns", [])}

    selected_patterns: list[str] = []
    selected_source_patterns: list[str] = []
    slide_specs = plan.get("slides", [])
    navigation_specs = [
        item
        for item in slide_specs
        if str(item.get("role", "content")) not in {"cover", "section", "closing"}
    ]
    navigation_titles = [str(item.get("title", "")) for item in navigation_specs]
    navigation_index = 0
    assets = plan.get("assets") if isinstance(plan.get("assets"), dict) else {}
    visual_records: list[dict[str, Any]] = []
    cleanup_records: list[dict[str, Any]] = []
    for generated_index, slide_spec in enumerate(slide_specs):
        layout_index = int(slide_spec.get("layout_index", 0))
        if layout_index < 0 or layout_index >= len(layouts):
            layout_index = 0
        pattern = patterns.get(str(slide_spec.get("pattern_id")), {})
        selected_source_patterns.append(str(slide_spec.get("pattern_id", "unassigned")))
        background = (
            pattern.get("background")
            if isinstance(pattern.get("background"), dict)
            else {}
        )
        context = {
            "slide": generated_index + 1,
            "role": str(slide_spec.get("role", "content")),
            "title": str(slide_spec.get("title", "")),
            "text": " ".join(
                [str(slide_spec.get("title", "")), str(slide_spec.get("body", ""))]
                + [str(item) for item in slide_spec.get("bullets", [])]
            ),
            "background": background.get("color"),
            "assets": assets,
            "records": visual_records,
            "cleanup": cleanup_records,
        }
        use_native_grid = (
            native_grid_source or pattern.get("source_kind") == "native_grid"
        )
        if use_native_grid:
            slide = prs.slides.add_slide(blank_layout)
            for shape in list(slide.shapes):
                slide.shapes._spTree.remove(shape.element)
            role = str(slide_spec.get("role", "content"))
            role_brand_source = (
                cover_brand_source
                if role in {"cover", "section"}
                else closing_brand_source
                if role == "closing"
                else brand_source
            )
            mode = _native_slide(
                slide,
                slide_spec,
                design_system,
                role_brand_source,
                repeated_image_hashes,
                slide_number=generated_index + 1,
                slide_count=len(slide_specs),
                navigation_titles=navigation_titles,
                navigation_index=navigation_index,
                context=context,
            )
            if role not in {"cover", "section", "closing"}:
                navigation_index += 1
            selected_patterns.append(mode)
            continue
        example_indices = [
            int(value)
            for value in pattern.get("example_slide_indices", [])
            if isinstance(value, int) and 1 <= value <= len(source_slides)
        ]
        if example_indices and pattern.get("source_kind") == "slide_exemplar":
            source_index = example_indices[generated_index % len(example_indices)] - 1
            slide = _clone_slide(prs, source_slides[source_index])
        else:
            slide = prs.slides.add_slide(layouts[layout_index])
        _fill_slide(slide, slide_spec, design_system, repeated_image_hashes, context)
        _contain_generated_text(slide, prs.slide_width, prs.slide_height)
        selected_patterns.append(
            str(slide_spec.get("pattern_id", f"layout-{layout_index}"))
        )

    _remove_slide_ids(prs, original_slide_ids)

    prs.core_properties.title = str(plan.get("title", "Generated presentation"))
    prs.core_properties.subject = "Template-adaptive presentation"
    prs.core_properties.keywords = "BrandDeck, template-adaptive, presentation"
    prs.core_properties.author = ""
    prs.core_properties.last_modified_by = ""
    prs.core_properties.comments = ""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    result = {
        "output": str(output_path.resolve()),
        "slide_count": len(prs.slides),
        "patterns_used": selected_patterns,
        "source_patterns": selected_source_patterns,
        "composition_mode": "native_grid" if native_grid_source else "template_layout",
        "visuals": _visual_summary(visual_records),
        "exemplar_cleanup": cleanup_records,
    }
    write_json(output_path.with_suffix(".manifest.json"), result)
    return result
