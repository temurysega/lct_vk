"""Conservative lexical coverage: unmatched paraphrases require human review."""

from __future__ import annotations

import hashlib
import json
import re
from copy import deepcopy
from pathlib import Path

from pptx import Presentation


def _normalized(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("•", " ")).strip().casefold()


def _contains(text: str, unit: str) -> bool:
    return (
        re.search(r"(?<!\w)" + re.escape(_normalized(unit)) + r"(?!\w)", text)
        is not None
    )


_NUMBER = re.compile(
    r"(?<![\w.,])\d{1,3}(?:[ \u00a0\u202f]\d{3})+(?:[.,]\d+)?|\d+(?:[.,]\d+)?"
)


def _numbers(text: str) -> set[str]:
    found = set()
    for match in _NUMBER.finditer(text):
        value = re.sub(r"[ \u00a0\u202f]", "", match.group()).replace(",", ".")
        found.add(value.rstrip("0").rstrip(".") if "." in value else value)
    return found


def _strings(value) -> list[str]:
    if isinstance(value, str):
        return [value]
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return [str(value)]
    if isinstance(value, dict):
        return [
            text
            for key, item in value.items()
            if key not in {"asset_id", "match", "match_score", "origin", "score"}
            for text in _strings(item)
        ]
    if isinstance(value, list):
        return [text for item in value for text in _strings(item)]
    return []


def visible_texts(slide: dict) -> list[str]:
    """Texts a slide shows; speaker notes are not visible."""
    return [
        str(text)
        for text in (
            slide.get("title", ""),
            slide.get("subtitle", ""),
            slide.get("body", ""),
            *slide.get("bullets", []),
            *_strings(slide.get("visual") or {}),
        )
    ]


_BUDGET = re.compile(r"\b(?:бюджет\w*|финансирован\w*|budget|funding)\b", re.IGNORECASE)
_BUDGET_APPROVAL = re.compile(
    r"\b(?:утверд\w*|одобр\w*|соглас\w*|выдел\w*|approve\w*|allocat\w*)\b",
    re.IGNORECASE,
)
_BUDGET_UNSPECIFIED = re.compile(
    r"\b(?:бюджет\w*|budget|funding)\b.{0,45}"
    r"(?:не указан\w*|не определ\w*|предстоит уточн\w*|not specified|unknown)",
    re.IGNORECASE,
)
_SECURITY = re.compile(
    r"\b(?:утеч\w*|риск\w*|безопасност\w*|защит\w*|leak\w*|risk\w*|security)\b",
    re.IGNORECASE,
)
_ABSOLUTE = re.compile(
    r"\b(?:исключ\w*|невозможн\w*|нулев\w*|полност\w*|гарантир\w*|"
    r"абсолют\w*|zero|eliminat\w*|guarantee\w*|impossible|never)\b|100\s*%",
    re.IGNORECASE,
)
_GUARANTEE = re.compile(r"\b(?:гарантир\w*|guarantee\w*)\b", re.IGNORECASE)
_DATA_CONTAINMENT = re.compile(
    r"\b(?:документ\w*|данн\w*|информаци\w*|documents?|data|information)\b"
    r".{0,85}?\b(?:не\s+(?:покида\w*|уход\w*|выход\w*)|"
    r"(?:never\s+)?leav\w*|remain\w*|stay\w*)\b",
    re.IGNORECASE,
)
_SELF_SERVICE = re.compile(
    r"\b(?:дол\w*|процент\w*|числ\w*|количеств\w*|share|rate|number)\b"
    r".{0,100}?\b(?:обращени\w*|вопрос\w*|запрос\w*|tickets?|requests?)\b"
    r".{0,100}?\b(?:без\s+(?:участи\w*\s+|привлечени\w*\s+)?специалист\w*|"
    r"without\s+(?:a\s+)?(?:specialist|agent|human))\b",
    re.IGNORECASE,
)
_DIRECTION = re.compile(
    r"\b(?P<down>сниж\w*|сократ\w*|уменьш\w*|падени\w*|"
    r"decreas\w*|reduc\w*|lower\w*)\b|"
    r"\b(?P<up>рост\w*|увелич\w*|повыш\w*|increas\w*|rais\w*)\b|"
    r"\b(?P<neutral>измер\w*|отслеж\w*|контрол\w*|оцен\w*|measure\w*|track\w*)\b",
    re.IGNORECASE,
)
_CURRENT_DURATION = re.compile(
    r"\b(?:до|up\s+to)\s*(?P<number>\d+(?:[.,]\d+)?)\s*"
    r"(?P<unit>минут\w*|час\w*|секунд\w*|minutes?|hours?|seconds?)\b",
    re.IGNORECASE,
)
_BASELINE_CUE = re.compile(
    r"\b(?:сейчас|сегодня|текущ\w*|занима\w*|трат\w*|"
    r"currently|today|now|takes?|spends?)\b",
    re.IGNORECASE,
)
_REDUCTION_WITH_DURATION = re.compile(
    r"\b(?:сокращ\w*|сократ\w*|сниж\w*|сниз\w*|уменьш\w*|"
    r"reduc\w*|decreas\w*|lower\w*)\b[^.!?;\n]{0,100}?"
    r"\b(?P<relation>до|с|to|from)\s*(?P<number>\d+(?:[.,]\d+)?)\s*"
    r"(?P<unit>минут\w*|час\w*|секунд\w*|minutes?|hours?|seconds?)\b",
    re.IGNORECASE,
)


def _duration_key(match: re.Match[str]) -> tuple[str, str]:
    number = match.group("number").replace(",", ".")
    number = number.rstrip("0").rstrip(".") if "." in number else number
    unit = match.group("unit").casefold()
    if unit.startswith(("минут", "minute")):
        return number, "minute"
    if unit.startswith(("секунд", "second")):
        return number, "second"
    return number, "hour"


def _baseline_durations(source: str) -> dict[tuple[str, str], str]:
    """Current durations and a short literal source phrase safe for replacement."""
    baselines: dict[tuple[str, str], str] = {}
    for match in _CURRENT_DURATION.finditer(source):
        start = max(source.rfind(char, 0, match.start()) for char in ".!?;\n") + 1
        before = source[start:match.start()]
        if not _BASELINE_CUE.search(before):
            continue
        end = min(
            (index for char in ".!?;\n" if (index := source.find(char, match.end())) >= 0),
            default=len(source),
        )
        phrase = source[start:end].strip(" \t,:—-")
        phrase = re.split(r",\s*(?:а|но|but|while)\b", phrase, maxsplit=1, flags=re.IGNORECASE)[0]
        baselines[_duration_key(match)] = phrase.strip(" \t,:—-")
    return baselines


def _baseline_as_target(source: str, text: str) -> bool:
    """A stated current duration must not become the same reduction target."""
    baselines = _baseline_durations(source)
    if not baselines:
        return False
    source_claims = {
        (_duration_key(match), match.group("relation").casefold())
        for match in _REDUCTION_WITH_DURATION.finditer(source)
    }
    for match in _REDUCTION_WITH_DURATION.finditer(text):
        if re.search(r"\b(?:не|без|not|without)\s+$", text[max(0, match.start() - 12):match.start()], re.IGNORECASE):
            continue
        key = _duration_key(match)
        relation = match.group("relation").casefold()
        if key in baselines and (key, relation) not in source_claims:
            return True
    return False
_TIME_WORDS = {
    "один": 1, "одна": 1, "одно": 1, "одного": 1,
    "два": 2, "две": 2, "двух": 2,
    "три": 3, "трех": 3, "трёх": 3,
    "четыре": 4, "четырех": 4, "четырёх": 4,
    "пять": 5, "пяти": 5,
    "шесть": 6, "шести": 6,
    "семь": 7, "семи": 7,
    "восемь": 8, "восьми": 8,
    "девять": 9, "девяти": 9,
    "десять": 10, "десяти": 10,
    "одиннадцать": 11, "одиннадцати": 11,
    "двенадцать": 12, "двенадцати": 12,
    "one": 1, "two": 2, "three": 3, "four": 4,
    "five": 5, "six": 6, "seven": 7, "eight": 8,
    "nine": 9, "ten": 10, "eleven": 11, "twelve": 12,
}
_TIME_QUANTITY = re.compile(
    r"(?<!\w)(?P<number>\d+|" + "|".join(_TIME_WORDS) + r")\s+"
    r"(?P<unit>дн\w*|недел\w*|месяц\w*|год\w*|лет|days?|weeks?|months?|years?)(?!\w)",
    re.IGNORECASE,
)
_DELIVERY_DEADLINE = (
    r"(?:запуст\w*|внедр\w*|разверн\w*|launch\w*|deploy\w*|release\w*|"
    r"начать|начн\w*|старт\w*|start\w*|begin\w*|"
    r"(?:для|до|к)\s+(?:запуск\w*|внедрен\w*|релиз\w*)|"
    r"запуск\w*.{0,40}(?:ожида\w*|запланир\w*|намеч\w*|долж\w*|состо\w*|произойд\w*))"
)
_DEADLINE_LINK = r"(?:за|в\s+течение|через|не\s+позднее|within|by|in)"
_PILOT_EVENT = re.compile(r"\b(?:пилот\w*|pilot|trial)\b", re.IGNORECASE)
_TEAM_DUTY = re.compile(
    r"\b(?:команд\w*|специалист\w*|аналитик\w*|инженер\w*|"
    r"разработчик\w*|ответственн\w*|менеджер\w*|"
    r"team|specialist|analyst|engineer|developer|owner|manager)\b.{0,85}?"
    r"\b(?:займ\w*|буд\w*|отвеч\w*|настро\w*|анализир\w*|"
    r"подготов\w*|разработ\w*|провед\w*|интегрир\w*|"
    r"подбир\w*|координир\w*|will|responsib\w*|"
    r"configur\w*|analy[sz]\w*|develop\w*)\b",
    re.IGNORECASE,
)
_DUTY_ACTION = re.compile(
    r"\b(?:настр\w*|анализ\w*|оцени\w*|разраб\w*|внедр\w*|"
    r"обуч\w*|подгот\w*|тестир\w*|провер\w*|монитор\w*|"
    r"исслед\w*|созда\w*|интегр\w*|подбир\w*|коорд\w*|"
    r"configur\w*|analy[sz]\w*|develop\w*|"
    r"implement\w*|train\w*|test\w*|monitor\w*|prepare\w*)\b",
    re.IGNORECASE,
)
_ROLE_KEYS = {
    "team": re.compile(r"\b(?:команд\w*|team)\b", re.IGNORECASE),
    "specialist": re.compile(r"\b(?:специалист\w*|specialist)\b", re.IGNORECASE),
    "analyst": re.compile(r"\b(?:аналитик\w*|analyst)\b", re.IGNORECASE),
    "engineer": re.compile(r"\b(?:инженер\w*|engineer)\b", re.IGNORECASE),
    "developer": re.compile(r"\b(?:разработчик\w*|developer)\b", re.IGNORECASE),
    "owner": re.compile(r"\b(?:ответственн\w*|owner)\b", re.IGNORECASE),
    "manager": re.compile(r"\b(?:менеджер\w*|manager)\b", re.IGNORECASE),
}
_BROAD_ROLLOUT = re.compile(
    r"\b(?:на\s+всю\s+(?:компани\w*|организаци\w*)|"
    r"для\s+всех\s+сотрудник\w*|полн\w*\s+внедрен\w*|"
    r"company[- ]wide|all\s+employees|full\s+rollout)\b",
    re.IGNORECASE,
)
_SOURCE_ROLLOUT = re.compile(
    r"\b(?:масштаб\w*|всю\s+(?:компани\w*|организаци\w*)|"
    r"всех\s+сотрудник\w*|полн\w*\s+внедрен\w*|rollout|"
    r"company[- ]wide|all\s+employees)\b",
    re.IGNORECASE,
)
_ROLLOUT_DECISION = re.compile(
    r"\b(?:реш\w*\s+о|рассмотр\w*|оцен\w*\s+(?:возможност\w*|"
    r"целесообразност\w*)|приня\w*\s+решен\w*|decide\s+whether|"
    r"consider|evaluate\s+whether)\b",
    re.IGNORECASE,
)
_PILOT_APPROVAL_REQUEST = re.compile(
    r"\b(?:просим|request|seek)\b[^.!?;\n]{0,85}?"
    r"\b(?:одобр\w*|утверд\w*|соглас\w*|approv\w*)\b"
    r"[^.!?;\n]{0,55}?\b(?:пилот\w*|pilot|trial)\b",
    re.IGNORECASE,
)
_PILOT_REPORTED_STATUS = re.compile(
    r"\b(?:пилот\w*|pilot|trial)\b(?P<between>[^.!?;\n]{0,65}?)"
    r"\b(?:запущен\w*|запустил\w*|стартовал\w*|начал(?:ся|ась|ись)|"
    r"заверш[её]н\w*|завершил(?:ся|ась|ись)|окончен\w*|"
    r"launched|started|completed|finished)\b",
    re.IGNORECASE,
)
_PILOT_STATUS_MODAL = re.compile(
    r"\b(?:не|будет|после|если|планиру\w*|ожида\w*|"
    r"not|will|would|could|after|if|planned)\b",
    re.IGNORECASE,
)


def _asserts_pilot_started_or_finished(text: str) -> bool:
    """Recognise an affirmative completed status, not a proposed next step."""
    return any(
        not _PILOT_STATUS_MODAL.search(match.group("between"))
        for match in _PILOT_REPORTED_STATUS.finditer(text)
    )


def _pilot_awaits_approval(source: str) -> bool:
    return bool(_PILOT_APPROVAL_REQUEST.search(source)) and not _asserts_pilot_started_or_finished(source)


def _time_unit(unit: str) -> str:
    unit = unit.casefold()
    if unit.startswith(("дн", "day")):
        return "day"
    if unit.startswith(("недел", "week")):
        return "week"
    if unit.startswith(("месяц", "month")):
        return "month"
    return "year"


def _time_claims(text: str) -> list[tuple[int, str, str]]:
    """Numeric duration and the event it describes, when the syntax is clear."""
    claims = []
    for match in _TIME_QUANTITY.finditer(text):
        token = match.group("number").casefold()
        number = int(token) if token.isdigit() else _TIME_WORDS[token]
        before = text[max(0, match.start() - 110) : match.start()]
        after = text[match.end() : match.end() + 65]
        deadline_before = re.search(
            rf"\b{_DELIVERY_DEADLINE}\b.{{0,75}}\b{_DEADLINE_LINK}\s*$",
            before,
            re.IGNORECASE,
        )
        deadline_after = re.search(
            rf"\b{_DEADLINE_LINK}\s*$", before, re.IGNORECASE
        ) and re.search(rf"^\W*.{{0,30}}\b{_DELIVERY_DEADLINE}\b", after, re.IGNORECASE)
        if deadline_before or deadline_after:
            relation = "delivery_deadline"
        elif _PILOT_EVENT.search(before + after):
            relation = "pilot_duration"
        else:
            relation = "unspecified"
        claims.append((number, _time_unit(match.group("unit")), relation))
    return claims


def _duty_stems(text: str) -> set[str]:
    return {match.group().casefold()[:5] for match in _DUTY_ACTION.finditer(text)}


def _role_keys(text: str) -> set[str]:
    return {key for key, pattern in _ROLE_KEYS.items() if pattern.search(text)}


def _source_team_duties(source: str) -> list[tuple[set[str], set[str]]]:
    return [
        (_role_keys(sentence), _duty_stems(sentence))
        for sentence in re.split(r"[.!?;\n]", source)
        if _TEAM_DUTY.search(sentence)
    ]


def _decreases_self_service(text: str) -> bool:
    for metric in _SELF_SERVICE.finditer(text):
        prefix = text[max(0, metric.start() - 45) : metric.start()]
        directions = list(_DIRECTION.finditer(prefix))
        if not directions or directions[-1].lastgroup != "down":
            continue
        before = prefix[: directions[-1].start()]
        if re.search(r"(?:\bне\s+|\bnot\s+)$", before, re.IGNORECASE):
            continue
        return True
    return False


def source_backed_kpi_fallback(source: str, slide: dict) -> dict | None:
    """Replace an inverted self-service claim with the brief's exact metric.

    This is intentionally limited to plain text fields and a source that only
    names the metric without a direction. Other claim types still need review.
    """
    metric = _SELF_SERVICE.search(source)
    if metric is None or _decreases_self_service(source):
        return None
    prefix = source[max(0, metric.start() - 45) : metric.start()]
    directions = list(_DIRECTION.finditer(prefix))
    if directions and directions[-1].lastgroup == "up":
        return None
    label = "Показатель: " if re.search(r"[А-Яа-яЁё]", source) else "Metric: "
    replacement = label + metric.group().strip()
    if len(replacement) > 120:
        return None
    revised = deepcopy(slide)
    changed = False
    for key in ("title", "subtitle", "body", "speaker_notes"):
        value = revised.get(key)
        if isinstance(value, str) and _decreases_self_service(value):
            revised[key] = replacement
            changed = True
    bullets = revised.get("bullets")
    if isinstance(bullets, list):
        revised["bullets"] = [
            replacement if isinstance(item, str) and _decreases_self_service(item) else item
            for item in bullets
        ]
        changed |= revised["bullets"] != bullets
    return revised if changed else None


def source_backed_baseline_fallback(source: str, slide: dict) -> dict | None:
    """Replace a current-duration-as-target claim with the source's words.

    This deliberately handles plain text only. A visual claim that cannot be
    repaired by the model remains a blocking grounding issue.
    """
    baselines = _baseline_durations(source)
    if not baselines:
        return None
    source_claims = {
        (_duration_key(match), match.group("relation").casefold())
        for match in _REDUCTION_WITH_DURATION.finditer(source)
    }

    def replacement(value: str) -> str:
        for match in _REDUCTION_WITH_DURATION.finditer(value):
            key = _duration_key(match)
            phrase = baselines.get(key)
            relation = match.group("relation").casefold()
            if phrase and (key, relation) not in source_claims and len(phrase) <= 100:
                return phrase
        return value

    revised = deepcopy(slide)
    changed = False
    for key in ("title", "subtitle", "body", "goal", "speaker_notes"):
        value = revised.get(key)
        if isinstance(value, str) and _baseline_as_target(source, value):
            revised[key] = replacement(value)
            changed |= revised[key] != value
    bullets = revised.get("bullets")
    if isinstance(bullets, list):
        revised["bullets"] = [
            replacement(item) if isinstance(item, str) and _baseline_as_target(source, item) else item
            for item in bullets
        ]
        changed |= revised["bullets"] != bullets
    return revised if changed else None


def source_backed_pilot_status_fallback(source: str, slide: dict) -> dict | None:
    """Correct only an explicit false launch of a pilot awaiting approval."""
    if not _pilot_awaits_approval(source):
        return None

    def replacement(value: str) -> str:
        return re.sub(
            r"\b(пилот(?:ный проект)?)\s+запущен\b",
            lambda match: f"{match.group(1)} планируется",
            value,
            flags=re.IGNORECASE,
        )

    revised = deepcopy(slide)
    changed = False
    for key in ("title", "subtitle", "body", "speaker_notes"):
        value = revised.get(key)
        if isinstance(value, str):
            revised[key] = replacement(value)
            changed |= revised[key] != value
    bullets = revised.get("bullets")
    if isinstance(bullets, list):
        revised["bullets"] = [
            replacement(item) if isinstance(item, str) else item for item in bullets
        ]
        changed |= revised["bullets"] != bullets
    return revised if changed else None


def grounding_findings(source: str, plan: dict) -> list[dict]:
    """High-confidence, exact-quote findings for risky brief extrapolations.

    These rules are deliberately narrow. They do not certify every claim.
    """
    source_budget = bool(_BUDGET.search(source))
    source_budget_approval = any(
        _BUDGET.search(part) and _BUDGET_APPROVAL.search(part)
        for part in re.split(r"[.!?\n]", source)
    )
    source_decreases_self_service = _decreases_self_service(source)
    source_time_claims = _time_claims(source)
    source_duties = _source_team_duties(source)
    source_allows_rollout = bool(_SOURCE_ROLLOUT.search(source))
    source_has_pilot = bool(_PILOT_EVENT.search(source))
    source_pilot_pending = _pilot_awaits_approval(source)
    normalized_source = _normalized(source)
    findings = []
    seen = set()
    for number, slide in enumerate(plan.get("slides", []), 1):
        # Notes are shipped in the PPTX even though they do not count towards
        # visual coverage, so their factual claims need the same grounding.
        texts = [*visible_texts(slide), str(slide.get("speaker_notes") or "")]
        visual = slide.get("visual") or {}
        if isinstance(visual, dict):
            for item in visual.get("items") or []:
                if isinstance(item, dict):
                    label = str(item.get("label") or "").strip()
                    detail = str(item.get("detail") or "").strip()
                    if label and detail:
                        texts.append(f"{label} — {detail}")
        for text in texts:
            if not text.strip() or _normalized(text) in normalized_source:
                continue
            reasons = []
            if _BUDGET.search(text) and not _BUDGET_UNSPECIFIED.search(text):
                if not source_budget:
                    reasons.append("В брифе нет бюджета: не добавляйте его в решение или запрос на одобрение.")
                elif _BUDGET_APPROVAL.search(text) and not source_budget_approval:
                    reasons.append("Бриф не запрашивает утверждение бюджета.")
            if _SECURITY.search(text) and _ABSOLUTE.search(text):
                reasons.append("Абсолютная гарантия безопасности не следует из брифа; используйте его точную формулировку.")
            if (
                _GUARANTEE.search(text)
                and _DATA_CONTAINMENT.search(text)
                and not _GUARANTEE.search(source)
            ):
                reasons.append(
                    "Бриф утверждает, что документы остаются внутри контура, "
                    "но не даёт основания обещать гарантию этого свойства."
                )
            if _decreases_self_service(text) and not source_decreases_self_service:
                reasons.append("Уменьшение доли обращений без специалиста меняет смысл метрики успеха.")
            if _baseline_as_target(source, text):
                reasons.append(
                    "В брифе текущее время указано как верхняя граница, а не точный "
                    "исходный уровень или целевой результат сокращения."
                )
            if source_pilot_pending and _asserts_pilot_started_or_finished(text):
                reasons.append("Бриф просит одобрить пилот и не сообщает, что он уже начался или завершился.")
            broad_scope = _BROAD_ROLLOUT.search(text)
            if broad_scope and source_has_pilot and not source_allows_rollout:
                before_scope = text[max(0, broad_scope.start() - 35) : broad_scope.start()]
                negated = re.search(r"\b(?:не|not)\s+(?:\w+\s+){0,2}$", before_scope, re.IGNORECASE)
                if not negated and not _ROLLOUT_DECISION.search(text):
                    reasons.append("Бриф ограничен пилотом и не обещает масштабирование на всю компанию.")
            for time_claim in _time_claims(text):
                quantity = time_claim[:2]
                matches = [claim for claim in source_time_claims if claim[:2] == quantity]
                if not matches:
                    reasons.append("Такой срок не указан в брифе; не придумывайте дату или длительность.")
                elif (
                    time_claim[2] != "unspecified"
                    and all(claim[2] != "unspecified" and claim[2] != time_claim[2] for claim in matches)
                ):
                    reasons.append("Число есть в брифе, но относится к другому событию: длительность не равна сроку запуска.")
            actions = _duty_stems(text)
            if (
                actions
                and _TEAM_DUTY.search(text)
                and not any(
                    _role_keys(text) & roles and actions <= supported
                    for roles, supported in source_duties
                )
            ):
                reasons.append("Бриф не назначает команде эти обязанности; не выводите их из численности команды.")
            for reason in reasons:
                key = (number, text, reason)
                if key in seen:
                    continue
                seen.add(key)
                findings.append(
                    {"slide": number, "code": "unsupported_claim", "quote": text[:160], "reason": reason}
                )
    return findings


def unsupported_numbers(
    source: str, plan: dict, *, include_notes: bool = False
) -> list[dict]:
    """Visible numbers of the plan that the source never states.

    Integers below 10 are skipped: they usually count the slide's own items.
    Speaker notes are excluded from visible coverage, but brief generation can
    opt into checking them because they are saved in the final PPTX.
    """
    known = _numbers(source)
    result = []
    for number, slide in enumerate(plan.get("slides", []), 1):
        texts = visible_texts(slide)
        if include_notes:
            texts.append(str(slide.get("speaker_notes") or ""))
        missing = sorted(
            value
            for value in _numbers("\n".join(texts)) - known
            if "." in value or float(value) >= 10
        )
        if missing:
            result.append({"slide": number, "numbers": missing})
    return result


def coverage_report(source: str, plan: dict, output: Path | None = None) -> dict:
    # Imported here because brief planning also uses grounding_findings.
    from .planner import _sections, _sentences

    units = []
    for heading, lines in _sections(source):
        if heading:
            units.append(heading)
        for line in lines:
            units.extend(_sentences(line))
    slides = plan.get("slides", [])
    planned = _normalized(
        "\n".join(
            "\n".join(
                [
                    slide.get("title", ""),
                    slide.get("subtitle", ""),
                    slide.get("body", ""),
                    *slide.get("bullets", []),
                    json.dumps(slide.get("visual") or {}, ensure_ascii=False),
                ]
            )
            for slide in slides
        )
    )
    rendered = None
    if output is not None:
        texts = []

        def visit(shapes):
            for shape in shapes:
                if hasattr(shape, "shapes"):
                    visit(shape.shapes)
                if shape.has_text_frame:
                    texts.append(shape.text)
                if shape.has_table:
                    texts.extend(
                        cell.text for row in shape.table.rows for cell in row.cells
                    )
                if shape.has_chart:
                    for series in shape.chart.series:
                        texts.extend([series.name, *map(str, series.values)])
                    texts.extend(
                        str(category.label)
                        for category in shape.chart.plots[0].categories
                    )

        for slide in Presentation(output).slides:
            visit(slide.shapes)
        rendered = _normalized("\n".join(texts))
    rows = [
        {
            "source_id": f"s{index:04}",
            "text": text,
            "in_plan": _contains(planned, text),
            "in_pptx": _contains(rendered, text) if rendered is not None else None,
        }
        for index, text in enumerate(units, 1)
    ]
    missing_plan = [r["source_id"] for r in rows if not r["in_plan"]]
    missing_pptx = [r["source_id"] for r in rows if r["in_pptx"] is False]
    return {
        "method": "literal_source_units_v1",
        "semantic_verification": "not_performed",
        "source_sha256": hashlib.sha256(source.encode()).hexdigest(),
        "status": "needs_review"
        if missing_plan or missing_pptx
        else "literal_coverage_complete",
        "units": rows,
        "missing_from_plan": missing_plan,
        "missing_from_pptx": missing_pptx,
        "unsupported_numbers": unsupported_numbers(source, plan),
        "grounding_findings": grounding_findings(source, plan),
        "note": "Unmatched paraphrases are not proof of omission; speaker notes do not count as slide content.",
    }
