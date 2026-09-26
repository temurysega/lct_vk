"""Template decoration that new content lies across or hides.

A large picture or band kept from a template slide frames the slide, so the
structural audit does not treat it as a colliding block, and layout art is
not on the slide at all. Either still spoils a slide when new content lies
across its edge: a title crossed by a decorative ring, a diagram cut by a
colour band, text over a logo. Decoration is judged by what it draws: the
opaque pixels of a picture that differ from the slide background, or the
outline of a filled shape. Content fully on a panel or fully beside it is
the template's intended composition.
"""

from __future__ import annotations

import io
import textwrap
from typing import Any

from PIL import Image, ImageChops, ImageStat
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

from .layout_geometry import (
    Box,
    area,
    box,
    has_text,
    intersection,
    is_generated,
    is_structural,
    is_visible_box,
)
from .typography import effective_font_size

CELL = 0.05  # inches per ink-grid cell
EMU = 914400
# Share of a content block lying on decoration: below it the block is beside
# the decoration, above it the block sits on a panel.
PARTIAL = (0.12, 0.88)
EXEMPT = ("Native", "Brand Asset", "Background")
SMALL_ART = 2.0  # square inches: a logo or ornament, not a panel
TEXT_GLYPH_EM = 0.5  # average rendered glyph width for the text extent
SKIPPED_TYPES = {MSO_SHAPE_TYPE.LINE, MSO_SHAPE_TYPE.PLACEHOLDER}


def _grid(canvas: tuple[float, float]) -> tuple[int, int]:
    return max(1, round(canvas[0] / CELL)), max(1, round(canvas[1] / CELL))


def _hex(value: str) -> tuple[int, int, int]:
    value = str(value or "FFFFFF").lstrip("#")[:6].ljust(6, "0")
    return tuple(int(value[index : index + 2], 16) for index in (0, 2, 4))


def _children(group: Any, zone: Box) -> list[tuple[Any, Box]]:
    """Group members with their boxes on the slide."""
    xfrm = group.element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
    try:
        off, ext = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
        origin_x, origin_y = int(off.get("x")) / EMU, int(off.get("y")) / EMU
        scale_x = zone[2] / max(int(ext.get("cx")) / EMU, 1e-6)
        scale_y = zone[3] / max(int(ext.get("cy")) / EMU, 1e-6)
    except (AttributeError, TypeError, ValueError):
        origin_x, origin_y, scale_x, scale_y = zone[0], zone[1], 1.0, 1.0
    result = []
    for child in group.shapes:
        x, y, w, h = box(child)
        result.append(
            (
                child,
                (
                    zone[0] + (x - origin_x) * scale_x,
                    zone[1] + (y - origin_y) * scale_y,
                    w * scale_x,
                    h * scale_y,
                ),
            )
        )
    return result


def _picture_ink(shape: Any, background: tuple[int, int, int]) -> Image.Image | None:
    """Opaque picture pixels that stand out from the slide background."""
    try:
        image = Image.open(io.BytesIO(shape.image.blob))
        image.load()
    except Exception:  # noqa: BLE001 - vector or unreadable media is not judged
        return None
    width, height = image.size
    crops = [
        max(0.0, float(getattr(shape, name, 0) or 0))
        for name in ("crop_left", "crop_top", "crop_right", "crop_bottom")
    ]
    frame = (
        round(crops[0] * width),
        round(crops[1] * height),
        round((1 - crops[2]) * width),
        round((1 - crops[3]) * height),
    )
    if frame[2] <= frame[0] or frame[3] <= frame[1]:
        return None
    image = image.crop(frame)
    image.thumbnail((480, 480))
    rgba = image.convert("RGBA")
    red, green, blue = ImageChops.difference(
        rgba.convert("RGB"), Image.new("RGB", rgba.size, background)
    ).split()
    distinct = ImageChops.lighter(ImageChops.lighter(red, green), blue)
    distinct = distinct.point(lambda value: 255 if value > 48 else 0)
    solid = rgba.getchannel("A").point(lambda value: 255 if value > 96 else 0)
    return ImageChops.multiply(distinct, solid)


def _fill_differs(shape: Any, background: tuple[int, int, int]) -> bool:
    """A filled shape in the background colour draws nothing visible."""
    try:
        color = shape.fill.fore_color.rgb
    except (AttributeError, TypeError, ValueError):
        return True  # theme or gradient fill: assume it shows
    distance = max(abs(a - b) for a, b in zip(_hex(str(color)), background))
    line = shape.element.find(qn("p:spPr")).find(qn("a:ln"))
    outlined = line is not None and any(
        line.find(qn(tag)) is not None for tag in ("a:solidFill", "a:gradFill")
    )
    return distance > 24 or outlined


def _paint(
    layer: Image.Image,
    shape: Any,
    zone: Box,
    background: tuple[int, int, int],
    *,
    erase: bool = False,
) -> None:
    """Draw a shape's ink; with ``erase`` a shape in the background colour hides ink."""
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child, child_zone in _children(shape, zone):
            _paint(layer, child, child_zone, background, erase=erase)
        return
    if shape.shape_type in SKIPPED_TYPES or zone[2] <= 0 or zone[3] <= 0:
        return
    size = (max(1, round(zone[2] / CELL)), max(1, round(zone[3] / CELL)))
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or shape.element.tag == qn("p:pic"):
        ink = _picture_ink(shape, background)
        if ink is None:
            return
        ink = ink.resize(size, Image.BILINEAR).point(
            lambda value: 255 if value >= 128 else 0
        )
    elif is_visible_box(shape) and _fill_differs(shape, background) or has_text(shape):
        ink = Image.new("L", size, 255)
    elif erase and is_visible_box(shape):
        left, top = round(zone[0] / CELL), round(zone[1] / CELL)
        layer.paste(
            0,
            (
                max(0, left),
                max(0, top),
                min(layer.width, left + size[0]),
                min(layer.height, top + size[1]),
            ),
        )
        return
    else:
        return
    layer.paste(ink, (round(zone[0] / CELL), round(zone[1] / CELL)), ink)


def _text_extent(shape: Any, zone: Box, default: float) -> Box:
    """The part of a text box its words occupy, from an estimated wrap."""
    frame = shape.text_frame
    sizes = [
        run.font.size.pt for p in frame.paragraphs for run in p.runs if run.font.size
    ]
    size = (
        max(sizes) if sizes else float(effective_font_size(shape, default) or default)
    )
    size = max(6.0, size)
    left, right = frame.margin_left / EMU, frame.margin_right / EMU
    top, bottom = frame.margin_top / EMU, frame.margin_bottom / EMU
    inner = max(0.1, zone[2] - left - right)
    glyph = size * TEXT_GLYPH_EM / 72
    per_line = max(1, int(inner / glyph))
    lines, widest = 0, 0.0
    for paragraph in frame.paragraphs:
        wrapped = textwrap.wrap(paragraph.text, per_line, break_on_hyphens=False) or [
            ""
        ]
        lines += len(wrapped)
        widest = max(widest, max(len(line) for line in wrapped) * glyph)
    width = min(inner, widest)
    # Fonts render taller than 1.2 em lines, and paragraphs keep spacing;
    # text may also run past a box it overflows.
    spacing = (len(frame.paragraphs) - 1) * size * 0.4 / 72
    height = lines * size * 1.3 / 72 + spacing + top + bottom
    alignment = frame.paragraphs[0].alignment
    if alignment == PP_ALIGN.CENTER:
        x = zone[0] + left + (inner - width) / 2
    elif alignment == PP_ALIGN.RIGHT:
        x = zone[0] + zone[2] - right - width
    else:
        x = zone[0] + left
    anchor = frame._txBody.find(qn("a:bodyPr")).get("anchor")
    if anchor == "ctr":
        y = zone[1] + (zone[3] - height) / 2
    elif anchor == "b":
        y = zone[1] + zone[3] - height
    else:
        y = zone[1]
    return (x, y, width, height)


def _marks(shape: Any, zone: Box, default: float) -> list[tuple[str, Box, bool]]:
    """Visible pieces of generated content: cards, pictures, text extents.

    The flag tells a drawn shape from text, so that a diagram can also be
    judged as one drawing.
    """
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [
            mark
            for child, child_zone in _children(shape, zone)
            for mark in _marks(child, child_zone, default)
        ]
    if shape.shape_type == MSO_SHAPE_TYPE.LINE or any(
        word in shape.name for word in EXEMPT
    ):
        return []
    if is_visible_box(shape):
        return [(shape.name, zone, True)]  # a card or segment draws its whole shape
    if has_text(shape):
        return [(shape.name, _text_extent(shape, zone, default), False)]
    return []


def _union(boxes: list[Box]) -> Box:
    left = min(x for x, _, _, _ in boxes)
    top = min(y for _, y, _, _ in boxes)
    right = max(x + w for x, _, w, _ in boxes)
    bottom = max(y + h for _, y, _, h in boxes)
    return (left, top, right - left, bottom - top)


def _decorations(slide: Any, canvas: tuple[float, float]) -> list[tuple[Any, bool]]:
    """Template shapes drawn under the new content, flagged when from the layout.

    Smaller template shapes on the slide are covered by the template-layer
    check; large ones frame the slide and are judged here, with the
    non-placeholder art of the layout and, when shown, of the master.
    """
    result = [
        (shape, False)
        for shape in slide.shapes
        if not is_generated(shape)
        and not has_text(shape)
        and not shape.is_placeholder
        and is_visible_box(shape)
        and is_structural(box(shape), canvas)
    ]
    layout = slide.slide_layout
    sources = [layout]
    if layout.element.get("showMasterSp") != "0":
        sources.append(layout.slide_master)
    for source in sources:
        result.extend(
            (shape, True) for shape in source.shapes if not shape.is_placeholder
        )
    return result


def _layout_visibility(
    slide: Any, canvas: tuple[float, float], background: tuple[int, int, int]
) -> Image.Image:
    """Where layout and master art shows: later shapes in the background colour mask it."""
    layer = Image.new("L", _grid(canvas), 0)
    layout = slide.slide_layout
    sources = [layout.slide_master] if layout.element.get("showMasterSp") != "0" else []
    for source in [*sources, layout]:
        for shape in source.shapes:
            if not shape.is_placeholder:
                _paint(layer, shape, box(shape), background, erase=True)
    return layer


def _ink(
    shape: Any,
    canvas: tuple[float, float],
    background: tuple[int, int, int],
    visible: Image.Image | None = None,
) -> tuple[Image.Image, float]:
    layer = Image.new("L", _grid(canvas), 0)
    _paint(layer, shape, box(shape), background)
    if visible is not None:
        layer = ImageChops.multiply(layer, visible)
    return layer, ImageStat.Stat(layer).sum[0] / 255 * CELL * CELL


def layout_art_boxes(slide: Any, design: dict[str, Any]) -> list[Box]:
    """Where the layout draws logos and ornaments that content must not hide."""
    canvas = (
        float(design["canvas"]["width_inches"]),
        float(design["canvas"]["height_inches"]),
    )
    background = _hex(design.get("brand", {}).get("background", "FFFFFF"))
    boxes = []
    visible = _layout_visibility(slide, canvas, background)
    for shape, from_layout in _decorations(slide, canvas):
        if not from_layout:
            continue
        layer, total = _ink(shape, canvas, background, visible)
        bounds = layer.getbbox()
        if bounds and 0.02 <= total <= SMALL_ART:
            left, top, right, bottom = (value * CELL for value in bounds)
            boxes.append((left, top, right - left, bottom - top))
    return boxes


def _layout_placed(shape: Any) -> bool:
    """A placeholder where its layout puts it: the layout composed it with its art.

    The height may have grown to fit the text; the anchor and width are kept.
    """
    if not shape.is_placeholder:
        return False
    try:
        layout_shape = shape.part.slide.slide_layout.placeholders.get(
            idx=shape.placeholder_format.idx
        )
    except (AttributeError, KeyError, ValueError):
        return False
    return layout_shape is not None and all(
        abs(getattr(shape, name) - getattr(layout_shape, name)) <= 0.05 * EMU
        for name in ("left", "top", "width")
    )


def _designed(mark: Box, template: list[Box], canvas: tuple[float, float]) -> bool:
    """A block placed where the template sample had one, up to canvas clipping."""
    for x, y, w, h in template:
        left, top = max(0.0, x), max(0.0, y)
        clipped = (left, top, min(canvas[0], x + w) - left, min(canvas[1], y + h) - top)
        if clipped[2] <= 0 or clipped[3] <= 0:
            continue
        if intersection(mark, clipped) >= 0.92 * max(area(mark), area(clipped)):
            return True
    return False


def decoration_overlap_issues(
    slide: Any,
    canvas: tuple[float, float],
    design: dict[str, Any] | None,
    default_size: float,
    template_shapes: set[tuple[Any, ...]] = frozenset(),
) -> list[dict[str, Any]]:
    """Decoration that new content lies across or hides.

    A picture or table placed where the template sample had a block follows
    the template's own composition, and so does a placeholder over the art
    of its own layout. A backdrop drawing over half the slide is the ground
    content is meant to sit on.
    """
    design = design or {}
    background = _hex(design.get("brand", {}).get("background", "FFFFFF"))
    template = [
        tuple(value / EMU for value in geometry) for geometry in template_shapes
    ]
    marks = []
    for shape in slide.shapes:
        if not is_generated(shape):
            continue
        zone = box(shape)
        if not has_text(shape) and _designed(zone, template, canvas):
            continue
        placed = _layout_placed(shape)
        pieces = _marks(shape, zone, default_size)
        drawn = [mark for _, mark, is_drawn in pieces if is_drawn]
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP and len(drawn) > 1:
            # A diagram cut by a band as a whole, not only piece by piece.
            pieces.append((shape.name, _union(drawn), True))
        marks.extend(
            (name, mark, placed) for name, mark, _ in pieces if area(mark) >= 0.05
        )
    if not marks:
        return []
    issues = []
    size = _grid(canvas)
    visible = _layout_visibility(slide, canvas, background)
    for decoration, from_layout in _decorations(slide, canvas):
        layer, total = _ink(
            decoration, canvas, background, visible if from_layout else None
        )
        if total < 0.02 or total > 0.5 * canvas[0] * canvas[1]:
            continue
        for name, mark, placed in marks:
            if placed and from_layout:
                continue
            left, top = max(0, round(mark[0] / CELL)), max(0, round(mark[1] / CELL))
            right = min(size[0], round((mark[0] + mark[2]) / CELL))
            bottom = min(size[1], round((mark[1] + mark[3]) / CELL))
            if right <= left or bottom <= top:
                continue
            share = ImageStat.Stat(layer.crop((left, top, right, bottom))).mean[0] / 255
            covered = share * (right - left) * (bottom - top) * CELL * CELL
            # Covering nearly all of a large drawing replaces it, not crosses it.
            across = PARTIAL[0] <= share <= PARTIAL[1] and covered < 0.85 * total
            hides = total <= SMALL_ART and covered >= max(0.01, 0.1 * total)
            if across or hides:
                zone = box(decoration)
                issues.append(
                    {
                        "severity": "warning",
                        "code": "decor_overlap",
                        "shape": f"{decoration.name} / {name}",
                        "ratio": round(share, 2),
                        "bounds": [
                            dict(zip("xywh", (round(value, 3) for value in zone)))
                        ],
                        "message": "Контент закрывает логотип или орнамент шаблона"
                        if hides
                        else "Контент лежит поперёк декоративного элемента шаблона",
                    }
                )
                break
    return issues
