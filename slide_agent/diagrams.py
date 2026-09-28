"""SmartArt-style diagrams, pictogram grids and illustrations from native shapes.

Every diagram is one named group of editable autoshapes, connectors, text
boxes and vector pictograms. Real OOXML SmartArt (``dgm`` parts) is not
emitted: LibreOffice renders SmartArt only from a cached drawing, so exported
PDF/HTML could silently differ from the PPTX. Colours come from the template
theme and are checked for WCAG contrast before text is placed on them.
"""

from __future__ import annotations

import math
import re
from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .pictograms import add_icon, add_polygon, assign_icons, choose_icon, resolve_icon
from .typography import ESTIMATED_GLYPH_WIDTH_EM, estimated_line_count
from .vector import drop_theme_style

DIAGRAM_TYPES = (
    "process",
    "cycle",
    "hierarchy",
    "pyramid",
    "funnel",
    "matrix",
    "icon_grid",
)
ITEM_LIMITS = {
    "process": (2, 6),
    "cycle": (3, 6),
    "hierarchy": (2, 5),
    "pyramid": (2, 5),
    "funnel": (2, 5),
    "matrix": (4, 4),
    "icon_grid": (2, 6),
}
_LABEL_SPLIT = re.compile(r"^(?P<head>[^:—–]{2,48}?(?:\s*[:—–]))(?P<tail>\s+\S.*)$")


# ---------------------------------------------------------------- colours ---
def _hex(value: Any, fallback: str) -> str:
    candidate = str(value or "").lstrip("#").upper()
    if len(candidate) == 8:
        candidate = candidate[-6:]
    return candidate if re.fullmatch(r"[0-9A-F]{6}", candidate) else fallback


def luminance(value: str) -> float:
    color = _hex(value, "000000")
    channels = [int(color[i : i + 2], 16) / 255 for i in (0, 2, 4)]
    linear = [
        c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4 for c in channels
    ]
    return 0.2126 * linear[0] + 0.7152 * linear[1] + 0.0722 * linear[2]


def contrast(first: str, second: str) -> float:
    a, b = sorted((luminance(first), luminance(second)), reverse=True)
    return (a + 0.05) / (b + 0.05)


def mix(first: str, second: str, ratio: float) -> str:
    ratio = min(1.0, max(0.0, ratio))
    left, right = _hex(first, "000000"), _hex(second, "FFFFFF")
    return "".join(
        f"{round(int(left[i : i + 2], 16) * (1 - ratio) + int(right[i : i + 2], 16) * ratio):02X}"
        for i in (0, 2, 4)
    )


def _saturation(value: str) -> float:
    color = _hex(value, "000000")
    channels = [int(color[i : i + 2], 16) / 255 for i in (0, 2, 4)]
    high, low = max(channels), min(channels)
    return 0.0 if high == 0 else (high - low) / high


def on_fill(
    fill: str, dark_text: str = "111111", *, prefer_shade: bool = True
) -> tuple[str, str]:
    """Return a fill and text colour pair with at least 4.5:1 contrast.

    A primary brand accent that is slightly too light for white text is
    deepened in small steps (like PowerPoint's theme shades) instead of
    switching to an unrelated colour. Tints keep their colour and use dark
    text when that already passes.
    """
    fill = _hex(fill, "0077FF")
    if contrast(fill, "FFFFFF") >= 4.5:
        return fill, "FFFFFF"
    if contrast(fill, dark_text) >= 4.5 and (luminance(fill) > 0.4 or not prefer_shade):
        return fill, dark_text
    for step in range(1, 9):
        shaded = mix(fill, "000000", step * 0.04)
        if contrast(shaded, "FFFFFF") >= 4.5:
            return shaded, "FFFFFF"
    return fill, (
        "FFFFFF" if contrast(fill, "FFFFFF") >= contrast(fill, dark_text) else dark_text
    )


def diagram_style(
    design: dict[str, Any], background: str | None = None, role: str = "content"
) -> dict[str, Any]:
    """Derive contrast-checked diagram tokens from the template design system."""
    theme = {
        str(item.get("role")): _hex(item.get("hex"), "")
        for item in design.get("colors", {}).get("theme", [])
    }
    brand = design.get("brand", {})
    prefix = (
        "cover"
        if role in {"cover", "section"}
        else "closing"
        if role == "closing"
        else "content"
    )
    background = _hex(
        background
        or brand.get(f"{prefix}_background")
        or brand.get("background")
        or theme.get("lt1"),
        "FFFFFF",
    )
    dark_background = contrast(background, "FFFFFF") > contrast(background, "000000")
    dark_text = _hex(theme.get("dk1"), "111111")
    if contrast(dark_text, "FFFFFF") < 7:
        dark_text = "111111"
    candidates = [theme.get(f"accent{index}") for index in range(1, 7)] + [
        theme.get("dk2"),
        brand.get("accent"),
    ]
    accents: list[str] = []
    for value in candidates:
        color = _hex(value, "")
        if not color or color in accents:
            continue
        if contrast(color, background) >= 2.2:
            accents.append(color)
    accents.sort(key=lambda color: 0 if _saturation(color) >= 0.35 else 1)
    accent = accents[0] if accents else ("FFFFFF" if dark_background else dark_text)
    text = "FFFFFF" if dark_background else dark_text
    for value in (brand.get(f"{prefix}_body"), brand.get("body")):
        color = _hex(value, "")
        if color and contrast(color, background) >= 4.5:
            text = color
            break
    heading = text
    for value in (brand.get(f"{prefix}_heading"), brand.get("heading")):
        color = _hex(value, "")
        if color and contrast(color, background) >= 4.5:
            heading = color
            break
    surface = mix(
        background,
        "FFFFFF" if dark_background else accent,
        0.1 if dark_background else 0.07,
    )
    if contrast(text, surface) >= 4.5:
        surface_text = text
    else:
        # Keep the pair together: on dark slides the card deepens for white text.
        surface, surface_text = on_fill(
            surface, dark_text, prefer_shade=dark_background
        )
    accent_fill, accent_text = on_fill(accent, dark_text)
    typography = design.get("typography", {})
    size = float(typography.get("body_size_pt", 16) or 16)
    scale = sorted(
        {
            float(value)
            for value in typography.get("observed_sizes_pt", [])
            if isinstance(value, (int, float)) and 9 <= float(value) <= 40
        }
    )
    return {
        "scale": scale,
        "background": background,
        "dark": dark_background,
        "accent": accent,
        "accents": accents or [accent],
        "accent_fill": accent_fill,
        "accent_text": accent_text,
        "text": text,
        "heading": heading,
        "surface": surface,
        "surface_text": surface_text,
        "line": mix(accent, background, 0.35),
        "dark_text": dark_text,
        "font": typography.get("primary_font") or "Arial",
        "heading_font": typography.get("heading_font")
        or typography.get("primary_font")
        or "Arial",
        "size": min(18.0, max(13.0, size)),
    }


def tint_series(
    style: dict[str, Any], count: int, *, spread: float = 0.55
) -> list[tuple[str, str]]:
    """Monochrome brand progression: darkest first, each with readable text."""
    result = []
    for index in range(count):
        ratio = spread * index / max(1, count - 1)
        fill = mix(style["accent_fill"], style["background"], ratio)
        result.append(on_fill(fill, style["dark_text"], prefer_shade=index == 0))
    return result


# ------------------------------------------------------------------ items ---
def normalize_item(item: Any) -> dict[str, Any]:
    """Return head/tail runs; string items keep their exact source wording."""
    if isinstance(item, dict):
        label = str(item.get("label") or item.get("title") or "").strip()
        detail = str(
            item.get("detail") or item.get("text") or item.get("description") or ""
        ).strip()
        value = str(item.get("value") or "").strip()
        if not label and not detail:
            label = value
            value = ""
        return {
            "head": label,
            "tail": detail,
            "inline": False,
            "value": value,
            "icon": item.get("icon"),
            "plain": " ".join(part for part in (label, detail) if part),
        }
    text = str(item or "").strip()
    match = _LABEL_SPLIT.match(text)
    if match and len(match.group("head").split()) <= 6:
        return {
            "head": match.group("head"),
            "tail": match.group("tail"),
            "inline": True,
            "value": "",
            "icon": None,
            "plain": text,
        }
    return {
        "head": "",
        "tail": text,
        "inline": False,
        "value": "",
        "icon": None,
        "plain": text,
    }


def visual_items(visual: dict[str, Any]) -> list[dict[str, Any]]:
    return [normalize_item(item) for item in visual.get("items", []) if item]


def _paragraphs(item: dict[str, Any] | str) -> list[str]:
    if isinstance(item, str):
        return [item] if item else []
    if item.get("inline"):
        return [item.get("head", "") + item.get("tail", "")]
    return [text for text in (item.get("head", ""), item.get("tail", "")) if text]


def text_height(paragraphs: list[str], width: float, size: float) -> float:
    """Estimated wrapped height in inches, using the same model as QA."""
    chars = max(1, int(max(0.1, width) * 72 / (size * ESTIMATED_GLYPH_WIDTH_EM)))
    lines = sum(estimated_line_count(text, chars) for text in paragraphs if text)
    gaps = max(0, len(paragraphs) - 1) * size * 0.25 / 72
    return lines * size * 1.22 / 72 + gaps + 0.08


def _sizes(size: float, minimum: float, scale: list[float]) -> list[float]:
    """Candidate sizes, largest first, preferring the template type scale."""
    allowed = [value for value in scale if minimum <= value <= size]
    if not allowed:
        # No template size in range: use the smallest legible template size.
        allowed = (
            [min(value for value in scale if value >= minimum)]
            if any(value >= minimum for value in scale)
            else []
        )
    if allowed:
        return sorted(allowed, reverse=True)
    steps, value = [], max(minimum, size)
    while value > minimum:
        steps.append(value)
        value -= 0.5
    return steps + [minimum]


def _fit(
    paragraphs: list[str],
    width: float,
    height: float,
    size: float,
    minimum: float,
    scale: list[float] | None = None,
) -> float:
    """Largest size whose estimated wrapped height fits, mirroring QA.

    The longest word must also fit the width: otherwise the office renderer
    breaks it mid-word ("Масштабировани / е").
    """
    candidates = _sizes(size, minimum, scale or [])
    longest = max((len(word) for text in paragraphs for word in text.split()), default=0)
    for candidate in candidates:
        word_width = longest * candidate * ESTIMATED_GLYPH_WIDTH_EM / 72
        if word_width <= width and text_height(paragraphs, width, candidate) <= height:
            return candidate
    return candidates[-1]


def _block_top(y: float, available: float, used: float) -> float:
    """Place a shorter block slightly above the optical centre of its zone."""
    return y + max(0.0, (available - used) * 0.38)


def _write(
    shape: Any,
    item: dict[str, Any] | str,
    style: dict[str, Any],
    *,
    color: str,
    size: float,
    minimum: float = 11.0,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
    bold: bool = False,
    margins: tuple[float, float, float, float] = (0.04, 0.02, 0.04, 0.02),
) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = anchor
    frame.margin_left, frame.margin_top, frame.margin_right, frame.margin_bottom = (
        Inches(value) for value in margins
    )
    if isinstance(item, str):
        item = {"head": "", "tail": item, "inline": False}
    head, tail = item.get("head", ""), item.get("tail", "")
    paragraphs = _paragraphs(item)
    if not paragraphs:
        return
    width = shape.width / 914400 - margins[0] - margins[2]
    height = shape.height / 914400 - margins[1] - margins[3]
    fitted = _fit(
        paragraphs, max(0.1, width), max(0.1, height), size, minimum, style.get("scale")
    )

    def run(paragraph: Any, text: str, strong: bool) -> None:
        piece = paragraph.add_run()
        piece.text = text
        piece.font.name = style["font"]
        piece.font.size = Pt(fitted)
        piece.font.bold = strong
        piece.font.color.rgb = RGBColor.from_string(color)

    first = frame.paragraphs[0]
    first.alignment = align
    first.line_spacing = 1.1
    if item.get("inline"):
        run(first, head, True)
        run(first, tail, bold)
        return
    if head:
        run(first, head, True)
        if tail:
            second = frame.add_paragraph()
            second.alignment = align
            second.line_spacing = 1.1
            second.space_before = Pt(fitted * 0.25)
            run(second, tail, bold)
    else:
        run(first, tail, bold)


# ----------------------------------------------------------------- shapes ---
def _box(
    shapes: Any,
    kind: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: str | None,
    name: str,
    *,
    line: str | None = None,
    line_pt: float = 0.0,
    radius: float | None = None,
) -> Any:
    shape = shapes.add_shape(
        kind, Inches(x), Inches(y), Inches(max(0.02, w)), Inches(max(0.02, h))
    )
    drop_theme_style(shape)
    shape.name = name
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = RGBColor.from_string(line)
        shape.line.width = Pt(line_pt or 1.0)
    if radius is not None:
        try:
            shape.adjustments[0] = radius
        except (IndexError, ValueError):
            pass
    return shape


def _text_box(shapes: Any, x: float, y: float, w: float, h: float, name: str) -> Any:
    shape = shapes.add_textbox(
        Inches(x), Inches(y), Inches(max(0.1, w)), Inches(max(0.1, h))
    )
    shape.name = name
    return shape


def _line(
    shapes: Any,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    color: str,
    name: str,
    width: float = 1.5,
) -> Any:
    connector = shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    connector.line.color.rgb = RGBColor.from_string(color)
    connector.line.width = Pt(width)
    connector.name = name
    return connector


def _badge(
    shapes: Any,
    cx: float,
    cy: float,
    diameter: float,
    style: dict[str, Any],
    *,
    icon: str | None,
    number: int,
    name: str,
    fill: str | None = None,
    text_color: str | None = None,
) -> None:
    fill, text_color = (
        (fill, text_color) if fill else (style["accent_fill"], style["accent_text"])
    )
    circle = _box(
        shapes,
        MSO_SHAPE.OVAL,
        cx - diameter / 2,
        cy - diameter / 2,
        diameter,
        diameter,
        fill,
        f"{name} Badge",
    )
    if icon:
        size = diameter * 0.56
        add_icon(shapes, icon, (cx - size / 2, cy - size / 2, size), text_color)
    else:
        _write(
            circle,
            f"{number:02d}" if number < 100 else str(number),
            style,
            color=text_color,
            size=max(10.0, min(16.0, diameter * 26)),
            minimum=9,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            bold=True,
            margins=(0.0, 0.0, 0.0, 0.0),
        )


def _describe(group: Any, text: str) -> None:
    """Alt text for screen readers and later contextual (VLM) audits."""
    properties = group._element.find(qn("p:nvGrpSpPr"))
    if properties is not None:
        properties.find(qn("p:cNvPr")).set("descr", text[:900])


# --------------------------------------------------------------- diagrams ---
def _grid_columns(count: int, w: float, h: float, gap: float) -> int:
    """Columns whose cards are closest to a slightly landscape shape."""
    best, best_penalty = 1, math.inf
    for columns in range(1, min(count, 4) + 1):
        rows = math.ceil(count / columns)
        empty = rows * columns - count
        card_w = (w - gap * (columns - 1)) / columns
        card_h = (h - gap * (rows - 1)) / rows
        if card_w < 1.5 or empty >= columns or card_h <= 0:
            continue
        penalty = abs(math.log(card_w / card_h / 1.25)) + 0.15 * empty
        if penalty < best_penalty:
            best, best_penalty = columns, penalty
    return best


def _icon_grid(shapes, items, zone, style, context, record, layout_variant: str) -> None:
    x, y, w, h = zone
    count = len(items)
    gap = min(0.24, max(0.12, w * 0.02))
    canvas_area = float(style.get("canvas_area") or 0)
    focus_h = min(h, 0.70 * canvas_area / w) if canvas_area and w > 0 else h
    wide_pair = count == 2 and w >= 7.5
    boxes: list[tuple[float, float, float, float]]
    narrow_stack = w < 5.6 and (h - gap * (count - 1)) / count >= 1.05
    if layout_variant == "focus" and wide_pair and focus_h >= 2.4:
        # A pair of equal cards would render exactly like the balanced mode.
        # Give the lead claim more width while both keep the same height.
        focus_h = min(focus_h, 3.0)
        y += (h - focus_h) / 2
        h = focus_h
        lead_w = w * 0.58
        boxes = [(x, y, lead_w, h), (x + lead_w + gap, y, w - lead_w - gap, h)]
        record["columns"] = 2
        record["emphasis"] = "first"
    elif (
        (layout_variant == "columns" and w >= 5.6) or narrow_stack
    ) and (h - gap * (count - 1)) / count >= 1.05:
        card_h = (h - gap * (count - 1)) / count
        boxes = [(x, y + index * (card_h + gap), w, card_h) for index in range(count)]
        record["columns"] = 1
        if narrow_stack:
            record["narrow_stack"] = True
        if narrow_stack and layout_variant == "focus" and count == 2:
            lead_h = (h - gap) * 0.6
            other_h = h - gap - lead_h
            if other_h >= 1.05 and all(
                text_height(_paragraphs(item), max(0.1, w - 1.0), 11)
                + 0.25
                <= card_h_candidate
                for item, card_h_candidate in zip(
                    items, (lead_h, other_h), strict=True
                )
            ):
                boxes = [(x, y, w, lead_h), (x, y + lead_h + gap, w, other_h)]
                record["emphasis"] = "first"
    elif (
        layout_variant == "focus"
        and count >= 3
        and w >= 7.5
        and focus_h >= 3.0
        and (focus_h - gap * (count - 2)) / (count - 1) >= 1.05
    ):
        # A full-height hero card plus its companions can cover over 75% of
        # a compact template. Bound the card group and centre it in the slot.
        y += (h - focus_h) / 2
        h = focus_h
        lead_w = w * 0.42
        side_x = x + lead_w + gap
        side_w = w - lead_w - gap
        side_h = (h - gap * (count - 2)) / (count - 1)
        boxes = [(x, y, lead_w, h)] + [
            (side_x, y + index * (side_h + gap), side_w, side_h)
            for index in range(count - 1)
        ]
        record["columns"] = 2
        record["emphasis"] = "first"
    else:
        layout_variant = "balanced"
        columns = 2 if count == 4 and w >= 6.0 else _grid_columns(count, w, h, gap)
        if count <= 4 and w >= 6 and max(len(item["plain"]) for item in items) > 48:
            columns = min(columns, 2)
        rows = math.ceil(count / columns)
        card_w = (w - gap * (columns - 1)) / columns
        available_h = (h - gap * (rows - 1)) / rows
        pad = min(0.24, max(0.12, card_w * 0.07))
        diameter = min(0.66, max(0.4, min(card_w, available_h) * 0.28))
        size = style["size"] + 2
        horizontal = wide_pair or (
            card_w / max(0.1, available_h) >= 1.9 and card_w >= 2.6
        )
        text_w = card_w - (diameter + pad * 2.6 if horizontal else pad * 1.4) - 0.08
        needed = max(text_height(_paragraphs(item), text_w, size) for item in items)
        if horizontal:
            wanted = max(1.1, needed + pad * 1.2, diameter + pad * 2)
        else:
            wanted = max(1.7, pad * 2.2 + diameter + needed)
        # Short claims still fill the native card area instead of making an
        # almost empty slide. Other variants use their full assigned zones.
        card_h = min(available_h, wanted + 0.12)
        if count == 2 and w >= 7.5 and not style.get("fill_zone"):
            # Two short claims need breathing room without a mostly empty
            # full-height card. Three inches still cover the required quarter
            # of a standard wide slide when both cards span the content zone.
            card_h = min(available_h, max(3.0, card_h))
        else:
            grid_area = w * (card_h * rows + gap * (rows - 1))
            if style.get("fill_zone") or grid_area < 0.27 * style.get("canvas_area", 0):
                card_h = available_h
        used_h = card_h * rows + gap * (rows - 1)
        top = y + (h - used_h) / 2 if count == 2 and w >= 7.5 else _block_top(y, h, used_h)
        boxes = [
            (
                x + (index % columns) * (card_w + gap),
                top + (index // columns) * (card_h + gap),
                card_w,
                card_h,
            )
            for index in range(count)
        ]
        record["columns"] = columns
    record["layout_variant"] = layout_variant
    hints = [
        resolve_icon(item.get("icon")) if item.get("icon") else None for item in items
    ]
    chosen = assign_icons([item["plain"] for item in items], context=context)
    icons = []
    for hint, automatic in zip(hints, chosen):
        icons.append(hint or automatic)
    record["icons"] = [icon or "circle-dot" for icon in icons]
    record["icon_matches"] = sum(bool(icon) for icon in icons)
    # Cards share one size: the largest at which every card's text fits, so
    # the asymmetric variant does not shrink only one statement.
    measures = []
    for _, _, card_w, card_h in boxes:
        pad = min(0.24, max(0.12, card_w * 0.07))
        diameter = min(0.66, max(0.4, min(card_w, card_h) * 0.28))
        horizontal = wide_pair or (
            card_w / max(0.1, card_h) >= 1.9 and card_w >= 2.6
        )
        if horizontal:
            box_w, box_h = card_w - diameter - pad * 2.6, card_h - pad * 1.2
        else:
            box_w = card_w - pad * 1.4
            box_h = card_h - pad * 1.5 - diameter - pad * 0.7
        measures.append((pad, diameter, horizontal, box_w, box_h))
    size = min(
        _fit(
            _paragraphs(item),
            max(0.1, box_w - 0.08),
            max(0.1, box_h - 0.04),
            style["size"] + 2,
            11.0,
            style.get("scale"),
        )
        for item, (_, _, _, box_w, box_h) in zip(items, measures)
    )
    for index, item in enumerate(items):
        cx, cy, card_w, card_h = boxes[index]
        pad, diameter, horizontal, _, _ = measures[index]
        _box(
            shapes,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            cx,
            cy,
            card_w,
            card_h,
            style["surface"],
            f"BrandDeck Diagram Card {index + 1}",
            radius=0.08,
        )
        icon = icons[index] or "circle-dot"
        if horizontal:
            _badge(
                shapes,
                cx + pad + diameter / 2,
                cy + card_h / 2,
                diameter,
                style,
                icon=icon,
                number=index + 1,
                name=f"BrandDeck Diagram Icon {index + 1}",
            )
            box = _text_box(
                shapes,
                cx + pad * 1.6 + diameter,
                cy + pad * 0.6,
                card_w - diameter - pad * 2.6,
                card_h - pad * 1.2,
                f"BrandDeck Diagram Text {index + 1}",
            )
            _write(
                box,
                item,
                style,
                color=style["surface_text"],
                size=size,
                anchor=MSO_ANCHOR.MIDDLE,
            )
        else:
            _badge(
                shapes,
                cx + pad + diameter / 2,
                cy + pad + diameter / 2,
                diameter,
                style,
                icon=icon,
                number=index + 1,
                name=f"BrandDeck Diagram Icon {index + 1}",
            )
            text_top = cy + pad + diameter + pad * 0.7
            box = _text_box(
                shapes,
                cx + pad * 0.7,
                text_top,
                card_w - pad * 1.4,
                cy + card_h - text_top - pad * 0.5,
                f"BrandDeck Diagram Text {index + 1}",
            )
            _write(box, item, style, color=style["surface_text"], size=size)


def _process(shapes, items, zone, style, context, record) -> None:
    x, y, w, h = zone
    count = len(items)
    if w >= 8.5 and w / max(0.1, h) >= 1.5 and count <= 5:
        arrow_h = min(0.9, max(0.62, h * 0.22))
        depth = arrow_h * 0.36
        joint = 0.05
        arrow_w = (w + (count - 1) * (depth - joint)) / count
        column_w = w / count
        fills = tint_series(style, count, spread=0.3)
        record["orientation"] = "horizontal"
        labels = [f"{index + 1:02d}" for index in range(count)]
        details = items
        size = style["size"]
        needed = max(
            text_height(_paragraphs(detail), column_w - 0.22, size)
            for detail in details
        )
        detail_h = min(h - arrow_h - 0.16, max(1.05, needed + 0.18))
        if style.get("fill_zone"):
            detail_h = h - arrow_h - 0.16
        top = _block_top(y, h, arrow_h + 0.12 + detail_h)
        for index, item in enumerate(items):
            ax = x + index * (arrow_w - depth + joint)
            kind = MSO_SHAPE.PENTAGON if index == 0 else MSO_SHAPE.CHEVRON
            fill, text_color = fills[index]
            arrow = _box(
                shapes,
                kind,
                ax,
                top,
                arrow_w,
                arrow_h,
                fill,
                f"BrandDeck Diagram Step {index + 1}",
            )
            try:
                arrow.adjustments[0] = 0.36
            except (IndexError, ValueError):
                pass
            inner = depth + 0.06
            _write(
                arrow,
                labels[index],
                style,
                color=text_color,
                size=min(style["size"], 16),
                minimum=10,
                align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE,
                bold=True,
                margins=(inner if index else 0.1, 0.02, inner, 0.02),
            )
            card_x = x + index * column_w + 0.04
            card_y = top + arrow_h + 0.12
            _box(
                shapes,
                MSO_SHAPE.ROUNDED_RECTANGLE,
                card_x,
                card_y,
                column_w - 0.08,
                detail_h,
                style["surface"],
                f"BrandDeck Diagram Detail Card {index + 1}",
                radius=0.12,
            )
            box = _text_box(
                shapes,
                card_x + 0.08,
                card_y + 0.08,
                column_w - 0.24,
                detail_h - 0.16,
                f"BrandDeck Diagram Text {index + 1}",
            )
            _write(
                box,
                details[index],
                style,
                color=style["surface_text"],
                size=size,
                anchor=MSO_ANCHOR.MIDDLE,
            )
        return
    record["orientation"] = "vertical"
    row_h = h / count
    diameter = min(0.56, max(0.34, row_h * 0.62))
    centre_x = x + diameter / 2
    if count > 1:
        _box(
            shapes,
            MSO_SHAPE.RECTANGLE,
            centre_x - 0.018,
            y + row_h / 2,
            0.036,
            row_h * (count - 1),
            style["line"],
            "BrandDeck Diagram Connector",
        )
    for index, item in enumerate(items):
        cy = y + row_h * index + row_h / 2
        _badge(
            shapes,
            centre_x,
            cy,
            diameter,
            style,
            icon=None,
            number=index + 1,
            name=f"BrandDeck Diagram Step {index + 1}",
        )
        _box(
            shapes,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            x + diameter + 0.12,
            y + row_h * index + 0.03,
            w - diameter - 0.12,
            row_h - 0.06,
            style["surface"],
            f"BrandDeck Diagram Detail Card {index + 1}",
            radius=0.12,
        )
        box = _text_box(
            shapes,
            x + diameter + 0.22,
            y + row_h * index + 0.03,
            w - diameter - 0.22,
            row_h - 0.06,
            f"BrandDeck Diagram Text {index + 1}",
        )
        _write(
            box,
            item,
            style,
            color=style["surface_text"],
            size=style["size"],
            anchor=MSO_ANCHOR.MIDDLE,
        )


def _cycle(shapes, items, zone, style, context, record) -> None:
    x, y, w, h = zone
    count = len(items)
    # The legend needs room for its longest word at a legible size; a narrow
    # zone gives the ring less width rather than squeezing the captions.
    longest = max((len(word) for item in items for word in item["plain"].split()), default=8)
    legend_need = 0.95 + longest * (style["size"] - 1) * ESTIMATED_GLYPH_WIDTH_EM / 72
    side = max(min(1.6, h), min(h, w * 0.46, w - legend_need))
    node = min(0.72, max(0.44, side * 0.19))
    radius = side / 2 - node / 2
    cx, cy = x + side / 2, y + h / 2
    ring = _box(
        shapes,
        MSO_SHAPE.OVAL,
        cx - radius,
        cy - radius,
        radius * 2,
        radius * 2,
        None,
        "BrandDeck Diagram Ring",
        line=style["line"],
        line_pt=2.25,
    )
    ring.fill.background()
    for index in range(count):
        theta = -math.pi / 2 + 2 * math.pi * (index + 0.5) / count
        size = min(0.22, node * 0.36)
        marker = _box(
            shapes,
            MSO_SHAPE.ISOSCELES_TRIANGLE,
            cx + radius * math.cos(theta) - size / 2,
            cy + radius * math.sin(theta) - size / 2,
            size,
            size,
            style["line"],
            f"BrandDeck Diagram Arrow {index + 1}",
        )
        marker.rotation = (math.degrees(theta) + 180) % 360
    for index in range(count):
        theta = -math.pi / 2 + 2 * math.pi * index / count
        _badge(
            shapes,
            cx + radius * math.cos(theta),
            cy + radius * math.sin(theta),
            node,
            style,
            icon=None,
            number=index + 1,
            name=f"BrandDeck Diagram Node {index + 1}",
        )
    legend_x = x + side + 0.35
    legend_w = w - side - 0.35
    row_h = min(h / count, 1.35)
    top = y + (h - row_h * count) / 2
    record["layout"] = "ring_with_legend"
    for index, item in enumerate(items):
        row_y = top + index * row_h
        small = min(0.34, row_h * 0.5)
        _badge(
            shapes,
            legend_x + small / 2,
            row_y + row_h / 2,
            small,
            style,
            icon=None,
            number=index + 1,
            name=f"BrandDeck Diagram Legend {index + 1}",
        )
        box = _text_box(
            shapes,
            legend_x + small + 0.16,
            row_y + 0.02,
            legend_w - small - 0.16,
            row_h - 0.04,
            f"BrandDeck Diagram Text {index + 1}",
        )
        _write(
            box,
            item,
            style,
            color=style["text"],
            size=style["size"] - 1,
            anchor=MSO_ANCHOR.MIDDLE,
        )


def _hierarchy(shapes, items, zone, style, context, record, root: str) -> None:
    x, y, w, h = zone
    count = len(items)
    root_h = min(0.8, max(0.5, h * 0.2))
    root_w = min(w * 0.42, 4.2)
    gap_v = min(0.6, max(0.34, h * 0.14))
    gap = min(0.24, max(0.12, w * 0.02))
    child_w = (w - gap * (count - 1)) / count
    size = style["size"]
    needed = max(text_height(_paragraphs(item), child_w - 0.28, size) for item in items)
    child_h = min(h - root_h - gap_v, max(1.1, needed + 0.42))
    if style.get("fill_zone"):
        child_h = h - root_h - gap_v
    y = _block_top(y, h, root_h + gap_v + child_h)
    root_x = x + (w - root_w) / 2
    root_shape = _box(
        shapes,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        root_x,
        y,
        root_w,
        root_h,
        style["accent_fill"],
        "BrandDeck Diagram Root",
        radius=0.18,
    )
    _write(
        root_shape,
        root,
        style,
        color=style["accent_text"],
        size=style["size"],
        minimum=11,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        bold=True,
        margins=(0.12, 0.04, 0.12, 0.04),
    )
    child_y = y + root_h + gap_v
    bus_y = y + root_h + gap_v / 2
    centres = [x + index * (child_w + gap) + child_w / 2 for index in range(count)]
    _line(
        shapes,
        x + w / 2,
        y + root_h,
        x + w / 2,
        bus_y,
        style["line"],
        "BrandDeck Diagram Link Root",
    )
    if count > 1:
        _line(
            shapes,
            centres[0],
            bus_y,
            centres[-1],
            bus_y,
            style["line"],
            "BrandDeck Diagram Link Bus",
        )
    for index, item in enumerate(items):
        left = x + index * (child_w + gap)
        _line(
            shapes,
            centres[index],
            bus_y,
            centres[index],
            child_y,
            style["line"],
            f"BrandDeck Diagram Link {index + 1}",
        )
        _box(
            shapes,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left,
            child_y,
            child_w,
            child_h,
            style["surface"],
            f"BrandDeck Diagram Card {index + 1}",
            radius=0.06,
        )
        _box(
            shapes,
            MSO_SHAPE.RECTANGLE,
            left + 0.14,
            child_y + 0.12,
            min(0.6, child_w * 0.3),
            0.05,
            style["accent"],
            f"BrandDeck Diagram Marker {index + 1}",
        )
        box = _text_box(
            shapes,
            left + 0.1,
            child_y + 0.26,
            child_w - 0.2,
            child_h - 0.36,
            f"BrandDeck Diagram Text {index + 1}",
        )
        _write(box, item, style, color=style["surface_text"], size=size)


def _stacked(shapes, items, zone, style, context, record, *, funnel: bool) -> None:
    x, y, w, h = zone
    count = len(items)
    # Two-level pyramids and funnels otherwise occupy a narrow strip in a
    # full-size content zone and fail the rendered quarter-slide fill check;
    # after an underfilled render (fill_zone) the levels take the whole zone.
    used = h if style.get("fill_zone") else min(h, max(count * 1.3, 3.2 if count == 2 else 0))
    y, h = _block_top(y, h, used), used
    region_w = min(w * 0.44, h * (1.5 if funnel else 1.25))
    cx = x + region_w / 2
    level_h = h / count
    gap = 0.06

    def width_at(yy: float) -> float:
        ratio = (yy - y) / h
        return region_w * (1 - 0.68 * ratio) if funnel else region_w * ratio

    fills = tint_series(style, count, spread=0.5)
    if funnel:
        fills = list(reversed(fills))
    record["layout"] = "funnel" if funnel else "pyramid"
    for index, item in enumerate(items):
        top = y + index * level_h + gap / 2
        bottom = y + (index + 1) * level_h - gap / 2
        top_w, bottom_w = max(0.05, width_at(top)), width_at(bottom)
        fill, text_color = fills[index]
        level = add_polygon(
            shapes,
            [
                (cx - top_w / 2, top),
                (cx + top_w / 2, top),
                (cx + bottom_w / 2, bottom),
                (cx - bottom_w / 2, bottom),
            ],
            fill,
            name=f"BrandDeck Diagram Level {index + 1}",
        )
        label = item["value"] or f"{index + 1:02d}"
        apex = not funnel and index == 0
        # The apex is a near-triangle: its label sits in the wide lower part.
        inner = bottom_w * 0.55 if apex else min(top_w, bottom_w)
        if inner >= (0.3 if apex else 0.5):
            _write(
                level,
                label,
                style,
                color=text_color,
                size=min(15.0, style["size"]),
                minimum=9,
                align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.BOTTOM if apex else MSO_ANCHOR.MIDDLE,
                bold=True,
                margins=(0.02, 0.0, 0.02, 0.08 if apex else 0.0),
            )
        box = _text_box(
            shapes,
            x + region_w + 0.3,
            top,
            w - region_w - 0.3,
            bottom - top,
            f"BrandDeck Diagram Text {index + 1}",
        )
        _write(
            box,
            item,
            style,
            color=style["text"],
            size=style["size"],
            anchor=MSO_ANCHOR.MIDDLE,
        )


def _matrix(shapes, items, zone, style, context, record, axes: dict[str, str]) -> None:
    x, y, w, h = zone
    # Axis captions stay horizontal: rotated boxes would defeat bounds checks.
    top_offset = 0.34 if axes.get("y") else 0.0
    bottom_offset = 0.34 if axes.get("x") else 0.0
    gx, gy, gw, gh = x, y + top_offset, w, h - top_offset - bottom_offset
    gap = 0.14
    cell_w, cell_h = (gw - gap) / 2, (gh - gap) / 2
    fills = tint_series(style, 4, spread=0.62)
    if axes.get("y"):
        axis = _text_box(
            shapes, gx, y, gw, top_offset - 0.04, "BrandDeck Diagram Axis Y"
        )
        _write(axis, axes["y"], style, color=style["text"], size=12, minimum=9)
    for index, item in enumerate(items[:4]):
        column, row = index % 2, index // 2
        fill, text_color = fills[index]
        cell = _box(
            shapes,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            gx + column * (cell_w + gap),
            gy + row * (cell_h + gap),
            cell_w,
            cell_h,
            fill,
            f"BrandDeck Diagram Quadrant {index + 1}",
            radius=0.06,
        )
        _write(
            cell,
            item,
            style,
            color=text_color,
            size=style["size"],
            margins=(0.18, 0.14, 0.18, 0.12),
        )
    if axes.get("x"):
        axis = _text_box(
            shapes,
            gx,
            gy + gh + 0.04,
            gw,
            bottom_offset - 0.04,
            "BrandDeck Diagram Axis X",
        )
        _write(
            axis,
            axes["x"],
            style,
            color=style["text"],
            size=12,
            minimum=9,
            align=PP_ALIGN.CENTER,
        )


def render_diagram(
    slide: Any,
    visual: dict[str, Any],
    zone: tuple[float, float, float, float],
    style: dict[str, Any],
    *,
    context: str = "",
) -> dict[str, Any]:
    """Draw ``visual`` into ``zone`` and return a manifest record."""
    kind = str(visual.get("type"))
    items = visual_items(visual)
    minimum, maximum = ITEM_LIMITS[kind]
    items = items[:maximum]
    if len(items) < minimum:
        raise ValueError(f"{kind} needs at least {minimum} items")
    group = slide.shapes.add_group_shape()
    group.name = f"BrandDeck Diagram {kind}"
    record: dict[str, Any] = {"type": kind, "items": len(items), "native": True}
    shapes = group.shapes
    if kind == "icon_grid":
        _icon_grid(
            shapes,
            items,
            zone,
            style,
            context,
            record,
            str(visual.get("layout_variant") or "balanced"),
        )
    elif kind == "process":
        _process(shapes, items, zone, style, context, record)
    elif kind == "cycle":
        _cycle(shapes, items, zone, style, context, record)
    elif kind == "hierarchy":
        root = str(visual.get("root") or context or "").strip()[:80]
        _hierarchy(shapes, items, zone, style, context, record, root)
    elif kind in {"pyramid", "funnel"}:
        _stacked(shapes, items, zone, style, context, record, funnel=kind == "funnel")
    elif kind == "matrix":
        axes = visual.get("axes") if isinstance(visual.get("axes"), dict) else {}
        _matrix(
            shapes,
            items,
            zone,
            style,
            context,
            record,
            {k: str(v)[:60] for k, v in axes.items() if k in {"x", "y"} and v},
        )
    _describe(group, f"{kind}: " + "; ".join(item["plain"] for item in items))
    return record


def render_illustration(
    slide: Any,
    zone: tuple[float, float, float, float],
    style: dict[str, Any],
    *,
    text: str,
    name: str = "BrandDeck Illustration",
) -> dict[str, Any]:
    """Native vector illustration used when no suitable picture is available."""
    x, y, w, h = zone
    icon, score = choose_icon(text)
    icon = icon or "image"
    group = slide.shapes.add_group_shape()
    group.name = name
    shapes = group.shapes
    panel_fill = mix(
        style["background"], style["accent"], 0.16 if style["dark"] else 0.1
    )
    _box(
        shapes,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        x,
        y,
        w,
        h,
        panel_fill,
        "BrandDeck Illustration Panel",
        radius=0.06,
    )
    side = min(w, h)
    big = side * 0.62
    _box(
        shapes,
        MSO_SHAPE.OVAL,
        x + w / 2 - big / 2,
        y + h / 2 - big / 2,
        big,
        big,
        mix(panel_fill, style["accent"], 0.22),
        "BrandDeck Illustration Halo",
    )
    small = side * 0.18
    _box(
        shapes,
        MSO_SHAPE.OVAL,
        x + w * 0.5 + big * 0.36,
        y + h * 0.5 - big * 0.5,
        small,
        small,
        mix(panel_fill, style["accent"], 0.45),
        "BrandDeck Illustration Accent",
    )
    dot = side * 0.08
    _box(
        shapes,
        MSO_SHAPE.OVAL,
        x + w * 0.5 - big * 0.55,
        y + h * 0.5 + big * 0.34,
        dot,
        dot,
        style["accent"],
        "BrandDeck Illustration Dot",
    )
    size = big * 0.52
    add_icon(
        shapes,
        icon,
        (x + w / 2 - size / 2, y + h / 2 - size / 2, size),
        style["accent_fill"],
        stroke_ratio=0.8,
    )
    _describe(group, f"illustration: {icon}")
    return {
        "type": "illustration",
        "icon": icon,
        "keyword_match": bool(score),
        "native": True,
    }
