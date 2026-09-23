"""Image assets for slides: uploads, pictures inside source documents, matching.

Images are normalised once into a content-addressed store and referenced from
the plan by id. Matching to slides is a transparent lexical comparison of the
image's alt text, caption, file name and surrounding document text with each
slide's title and text; the decision and its score are recorded in the plan.
Nothing is downloaded: remote image URLs in documents are ignored.
"""

from __future__ import annotations

import hashlib
import io
import math
import re
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from lxml import etree
from PIL import Image, ImageOps, UnidentifiedImageError

IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
MAX_SOURCE_BYTES = 25 * 1024 * 1024
MAX_PIXELS = 40_000_000
MAX_SIDE = 2400
MIN_SIDE = 96
MAX_IMAGES = 40

_MARKDOWN_IMAGE = re.compile(
    r"!\[(?P<alt>[^\]]*)\]\(\s*<?(?P<src>[^)\s>]+)>?(?:\s+[\"'](?P<title>[^\"']*)[\"'])?\s*\)"
)
_HTML_IMAGE = re.compile(r"<img\b[^>]*>", re.IGNORECASE)
_ATTRIBUTE = re.compile(r"(\w+)\s*=\s*[\"']([^\"']*)[\"']")
_WORD = re.compile(r"[0-9a-zа-я]+")
_CAMERA_NAME = re.compile(
    r"^(img|dsc|dscn|photo|image|pic|screenshot|снимок|изображение|фото|scan)?[\s_\-]*[\d\s_\-]*$",
    re.IGNORECASE,
)
_STOPWORDS = {
    "этот",
    "этого",
    "также",
    "более",
    "менее",
    "может",
    "могут",
    "будет",
    "которые",
    "который",
    "которая",
    "после",
    "перед",
    "между",
    "через",
    "всех",
    "всего",
    "есть",
    "нужна",
    "нужен",
    "нужно",
    "только",
    "каждый",
    "каждого",
    "очень",
    "with",
    "from",
    "that",
    "this",
    "have",
    "will",
    "into",
    "your",
    "their",
    "about",
    "slide",
    "слайд",
    "image",
    "picture",
    "photo",
    "рисунок",
    "изображение",
    "фото",
}


@dataclass
class ImageAsset:
    asset_id: str
    path: str
    source: str
    label: str
    context: str
    width: int
    height: int
    origin: str


def normalize_image(data: bytes) -> tuple[bytes, str, tuple[int, int]]:
    """Decode, bound, orient and re-encode an image; raise ``ValueError``."""
    if not data or len(data) > MAX_SOURCE_BYTES:
        raise ValueError("Image is empty or exceeds 25 MB")
    try:
        with Image.open(io.BytesIO(data)) as probe:
            width, height = probe.size
            source_format = (probe.format or "").upper()
            if width * height > MAX_PIXELS:
                raise ValueError("Image has too many pixels")
            probe.verify()
        with Image.open(io.BytesIO(data)) as image:
            image.seek(0)
            image = ImageOps.exif_transpose(image)
            has_alpha = image.mode in {"RGBA", "LA", "PA"} or (
                image.mode == "P" and "transparency" in image.info
            )
            image = image.convert("RGBA" if has_alpha else "RGB")
            if min(image.size) < MIN_SIDE:
                raise ValueError("Image is too small for a slide")
            image.thumbnail((MAX_SIDE, MAX_SIDE), Image.Resampling.LANCZOS)
            output = io.BytesIO()
            if has_alpha or source_format in {"PNG", "GIF", "BMP"}:
                image.save(output, format="PNG", optimize=True)
                suffix = "png"
            else:
                image.save(output, format="JPEG", quality=88, optimize=True)
                suffix = "jpg"
            return output.getvalue(), suffix, image.size
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
        raise ValueError(f"Unsupported or damaged image: {exc}") from exc


def label_from_filename(name: str) -> str:
    stem = Path(name).stem
    text = re.sub(r"[_\-.]+", " ", stem).strip()
    return "" if _CAMERA_NAME.match(text) else text


class ImageLibrary:
    """Content-addressed image store shared by all variants of one request."""

    def __init__(self, root: Path) -> None:
        self.root = Path(root)
        self.assets: dict[str, ImageAsset] = {}
        self.rejected: list[dict[str, str]] = []

    def __len__(self) -> int:
        return len(self.assets)

    def add(
        self,
        data: bytes,
        *,
        source: str,
        label: str = "",
        context: str = "",
        origin: str = "",
    ) -> ImageAsset | None:
        if len(self.assets) >= MAX_IMAGES:
            self.rejected.append({"origin": origin, "reason": "image limit reached"})
            return None
        try:
            payload, suffix, (width, height) = normalize_image(data)
        except ValueError as exc:
            self.rejected.append({"origin": origin, "reason": str(exc)[:200]})
            return None
        asset_id = hashlib.sha256(payload).hexdigest()[:16]
        if asset_id in self.assets:
            existing = self.assets[asset_id]
            existing.label = existing.label or label.strip()[:200]
            existing.context = existing.context or context.strip()[:600]
            return existing
        self.root.mkdir(parents=True, exist_ok=True)
        target = self.root / f"{asset_id}.{suffix}"
        if not target.exists():
            target.write_bytes(payload)
        asset = ImageAsset(
            asset_id=asset_id,
            path=str(target.resolve()),
            source=source,
            label=label.strip()[:200],
            context=re.sub(r"\s+", " ", context).strip()[:600],
            width=width,
            height=height,
            origin=origin[:200],
        )
        self.assets[asset_id] = asset
        return asset

    def add_path(
        self, path: str | Path, *, source: str = "upload", label: str = ""
    ) -> ImageAsset | None:
        candidate = Path(path)
        if candidate.suffix.lower() not in IMAGE_SUFFIXES:
            self.rejected.append(
                {"origin": candidate.name, "reason": "unsupported extension"}
            )
            return None
        try:
            data = candidate.read_bytes()
        except OSError as exc:
            self.rejected.append({"origin": candidate.name, "reason": str(exc)[:200]})
            return None
        return self.add(
            data,
            source=source,
            label=label or label_from_filename(candidate.name),
            origin=candidate.name,
        )

    def catalog(self) -> list[dict[str, str]]:
        """Bounded description for a planning model: ids, labels and context only."""
        return [
            {"id": asset.asset_id, "label": asset.label, "context": asset.context[:200]}
            for asset in self.assets.values()
        ]

    def plan_assets(self) -> dict[str, dict[str, Any]]:
        return {asset.asset_id: asdict(asset) for asset in self.assets.values()}


# ------------------------------------------------------ document extraction ---
def _context_before(text: str, position: int) -> str:
    before = text[:position]
    heading = ""
    for line in reversed(before.splitlines()):
        match = re.match(r"^#{1,6}\s+(.+)$", line.strip())
        if match:
            heading = match.group(1)
            break
    paragraph = before.rsplit("\n\n", 1)[-1]
    return f"{heading}. {paragraph}".strip(" .")


def extract_markdown_images(
    text: str, base_dir: Path | None
) -> tuple[str, list[dict[str, Any]]]:
    """Remove image markup from text; return local references inside ``base_dir``."""
    references: list[dict[str, Any]] = []
    root = base_dir.resolve() if base_dir else None

    def local(src: str) -> Path | None:
        if root is None or re.match(r"^[a-z][a-z0-9+.-]*:", src, re.IGNORECASE):
            return None
        candidate = (root / src).resolve()
        try:
            candidate.relative_to(root)
        except ValueError:
            return None
        return candidate if candidate.is_file() else None

    def replace(match: re.Match[str], src: str, alt: str) -> str:
        references.append(
            {
                "path": local(src),
                "src": src,
                "label": alt,
                "context": _context_before(text, match.start()),
            }
        )
        return ""

    def markdown(match: re.Match[str]) -> str:
        alt = match.group("alt") or match.group("title") or ""
        return replace(match, match.group("src"), alt)

    def html(match: re.Match[str]) -> str:
        attributes = {
            key.lower(): value for key, value in _ATTRIBUTE.findall(match.group())
        }
        if not attributes.get("src"):
            return match.group()
        return replace(
            match,
            attributes["src"],
            attributes.get("alt", "") or attributes.get("title", ""),
        )

    cleaned = _MARKDOWN_IMAGE.sub(markdown, text)
    cleaned = _HTML_IMAGE.sub(html, cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned, references


def _pptx_images(path: Path) -> list[dict[str, Any]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(path)
    area = max(1, prs.slide_width * prs.slide_height)
    found: list[dict[str, Any]] = []
    counts: dict[str, int] = {}

    def visit(shapes: Any, slide_number: int, context: str) -> None:
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                visit(shape.shapes, slide_number, context)
                continue
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            try:
                blob, digest = shape.image.blob, shape.image.sha1
            except (AttributeError, ValueError, KeyError):
                continue
            counts[digest] = counts.get(digest, 0) + 1
            if shape.width * shape.height / area < 0.015:
                continue
            properties = shape._element.xpath("./p:nvPicPr/p:cNvPr")
            label = ""
            if properties:
                label = properties[0].get("descr") or properties[0].get("title") or ""
            found.append(
                {
                    "data": blob,
                    "digest": digest,
                    "label": label,
                    "context": context,
                    "origin": f"{path.name}#slide-{slide_number}",
                }
            )

    for number, slide in enumerate(prs.slides, 1):
        texts = [
            shape.text.strip()
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text.strip()
        ]
        visit(slide.shapes, number, " ".join(texts)[:600])
    slides = max(1, len(prs.slides))
    # Pictures repeated across slides are logos and decoration, not content.
    return [
        item
        for item in found
        if counts[item["digest"]] < max(3, math.ceil(slides * 0.4))
    ]


def _docx_images(path: Path) -> list[dict[str, Any]]:
    namespaces = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    found: list[dict[str, Any]] = []
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        if "word/document.xml" not in names:
            return found
        parser = etree.XMLParser(
            resolve_entities=False, no_network=True, huge_tree=False
        )
        document = etree.fromstring(archive.read("word/document.xml"), parser)
        relationships = {}
        if "word/_rels/document.xml.rels" in names:
            rels = etree.fromstring(
                archive.read("word/_rels/document.xml.rels"), parser
            )
            for rel in rels.findall("rel:Relationship", namespaces):
                if rel.get("TargetMode") != "External":
                    relationships[rel.get("Id")] = rel.get("Target", "")
        paragraphs = document.findall(".//w:body/w:p", namespaces)
        texts = ["".join(node.itertext()).strip() for node in paragraphs]
        for index, paragraph in enumerate(paragraphs):
            for drawing in paragraph.iter(f"{{{namespaces['w']}}}drawing"):
                blip = drawing.find(".//a:blip", namespaces)
                if blip is None:
                    continue
                target = relationships.get(blip.get(f"{{{namespaces['r']}}}embed"), "")
                member = "word/" + target.lstrip("/").removeprefix("word/")
                if member not in names:
                    continue
                properties = drawing.find(".//wp:docPr", namespaces)
                label = ""
                if properties is not None:
                    label = properties.get("descr") or properties.get("title") or ""
                nearby = [
                    texts[i]
                    for i in range(max(0, index - 2), min(len(texts), index + 2))
                    if texts[i]
                ]
                found.append(
                    {
                        "data": archive.read(member),
                        "label": label,
                        "context": " ".join(nearby)[:600],
                        "origin": f"{path.name}#{Path(member).name}",
                    }
                )
    return found


def _pdf_images(path: Path) -> list[dict[str, Any]]:
    try:
        import pymupdf
    except ImportError:
        return []
    found: list[dict[str, Any]] = []
    with pymupdf.open(path) as document:
        for number, page in enumerate(document, 1):
            context = page.get_text("text")[:600]
            for info in page.get_images(full=True)[:6]:
                try:
                    extracted = document.extract_image(info[0])
                except (RuntimeError, ValueError):
                    continue
                if (
                    extracted
                    and min(extracted.get("width", 0), extracted.get("height", 0))
                    >= MIN_SIDE * 2
                ):
                    found.append(
                        {
                            "data": extracted["image"],
                            "label": "",
                            "context": context,
                            "origin": f"{path.name}#page-{number}",
                        }
                    )
    return found


def collect_content_images(
    content: str | Path, source_text: str, library: ImageLibrary
) -> str:
    """Add pictures from a content document; return text without image markup."""
    candidate = Path(str(content)) if not isinstance(content, Path) else content
    try:
        is_file = candidate.is_file()
    except (OSError, ValueError):
        is_file = False
    suffix = candidate.suffix.lower() if is_file else ""
    extracted: list[dict[str, Any]] = []
    try:
        if suffix == ".pptx":
            extracted = _pptx_images(candidate)
        elif suffix == ".docx":
            extracted = _docx_images(candidate)
        elif suffix == ".pdf":
            extracted = _pdf_images(candidate)
    except (
        zipfile.BadZipFile,
        etree.XMLSyntaxError,
        KeyError,
        ValueError,
        OSError,
    ) as exc:
        library.rejected.append(
            {"origin": candidate.name, "reason": f"extraction failed: {exc}"[:200]}
        )
    for item in extracted:
        library.add(
            item["data"],
            source="document",
            label=item.get("label", ""),
            context=item.get("context", ""),
            origin=item.get("origin", ""),
        )
    if suffix in {".pptx", ".docx", ".pdf", ".json", ".csv"}:
        return source_text
    cleaned, references = extract_markdown_images(
        source_text, candidate.parent if is_file else None
    )
    for reference in references:
        if reference["path"] is None:
            library.rejected.append(
                {
                    "origin": reference["src"][:200],
                    "reason": "not a local file next to the content",
                }
            )
            continue
        asset = library.add_path(
            reference["path"], source="document", label=reference["label"]
        )
        if asset is not None and not asset.context:
            asset.context = reference["context"][:600]
    return cleaned


# --------------------------------------------------------------- matching ---
def stems(text: str) -> set[str]:
    result = set()
    for word in _WORD.findall(str(text).lower().replace("ё", "е")):
        if len(word) < 4 or word in _STOPWORDS or word.isdigit():
            continue
        result.add(word[:6])
    return result


def _slide_score(asset: ImageAsset, slide: dict[str, Any]) -> float:
    image_terms = stems(f"{asset.label} {asset.context}")
    if not image_terms:
        return 0.0
    title = stems(slide.get("title", ""))
    body = stems(
        " ".join(
            [str(slide.get("body", ""))] + [str(b) for b in slide.get("bullets", [])]
        )
    )
    score = 2.0 * len(image_terms & title) + len(image_terms & (body - title))
    return score / math.sqrt(max(1, len(image_terms)) / 4 + 1)


def _replaceable(slide: dict[str, Any]) -> bool:
    if str(slide.get("role", "content")) in {"cover", "closing", "section"}:
        return False
    visual = slide.get("visual")
    if not visual:
        return True
    if visual.get("type") == "image" and not visual.get("asset_id"):
        return True
    # Heuristic pictogram grids are decoration; supplied pictures take priority.
    return visual.get("type") == "icon_grid" and visual.get("origin") == "offline_rules"


def _use(slide: dict[str, Any], asset: ImageAsset, match: str, score: float) -> None:
    visual = slide.get("visual") or {}
    if visual.get("type") == "icon_grid" and visual.get("origin") == "offline_rules":
        slide["bullets"] = [str(item) for item in visual.get("items", [])]
    slide["visual"] = {
        "type": "image",
        "asset_id": asset.asset_id,
        "match": match,
        "match_score": round(score, 2),
        **({"caption": visual["caption"]} if visual.get("caption") else {}),
    }


def attach_images(
    plan: dict[str, Any], library: ImageLibrary, *, minimum_score: float = 1.8
) -> dict[str, Any]:
    """Assign each image to at most one slide; record every decision."""
    slides = plan.get("slides", [])
    requested = {
        str(slide["visual"]["asset_id"])
        for slide in slides
        if isinstance(slide.get("visual"), dict) and slide["visual"].get("asset_id")
    }
    free = [
        asset for asset in library.assets.values() if asset.asset_id not in requested
    ]
    candidates = [index for index, slide in enumerate(slides) if _replaceable(slide)]
    pairs = sorted(
        (
            (_slide_score(asset, slides[index]), asset.asset_id, index)
            for asset in free
            for index in candidates
        ),
        key=lambda item: (-item[0], item[1], item[2]),
    )
    used_assets: set[str] = set()
    used_slides: set[int] = set()
    placements = []
    for score, asset_id, index in pairs:
        if score < minimum_score or asset_id in used_assets or index in used_slides:
            continue
        _use(slides[index], library.assets[asset_id], "lexical", score)
        used_assets.add(asset_id)
        used_slides.add(index)
        placements.append(
            {
                "asset_id": asset_id,
                "slide": index + 1,
                "match": "lexical",
                "score": round(score, 2),
            }
        )
    # Explicit uploads without descriptive names are placed in order on the
    # least text-heavy remaining slides; the manifest marks them for review.
    unlabeled = [
        asset
        for asset in free
        if asset.asset_id not in used_assets
        and asset.source == "upload"
        and not stems(asset.label)
    ]
    remaining = [index for index in candidates if index not in used_slides]
    remaining.sort(
        key=lambda index: (
            len(" ".join(map(str, slides[index].get("bullets", []))))
            + len(str(slides[index].get("body", "")))
        )
    )
    chosen = sorted(remaining[: len(unlabeled)])
    for asset, index in zip(unlabeled, chosen):
        _use(slides[index], asset, "order", 0.0)
        used_assets.add(asset.asset_id)
        placements.append(
            {
                "asset_id": asset.asset_id,
                "slide": index + 1,
                "match": "order",
                "score": 0.0,
            }
        )
    unplaced = [
        {
            "asset_id": asset.asset_id,
            "origin": asset.origin,
            "reason": "no matching slide",
        }
        for asset in free
        if asset.asset_id not in used_assets
    ]
    plan["assets"] = library.plan_assets()
    summary = {
        "available": len(library),
        "placed": len(placements) + len(requested & set(library.assets)),
        "placements": placements,
        "unplaced": unplaced,
        "rejected": library.rejected,
    }
    plan["image_matching"] = summary
    return summary


def crop_to_fill(
    image_size: tuple[int, int], box: tuple[float, float]
) -> tuple[float, float, float, float]:
    """Crops (left, top, right, bottom) that fill ``box`` without distortion."""
    image_w, image_h = image_size
    box_w, box_h = box
    if min(image_w, image_h, box_w, box_h) <= 0:
        return 0.0, 0.0, 0.0, 0.0
    image_ratio, box_ratio = image_w / image_h, box_w / box_h
    if abs(image_ratio - box_ratio) < 1e-6:
        return 0.0, 0.0, 0.0, 0.0
    if image_ratio > box_ratio:
        excess = 1 - box_ratio / image_ratio
        return excess / 2, 0.0, excess / 2, 0.0
    excess = 1 - image_ratio / box_ratio
    # Keep more of the upper part of tall pictures, where subjects usually are.
    return 0.0, excess * 0.35, 0.0, excess * 0.65
