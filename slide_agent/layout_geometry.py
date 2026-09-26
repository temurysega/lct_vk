"""Shape geometry shared by the composer cleanup and the structural audit."""

from __future__ import annotations

from typing import Any

from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn

Box = tuple[float, float, float, float]

OPAQUE_TYPES = {
    MSO_SHAPE_TYPE.PICTURE,
    MSO_SHAPE_TYPE.GROUP,
    MSO_SHAPE_TYPE.CHART,
    MSO_SHAPE_TYPE.TABLE,
    MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT,
}


def is_opaque(shape: Any) -> bool:
    """Pictures, diagrams, charts and tables, including filled placeholders."""
    return shape.shape_type in OPAQUE_TYPES or shape.element.tag in {
        qn("p:pic"),
        qn("p:graphicFrame"),
    }


def box(shape: Any) -> Box:
    return (
        (shape.left or 0) / 914400,
        (shape.top or 0) / 914400,
        (shape.width or 0) / 914400,
        (shape.height or 0) / 914400,
    )


def intersection(first: Box, second: Box) -> float:
    x1, y1, w1, h1 = first
    x2, y2, w2, h2 = second
    width = min(x1 + w1, x2 + w2) - max(x1, x2)
    height = min(y1 + h1, y2 + h2) - max(y1, y2)
    return width * height if width > 0 and height > 0 else 0.0


def gap(first: Box, second: Box) -> float:
    """Distance between box edges; 0 when the boxes touch or intersect."""
    x1, y1, w1, h1 = first
    x2, y2, w2, h2 = second
    dx = max(0.0, x2 - (x1 + w1), x1 - (x2 + w2))
    dy = max(0.0, y2 - (y1 + h1), y1 - (y2 + h2))
    return max(dx, dy)


def area(value: Box) -> float:
    return max(0.0, value[2]) * max(0.0, value[3])


def has_text(shape: Any) -> bool:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return any(has_text(child) for child in shape.shapes)
    return bool(getattr(shape, "has_text_frame", False) and shape.text.strip())


def is_structural(value: Box, canvas: tuple[float, float]) -> bool:
    """Bands, panels and backgrounds frame a slide rather than one block."""
    width, height = canvas
    return (
        value[2] >= width * 0.85
        or value[3] >= height * 0.85
        or area(value) >= width * height * 0.3
    )


def is_framing(
    shape: Any,
    zone: Box,
    canvas: tuple[float, float],
    content: list[tuple[Any, Box]],
    held_sample: bool,
) -> bool:
    """Bands, panels and cards that frame new content rather than compete.

    A large shape frames the slide unless it held sample text: then it was
    exemplar content (a chart drawn with shapes, a quote panel). A card frames
    the new content it lies under; a group is never such a card, because new
    content is drawn over a group, not inside it.
    """
    if is_structural(zone, canvas) and not held_sample:
        return True
    return shape.shape_type != MSO_SHAPE_TYPE.GROUP and any(
        hosts(zone, other) for _, other in content
    )


def hosts(container: Box, content: Box) -> bool:
    """A card behind text: it covers most of the content box."""
    return intersection(container, content) >= 0.8 * max(area(content), 1e-6)


def covered_share(content: Any) -> float:
    """Share of a template shape that content must cover to collide with it.

    Pictures and diagrams are opaque; a text box is often wider than its
    text, so a sliver under its edge is not a visible collision.
    """
    return 0.1 if is_opaque(content) else 0.25


def collides(
    shape: Any, zone: Box, content: Any, other: Box, template_origin: bool
) -> bool:
    """Whether new content covers a template shape.

    A line collides when it crosses non-title content. Another shape collides
    when enough of it lies under the content; decoration inside a visible
    template card (an icon in its corner) is the card's own design, while a
    bar inside an invisible text container sits in the flow of the new text.
    """
    if shape.shape_type == MSO_SHAPE_TYPE.LINE:
        return not is_title(content) and intersection(_thick(zone), other) > 0
    covered = intersection(zone, other)
    if covered < covered_share(content) * max(area(zone), 1e-6):
        return False
    inside_card = (
        template_origin
        and covered >= 0.95 * max(area(zone), 1e-6)
        and is_visible_box(content)
    )
    return not inside_card


def _thick(zone: Box) -> Box:
    """Give a horizontal or vertical line a width so that it can intersect."""
    x, y, w, h = zone
    return (x - 0.01, y - 0.01, w + 0.02, h + 0.02)


def is_title(shape: Any) -> bool:
    if not getattr(shape, "is_placeholder", False):
        return False
    try:
        return shape.placeholder_format.type in {
            PP_PLACEHOLDER.TITLE,
            PP_PLACEHOLDER.CENTER_TITLE,
        }
    except ValueError:
        return False


def is_generated(shape: Any) -> bool:
    return str(shape.name).startswith("BrandDeck")


def is_visible_box(shape: Any) -> bool:
    """Whether a text-free shape still draws something (fill, outline, art)."""
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return any(is_visible_box(child) for child in shape.shapes)
    if is_opaque(shape):
        return True
    properties = shape.element.find(qn("p:spPr"))
    if properties is None:
        return False
    fills = ("a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill")
    if any(properties.find(qn(tag)) is not None for tag in fills):
        return True
    line = properties.find(qn("a:ln"))
    no_line = line is not None and line.find(qn("a:noFill")) is not None
    if line is not None and any(line.find(qn(tag)) is not None for tag in fills):
        return True
    # Without explicit properties the theme style (p:style) draws the shape.
    style = shape.element.find(qn("p:style"))
    if style is None:
        return False
    fill_ref = style.find(qn("a:fillRef"))
    line_ref = style.find(qn("a:lnRef"))
    styled_fill = (
        fill_ref is not None
        and fill_ref.get("idx", "0") != "0"
        and properties.find(qn("a:noFill")) is None
    )
    styled_line = line_ref is not None and line_ref.get("idx", "0") != "0" and not no_line
    return styled_fill or styled_line


FRAME_GEOMETRIES = {"ellipse", "roundRect"}


def is_photo_frame(
    shape: Any,
    zone: Box,
    content: list[tuple[Any, Box]],
    occupants: list[Box] = (),
) -> bool:
    """An empty avatar circle or icon tile left of a caption.

    The speaker photo or the icon it framed is gone: nothing else is centred
    inside it and its own fill is not a picture.
    """
    if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
        return False
    properties = shape.element.find(qn("p:spPr"))
    geometry = properties.find(qn("a:prstGeom")) if properties is not None else None
    if geometry is None or geometry.get("prst") not in FRAME_GEOMETRIES:
        return False
    if properties.find(qn("a:blipFill")) is not None:
        return False  # a real photo
    x, y, w, h = zone
    if not (0.5 <= min(w, h) and max(w, h) <= 1.6 and 0.8 <= w / max(h, 0.01) <= 1.25):
        return False
    if any(
        x <= ox + ow / 2 <= x + w and y <= oy + oh / 2 <= y + h
        for ox, oy, ow, oh in occupants
    ):
        return False
    middle = y + h / 2
    return any(
        has_text(item)
        and 0 <= other[0] - (x + w) <= 0.5
        and other[1] <= middle <= other[1] + other[3]
        for item, other in content
    )
