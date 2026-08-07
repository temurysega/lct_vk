from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Inches, Pt

MAX_PDF_VECTORS_PER_PAGE = 300
MAX_PDF_IMAGES_PER_PAGE = 150
MIN_PDF_VECTOR_AREA_INCHES = 0.003
MIN_PDF_IMAGE_PIXELS = 5


def _rgb_tuple(value: Any, fallback: str = "000000") -> RGBColor:
    if not isinstance(value, (list, tuple)) or len(value) < 3:
        return RGBColor.from_string(fallback)
    channels = [max(0, min(255, round(float(channel) * 255))) for channel in value[:3]]
    return RGBColor(*channels)


def _pdf_int_color(value: int | None, fallback: str = "000000") -> RGBColor:
    if not isinstance(value, int):
        return RGBColor.from_string(fallback)
    return RGBColor((value >> 16) & 255, (value >> 8) & 255, value & 255)


def _clean_font_name(value: str) -> str:
    value = re.sub(r"^[A-Z]{6}\+", "", value or "")
    value = value.replace("-Bold", "").replace("-Regular", "")
    aliases = {
        "ArialMT": "Arial",
        "Helvetica": "Arial",
        "HelveticaNeue": "Arial",
    }
    return aliases.get(value, value) or "Arial"


def _to_inches(value: float, page_points: float, slide_inches: float) -> float:
    return float(value) / float(page_points) * slide_inches


def _add_drawings(
    slide: Any, page: Any, slide_width: float, slide_height: float
) -> int:
    added = 0
    candidates: list[tuple[float, Any]] = []
    for drawing in page.get_drawings():
        fill = drawing.get("fill")
        rect = drawing.get("rect")
        if fill is None or rect is None or rect.width <= 0 or rect.height <= 0:
            continue
        area = (
            _to_inches(rect.width, page.rect.width, slide_width)
            * _to_inches(rect.height, page.rect.height, slide_height)
        )
        if area < MIN_PDF_VECTOR_AREA_INCHES:
            continue
        candidates.append((area, drawing))
    candidates.sort(key=lambda item: item[0], reverse=True)
    for _, drawing in candidates[:MAX_PDF_VECTORS_PER_PAGE]:
        fill = drawing["fill"]
        rect = drawing["rect"]
        x = _to_inches(rect.x0, page.rect.width, slide_width)
        y = _to_inches(rect.y0, page.rect.height, slide_height)
        w = _to_inches(rect.width, page.rect.width, slide_width)
        h = _to_inches(rect.height, page.rect.height, slide_height)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
        )
        shape.name = f"PDF vector {added + 1}"
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb_tuple(fill, "FFFFFF")
        line_color = drawing.get("color")
        if line_color is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = _rgb_tuple(line_color)
            shape.line.width = Pt(max(0.1, float(drawing.get("width", 0.5))))
        added += 1
    return added


def _add_images(
    slide: Any,
    page: Any,
    document: Any,
    slide_width: float,
    slide_height: float,
) -> int:
    added = 0
    seen_rectangles: set[tuple[float, float, float, float]] = set()
    for image in page.get_images(full=True):
        xref = int(image[0])
        width_px = int(image[2])
        height_px = int(image[3])
        if width_px < MIN_PDF_IMAGE_PIXELS or height_px < MIN_PDF_IMAGE_PIXELS:
            continue
        try:
            extracted = document.extract_image(xref)
            payload = extracted.get("image")
            if not payload:
                continue
            rectangles = page.get_image_rects(xref)
        except (RuntimeError, ValueError):
            continue
        for rect in rectangles:
            key = tuple(round(value, 2) for value in rect)
            if key in seen_rectangles or rect.width <= 0 or rect.height <= 0:
                continue
            seen_rectangles.add(key)
            x = _to_inches(rect.x0, page.rect.width, slide_width)
            y = _to_inches(rect.y0, page.rect.height, slide_height)
            w = _to_inches(rect.width, page.rect.width, slide_width)
            h = _to_inches(rect.height, page.rect.height, slide_height)
            try:
                picture = slide.shapes.add_picture(
                    io.BytesIO(payload), Inches(x), Inches(y), Inches(w), Inches(h)
                )
                picture.name = f"PDF image {added + 1}"
                added += 1
                if added >= MAX_PDF_IMAGES_PER_PAGE:
                    return added
            except (OSError, ValueError):
                continue
    return added


def _add_text(slide: Any, page: Any, slide_width: float, slide_height: float) -> int:
    added = 0
    text_dict = page.get_text("dict", flags=11)
    for block in text_dict.get("blocks", []):
        if block.get("type") != 0 or not block.get("lines"):
            continue
        x0, y0, x1, y1 = block["bbox"]
        x = _to_inches(x0, page.rect.width, slide_width)
        y = _to_inches(y0, page.rect.height, slide_height)
        w = max(0.05, _to_inches(x1 - x0, page.rect.width, slide_width))
        h = max(0.05, _to_inches(y1 - y0, page.rect.height, slide_height))
        textbox = slide.shapes.add_textbox(
            Inches(x), Inches(y), Inches(w + 0.03), Inches(h + 0.03)
        )
        textbox.name = f"PDF text block {added + 1}"
        frame = textbox.text_frame
        frame.clear()
        frame.word_wrap = False
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        frame.vertical_anchor = MSO_ANCHOR.TOP
        frame.margin_left = 0
        frame.margin_right = 0
        frame.margin_top = 0
        frame.margin_bottom = 0
        for line_index, line in enumerate(block["lines"]):
            paragraph = (
                frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
            )
            for span in line.get("spans", []):
                if not span.get("text"):
                    continue
                run = paragraph.add_run()
                run.text = span["text"]
                run.font.name = _clean_font_name(span.get("font", ""))
                run.font.size = Pt(max(4.0, float(span.get("size", 10))))
                run.font.bold = bool(int(span.get("flags", 0)) & 16)
                run.font.italic = bool(int(span.get("flags", 0)) & 2)
                run.font.color.rgb = _pdf_int_color(span.get("color"))
        if not textbox.text.strip():
            slide.shapes._spTree.remove(textbox.element)
            continue
        added += 1
    return added


def pdf_to_pptx(
    pdf_path: str | Path,
    output_path: str | Path,
    *,
    max_pages: int | None = None,
) -> dict[str, Any]:
    """Convert a vector PDF deck into an editable PPTX reference fixture."""
    try:
        import pymupdf
    except ImportError as exc:
        raise RuntimeError(
            "Install PDF reference support: pip install -e '.[pdf-reference]'"
        ) from exc

    source = Path(pdf_path).expanduser().resolve()
    target = Path(output_path).expanduser().resolve()
    document = pymupdf.open(source)
    if document.page_count == 0:
        raise ValueError(f"PDF has no pages: {source}")
    first_page = document[0]
    slide_width = 13.333
    slide_height = slide_width * first_page.rect.height / first_page.rect.width
    if slide_height > 8.5:
        slide_height = 7.5
        slide_width = slide_height * first_page.rect.width / first_page.rect.height

    prs = Presentation()
    prs.slide_width = Inches(slide_width)
    prs.slide_height = Inches(slide_height)
    blank = prs.slide_layouts[6]
    counts = {"vectors": 0, "images": 0, "text_blocks": 0}
    page_limit = min(document.page_count, max_pages or document.page_count)
    for page_number in range(page_limit):
        page = document[page_number]
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        counts["vectors"] += _add_drawings(slide, page, slide_width, slide_height)
        counts["images"] += _add_images(
            slide, page, document, slide_width, slide_height
        )
        counts["text_blocks"] += _add_text(slide, page, slide_width, slide_height)

    prs.core_properties.title = source.stem
    prs.core_properties.subject = "Editable reference imported from a PDF template"
    target.parent.mkdir(parents=True, exist_ok=True)
    prs.save(target)
    document.close()
    return {
        "source": str(source),
        "output": str(target),
        "slide_count": page_limit,
        **counts,
    }
