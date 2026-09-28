from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Any

from pptx import Presentation

from .brief import plan_from_brief, use_brief_mode
from .diagrams import DIAGRAM_TYPES
from .llm import InferenceClient, InferenceError
from .model_context import compact_patterns
from .prompt_config import load_prompt
from .visual_planning import (
    add_offline_visuals,
    sanitize_diagram,
    sanitize_image,
    visual_text,
)

PLANNER_SYSTEM = load_prompt("planner")
# Target deck size of the brief when the user does not set a number.
DEFAULT_SLIDE_RANGE = (10, 15)


def load_content(value: str | Path) -> tuple[str, str]:
    """Return ``(content, source_label)`` from a path or inline text."""
    candidate = Path(str(value)).expanduser()
    try:
        exists = candidate.is_file()
    except OSError:
        exists = False
    if not exists:
        return str(value).strip(), "inline"

    suffix = candidate.suffix.lower()
    if suffix in {".md", ".txt", ".html"}:
        return candidate.read_text(encoding="utf-8"), candidate.name
    if suffix == ".json":
        data = json.loads(candidate.read_text(encoding="utf-8"))
        return json.dumps(data, ensure_ascii=False, indent=2), candidate.name
    if suffix == ".csv":
        with candidate.open("r", encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        return "\n".join(" | ".join(row) for row in rows), candidate.name
    if suffix == ".pptx":
        prs = Presentation(candidate)
        blocks = []
        for index, slide in enumerate(prs.slides, 1):
            texts = [
                shape.text.strip()
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text.strip()
            ]
            blocks.append(f"## Slide {index}\n" + "\n".join(texts))
        return "\n\n".join(blocks), candidate.name

    try:
        from markitdown import MarkItDown

        result = MarkItDown().convert(str(candidate))
        return result.text_content, candidate.name
    except ImportError as exc:
        raise ValueError(
            f"{suffix or 'This'} document type requires the optional 'documents' dependencies"
        ) from exc


def _detect_russian(text: str) -> bool:
    cyrillic = len(re.findall(r"[А-Яа-яЁё]", text))
    latin = len(re.findall(r"[A-Za-z]", text))
    return cyrillic > latin


def _sentences(text: str) -> list[str]:
    cleaned = re.sub(r"\s+", " ", text).strip()
    if not cleaned:
        return []
    # A semicolon separates list items only when a new statement follows;
    # "…специалисту; ссылку можно проверить" is one sentence.
    parts = re.split(r"(?<=[.!?])\s+|\s*•\s*|\s*;\s*(?=[A-ZА-ЯЁ0-9])", cleaned)
    return [part.strip(" -\t") for part in parts if part.strip(" -\t")]


def _sections(text: str) -> list[tuple[str, list[str]]]:
    sections: list[tuple[str, list[str]]] = []
    title = ""
    content: list[str] = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        heading = re.match(r"^#{1,6}\s+(.+)$", line)
        if heading:
            if title or content:
                sections.append((title, content))
            title, content = heading.group(1).strip(), []
            continue
        if line:
            content.append(re.sub(r"^[-*+]\s+", "", line))
    if title or content:
        sections.append((title, content))
    return sections


def _fallback_plan(text: str, slide_count: int | None) -> dict[str, Any]:
    russian = _detect_russian(text)
    sections = _sections(text)
    if not sections:
        sections = [("", _sentences(text))]

    first_title = next((title for title, _ in sections if title), "")
    if not first_title:
        words = re.sub(r"\s+", " ", text).strip().split()
        first_title = " ".join(words[:8]) or (
            "Презентация" if russian else "Presentation"
        )

    # The service targets 10-15 slides unless the user sets a number; thin
    # material gives fewer slides rather than one statement per slide.
    low, high = DEFAULT_SLIDE_RANGE
    desired = max(3, slide_count or min(high, max(low, len(sections) + 2)))
    content_slots = desired - 2
    material: list[tuple[str, list[str]]] = []
    for index, (heading, lines) in enumerate(sections):
        if index == 0 and heading == first_title and len(sections) > 1 and not lines:
            continue
        bullets = []
        for line in lines:
            bullets.extend(_sentences(line) or [line])
        material.append((heading, bullets))
    if not material:
        material = [("", _sentences(text))]

    # Merge adjacent sections instead of dropping the tail of the document.
    while len(material) > content_slots:
        index = min(
            range(len(material) - 1),
            key=lambda i: sum(
                len(item) for _, items in material[i : i + 2] for item in items
            ),
        )
        first, second = material[index : index + 2]
        material[index : index + 2] = [
            (
                " / ".join(heading for heading in (first[0], second[0]) if heading),
                first[1] + second[1],
            )
        ]

    while len(material) < content_slots:
        split_at = max(range(len(material)), key=lambda idx: len(material[idx][1]))
        heading, bullets = material[split_at]
        if len(bullets) < (2 if slide_count is not None else 4):
            break
        midpoint = max(1, len(bullets) // 2)
        material[split_at : split_at + 1] = [
            (heading, bullets[:midpoint]),
            (heading, bullets[midpoint:]),
        ]

    if len(material) < content_slots:
        if slide_count is not None:
            raise ValueError(
                "Not enough source material for the requested slide count; request fewer slides"
            )
        desired = len(material) + 2

    slides: list[dict[str, Any]] = [
        {
            "title": first_title,
            "role": "cover",
            # Service captions ("generated automatically") are slide noise
            # (Appendix 1, question 7); the template's subtitle stays empty.
            "subtitle": "",
            "body": "",
            "bullets": [],
            "visual": None,
            "speaker_notes": "",
        }
    ]
    for index, (heading, bullets) in enumerate(material[:content_slots], 1):
        default_title = f"Раздел {index}" if russian else f"Section {index}"
        slide_bullets = list(bullets)
        slides.append(
            {
                "title": heading or default_title,
                "role": "content",
                "subtitle": "",
                "body": "",
                "bullets": slide_bullets,
                "visual": None,
                "speaker_notes": "",
            }
        )
    slides.append(
        {
            "title": "Спасибо" if russian else "Thank you",
            "role": "closing",
            "subtitle": "Вопросы?" if russian else "Questions?",
            "body": "",
            "bullets": [],
            "visual": None,
            "speaker_notes": "",
        }
    )
    add_offline_visuals(slides)
    return {
        "language": "ru" if russian else "en",
        "title": first_title,
        "slides": slides,
    }


def _sanitize_visual(
    value: Any, asset_ids: set[str] | frozenset[str] = frozenset()
) -> dict[str, Any] | None:
    if not isinstance(value, dict):
        return None
    if value.get("type") in DIAGRAM_TYPES:
        return sanitize_diagram(value)
    if value.get("type") == "image":
        return sanitize_image(value, asset_ids)
    supported = {"metric_cards", "bar_chart", "pie_chart", "table", "timeline"}
    if value.get("type") not in supported:
        return None
    return value


def normalize_plan(
    plan: dict[str, Any],
    *,
    source_text: str,
    slide_count: int | None,
    asset_ids: set[str] | frozenset[str] = frozenset(),
) -> dict[str, Any]:
    raw_slides = plan.get("slides")
    if not isinstance(raw_slides, list) or not raw_slides:
        return _fallback_plan(source_text, slide_count)

    allowed_roles = {
        "cover",
        "section",
        "content",
        "two_column",
        "comparison",
        "data",
        "image",
        "closing",
    }
    slides: list[dict[str, Any]] = []
    for raw in raw_slides:
        if not isinstance(raw, dict):
            continue
        title = str(raw.get("title") or "Untitled").strip()
        bullets = raw.get("bullets") if isinstance(raw.get("bullets"), list) else []
        bullets = [str(item).strip() for item in bullets if str(item).strip()]
        role = str(raw.get("role", "content")).lower()
        slides.append(
            {
                "title": title,
                "role": role if role in allowed_roles else "content",
                "subtitle": str(raw.get("subtitle") or "").strip(),
                "body": str(raw.get("body") or "").strip(),
                "bullets": bullets,
                "visual": _sanitize_visual(raw.get("visual"), asset_ids),
                "speaker_notes": str(raw.get("speaker_notes") or "").strip(),
            }
        )

    if len(slides) < 2:
        return _fallback_plan(source_text, slide_count)
    slides[0]["role"] = "cover"
    slides[-1]["role"] = "closing"
    if slide_count and len(slides) != slide_count:
        rebuilt = _fallback_plan(source_text, slide_count)
        rebuilt["replanned_reason"] = "slide_count_mismatch"
        return rebuilt
    return {
        "language": str(
            plan.get("language") or ("ru" if _detect_russian(source_text) else "en")
        ),
        "title": str(plan.get("title") or slides[0]["title"]),
        "slides": slides,
    }


def plan_deck(
    content: str,
    *,
    pattern_catalog: dict[str, Any],
    design_system: dict[str, Any],
    slide_count: int | None = None,
    client: InferenceClient | None = None,
    image_catalog: list[dict[str, str]] | None = None,
    mode: str = "auto",
    purpose: str | None = None,
) -> dict[str, Any]:
    asset_ids = {str(item["id"]) for item in image_catalog or []}
    brief_mode = use_brief_mode(content, mode, client)
    if brief_mode:
        assert client is not None
        try:
            plan = plan_from_brief(
                content,
                client=client,
                slide_count=slide_count,
                purpose=purpose,
                image_catalog=image_catalog,
            )
        except InferenceError as exc:
            # Expanding a brief is the model's job; a deterministic fallback
            # would only produce a few near-empty slides.
            raise ValueError(f"Could not build a deck from the brief: {exc}") from exc
        planner_mode = "inference"
    elif client is None:
        plan = _fallback_plan(content, slide_count)
        planner_mode = "offline"
    else:
        payload = {
            "requested_slide_count": slide_count,
            "source": content,
            "available_template_patterns": compact_patterns(pattern_catalog),
            "design_tone": {
                "font": design_system.get("typography", {}).get("primary_font"),
                "palette": design_system.get("colors", {}).get("theme", []),
            },
        }
        if image_catalog:
            payload["available_images"] = image_catalog[:20]
        user = json.dumps(payload, ensure_ascii=False)
        # A wrong slide count discards the whole plan in normalize_plan, so the
        # full model gets it as a grammar constraint, not only as a request.
        schema = (
            None
            if not slide_count
            else {
                "type": "object",
                "properties": {
                    "language": {"type": "string"},
                    "title": {"type": "string"},
                    "slides": {
                        "type": "array",
                        "items": {"type": "object"},
                        "minItems": slide_count,
                        "maxItems": slide_count,
                    },
                },
                "required": ["language", "title", "slides"],
            }
        )
        try:
            plan = client.chat_json(
                system=PLANNER_SYSTEM,
                user=user,
                max_tokens=6000,
                schema=schema,
            )
            planner_mode = "inference"
        except InferenceError as exc:
            plan = _fallback_plan(content, slide_count)
            planner_mode = "fallback"
            plan["inference_error"] = str(exc)[:400]
    if planner_mode == "inference":
        raw_slides = plan.get("slides") if isinstance(plan, dict) else None
        if (
            not isinstance(raw_slides, list)
            or sum(isinstance(item, dict) for item in raw_slides) < 2
        ):
            planner_mode = "fallback"
    normalized = normalize_plan(
        plan, source_text=content, slide_count=slide_count, asset_ids=asset_ids
    )
    if normalized.get("replanned_reason"):
        planner_mode = "fallback"
    normalized["planner"] = {
        "mode": planner_mode,
        "input": "brief" if brief_mode else "source",
    }
    if brief_mode:
        normalized["brief"] = plan["brief"]
    if plan.get("inference_error"):
        normalized["planner"]["reason"] = plan["inference_error"]
    return normalized


def _data_text(visual: dict[str, Any]) -> list[str]:
    """Text of cards, tables, timelines and charts.

    Counting it lets capacity scoring reject layouts without a usable zone,
    which would otherwise squeeze the visual into a decorative label.
    """
    kind = visual.get("type")
    if kind in {"metric_cards", "timeline"}:
        return [
            str(item.get(key))
            for item in visual.get("items", [])
            if isinstance(item, dict)
            for key in ("label", "value", "detail")
            if item.get(key)
        ]
    if kind == "table":
        cells = list(visual.get("headers", []))
        for row in visual.get("rows", []):
            cells.extend(row if isinstance(row, list) else [])
        return [str(cell) for cell in cells]
    if kind in {"bar_chart", "pie_chart"}:
        return [str(item) for item in visual.get("categories", [])]
    return []


def _slide_requirements(slide: dict[str, Any]) -> dict[str, Any]:
    # Diagrams are described as the list of statements they carry.
    visual = slide.get("visual") or {}
    diagram_text = visual_text(visual)
    body_chars = (
        len(str(slide.get("body", "")))
        + sum(len(str(item)) for item in slide.get("bullets", []))
        + sum(len(text) for text in diagram_text)
        + sum(len(text) for text in _data_text(visual))
    )
    bullet_count = len(slide.get("bullets", [])) or len(
        visual.get("items", []) if diagram_text else []
    )
    role = str(slide.get("role", "content"))
    if role in {"comparison", "two_column"}:
        rhetorical_pattern = "two_column"
    elif role == "data" or visual.get("type") in {
        "bar_chart",
        "pie_chart",
        "table",
        "metric_cards",
    }:
        rhetorical_pattern = "data"
    elif bullet_count >= 2:
        rhetorical_pattern = "bullet_list"
    elif role in {"cover", "closing", "section"}:
        rhetorical_pattern = "hero" if role == "cover" else role
    else:
        rhetorical_pattern = "content"
    return {
        "role": role,
        "title_chars": len(str(slide.get("title", ""))),
        "body_chars": body_chars,
        "bullet_count": bullet_count,
        "needs_data": rhetorical_pattern == "data",
        "needs_image": role == "image" or visual.get("type") == "image",
        "rhetorical_pattern": rhetorical_pattern,
        "has_visual": bool(visual),
    }


def _score_pattern(
    pattern: dict[str, Any],
    requirements: dict[str, Any],
    *,
    reuse_count: int,
    avoided: bool,
) -> tuple[float, list[str], list[str]]:
    roles = set(pattern.get("roles", []))
    capacity = pattern.get("capacity", {})
    role = requirements["role"]
    reasons: list[str] = []
    risks: list[str] = []
    score = 0.0

    if role in roles:
        score += 10.0
        reasons.append(f"роль {role} поддерживается")
    if role == "content" and roles.intersection({"cover", "closing", "section"}):
        score -= 5.0
    if role not in {"cover", "closing", "section"} and "closing" in roles:
        score -= 9.0
    if role != "cover" and "cover" in roles:
        score -= 6.0

    requested_pattern = requirements["rhetorical_pattern"]
    available_pattern = capacity.get("rhetorical_pattern", "content")
    compatible_patterns = {
        ("content", "bullet_list"),
        ("bullet_list", "content"),
        ("data", "multi_card"),
        ("two_column", "comparison"),
    }
    if requested_pattern == available_pattern:
        score += 5.0
        reasons.append(f"риторический паттерн {available_pattern}")
    elif (requested_pattern, available_pattern) in compatible_patterns:
        score += 2.0
    elif requested_pattern in {"data", "two_column"}:
        score -= 3.0
        risks.append(f"нужен {requested_pattern}, доступен {available_pattern}")

    title_capacity = max(1, int(capacity.get("title_chars", 80) or 80))
    if requirements["title_chars"] <= title_capacity:
        score += 1.0
    else:
        score -= min(5.0, (requirements["title_chars"] / title_capacity - 1) * 4)
        risks.append("заголовок длиннее визуальной ёмкости")

    body_capacity = max(1, int(capacity.get("body_chars", 500) or 500))
    if requirements["body_chars"]:
        if capacity.get("has_title_zone") is False:
            score -= 12
            risks.append("нет отдельной зоны заголовка")
        if float(capacity.get("title_top_ratio", 0)) > 0.28:
            score -= 12
            risks.append("центральный заголовок оставляет мало места для текста")
        score -= int(capacity.get("out_of_canvas_body_zones", 0)) * 12
        if not requirements["has_visual"]:
            score -= min(16, int(capacity.get("specialized_labels", 0)) * 3)
            if 0.25 <= float(capacity.get("body_area_ratio", 0)) <= 0.7:
                score += 4
                reasons.append("полноценная область для чтения")
        ratio = requirements["body_chars"] / body_capacity
        if ratio <= 1:
            score += 3.0
            reasons.append("текст помещается по геометрической ёмкости")
        else:
            score -= min(10.0, (ratio - 1) * 7)
            risks.append(f"текст/ёмкость: {ratio:.1f}×")
        if capacity.get("usable_body_zones") == 0:
            score -= 25
            risks.append("нет полноценной текстовой зоны")

    bullet_capacity = max(
        1,
        int(capacity.get("body_zones", 1) or 1)
        * int(capacity.get("bullets_per_zone", 4) or 4),
    )
    if requirements["bullet_count"]:
        usable_zones = int(capacity.get("usable_body_zones", 1))
        if usable_zones > requirements["bullet_count"]:
            score -= (usable_zones - requirements["bullet_count"]) * 4
            risks.append("часть текстовых блоков останется пустой")
        if requirements["bullet_count"] <= bullet_capacity:
            score += 2.0
        else:
            score -= (requirements["bullet_count"] - bullet_capacity) * 1.5
            risks.append("слишком много тезисов для числа слотов")

    if requirements["needs_data"]:
        if capacity.get("supports_data"):
            score += 4.0
            reasons.append("есть структура для данных")
        else:
            score -= 4.0
            risks.append("нет нативной структуры для данных")
    if requirements["needs_image"]:
        # Only picture placeholders and exemplar photos can be replaced;
        # pictures baked into a layout would stay next to the new image.
        slot_area = float(capacity.get("image_slot_area_ratio", 0) or 0)
        content_image_area = float(capacity.get("content_image_area_ratio", 0) or 0)
        if slot_area >= 0.03 or (
            pattern.get("source_kind") == "slide_exemplar"
            and content_image_area >= 0.04
        ):
            score += 5.0
            reasons.append("заменяемый слот изображения")
        else:
            # Without a slot the picture competes with exemplar decoration
            # for the text zone; any slotted layout of the template is better.
            score -= 6.0 + content_image_area * 20.0
            risks.append("нет слота изображения")
            if content_image_area:
                risks.append("изображения макета останутся рядом с новым")
    else:
        # Charts and diagrams do not fill picture slots either.
        content_image_area = float(capacity.get("content_image_area_ratio", 0) or 0)
        image_slot_area = float(capacity.get("image_slot_area_ratio", 0) or 0)
        score -= content_image_area * 20.0 + image_slot_area * 24.0
        if image_slot_area > 0.1:
            risks.append("слот изображения останется пустым")
        if content_image_area >= 0.35:
            risks.append("крупное исходное изображение не связано с новым текстом")
    backdrop = float(capacity.get("backdrop_area_ratio", 0) or 0)
    carries_text = role not in {"cover", "closing", "section"} or (
        role == "closing" and requirements["body_chars"]
    )
    if backdrop and carries_text:
        # A colour block reserved for an object stays empty on a generated
        # content slide: nothing in the plan is placed there.
        score -= backdrop * 24.0
        risks.append("пустая область макета под объект")
    if requirements["has_visual"] and capacity.get("body_zones", 0) >= 1:
        score += 1.0

    complexity = float(capacity.get("complexity", 0) or 0)
    if (
        requested_pattern in {"content", "bullet_list"}
        and not requirements["has_visual"]
    ):
        score -= complexity * 4.0
        if complexity >= 0.75:
            risks.append("сложный декоративный exemplar для простого текста")

    if pattern.get("source_kind") == "slide_exemplar":
        score += 1.5
        reasons.append("проверенный exemplar-слайд")
    elif pattern.get("source_kind") == "native_grid":
        score += 4.0
        reasons.append("чистая нативная сетка без PDF-фрагментов")
    score += min(len(pattern.get("example_slide_indices", [])), 2) * 0.2
    reuse_penalty = 0.5
    score -= reuse_count * reuse_penalty
    if avoided:
        score -= 50.0
        risks.append("паттерн исключён предыдущим QA")
    return score, reasons, risks


def _side_title_column(pattern: dict[str, Any]) -> bool:
    """The title has its own column beside the main content zone.

    Below such a title the column stays empty unless the slide has text for
    it, so a slide carried by one diagram leaves much of the canvas blank.
    """
    zones = pattern.get("placeholders") or []
    titles = [
        zone for zone in zones
        if str(zone.get("type", "")).startswith(("title", "center_title"))
    ]
    bodies = [
        zone for zone in zones
        if str(zone.get("type", "")).startswith(("body", "object"))
        and float(zone.get("h", 0) or 0) >= 1.0
    ]
    if len(titles) != 1 or not bodies:
        return False
    title = titles[0]
    main = max(bodies, key=lambda zone: float(zone["w"]) * float(zone["h"]))
    title_bottom = float(title["y"]) + float(title["h"])
    return (
        float(title["x"]) + float(title["w"]) <= float(main["x"]) + 0.1
        and float(main["y"]) < title_bottom
        and float(main["y"]) + float(main["h"]) - title_bottom >= 1.5
    )


def assign_patterns(
    plan: dict[str, Any],
    pattern_catalog: dict[str, Any],
    *,
    avoid_by_slide: dict[int, set[str]] | None = None,
    layout_strategy: str = "balanced",
    previous_by_slide: dict[int, set[str]] | None = None,
) -> dict[str, Any]:
    if layout_strategy not in {"balanced", "columns", "focus"}:
        raise ValueError("Unknown layout strategy")
    patterns = pattern_catalog.get("patterns", [])
    if not patterns:
        raise ValueError("Template has no usable slide layouts")
    usage: dict[str, int] = {}
    for slide_index, slide in enumerate(plan.get("slides", []), 1):
        # Keep the mode on the slide too: a sparse text slide may become a
        # native icon grid only during composition, after pattern assignment.
        slide["layout_variant"] = layout_strategy
        requirements = _slide_requirements(slide)
        diagram = (slide.get("visual") or {}).get("type") in DIAGRAM_TYPES
        scored: list[tuple[float, dict[str, Any], list[str], list[str]]] = []
        for pattern in patterns:
            score, reasons, risks = _score_pattern(
                pattern,
                requirements,
                reuse_count=usage.get(pattern["id"], 0),
                avoided=pattern["id"] in (avoid_by_slide or {}).get(slide_index, set()),
            )
            if diagram:
                # Diagrams need one calm, large content area.
                capacity = pattern.get("capacity", {})
                area = float(capacity.get("body_area_ratio", 0) or 0)
                score += min(area, 0.45) * 24 - 5
                score -= float(capacity.get("complexity", 0) or 0) * 4
                if area >= 0.25:
                    reasons.append("крупная зона для схемы")
                else:
                    risks.append("схема займёт общую зону контента")
                if (
                    slide.get("remap_underfilled")
                    and not slide.get("bullets")
                    and not slide.get("body")
                ):
                    # The render showed this diagram below a quarter of the
                    # slide: the size of the new zone now matters most, and
                    # the column under a side title would stay empty again.
                    score += area * 40
                    if _side_title_column(pattern):
                        score -= 12.0
                        risks.append("колонка под заголовком останется пустой")
            if requirements["body_chars"]:
                zones = int(pattern.get("capacity", {}).get("usable_body_zones", 1))
                if layout_strategy == "columns" and 2 <= zones <= 4:
                    score += 4
                    reasons.append("вариант с несколькими текстовыми блоками")
                elif layout_strategy == "focus" and zones == 1:
                    score += 6
                    reasons.append("вариант с единым текстовым блоком")
                if pattern["id"] in (previous_by_slide or {}).get(slide_index, set()):
                    score -= 3
            scored.append((score, pattern, reasons, risks))
        scored.sort(key=lambda item: (-item[0], item[1]["id"]))
        selected_score, selected, reasons, risks = scored[0]
        slide["pattern_id"] = selected["id"]
        slide["layout_index"] = selected["layout_index"]
        slide["master_index"] = selected["master_index"]
        slide["pattern_selection"] = {
            "selector": {"mode": "heuristic"},
            "score": round(selected_score, 2),
            "layout_pattern": selected.get("capacity", {}).get(
                "rhetorical_pattern", "content"
            ),
            "why_fit": "; ".join(reasons[:4]) or "лучший доступный компромисс",
            "risk": "; ".join(risks[:3]) or "низкий",
            "alternatives": [
                {"pattern_id": item[1]["id"], "score": round(item[0], 2)}
                for item in [s for s in scored if s[1]["id"] != selected["id"]][:3]
            ],
        }
        usage[selected["id"]] = usage.get(selected["id"], 0) + 1
    return plan
