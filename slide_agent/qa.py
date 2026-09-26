from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .decor_audit import decoration_overlap_issues
from .layout_geometry import (
    box,
    collides,
    has_text,
    hosts,
    is_framing,
    is_generated,
    is_opaque,
    is_photo_frame,
    is_visible_box,
)
from .style_audit import inspect_style
from .typography import (
    ESTIMATED_GLYPH_WIDTH_EM,
    effective_font_size,
    estimated_line_count,
)
from .utils import write_json


def _inches(value: int) -> float:
    return value / 914400


def _shape_font_size(shape: Any, default: float) -> float:
    sizes: list[float] = []
    if not getattr(shape, "has_text_frame", False):
        return default
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size:
                sizes.append(run.font.size.pt)
    return min(sizes) if sizes else float(effective_font_size(shape, default))


def _overflow_ratio(shape: Any, default_size: float) -> float:
    text = getattr(shape, "text", "").strip()
    if not text:
        return 0.0
    frame = shape.text_frame
    width = max(0.01, _inches(shape.width - frame.margin_left - frame.margin_right))
    height = max(0.01, _inches(shape.height - frame.margin_top - frame.margin_bottom))
    font_size = max(6.0, _shape_font_size(shape, default_size))
    chars_per_line = max(1, int(width * 72 / (font_size * ESTIMATED_GLYPH_WIDTH_EM)))
    lines = 0
    for logical_line in text.splitlines() or [text]:
        lines += estimated_line_count(logical_line, chars_per_line)
    required_height = lines * font_size * 1.22 / 72 + 0.08
    return required_height / height


def _inspected_shapes(shapes: Any) -> list[tuple[Any, bool]]:
    """Top-level shapes plus the children of generated (BrandDeck) groups.

    Template groups are left as authored; generated diagrams are groups, and
    their text must pass the same overflow and overlap checks as plain text.
    """
    result: list[tuple[Any, bool]] = []

    def visit(collection: Any, nested: bool) -> None:
        for shape in collection:
            result.append((shape, nested))
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP and str(shape.name).startswith(
                "BrandDeck"
            ):
                visit(shape.shapes, True)

    visit(shapes, False)
    return result


def _geometry(shape: Any) -> tuple[Any, ...]:
    return (shape.left, shape.top, shape.width, shape.height)


def _template_reference(
    template_path: str | Path | None,
) -> tuple[set[tuple[Any, ...]], set[tuple[Any, ...]], set[str]]:
    """Template shape geometry, shapes holding sample text, layout names.

    Filled template blocks are renamed by the composer but keep their
    geometry, so geometry identifies template-origin shapes in the output.
    """
    if not template_path or not Path(template_path).exists():
        return set(), set(), set()
    try:
        template = Presentation(template_path)
    except Exception:  # noqa: BLE001 - the template reference is optional
        return set(), set(), set()
    shapes = [shape for slide in template.slides for shape in slide.shapes]
    layouts = {
        layout.name for master in template.slide_masters for layout in master.slide_layouts
    }
    return (
        {_geometry(shape) for shape in shapes},
        {(str(shape.name), *_geometry(shape)) for shape in shapes if has_text(shape)},
        layouts,
    )


def _template_layer_issues(
    slide: Any,
    canvas: tuple[float, float],
    template_shapes: set[tuple[Any, ...]],
    sample_keys: set[tuple[Any, ...]],
) -> list[dict[str, Any]]:
    """Visible template shapes that collide with, or lost, the new content.

    Hosting cards, full-width bands and backgrounds are expected around new
    content. A smaller template shape is a defect when new content covers
    part of it, when its sample text was erased and it stayed as an empty
    block, or when it is an avatar circle without a photo.
    """
    shapes = list(slide.shapes)
    content = [
        (shape, box(shape))
        for shape in shapes
        if is_generated(shape)
        and (has_text(shape) or is_opaque(shape))
    ]
    issues: list[dict[str, Any]] = []
    for shape in shapes:
        if has_text(shape) or not is_visible_box(shape):
            continue
        zone = box(shape)
        if is_generated(shape):
            # A template text block renamed on filling, then left empty.
            if (
                shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
                and _geometry(shape) in template_shapes
                and getattr(shape, "has_text_frame", False)
                and not any(hosts(zone, other) for _, other in content)
            ):
                issues.append(
                    {
                        "severity": "warning",
                        "code": "emptied_template_block",
                        "shape": shape.name,
                        "bounds": [dict(zip("xywh", (round(v, 3) for v in zone)))],
                        "message": "Блок шаблона остался пустым после удаления образца текста",
                    }
                )
            continue
        sample = (str(shape.name), *_geometry(shape)) in sample_keys
        if is_framing(shape, zone, canvas, content, sample):
            continue
        bounds = dict(zip("xywh", (round(value, 3) for value in zone)))
        covered = [
            item.name
            for item, other in content
            if collides(
                shape,
                zone,
                item,
                other,
                item.is_placeholder or _geometry(item) in template_shapes,
            )
        ]
        if covered:
            issues.append(
                {
                    "severity": "warning",
                    "code": "template_overlap",
                    "shape": f"{shape.name} / {covered[0]}",
                    "bounds": [bounds],
                    "message": "Новый контент перекрывает элемент шаблона",
                }
            )
        elif shape.shape_type == MSO_SHAPE_TYPE.LINE:
            continue
        elif sample:
            issues.append(
                {
                    "severity": "warning",
                    "code": "emptied_template_block",
                    "shape": shape.name,
                    "bounds": [bounds],
                    "message": "Блок шаблона остался пустым после удаления образца текста",
                }
            )
        elif is_photo_frame(
            shape, zone, content, [box(other) for other in shapes if other is not shape]
        ):
            issues.append(
                {
                    "severity": "warning",
                    "code": "empty_photo_frame",
                    "shape": shape.name,
                    "bounds": [bounds],
                    "message": "Пустая рамка под фото или иконку рядом с подписью",
                }
            )
    return issues


def _picture_distortion(shape: Any) -> float:
    try:
        pixel_w, pixel_h = shape.image.size
    except (AttributeError, ValueError, KeyError, TypeError):
        return 0.0
    visible_w = pixel_w * (1 - (shape.crop_left or 0) - (shape.crop_right or 0))
    visible_h = pixel_h * (1 - (shape.crop_top or 0) - (shape.crop_bottom or 0))
    if min(visible_w, visible_h, shape.width or 0, shape.height or 0) <= 0:
        return 0.0
    return abs((visible_w / visible_h) / (shape.width / shape.height) - 1)


def inspect_presentation(
    output_path: str | Path,
    *,
    design_system: dict[str, Any] | None = None,
    expected_slide_count: int | None = None,
    report_path: str | Path | None = None,
    template_path: str | Path | None = None,
) -> dict[str, Any]:
    path = Path(output_path).resolve()
    issues: list[dict[str, Any]] = []
    if not path.exists():
        report = {
            "status": "failed",
            "score": 0,
            "output": str(path),
            "issues": [
                {
                    "severity": "error",
                    "code": "missing_file",
                    "message": "PPTX file does not exist",
                }
            ],
        }
        if report_path:
            write_json(Path(report_path), report)
        return report

    try:
        with zipfile.ZipFile(path) as archive:
            bad_member = archive.testzip()
            if bad_member:
                issues.append(
                    {
                        "severity": "error",
                        "code": "corrupt_zip",
                        "message": f"Corrupt package member: {bad_member}",
                    }
                )
    except zipfile.BadZipFile:
        issues.append(
            {
                "severity": "error",
                "code": "invalid_pptx",
                "message": "File is not a valid ZIP/OOXML package",
            }
        )

    try:
        prs = Presentation(path)
    except Exception as exc:  # noqa: BLE001 - QA must report every malformed-package failure
        issues.append({"severity": "error", "code": "open_failed", "message": str(exc)})
        report = {"status": "failed", "score": 0, "output": str(path), "issues": issues}
        if report_path:
            write_json(Path(report_path), report)
        return report

    slide_width = _inches(prs.slide_width)
    slide_height = _inches(prs.slide_height)
    default_body = float(
        (design_system or {}).get("typography", {}).get("body_size_pt", 18)
    )
    explicit_fonts: set[str] = set()
    slide_reports: list[dict[str, Any]] = []
    expects_native_grid = (design_system or {}).get("source_model", {}).get(
        "composition_mode"
    ) == "native_grid"
    template_shapes, sample_keys, layout_names = _template_reference(template_path)
    for slide_index, slide in enumerate(prs.slides, 1):
        slide_issues: list[dict[str, Any]] = []
        text_count = 0
        text_shapes: list[Any] = []
        for shape, nested in _inspected_shapes(slide.shapes):
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                distortion = _picture_distortion(shape)
                if distortion > 0.04:
                    generated_picture = str(shape.name).startswith("BrandDeck")
                    slide_issues.append(
                        {
                            "severity": "warning" if generated_picture else "info",
                            "code": "image_distorted",
                            "shape": shape.name,
                            "ratio": round(distortion, 3),
                            "message": "Picture aspect ratio is distorted by more than 4%",
                        }
                    )
            if nested:
                # Group bounds are checked once at top level; children only
                # contribute text checks.
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    text_count += 1
                    text_shapes.append(shape)
                    ratio = _overflow_ratio(shape, default_body)
                    if ratio > 1.45:
                        slide_issues.append(
                            {
                                "severity": "error" if ratio > 2.0 else "warning",
                                "code": "text_overflow_risk",
                                "shape": shape.name,
                                "ratio": round(ratio, 2),
                                "message": "Estimated text height exceeds the available box",
                            }
                        )
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if run.font.name:
                                explicit_fonts.add(run.font.name)
                continue
            if expects_native_grid and str(shape.name).lower().startswith("pdf "):
                slide_issues.append(
                    {
                        "severity": "error",
                        "code": "source_fragment_leak",
                        "shape": shape.name,
                        "message": "A PDF-conversion fragment leaked into a native slide",
                    }
                )
            x, y, w, h = map(
                _inches, (shape.left, shape.top, shape.width, shape.height)
            )
            selectable_background = (
                shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                and w >= slide_width * 0.8
                and h >= slide_height * 0.8
            )
            if expects_native_grid and (
                str(shape.name) == "BrandDeck Background Asset" or selectable_background
            ):
                slide_issues.append(
                    {
                        "severity": "error",
                        "code": "selectable_background",
                        "shape": shape.name,
                        "message": "A full-slide raster must be stored as a slide background",
                    }
                )
            if (
                x < -0.02
                or y < -0.02
                or x + w > slide_width + 0.02
                or y + h > slide_height + 0.02
            ):
                generated_shape = str(shape.name).startswith("BrandDeck")
                brand_bleed = shape.name in {
                    "BrandDeck Background Asset",
                    "BrandDeck Brand Asset",
                }
                slide_issues.append(
                    {
                        "severity": (
                            "info" if brand_bleed or not generated_shape else "warning"
                        ),
                        "code": (
                            "brand_asset_bleed"
                            if brand_bleed
                            else "out_of_bounds"
                            if generated_shape
                            else "template_bleed"
                        ),
                        "shape": shape.name,
                        "message": f"Shape bounds ({x:.2f}, {y:.2f}, {w:.2f}, {h:.2f}) exceed canvas",
                    }
                )
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                text_count += 1
                text_shapes.append(shape)
                ratio = _overflow_ratio(shape, default_body)
                if ratio > 1.45:
                    slide_issues.append(
                        {
                            "severity": "error" if ratio > 2.0 else "warning",
                            "code": "text_overflow_risk",
                            "shape": shape.name,
                            "ratio": round(ratio, 2),
                            "message": "Estimated text height exceeds the available box",
                        }
                    )
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.name:
                            explicit_fonts.add(run.font.name)
                        if (
                            expects_native_grid
                            and str(shape.name).startswith("BrandDeck Native")
                            and not any(
                                token in str(shape.name)
                                for token in (
                                    "Footer",
                                    "Badge",
                                    "Number",
                                    "Navigation Label",
                                )
                            )
                            and run.font.size
                            and run.font.size.pt < 14
                        ):
                            slide_issues.append(
                                {
                                    "severity": "warning",
                                    "code": "native_text_too_small",
                                    "shape": shape.name,
                                    "message": "Native slide text is smaller than 14 pt",
                                }
                            )
        for first_index, first in enumerate(text_shapes):
            for second in text_shapes[first_index + 1 :]:
                if not (
                    str(first.name).startswith("BrandDeck")
                    or str(second.name).startswith("BrandDeck")
                ):
                    continue
                left = max(first.left, second.left)
                top = max(first.top, second.top)
                right = min(first.left + first.width, second.left + second.width)
                bottom = min(first.top + first.height, second.top + second.height)
                if right <= left or bottom <= top:
                    continue
                intersection = (right - left) * (bottom - top)
                smaller_area = max(
                    1, min(first.width * first.height, second.width * second.height)
                )
                ratio = intersection / smaller_area
                if ratio > 0.15:
                    slide_issues.append(
                        {
                            "severity": "error" if ratio > 0.4 else "warning",
                            "code": "text_overlap",
                            "shape": f"{first.name} / {second.name}",
                            "ratio": round(ratio, 2),
                            "message": "Visible text boxes overlap",
                        }
                    )
        slide_issues.extend(
            _template_layer_issues(
                slide, (slide_width, slide_height), template_shapes, sample_keys
            )
        )
        slide_issues.extend(
            decoration_overlap_issues(
                slide,
                (slide_width, slide_height),
                design_system,
                default_body,
                template_shapes,
            )
        )
        slide_issues.extend(
            inspect_style(
                slide,
                (slide_width, slide_height),
                design_system,
                template_shapes,
                layout_names,
            )
        )
        if text_count == 0:
            slide_issues.append(
                {
                    "severity": "warning",
                    "code": "empty_slide",
                    "message": "No visible text detected",
                }
            )
        if expects_native_grid and not any(
            str(shape.name).startswith("BrandDeck Native") for shape in slide.shapes
        ):
            slide_issues.append(
                {
                    "severity": "error",
                    "code": "native_grid_missing",
                    "message": "Fragmented source was not rebuilt with native objects",
                }
            )
        for item in slide_issues:
            item["slide"] = slide_index
        issues.extend(slide_issues)
        slide_reports.append(
            {
                "slide": slide_index,
                "shape_count": len(slide.shapes),
                "text_shape_count": text_count,
            }
        )

    if expected_slide_count is not None and len(prs.slides) != expected_slide_count:
        issues.append(
            {
                "severity": "error",
                "code": "slide_count_mismatch",
                "message": f"Expected {expected_slide_count} slides, found {len(prs.slides)}",
            }
        )

    allowed_fonts = set((design_system or {}).get("typography", {}).get("families", []))
    primary = (design_system or {}).get("typography", {}).get("primary_font")
    if primary:
        allowed_fonts.add(primary)
    foreign_fonts = sorted(explicit_fonts - allowed_fonts) if allowed_fonts else []
    if foreign_fonts:
        issues.append(
            {
                "severity": "warning",
                "code": "foreign_fonts",
                "message": "Explicit fonts outside the extracted design system",
                "fonts": foreign_fonts,
            }
        )

    errors = sum(item["severity"] == "error" for item in issues)
    warnings = sum(item["severity"] == "warning" for item in issues)
    score = max(0, 100 - errors * 18 - warnings * 4)
    status = "failed" if errors else ("warning" if warnings else "passed")
    report = {
        "status": status,
        "score": score,
        "output": str(path),
        "file_size_bytes": path.stat().st_size,
        "slide_count": len(prs.slides),
        "canvas": {"width_inches": slide_width, "height_inches": slide_height},
        "explicit_fonts": sorted(explicit_fonts),
        "issues": issues,
        "slides": slide_reports,
    }
    if report_path:
        write_json(Path(report_path), report)
    return report


def compact_plan(plan: dict[str, Any]) -> dict[str, Any]:
    """Preserve content during remapping; shortening requires a semantic editor."""
    import copy

    return copy.deepcopy(plan)
