"""Audit the actual office render: text contrast (with repair) and slide fill."""

from __future__ import annotations

import hashlib
from collections import Counter
from pathlib import Path
from typing import Any

import pymupdf
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor

from .layout_geometry import box, has_text, is_generated, is_opaque, is_visible_box
from .qa import _inspected_shapes


def contrast_ratio(
    foreground: tuple[int, int, int], background: tuple[int, int, int]
) -> float:
    def luminance(color: tuple[int, int, int]) -> float:
        channels = [component / 255 for component in color]
        linear = [
            value / 12.92 if value <= 0.04045 else ((value + 0.055) / 1.055) ** 2.4
            for value in channels
        ]
        return sum(
            weight * channel
            for weight, channel in zip((0.2126, 0.7152, 0.0722), linear)
        )

    first, second = sorted((luminance(foreground), luminance(background)), reverse=True)
    return (first + 0.05) / (second + 0.05)


def _hex(color: tuple[int, int, int]) -> str:
    return f"#{color[0]:02X}{color[1]:02X}{color[2]:02X}"


def _background(
    image: Image.Image, bounds: tuple[float, float, float, float], scale: float
) -> tuple[tuple[int, int, int], float] | None:
    x0, y0, x1, y1 = bounds
    box = (
        max(0, int(x0 * scale)),
        max(0, int(y0 * scale)),
        min(image.width, int(x1 * scale + 1)),
        min(image.height, int(y1 * scale + 1)),
    )
    if box[2] <= box[0] or box[3] <= box[1]:
        return None
    crop = image.crop(box)
    crop.thumbnail((250, 90))
    pixels = crop.tobytes()
    counts = Counter(zip(pixels[0::3], pixels[1::3], pixels[2::3]))
    color, count = counts.most_common(1)[0]
    return color, count / sum(counts.values())


def inspect_rendered_contrast(pdf_path: Path, pptx_path: Path) -> dict[str, Any]:
    """Audit text spans against the dominant pixels behind them in the final PDF.

    Complex image backgrounds are left for visual review when no dominant color
    can be established. One issue is retained per rendered text line.
    """
    prs = Presentation(pptx_path)
    issues: list[dict[str, Any]] = []
    checked = 0
    uncertain = 0
    with pymupdf.open(pdf_path) as document:
        for slide_number, page in enumerate(document, 1):
            scale = 2.0
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
            image = Image.frombytes(
                "RGB", (pixmap.width, pixmap.height), pixmap.samples
            )
            width_scale = prs.slide_width / 914400 / page.rect.width
            height_scale = prs.slide_height / 914400 / page.rect.height
            for block_index, block in enumerate(page.get_text("dict")["blocks"]):
                for line_index, line in enumerate(block.get("lines", [])):
                    worst: dict[str, Any] | None = None
                    for span in line.get("spans", []):
                        value = span.get("text", "").strip()
                        if len(value) < 4 or not any(char.isalpha() for char in value):
                            continue
                        sample = _background(image, span["bbox"], scale)
                        if not sample or sample[1] < 0.30:
                            uncertain += 1
                            continue
                        checked += 1
                        color = int(span["color"])
                        foreground = (
                            (color >> 16) & 255,
                            (color >> 8) & 255,
                            color & 255,
                        )
                        background = sample[0]
                        ratio = contrast_ratio(foreground, background)
                        if ratio >= 4.5 or (worst and ratio >= worst["ratio"]):
                            continue
                        x0, y0, x1, y1 = span["bbox"]
                        worst = {
                            "slide": slide_number,
                            "code": "rendered_low_contrast",
                            "severity": "warning",
                            "check_type": "deterministic",
                            "ratio": round(ratio, 2),
                            "text": value[:100],
                            "foreground": _hex(foreground),
                            "background": _hex(background),
                            "bounds": [
                                {
                                    "x": round(x0 * width_scale, 4),
                                    "y": round(y0 * height_scale, 4),
                                    "w": round((x1 - x0) * width_scale, 4),
                                    "h": round((y1 - y0) * height_scale, 4),
                                }
                            ],
                            "message": f"Контраст текста {ratio:.2f}:1 ниже 4.5:1 после офисного рендера",
                            "repairable": False,
                        }
                    if worst:
                        identity = f"{slide_number}:{block_index}:{line_index}:{worst['text']}"
                        worst["id"] = hashlib.sha256(identity.encode()).hexdigest()[:16]
                        issues.append(worst)
    return {
        "status": "warning" if issues else "passed",
        "checked_spans": checked,
        "uncertain_spans": uncertain,
        "issues": issues,
    }


def inspect_rendered_fill(
    pdf_path: Path, pptx_path: Path, roles: list[str]
) -> dict[str, Any]:
    """Share of each content slide occupied by the new content (Appendix 1).

    Content is the rendered text lines inside generated blocks, visible cards
    carrying that text and generated pictures, diagrams, charts and tables;
    template decoration is not content. Below a quarter or above three
    quarters of the slide is reported. Cover, section and closing slides are
    sparse by design and are not measured.
    """
    prs = Presentation(pptx_path)
    width = prs.slide_width / 914400
    height = prs.slide_height / 914400
    scale = 20  # mask pixels per inch
    issues: list[dict[str, Any]] = []
    measured: list[dict[str, Any]] = []
    with pymupdf.open(pdf_path) as document:
        for number, (page, slide) in enumerate(zip(document, prs.slides), 1):
            role = roles[number - 1] if number <= len(roles) else "content"
            if role in {"cover", "section", "closing"}:
                continue
            sx = width / page.rect.width
            sy = height / page.rect.height
            lines = [
                (x0 * sx, y0 * sy, x1 * sx, y1 * sy)
                for block in page.get_text("dict")["blocks"]
                for line in block.get("lines", [])
                if "".join(span.get("text", "") for span in line["spans"]).strip()
                for x0, y0, x1, y1 in [line["bbox"]]
            ]
            boxes: list[tuple[float, float, float, float]] = []
            for shape in slide.shapes:
                if not is_generated(shape):
                    continue
                x, y, w, h = box(shape)
                if is_opaque(shape) or (
                    has_text(shape) and is_visible_box(shape)
                ):
                    boxes.append((x, y, x + w, y + h))
                elif has_text(shape):
                    inside = [
                        line
                        for line in lines
                        if x - 0.05 <= (line[0] + line[2]) / 2 <= x + w + 0.05
                        and y - 0.05 <= (line[1] + line[3]) / 2 <= y + h + 0.05
                    ]
                    if inside:
                        boxes.append(
                            (
                                min(line[0] for line in inside),
                                min(line[1] for line in inside),
                                max(line[2] for line in inside),
                                max(line[3] for line in inside),
                            )
                        )
            mask = Image.new("L", (round(width * scale), round(height * scale)), 0)
            draw = ImageDraw.Draw(mask)
            for x0, y0, x1, y1 in boxes:
                draw.rectangle(
                    (x0 * scale, y0 * scale, x1 * scale, y1 * scale), fill=255
                )
            coverage = mask.histogram()[255] / (mask.width * mask.height)
            measured.append({"slide": number, "coverage": round(coverage, 3)})
            if 0.25 <= coverage <= 0.75:
                continue
            code = "slide_underfilled" if coverage < 0.25 else "slide_overfilled"
            issues.append(
                {
                    "id": hashlib.sha256(f"{number}:{code}".encode()).hexdigest()[:16],
                    "slide": number,
                    "code": code,
                    "severity": "warning",
                    "check_type": "deterministic",
                    "coverage": round(coverage, 3),
                    "bounds": [],
                    "message": (
                        f"Контент занимает {coverage:.0%} слайда"
                        + (" (меньше четверти)" if coverage < 0.25 else " (больше трёх четвертей)")
                    ),
                    "repairable": True,
                }
            )
    return {
        "status": "warning" if issues else "passed",
        "slides": measured,
        "issues": issues,
    }


def _parse_colors(values: list[str] | None) -> list[tuple[int, int, int]]:
    colors = []
    for value in values or []:
        cleaned = str(value).lstrip("#")
        if len(cleaned) == 6:
            try:
                colors.append(tuple(int(cleaned[i : i + 2], 16) for i in (0, 2, 4)))
            except ValueError:
                continue
    return colors


def repair_rendered_contrast(
    pptx_path: Path,
    issues: list[dict[str, Any]],
    palette: list[str] | None = None,
    accents: list[str] | None = None,
) -> list[dict[str, Any]]:
    """Recolor only generated editable text whose box contains a flagged span.

    Text takes the nearest readable text color of the template (``palette``:
    dk1, lt1, dk2, lt2, plus black and white); accents are used only when no
    text color reads on that background, so body text does not turn into an
    accent color.
    """
    prs = Presentation(pptx_path)
    repaired: list[dict[str, Any]] = []
    seen: set[tuple[int, int]] = set()
    for issue in issues:
        number = issue["slide"]
        rect = issue["bounds"][0]
        cx, cy = rect["x"] + rect["w"] / 2, rect["y"] + rect["h"] / 2
        background = tuple(int(issue["background"][i : i + 2], 16) for i in (1, 3, 5))
        original = tuple(int(issue["foreground"][i : i + 2], 16) for i in (1, 3, 5))
        text_colors = [(0, 0, 0), (255, 255, 255), *_parse_colors(palette)]
        readable = [
            color for color in text_colors if contrast_ratio(color, background) >= 4.5
        ] or [
            color
            for color in _parse_colors(accents)
            if contrast_ratio(color, background) >= 4.5
        ]
        replacement = min(
            readable,
            key=lambda color: sum((a - b) ** 2 for a, b in zip(color, original)),
        )
        for shape, nested in _inspected_shapes(prs.slides[number - 1].shapes):
            if not str(shape.name).startswith("BrandDeck") or not getattr(
                shape, "has_text_frame", False
            ):
                continue
            x, y, w, h = (
                value / 914400
                for value in (shape.left, shape.top, shape.width, shape.height)
            )
            if not (x - 0.05 <= cx <= x + w + 0.05 and y - 0.05 <= cy <= y + h + 0.05):
                continue
            key = (number, shape.shape_id)
            if key in seen:
                continue
            seen.add(key)
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(*replacement)
            repaired.append(
                {
                    "slide": number,
                    "shape": shape.name,
                    "color": _hex(replacement),
                    "issue_id": issue["id"],
                }
            )
            break
    if repaired:
        prs.save(pptx_path)
    return repaired
