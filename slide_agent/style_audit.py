"""Template-conformance checks of Appendix 1 for generated content.

Deterministic checks on the saved PPTX: readable and template-scale type
sizes, colours derived from the template palette, generated blocks kept out
of the edge margins, labelled charts and slides built on template layouts.
Template shapes keep their authored values and are not judged here.
"""

from __future__ import annotations

from typing import Any

from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .layout_geometry import box, has_text, is_generated

EDGE_INCHES = 0.15
PALETTE_DISTANCE = 16.0
# Chrome drawn by the native-grid composer follows its own tokens and is
# covered by the native checks.
EXEMPT_NAMES = ("Native", "Brand Asset", "Background")
PIE_TYPES = {XL_CHART_TYPE.PIE, XL_CHART_TYPE.PIE_EXPLODED, XL_CHART_TYPE.DOUGHNUT}


def _hex_rgb(value: str) -> tuple[int, int, int]:
    value = value.lstrip("#")
    return tuple(int(value[index : index + 2], 16) for index in (0, 2, 4))


def _palette(design: dict[str, Any]) -> list[tuple[int, int, int]]:
    colors = design.get("colors", {})
    values = [item.get("hex") for item in colors.get("theme", [])]
    values += [item.get("hex") for item in colors.get("observed", [])]
    values += [value for value in design.get("brand", {}).values() if isinstance(value, str)]
    palette = []
    for value in values + ["000000", "FFFFFF"]:
        text = str(value or "").lstrip("#")
        if len(text) == 6:
            try:
                palette.append(_hex_rgb(text))
            except ValueError:
                continue
    return palette


def _derived(color: tuple[int, int, int], palette: list[tuple[int, int, int]]) -> bool:
    """A palette colour or a mix of two (tints, shades, accent over background)."""
    unique = list(dict.fromkeys(palette))
    for index, base in enumerate(unique):
        for pole in unique[index:]:
            for step in range(21):
                ratio = step / 20
                mixed = [a + (b - a) * ratio for a, b in zip(base, pole)]
                if sum((m - c) ** 2 for m, c in zip(mixed, color)) ** 0.5 <= PALETTE_DISTANCE:
                    return True
    return False


def _explicit_colors(shape: Any) -> list[str]:
    colors = []
    try:
        if shape.fill.type == MSO_FILL.SOLID and shape.fill.fore_color.type == MSO_COLOR_TYPE.RGB:
            colors.append(str(shape.fill.fore_color.rgb))
    except (AttributeError, TypeError, ValueError):
        pass
    try:
        line = shape.line
        if line.fill.type == MSO_FILL.SOLID and line.color.type == MSO_COLOR_TYPE.RGB:
            colors.append(str(line.color.rgb))
    except (AttributeError, TypeError, ValueError):
        pass
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                try:
                    if run.font.color.type == MSO_COLOR_TYPE.RGB:
                        colors.append(str(run.font.color.rgb))
                except (AttributeError, TypeError, ValueError):
                    continue
    return colors


def _generated_parts(shapes: Any) -> list[tuple[Any, bool]]:
    """Generated top-level shapes plus the children of generated groups."""
    parts: list[tuple[Any, bool]] = []
    for shape in shapes:
        if not is_generated(shape) or any(name in shape.name for name in EXEMPT_NAMES):
            continue
        parts.append((shape, False))
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            parts.extend((child, True) for child in shape.shapes)
    return parts


def _chart_labelled(chart: Any) -> bool:
    plot = chart.plots[0] if len(chart.plots) else None
    if plot is None:
        return False
    series = list(plot.series)
    if chart.chart_type in PIE_TYPES:
        return plot.has_data_labels or chart.has_legend
    if len(series) > 1 and not chart.has_legend:
        return False
    try:
        axis_title = chart.value_axis.has_title
    except (AttributeError, ValueError):
        axis_title = False
    return plot.has_data_labels or axis_title


def _chart_readability_issues(shape: Any) -> list[dict[str, Any]]:
    """Catch pie layouts whose legend squeezes the plot or duplicates labels."""
    chart = shape.chart
    if chart.chart_type not in PIE_TYPES or not chart.plots:
        return []
    plot = chart.plots[0]
    issues: list[dict[str, Any]] = []
    width = shape.width / 914400
    height = shape.height / 914400
    if width < 2.8 or height < 2.4:
        issues.append(
            {
                "severity": "warning",
                "code": "chart_too_small",
                "shape": shape.name,
                "message": "Круговая диаграмма слишком мала для читаемых секторов и подписей",
            }
        )
    if chart.has_legend and chart.legend.position == XL_LEGEND_POSITION.RIGHT:
        labels = [str(category.label) for category in plot.categories]
        longest = max((len(label) for label in labels), default=0)
        # Leave a usable pie and space for the longest category. A right
        # legend in a narrower frame produced one-letter lines in Office.
        required = 2.8 + min(longest, 24) * 0.11
        if width < required:
            issues.append(
                {
                    "severity": "warning",
                    "code": "chart_legend_too_narrow",
                    "shape": shape.name,
                    "message": "Легенда круговой диаграммы сжимает подписи и сектора",
                }
            )
    if plot.has_data_labels:
        labels = plot.data_labels
        if labels.show_value and labels.show_percentage:
            issues.append(
                {
                    "severity": "warning",
                    "code": "chart_redundant_labels",
                    "shape": shape.name,
                    "message": "Подписи круговой диаграммы смешивают значения и доли",
                }
            )
    return issues


def inspect_style(
    slide: Any,
    canvas: tuple[float, float],
    design: dict[str, Any] | None,
    template_shapes: set[tuple[Any, ...]],
    layout_names: set[str],
) -> list[dict[str, Any]]:
    design = design or {}
    width, height = canvas
    floor = max(7.0, 10.0 * width / 13.333)
    scale = sorted(
        float(value)
        for value in design.get("typography", {}).get("observed_sizes_pt", [])
        if 6 <= float(value) <= 96
    )
    palette = _palette(design)
    issues: list[dict[str, Any]] = []
    if layout_names and slide.slide_layout.name not in layout_names:
        issues.append(
            {
                "severity": "error",
                "code": "foreign_layout",
                "message": "Слайд собран не на макете шаблона",
            }
        )
    for shape, nested in _generated_parts(slide.shapes):
        sizes = []
        if has_text(shape) and getattr(shape, "has_text_frame", False):
            sizes = [
                run.font.size.pt
                for paragraph in shape.text_frame.paragraphs
                for run in paragraph.runs
                if run.font.size and run.text.strip()
            ]
        if sizes and min(sizes) < floor - 0.05:
            issues.append(
                {
                    "severity": "warning",
                    "code": "text_too_small",
                    "shape": shape.name,
                    "size_pt": round(min(sizes), 1),
                    "message": f"Кегль {min(sizes):.1f} pt меньше читаемого минимума {floor:.1f} pt",
                }
            )
        off = [size for size in sizes if scale and min(abs(size - s) for s in scale) > 0.6]
        if off:
            issues.append(
                {
                    "severity": "warning",
                    "code": "font_off_scale",
                    "shape": shape.name,
                    "size_pt": round(off[0], 1),
                    "message": f"Кегль {off[0]:.1f} pt не из типографической шкалы шаблона",
                }
            )
        foreign = [
            color
            for color in _explicit_colors(shape)
            if palette and not _derived(_hex_rgb(color), palette)
        ]
        if foreign:
            issues.append(
                {
                    "severity": "warning",
                    "code": "off_palette_color",
                    "shape": shape.name,
                    "color": f"#{foreign[0]}",
                    "message": f"Цвет #{foreign[0]} не выводится из палитры шаблона",
                }
            )
        if nested:
            continue
        if getattr(shape, "has_chart", False):
            if not _chart_labelled(shape.chart):
                issues.append(
                    {
                        "severity": "warning",
                        "code": "chart_unlabeled",
                        "shape": shape.name,
                        "message": "У диаграммы нет подписей значений, единиц или легенды",
                    }
                )
            issues.extend(_chart_readability_issues(shape))
        placed_by_composer = not getattr(shape, "is_placeholder", False) and (
            (shape.left, shape.top, shape.width, shape.height) not in template_shapes
        )
        visible = has_text(shape) or shape.shape_type in {
            MSO_SHAPE_TYPE.GROUP,
            MSO_SHAPE_TYPE.TABLE,
            MSO_SHAPE_TYPE.CHART,
        } or getattr(shape, "has_table", False) or getattr(shape, "has_chart", False)
        if placed_by_composer and visible:
            x, y, w, h = box(shape)
            if (
                x < EDGE_INCHES
                or y < EDGE_INCHES
                or x + w > width - EDGE_INCHES
                or y + h > height - EDGE_INCHES
            ) and not (x < -0.02 or y < -0.02 or x + w > width + 0.02 or y + h > height + 0.02):
                issues.append(
                    {
                        "severity": "warning",
                        "code": "margin_intrusion",
                        "shape": shape.name,
                        "message": "Сгенерированный блок заходит в поля у края слайда",
                    }
                )
    return issues
