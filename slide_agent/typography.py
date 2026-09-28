"""Resolve text size through OOXML placeholder inheritance without changing it."""

from __future__ import annotations

import textwrap
from typing import Any

from pptx.enum.shapes import PP_PLACEHOLDER

# Conservative fallback when exact font metrics are unavailable. Corporate
# Cyrillic fonts can be wider than the old 0.52-em Latin approximation.
ESTIMATED_GLYPH_WIDTH_EM = 0.6


def needs_cyrillic_font_fallback(font_name: str | None) -> bool:
    """East Asian template fonts can render Cyrillic with broken spacing."""
    name = str(font_name or "").casefold()
    return bool(
        any("\u3040" <= char <= "\u9fff" or "\uac00" <= char <= "\ud7af" for char in name)
        or any(
            token in name
            for token in ("yu gothic", "yu mincho", "meiryo", "malgun gothic", "noto sans cjk")
        )
    )


def readable_font(font_name: str | None, text: str) -> str | None:
    """Use a Cyrillic-capable sans font only when the template font is CJK."""
    if needs_cyrillic_font_fallback(font_name) and any(
        "\u0400" <= char <= "\u052f" for char in text
    ):
        return "Arial"
    return font_name


def estimated_line_count(text: str, characters_per_line: int) -> int:
    """Approximate wrapping at word boundaries, including explicit newlines."""
    return sum(
        max(
            1,
            len(
                textwrap.wrap(
                    line, width=max(1, characters_per_line), break_on_hyphens=False
                )
            ),
        )
        for line in str(text).split("\n")
    )


def effective_font_size(shape: Any, default: float | None = None) -> float | None:
    """Resolve first-paragraph size; mixed explicit runs use their maximum.

    Composition writes uniform paragraphs. This is not a full shaping engine.
    """
    if not getattr(shape, "has_text_frame", False):
        return default
    level = shape.text_frame.paragraphs[0].level + 1

    def local_size(element: Any) -> float | None:
        paths = [
            "./p:txBody/a:p[1]/a:r/a:rPr/@sz",
            "./p:txBody/a:p[1]/a:pPr/a:defRPr/@sz",
            f"./p:txBody/a:lstStyle/a:lvl{level}pPr/a:defRPr/@sz",
            "./p:txBody/a:lstStyle/a:defPPr/a:defRPr/@sz",
        ]
        for path in paths:
            values = element.xpath(path)
            if values:
                return max(float(value) / 100 for value in values)
        return None

    own = local_size(shape.element)
    if own is not None:
        return own
    slide = getattr(shape.part, "slide", None)
    if slide is None:
        return default
    kind = None
    if shape.is_placeholder:
        kind = shape.placeholder_format.type
        layout_shape = slide.slide_layout.placeholders.get(
            idx=shape.placeholder_format.idx
        )
        if layout_shape is not None:
            value = local_size(layout_shape.element)
            if value is not None:
                return value
        master_kind = (
            PP_PLACEHOLDER.TITLE if kind == PP_PLACEHOLDER.CENTER_TITLE else kind
        )
        for master_shape in slide.slide_layout.slide_master.placeholders:
            if master_shape.placeholder_format.type == master_kind:
                value = local_size(master_shape.element)
                if value is not None:
                    return value
    title_types = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
    body_types = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE}
    style = (
        "titleStyle"
        if kind in title_types
        else "bodyStyle"
        if kind in body_types
        else "otherStyle"
    )
    values = slide.slide_layout.slide_master.element.xpath(
        f"./p:txStyles/p:{style}/a:lvl{level}pPr/a:defRPr/@sz"
    )
    return float(values[0]) / 100 if values else default
