import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from slide_agent.composer import (
    _CYRILLIC_THEME_FALLBACK,
    _add_body_text,
    _add_pie_chart,
    _constrain_title_width,
    _fill_slide,
    _set_text_frame,
    _sparse_statement_cards,
    _whole_visual_zone,
)
from slide_agent.typography import (
    effective_font_size,
    estimated_line_count,
    readable_font,
)


def test_east_asian_template_font_is_replaced_for_cyrillic_text():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    explicit = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    explicit.text = "Образец"
    explicit.text_frame.paragraphs[0].runs[0].font.name = "游ゴシック"
    _set_text_frame(explicit, ["Русский текст"])
    assert explicit.text_frame.paragraphs[0].runs[0].font.name == "Arial"

    inherited = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
    token = _CYRILLIC_THEME_FALLBACK.set("Arial")
    try:
        _set_text_frame(inherited, ["Русский текст"])
    finally:
        _CYRILLIC_THEME_FALLBACK.reset(token)
    assert inherited.text_frame.paragraphs[0].runs[0].font.name == "Arial"
    assert readable_font("Play", "Русский текст") == "Play"
    assert readable_font("游ゴシック", "English only") == "游ゴシック"


def _design():
    return {
        "canvas": {"width_inches": 10, "height_inches": 7.5},
        "spacing": {},
        "typography": {"primary_font": "Arial", "body_size_pt": 18},
    }


def test_body_skips_caption_slot_and_preserves_every_paragraph():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[3])
    bodies = list(slide.placeholders)[1:]
    caption, body = bodies
    for shape in bodies:
        shape.width = Inches(4.3)
    caption.height = Inches(0.27)
    body.height = Inches(4)
    paragraphs = ["Подтвержденный факт из исходного документа." for _ in range(5)]
    geometry = [(s.left, s.top, s.width, s.height) for s in bodies]

    _add_body_text(slide, {"bullets": paragraphs}, bodies, (1, 1, 8, 5), _design())

    assert caption.text == ""
    assert body.text.splitlines() == paragraphs
    assert [(s.left, s.top, s.width, s.height) for s in bodies] == geometry


def test_body_uses_more_than_two_slots_without_reordering():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[3])
    bodies = list(slide.placeholders)[1:]
    # Clone a content placeholder to exercise a three-column template.
    from copy import deepcopy

    slide.shapes._spTree.insert_element_before(deepcopy(bodies[0].element), "p:extLst")
    bodies = list(slide.placeholders)[1:]
    for index, shape in enumerate(bodies):
        shape.left = Inches(0.5 + index * 3)
        shape.width = Inches(2.5)
        shape.height = Inches(0.8)
    paragraphs = [f"Факт номер {i}" for i in range(6)]

    _add_body_text(slide, {"bullets": paragraphs}, bodies, (1, 1, 8, 5), _design())

    assert all(shape.text for shape in bodies)
    assert [line for shape in bodies for line in shape.text.splitlines()] == paragraphs


def test_visual_slot_does_not_contain_body_text():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[3])
    bodies = list(slide.placeholders)[1:]
    for shape in bodies:
        shape.width = Inches(4)
    bodies[0].height = Inches(2)
    bodies[1].height = Inches(5)
    zone = _add_body_text(
        slide,
        {"bullets": ["Подтвержденный факт"], "visual": {"type": "bar_chart"}},
        bodies,
        (1, 1, 8, 5),
        _design(),
    )
    assert bodies[0].text == "Подтвержденный факт"
    assert bodies[1].text == ""
    assert zone == tuple(
        value / 914400
        for value in (bodies[1].left, bodies[1].top, bodies[1].width, bodies[1].height)
    )


def test_cover_reuses_body_placeholder_below_title():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title, body = list(slide.placeholders)
    title.top, title.height = Inches(3.2), Inches(1.82)
    title.width = Inches(9)
    body.top, body.height = Inches(5.32), Inches(0.61)
    body.width = Inches(5.48)
    original = (body.left, body.top, body.width, body.height)

    _fill_slide(
        slide,
        {
            "role": "cover",
            "title": "Пилот",
            "subtitle": "Описание проекта",
            "bullets": [],
            "body": "",
        },
        _design(),
        set(),
    )

    assert body.text == "Описание проекта"
    assert (body.left, body.top, body.width, body.height) == original
    assert not any(s.name == "BrandDeck Subtitle" for s in slide.shapes)


@pytest.mark.parametrize(
    "heading, caption, wrapped, latest_caption_bottom",
    [
        (
            "Ускорение работы сотрудников через ИИ",
            "Инициатива по внедрению ИИ-помощника во внутренний портал",
            False,
            2.5,
        ),
        (
            "Снижение времени поиска инструкций и нагрузки на поддержку",
            "Внедрение ИИ-помощника для внутренних регламентов",
            True,
            2.8,
        ),
    ],
)
def test_long_centered_cover_uses_wide_text_band_above_art(
    heading, caption, wrapped, latest_caption_bottom
):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title, subtitle = list(slide.placeholders)
    for shape, size, value in (
        (title, 48, "Короткий образец"),
        (subtitle, 16, "Подпись образца"),
    ):
        shape.left = Inches(2.71)
        shape.width = Inches(4.69)
        shape.text = value
        shape.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    title.top, title.height = Inches(1.5), Inches(0.87)
    subtitle.top, subtitle.height = Inches(2.41), Inches(0.31)
    design = {
        **_design(),
        "canvas": {"width_inches": 10, "height_inches": 5.625},
    }

    _fill_slide(
        slide,
        {
            "role": "cover",
            "title": heading,
            "subtitle": caption,
            "bullets": [],
            "body": "",
        },
        design,
        set(),
    )

    assert title.width >= Inches(7.5)
    assert title.top < Inches(1.5)
    assert subtitle.top >= title.top + title.height
    assert subtitle.top + subtitle.height < Inches(latest_caption_bottom)
    assert title.text_frame.paragraphs[0].runs[0].font.size.pt >= 23
    assert subtitle.text_frame.paragraphs[0].runs[0].font.size.pt >= 13
    assert ("\v" in title.text) is wrapped
    assert title.text.replace("\v", " ") == heading
    assert all(run.font.size is not None for run in title.text_frame.paragraphs[0].runs)


def test_replacing_exemplar_text_preserves_explicit_style():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    shape.text = "Исходный заголовок"
    run = shape.text_frame.paragraphs[0].runs[0]
    run.font.name = "VK Sans Display"
    run.font.size = Pt(32)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 211, 227)

    _set_text_frame(shape, ["Новый заголовок"])

    output_run = shape.text_frame.paragraphs[0].runs[0]
    assert output_run.font.name == "VK Sans Display"
    assert output_run.font.size.pt == 32
    assert output_run.font.bold is True
    assert output_run.font.color.rgb == RGBColor(0, 211, 227)


def test_inherited_layout_font_is_fitted_without_changing_template():
    prs = Presentation()
    layout = prs.slide_layouts[1]
    layout.placeholders[0].text_frame.paragraphs[0].font.size = Pt(48)
    slide = prs.slides.add_slide(layout)
    title = slide.shapes.title
    title.width, title.height = Inches(4.7), Inches(0.875)
    assert effective_font_size(title) == 48
    layout_xml = layout.element.xml
    text = "Пилот сервиса поддержки сотрудников"

    _set_text_frame(title, [text])

    assert title.text == text
    assert 9 <= title.text_frame.paragraphs[0].runs[0].font.size.pt < 48
    assert layout.element.xml == layout_xml


def test_generated_text_overrides_vertical_layout_writing_direction(tmp_path):
    prs = Presentation()
    layout = prs.slide_layouts[1]
    body_layout = layout.placeholders[1]
    body_layout.text_frame._txBody.bodyPr.set("vert", "eaVert")
    slide = prs.slides.add_slide(layout)
    body = slide.placeholders[1]

    _set_text_frame(body, ["Горизонтальный русский текст"])
    output = tmp_path / "vertical-template.pptx"
    prs.save(output)
    restored = Presentation(output)

    assert restored.slides[0].placeholders[1].text_frame._txBody.bodyPr.get("vert") == "horz"
    assert restored.slide_layouts[1].placeholders[1].text_frame._txBody.bodyPr.get("vert") == "eaVert"


def test_large_page_number_does_not_replace_real_exemplar_title():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    number = slide.shapes.add_textbox(Inches(8.1), Inches(0.3), Inches(1.2), Inches(1.4))
    number.text = "02"
    number.text_frame.paragraphs[0].runs[0].font.size = Pt(72)
    heading = slide.shapes.add_textbox(Inches(0.6), Inches(0.9), Inches(6), Inches(0.65))
    heading.text = "Настоящий заголовок"
    heading.text_frame.paragraphs[0].runs[0].font.size = Pt(30)
    body = slide.shapes.add_textbox(Inches(0.6), Inches(2.1), Inches(6), Inches(2.5))
    body.text = "Исходный абзац."

    _fill_slide(
        slide,
        {"role": "content", "title": "Новая тема", "body": "Подтверждённое содержание", "bullets": []},
        _design(),
        set(),
    )

    assert heading.text == "Новая тема"
    assert number.text == ""
    assert any("Подтверждённое содержание" in shape.text for shape in slide.shapes if shape.has_text_frame)


def test_pie_chart_expands_narrow_exemplar_slot_and_uses_readable_labels():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(8.8), Inches(0.8))
    title.name = "BrandDeck Title"
    title.text = "Доли участников"

    _add_pie_chart(
        slide,
        {"categories": ["Продажи", "Поддержка"], "values": [60, 90]},
        (3.0, 1.45, 3.0, 5.3),
        _design(),
    )

    shape = next(shape for shape in slide.shapes if shape.has_chart)
    chart = shape.chart
    labels = chart.plots[0].data_labels
    assert shape.name == "BrandDeck Pie Chart"
    assert shape.width >= Inches(6)
    assert chart._chartSpace.xpath(".//c:autoTitleDeleted")[0].get("val") == "1"
    assert labels.show_percentage is True
    assert labels.show_value is False
    assert labels.font.size.pt == 12


def test_explicit_font_overrides_layout_size():
    prs = Presentation()
    layout = prs.slide_layouts[1]
    layout.placeholders[0].text_frame.paragraphs[0].font.size = Pt(48)
    title = prs.slides.add_slide(layout).shapes.title
    title.text = "Заголовок"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    assert effective_font_size(title) == 24


def test_word_wrapping_reserves_full_lines_for_long_words():
    assert estimated_line_count("aaaaaa bbbbbb cccccc", 10) == 3
    assert estimated_line_count("aaa\n\nbbb", 10) == 3


def test_title_avoids_inherited_artwork_without_moving_it(tmp_path):
    from copy import deepcopy

    from PIL import Image

    prs = Presentation()
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    title = slide.shapes.title
    title.left, title.top = Inches(0.7), Inches(3.2)
    title.width, title.height = Inches(8), Inches(1.8)
    picture_path = tmp_path / "art.png"
    Image.new("RGB", (20, 20), "blue").save(picture_path)
    art = slide.shapes.add_picture(
        str(picture_path), Inches(6.6), Inches(0.7), Inches(3), Inches(6)
    )
    layout.shapes._spTree.insert_element_before(deepcopy(art.element), "p:extLst")
    slide.shapes._spTree.remove(art.element)
    before = layout.element.xml
    _constrain_title_width(slide, title)
    assert title.left + title.width < Inches(6.6)
    assert title.top == Inches(3.2)
    assert title.height == Inches(1.8)
    assert layout.element.xml == before


def test_fitting_accounts_for_word_boundaries():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2.1), Inches(1.4))
    text = "Ответственный подтверждает актуальность документации"
    _set_text_frame(shape, [text], font_size=30)
    assert shape.text == text
    assert shape.text_frame.paragraphs[0].runs[0].font.size.pt < 30


def test_exemplar_body_zone_does_not_extend_above_reserved_title():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    source = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(7), Inches(4))
    source.text = "Исходное содержание"
    _add_body_text(
        slide, {"bullets": ["Новый факт"]}, [source], (1, 2, 7, 4), _design()
    )
    generated = next(s for s in slide.shapes if s.name == "BrandDeck Body")
    assert generated.top >= Inches(2)
    assert "Новый факт" in generated.text


def test_every_paragraph_keeps_the_exemplar_bullet_and_indent():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(3))
    properties = shape.text_frame.paragraphs[0]._p.get_or_add_pPr()
    properties.set("marL", "255588")
    properties.set("indent", "-255588")
    properties.append(properties.makeelement(qn("a:buChar"), {"char": "•"}))
    shape.text_frame.paragraphs[0].text = "Образец"

    _set_text_frame(shape, ["Первый тезис.", "Второй тезис.", "Третий тезис."])

    for paragraph in shape.text_frame.paragraphs:
        pPr = paragraph._p.pPr
        assert pPr.find(qn("a:buChar")) is not None
        assert (pPr.get("marL"), pPr.get("indent")) == ("255588", "-255588")


def test_few_short_statements_in_a_large_empty_zone_become_cards():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    _fill_slide(
        slide,
        {
            "role": "content",
            "title": "Границы пилота",
            "subtitle": "",
            "body": "",
            "bullets": ["Два подразделения.", "Шесть недель.", "Ответы и передача."],
        },
        _design(),
        set(),
    )
    grids = [shape for shape in slide.shapes if shape.name == "BrandDeck Diagram icon_grid"]
    assert len(grids) == 1
    texts = [child.text for child in grids[0].shapes if child.has_text_frame]
    assert {"Два подразделения.", "Шесть недель.", "Ответы и передача."} <= set(texts)
    assert not any(
        shape.is_placeholder and shape.text.strip() == "Шесть недель." for shape in slide.shapes
    )


def test_sparse_cards_follow_variant_layout_without_changing_claims():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    claims = ["Два подразделения.", "Шесть недель.", "Ответы по регламенту.", "Передача специалисту."]
    _fill_slide(
        slide,
        {
            "role": "content",
            "title": "Границы пилота",
            "bullets": claims,
            "layout_variant": "columns",
        },
        _design(),
        set(),
    )

    grid = next(shape for shape in slide.shapes if shape.name == "BrandDeck Diagram icon_grid")
    cards = [shape for shape in grid.shapes if shape.name.startswith("BrandDeck Diagram Card")]
    assert len(cards) == len(claims)
    assert len({shape.left for shape in cards}) == 1
    assert [
        shape.text for shape in grid.shapes if shape.name.startswith("BrandDeck Diagram Text")
    ] == claims


def test_cards_are_not_forced_on_template_cards_or_dense_text():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2), Inches(8), Inches(4)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0xEE, 0xF2, 0xFF)
    short = {"role": "content", "bullets": ["Первый факт.", "Второй факт."]}
    dense = {"role": "content", "bullets": ["Слово " * 20] * 6}

    assert _sparse_statement_cards(short, [], [card], _design()) is None
    assert _sparse_statement_cards(dense, [], [], _design()) is None
    assert _sparse_statement_cards({**short, "role": "closing"}, [], [], _design()) is None
    cards = _sparse_statement_cards(short, [], [], _design())
    assert cards["type"] == "icon_grid" and cards["origin"] == "sparse_text"
    assert cards["estimated_fill"] < 0.28


def test_rendered_fill_hint_turns_statements_into_cards():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2), Inches(8), Inches(4)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0xEE, 0xF2, 0xFF)
    dense = {"role": "content", "bullets": ["Слово " * 20] * 3, "layout_hint": "cards"}
    cards = _sparse_statement_cards(dense, [], [card], _design())
    assert cards["type"] == "icon_grid" and cards["origin"] == "rendered_fill"


def test_short_strip_of_exemplar_labels_extends_to_the_content_area():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    left = slide.shapes.add_textbox(Inches(0.3), Inches(1.9), Inches(4.4), Inches(1.8))
    right = slide.shapes.add_textbox(Inches(4.8), Inches(1.9), Inches(5.0), Inches(1.8))
    # Typical margins of a template with right-side art leave a narrow zone.
    zone = _whole_visual_zone(slide, [left, right], (1.1, 1.5, 5.0, 5.5), _design())
    assert zone == pytest.approx((0.3, 1.9, 9.5, 5.1))


def test_exemplar_visual_zone_stops_before_reserved_footer():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    left = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(5), Inches(5.2))
    right = slide.shapes.add_textbox(Inches(5.6), Inches(2), Inches(5), Inches(5.2))
    content = (0.48, 1.64, 12.37, 5.41)
    design = {**_design(), "canvas": {"width_inches": 13.333, "height_inches": 7.5}}

    result = _whole_visual_zone(slide, [left, right], content, design)

    assert result[1] + result[3] == pytest.approx(6.97)


@pytest.mark.parametrize(
    "heading",
    ["Поиск инструкции занимает до 30 минут", "Сотрудники теряют до 30 минут на поиск инструкций"],
)
def test_long_cover_stack_fits_the_overflow_estimate_at_template_sizes(heading):
    from slide_agent.composer import _TYPE_SCALE
    from slide_agent.qa import _overflow_ratio

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title, subtitle = list(slide.placeholders)
    for shape, size, value in ((title, 48, "Образец"), (subtitle, 16, "Подпись")):
        shape.left, shape.width = Inches(2.71), Inches(4.69)
        shape.text = value
        shape.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    title.top, title.height = Inches(1.5), Inches(0.87)
    subtitle.top, subtitle.height = Inches(2.41), Inches(0.31)
    caption = "Инициатива по внедрению ИИ-помощника для ускорения работы с регламентами"
    # A template scale without the fitted sizes must not round text up.
    token = _TYPE_SCALE.set([48.0, 24.4, 24.0, 18.0, 16.0, 13.2, 12.2, 12.0, 11.0])
    try:
        _fill_slide(
            slide,
            {"role": "cover", "title": heading, "subtitle": caption,
             "bullets": [], "body": ""},
            {**_design(), "canvas": {"width_inches": 10, "height_inches": 5.625}},
            set(),
        )
    finally:
        _TYPE_SCALE.reset(token)
    assert subtitle.text == caption
    assert title.text.replace("\v", " ") == heading
    assert _overflow_ratio(title, 24) <= 1.0
    assert _overflow_ratio(subtitle, 16) <= 1.0
    assert title.text_frame.paragraphs[0].runs[0].font.size.pt >= 18
    assert subtitle.text_frame.paragraphs[0].runs[0].font.size.pt >= 11
