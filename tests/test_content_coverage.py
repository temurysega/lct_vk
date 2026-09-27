from copy import deepcopy

import pytest
from pptx import Presentation
from pptx.util import Inches

from slide_agent.coverage import (
    coverage_report,
    grounding_findings,
    unsupported_numbers,
)
from slide_agent.planner import _fallback_plan, _sentences, normalize_plan
from slide_agent.qa import compact_plan


def test_last_sections_and_first_section_body_survive_small_slide_count():
    source = "# Проект\nВводный факт 40%.\n" + "\n".join(
        f"## Раздел {i}\nПодтвержденный факт {i}." for i in range(12)
    )
    plan = normalize_plan(_fallback_plan(source, 5), source_text=source, slide_count=5)
    assert len(plan["slides"]) == 5
    report = coverage_report(source, plan)
    assert report["missing_from_plan"] == []
    assert any("Раздел 11" in slide["title"] for slide in plan["slides"])


def test_normalization_does_not_cut_long_text_or_sixth_bullet():
    text = "Подтвержденный факт " * 100
    plan = {
        "slides": [
            {"title": "Начало"},
            {
                "title": "Данные",
                "body": text,
                "bullets": [f"Факт {i}" for i in range(8)],
            },
            {"title": "Конец"},
        ]
    }
    normalized = normalize_plan(plan, source_text=text, slide_count=3)
    assert normalized["slides"][1]["body"] == text.strip()
    assert len(normalized["slides"][1]["bullets"]) == 8


def test_retry_preserves_chart_arrays_and_every_fact():
    plan = {
        "slides": [
            {
                "body": "Факт " * 200,
                "bullets": list(map(str, range(9))),
                "visual": {
                    "type": "bar_chart",
                    "categories": list(map(str, range(10))),
                    "series": [{"values": list(range(10))}],
                },
            }
        ]
    }
    original = deepcopy(plan)
    retry = compact_plan(plan)
    assert retry == original
    retry["slides"][0]["visual"]["categories"].clear()
    assert plan == original


def test_pptx_coverage_detects_missing_fact_even_if_it_is_in_plan(tmp_path):
    source = "# Проект\nБюджет 40 млн рублей.\nСрок 4 недели."
    plan = _fallback_plan(source, 3)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(8), Inches(2)
    ).text = "Проект\nБюджет 40 млн рублей."
    output = tmp_path / "deck.pptx"
    prs.save(output)
    report = coverage_report(source, plan, output)
    assert not report["missing_from_plan"]
    assert report["missing_from_pptx"]
    assert report["status"] == "needs_review"


def test_paraphrase_and_speaker_notes_are_not_claimed_as_verified():
    source = "Срок 4 недели."
    report = coverage_report(
        source, {"slides": [{"body": "Завершим за месяц.", "speaker_notes": source}]}
    )
    assert report["status"] == "needs_review"
    assert report["semantic_verification"] == "not_performed"


def test_image_match_score_is_not_treated_as_slide_claim():
    plan = {
        "slides": [
            {
                "title": "Сотрудники",
                "visual": {
                    "type": "image",
                    "asset_id": "abc123",
                    "match": "lexical",
                    "match_score": 4.54,
                },
            }
        ]
    }
    assert unsupported_numbers("Сотрудники используют помощника.", plan) == []


def test_grounding_rules_find_budget_guarantee_and_reversed_success_metric():
    source = (
        "Документы не уходят наружу. Успех: время ответа и доля обращений, "
        "закрытых без специалиста. Просим одобрить пилот и команду."
    )
    plan = {"slides": [
        {"title": "Снижение доли обращений, закрытых без специалиста"},
        {"title": "Риск утечки исключён"},
        {"title": "Просим утвердить бюджет пилота"},
    ]}
    findings = grounding_findings(source, plan)
    assert [(item["slide"], item["quote"]) for item in findings] == [
        (1, "Снижение доли обращений, закрытых без специалиста"),
        (2, "Риск утечки исключён"),
        (3, "Просим утвердить бюджет пилота"),
    ]
    assert coverage_report(source, plan)["grounding_findings"] == findings


def test_grounding_rules_keep_supported_and_neutral_wording():
    source = (
        "Документы не уходят наружу. Измеряем долю обращений, закрытых "
        "без специалиста. Просим утвердить бюджет пилота."
    )
    plan = {"slides": [
        {"title": "Рост доли обращений, закрытых без специалиста"},
        {"title": "Измеряем долю обращений, закрытых без специалиста"},
        {"title": "Документы не уходят наружу"},
        {"title": "Просим утвердить бюджет пилота"},
    ]}
    assert grounding_findings(source, plan) == []


def test_too_little_content_does_not_create_empty_slides():
    with pytest.raises(ValueError, match="fewer slides"):
        _fallback_plan("Один факт.", 10)


def test_semicolon_splits_only_before_a_new_statement():
    assert _sentences("Вопрос передаётся специалисту; ссылку можно проверить.") == [
        "Вопрос передаётся специалисту; ссылку можно проверить."
    ]
    assert _sentences("Первый пункт; Второй пункт") == ["Первый пункт", "Второй пункт"]


def test_offline_cover_has_no_service_caption():
    source = "# Пилот\n\n## Проблема\nПоиск занимает время.\n\n## Решение\nПомощник отвечает."
    cover = _fallback_plan(source, 4)["slides"][0]
    assert cover["role"] == "cover" and cover["subtitle"] == ""
