import ast

import pytest

from training.curation import is_page_marker, repeated_marginal_text, select_content
from training.data_utils import check_split
from training.prepare_curated_ru import build_notebook, filtered_copy
from training.render_quality import audit_page, contrast_ratio, render_pdf


def element(text, size, top, **extra):
    return {
        "left": 1,
        "top": top,
        "paragraphs": [{"text": text, "font": {"size_pt": size}}],
        **extra,
    }


def test_page_number_never_wins_over_title_and_glyph_is_not_body():
    slide = {
        "text_elements": [
            element("P · 02", 9, 0.1),
            element("«", 80, 2),
            element("Настоящий заголовок", 36, 1),
            element("Полезный текст", 20, 3),
            element("02 / 14", 9, 7),
        ]
    }
    title, bullets = select_content(slide, 7.5, set())
    assert title == "Настоящий заголовок"
    assert bullets == ["Полезный текст"]


@pytest.mark.parametrize("text", ["P · 02", "02 / 14", "Стр. 3", "Слайд 2 из 9", "‹#›"])
def test_page_markers(text):
    assert is_page_marker(text)
    assert not is_page_marker("Выручка выросла на 20%")


def test_repeated_footer_removed_but_nonrepeated_body_retained():
    context = {
        "presentation": {"slide_height_inches": 7.5},
        "slides": [
            {
                "text_elements": [
                    element("Мой университет", 12, 7),
                    element(f"Раздел {i}", 32, 1),
                ]
            }
            for i in range(4)
        ],
    }
    repeated = repeated_marginal_text(context)
    assert repeated == {"мой университет"}
    assert select_content(context["slides"][0], 7.5, repeated) == ("Раздел 0", [])


def test_explicit_title_placeholder_beats_large_body_quote():
    slide = {
        "text_elements": [
            element("Подзаголовок", 40, 2, placeholder_type="SUBTITLE"),
            element("Заголовок", 24, 1, placeholder_type="TITLE"),
        ]
    }
    assert select_content(slide, 7.5, set())[0] == "Заголовок"


def test_pdf_filter_detects_low_contrast_and_overlap():
    import pymupdf as fitz

    doc = fitz.open()
    good = doc.new_page(width=720, height=400)
    good.insert_text((30, 60), "Readable heading", fontsize=28, color=(0, 0, 0))
    assert audit_page(good)["status"] == "pass"
    bad = doc.new_page(width=720, height=400)
    bad.draw_rect(bad.rect, color=(1, 0.5, 0.5), fill=(1, 0.5, 0.5))
    bad.insert_text((30, 60), "Low contrast", fontsize=28, color=(0.8, 0.3, 0.3))
    assert audit_page(bad)["issues"]["low_contrast"]
    overlap = doc.new_page(width=720, height=400)
    overlap.insert_text((30, 60), "Large heading", fontsize=28)
    overlap.insert_text((40, 63), "Overlapping subtitle", fontsize=22)
    assert audit_page(overlap)["issues"]["text_overlap"]
    assert contrast_ratio((0, 0, 0), (255, 255, 255)) == pytest.approx(21)


def test_filtered_pptx_preserves_original_and_hidden_slide_mapping(tmp_path):
    import pymupdf as fitz
    from pptx import Presentation
    from pptx.util import Inches

    from slide_agent.exporter import find_libreoffice

    if not find_libreoffice():
        pytest.skip("LibreOffice unavailable")
    prs = Presentation()
    for i in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(5), Inches(1)
        ).text = f"Slide {i + 1}"
    prs.slides[1]._element.set("show", "0")
    source, cleaned = tmp_path / "original.pptx", tmp_path / "clean.pptx"
    prs.save(source)
    before = source.read_bytes()
    filtered_copy(source, cleaned, [2, 3])
    assert source.read_bytes() == before
    render_pdf(cleaned, tmp_path / "render.pdf")
    with fitz.open(tmp_path / "render.pdf") as pdf:
        assert len(pdf) == 2
        assert "Slide 2" in pdf[0].get_text()
        assert "Slide 3" in pdf[1].get_text()


def test_family_leakage_is_checked_against_final_test():
    def row(group, family):
        return {
            "provenance": {
                "group": group,
                "template_sha256": group,
                "design_family": family,
            }
        }

    check_split([row("a", "brand-a")], [row("b", "brand-b")], [row("c", "brand-c")])
    with pytest.raises(ValueError, match="design_family"):
        check_split([row("a", "brand-a")], [row("b", "brand-b")], [row("c", "brand-a")])


def test_ru_notebook_uses_new_folder_and_keeps_test_after_export():
    document = build_notebook("lct_cpu_ru_clean")
    code = "\n".join(
        "".join(c["source"]) for c in document["cells"] if c["cell_type"] == "code"
    )
    ast.parse(code)
    assert "/content/drive/MyDrive/lct_cpu_ru_clean" in code
    assert "test_report.json" in code
    assert "lct-ru-clean-work" in code
    assert "RESUME = 'auto'" in code


def test_schematic_preview_keeps_vector_asset_when_decoder_is_unavailable(tmp_path):
    from PIL import Image

    from slide_agent.renderer import _draw_image

    asset = tmp_path / "diagram.x-emf"
    asset.write_bytes(b"vector-format-without-pillow-decoder")
    original = asset.read_bytes()
    canvas = Image.new("RGB", (100, 100), "white")
    _draw_image(
        canvas,
        {"path": str(asset), "left": 1, "top": 1, "width": 3, "height": 2},
        10,
        tmp_path,
    )
    assert asset.read_bytes() == original
    assert canvas.getpixel((20, 20)) != (255, 255, 255)


def test_extracted_asset_names_are_portable_and_distinct():
    from types import SimpleNamespace

    from shared.scripts.extract_template import _image_asset_name

    shape = SimpleNamespace(
        name="Что будет в случае сбоя? «Сеть» / устройство: 1", shape_id=5
    )
    name = _image_asset_name("slide25", shape, "png")
    assert len(name) < 50
    assert not any(c in name for c in '<>:"/\\|?*')
    assert name != _image_asset_name(
        "slide25", SimpleNamespace(name=shape.name, shape_id=6), "png"
    )
