"""Template decoration that new content lies across or hides."""

import copy
import io

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from slide_agent.composer import _clear_layout_art
from slide_agent.decor_audit import (
    _is_fullscreen_backdrop,
    decoration_overlap_issues,
    layout_art_boxes,
)

CANVAS = (13.333, 7.5)
DESIGN = {
    "canvas": {"width_inches": CANVAS[0], "height_inches": CANVAS[1]},
    "brand": {"background": "FFFFFF"},
}


def _deck() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(CANVAS[0]), Inches(CANVAS[1])
    return prs


def _rect(shapes, name, box, color="1E6BFF"):
    shape = shapes.add_shape(MSO_SHAPE.RECTANGLE, *(Inches(value) for value in box))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    return shape


def _text(slide, box, text, size=18):
    shape = slide.shapes.add_textbox(*(Inches(value) for value in box))
    shape.name = "BrandDeck Body"
    shape.text_frame.text = text
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)
    return shape


def _codes(slide, template_shapes=frozenset()):
    return [
        issue["message"]
        for issue in decoration_overlap_issues(
            slide, CANVAS, DESIGN, 16, template_shapes
        )
    ]


def _to_layout(layout, shape):
    """Move a shape drawn on a scratch slide into the layout."""
    layout.shapes._spTree.append(copy.deepcopy(shape.element))
    shape.element.getparent().remove(shape.element)


def test_a_band_across_text_is_reported_and_text_beside_or_on_it_is_not():
    prs = _deck()
    lines = "Первая строка\nВторая строка\nТретья строка"
    for top, expected in ((4.7, 1), (2.0, 0), (5.6, 0)):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _rect(slide.shapes, "Google Shape;679", (0, 5.2, CANVAS[0], 2.3))
        _text(slide, (1, top, 6, 1.5), lines)
        assert len(_codes(slide)) == expected, top


def test_text_across_a_transparent_ring_picture_is_reported():
    ring = Image.new("RGBA", (400, 360), (0, 0, 0, 0))
    ImageDraw.Draw(ring).ellipse((0, 0, 399, 359), outline=(40, 90, 255, 255), width=58)
    data = io.BytesIO()
    ring.save(data, format="PNG")
    prs = _deck()
    for left, expected in ((3.0, 1), (0.3, 0)):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        picture = slide.shapes.add_picture(
            io.BytesIO(data.getvalue()),
            Inches(5),
            Inches(0),
            Inches(8.333),
            Inches(7.5),
        )
        picture.name = "Google Shape;727"
        _text(slide, (left, 3.5, 6, 0.6), "Время поиска до 30 минут на инструкцию")
        assert len(_codes(slide)) == expected, left


def test_only_fullscreen_pictures_behind_layout_text_are_backdrops():
    image = Image.new("RGBA", (1200, 675), (0, 0, 0, 0))
    ImageDraw.Draw(image).rectangle((0, 250, 1199, 400), fill=(240, 0, 0, 255))
    payload = io.BytesIO()
    image.save(payload, format="PNG")
    for behind in (True, False):
        prs = _deck()
        layout = prs.slide_layouts[1]
        scratch = prs.slides.add_slide(layout)
        picture = scratch.shapes.add_picture(
            io.BytesIO(payload.getvalue()), 0, 0, prs.slide_width, prs.slide_height
        )
        element = copy.deepcopy(picture.element)
        if behind:
            layout.shapes._spTree.insert(2, element)
        else:
            layout.shapes._spTree.insert_element_before(element, "p:extLst")
        slide = prs.slides.add_slide(layout)
        layout_picture = next(
            shape for shape in layout.shapes if shape.shape_type == picture.shape_type
        )
        assert _is_fullscreen_backdrop(layout_picture, slide, CANVAS) is behind


def test_a_hidden_layout_logo_is_reported_and_the_visual_zone_steps_aside():
    prs = _deck()
    layout = prs.slide_layouts[6]
    scratch = prs.slides.add_slide(layout)
    _to_layout(layout, _rect(scratch.shapes, "Logo", (0.5, 6.8, 1.5, 0.4)))
    slide = prs.slides.add_slide(layout)
    card = _rect(slide.shapes, "BrandDeck Diagram Card", (0.4, 4, 8, 3.2), "F2F2F2")
    assert any("логотип" in message for message in _codes(slide))
    # A block where the template sample had one is the template's own composition.
    geometry = (card.left, card.top, card.width, card.height)
    assert _codes(slide, {geometry}) == []

    art = layout_art_boxes(slide, DESIGN)
    assert len(art) == 1
    x, y, w, h = _clear_layout_art((0.4, 4, 8, 3.2), art)
    assert (x, y, w) == (0.4, 4, 8) and y + h <= 6.8 - 0.05
    # Stepping aside would cost half of a small zone: it stays and QA reports it.
    assert _clear_layout_art((0.4, 6.0, 3, 1.4), art) == (0.4, 6.0, 3, 1.4)


def test_layout_art_masked_by_a_background_shape_is_not_decoration():
    prs = _deck()
    layout = prs.slide_layouts[6]
    scratch = prs.slides.add_slide(layout)
    _to_layout(layout, _rect(scratch.shapes, "Dots", (0.5, 6.8, 1.5, 0.4)))
    _to_layout(layout, _rect(scratch.shapes, "Mask", (0, 6.5, 3, 1), "FFFFFF"))
    slide = prs.slides.add_slide(layout)
    _rect(slide.shapes, "BrandDeck Diagram Card", (0.4, 4, 8, 3.2), "F2F2F2")
    assert layout_art_boxes(slide, DESIGN) == []
    assert _codes(slide) == []


def test_a_diagram_cut_by_a_band_is_judged_as_one_drawing():
    prs = _deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide.shapes, "Google Shape;679", (0, 5.19, CANVAS[0], 2.31))
    group = slide.shapes.add_group_shape()
    group.name = "BrandDeck Diagram pyramid"
    for index, (top, height) in enumerate(((2.6, 1.1), (3.8, 1.1), (5.1, 1.2)), 1):
        _rect(
            group.shapes,
            f"BrandDeck Diagram Level {index}",
            (1, top, 3, height),
            "0B3D91",
        )
    # The lowest level alone lies almost wholly on the band; the pyramid does not.
    assert [
        issue["shape"] for issue in decoration_overlap_issues(slide, CANVAS, DESIGN, 16)
    ] == ["Google Shape;679 / BrandDeck Diagram pyramid"]
