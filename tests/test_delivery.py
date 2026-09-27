from __future__ import annotations

import copy
import hashlib
import re
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from examples.create_demo_assets import create_template
from slide_agent.analyzer import analyze_template
from slide_agent.api import app
from slide_agent.audit import enrich_audit
from slide_agent.cli import main
from slide_agent.exporter import export_presentation, find_libreoffice
from slide_agent.jobs import create_job, get_job, run_generation_job
from slide_agent.qa import inspect_presentation
from slide_agent.service import generate_deck, generate_variants, repair_presentation
from slide_agent.utils import read_json, write_json

CONTENT = """# Product
## Problem
- Reports take 10 hours.
- Three systems need checking.
## Solution
- One dashboard collects the reports.
- Two teams participate in the pilot.
## Plan
- Connect the systems.
- Measure the results.
"""


@pytest.fixture
def template(tmp_path):
    path = tmp_path / "brand.pptx"
    create_template(path, "clean-blue")
    return path


def test_variants_share_content_and_preserve_source(template, tmp_path):
    before = hashlib.sha256(template.read_bytes()).hexdigest()
    result = generate_variants(
        template=template,
        content=CONTENT,
        workspace=tmp_path / "w",
        slide_count=5,
        offline=True,
        export_formats=(),
    )
    assert len(result["variants"]) == 3
    content = []
    for deck in result["variants"]:
        plan = read_json(Path(deck["presentation_dir"]) / "deck_plan.final.json")
        content.append(
            [
                {
                    k: v
                    for k, v in s.items()
                    if k
                    not in {
                        "pattern_id",
                        "layout_index",
                        "master_index",
                        "pattern_selection",
                    }
                }
                for s in plan["slides"]
            ]
        )
        assert len(Presentation(deck["output"]).slides) == 5
        assert deck["workflow"]["sha256"]["planner"]
    assert content[0] == content[1] == content[2]
    assert hashlib.sha256(template.read_bytes()).hexdigest() == before


def test_selective_repair_preserves_parent_and_unselected_slides(template, tmp_path, monkeypatch):
    workspace = tmp_path / "w"
    deck = generate_deck(
        template=template,
        content=CONTENT,
        workspace=workspace,
        slide_count=5,
        offline=True,
    )
    directory = Path(deck["presentation_dir"])
    original = Path(deck["output"]).read_bytes()
    plan = read_json(directory / "deck_plan.final.json")
    qa = read_json(directory / "qa_report.json")
    qa["issues"].append(
        {
            "id": "selected",
            "slide": 2,
            "repairable": True,
            "code": "text_overflow_risk",
            "severity": "warning",
        }
    )
    write_json(directory / "qa_report.json", qa)
    fixed = repair_presentation(
        deck["presentation_id"], ["selected"], workspace=workspace
    )
    updated = read_json(Path(fixed["presentation_dir"]) / "deck_plan.final.json")
    assert Path(deck["output"]).read_bytes() == original
    assert fixed["presentation_id"] != deck["presentation_id"]
    assert updated["slides"][1]["pattern_id"] != plan["slides"][1]["pattern_id"]
    for n in (0, 2, 3, 4):
        assert plan["slides"][n] == updated["slides"][n]
        assert (
            Presentation(deck["output"]).slides[n].element.xml
            == Presentation(fixed["output"]).slides[n].element.xml
        )
    with pytest.raises(ValueError):
        repair_presentation(deck["presentation_id"], ["unknown"], workspace=workspace)
    with pytest.raises(ValueError):
        repair_presentation("../outside", ["selected"], workspace=workspace)

    # A failed parent export must not silently turn into a skipped export in
    # the revision. A rebuilt deck must also report when the chosen defect
    # survives, even if a different issue ID was assigned to it.
    original_manifest = read_json(directory / "manifest.json")
    original_manifest["requested_export_formats"] = ["pdf", "html"]
    original_manifest["exports"] = {"status": "failed", "artifacts": {}}
    write_json(directory / "manifest.json", original_manifest)
    captured = {}

    def fake_generate_from_plan(**kwargs):
        captured["formats"] = kwargs["export_formats"]
        target = tmp_path / "fake-repair"
        target.mkdir(exist_ok=True)
        return {
            "presentation_dir": str(target),
            "revision": kwargs["revision"],
            "qa": {"issues": [{"id": "new-id", "slide": 2, "code": "text_overflow_risk"}]},
        }

    monkeypatch.setattr("slide_agent.service._generate_from_plan", fake_generate_from_plan)
    unresolved = repair_presentation(
        deck["presentation_id"], ["selected"], workspace=workspace
    )
    assert captured["formats"] == ("pdf", "html")
    assert unresolved["revision"]["repair_result"]["status"] == "unresolved"
    assert unresolved["revision"]["repair_result"]["unresolved_issue_ids"] == ["selected"]
    assert read_json(tmp_path / "fake-repair" / "manifest.json") == unresolved
    original_manifest.pop("requested_export_formats")
    write_json(directory / "manifest.json", original_manifest)
    repair_presentation(deck["presentation_id"], ["selected"], workspace=workspace)
    assert captured["formats"] == ("pdf", "html")


def test_failed_qa_does_not_complete_job(monkeypatch, tmp_path):
    record = create_job(
        template_id="template", slide_count=5, offline=True, workspace=tmp_path
    )
    monkeypatch.setattr(
        "slide_agent.jobs.generate_deck",
        lambda **kwargs: {
            "status": "failed",
            "presentation_id": "bad",
            "output": "bad.pptx",
            "qa": {"status": "failed", "score": 0},
        },
    )
    run_generation_job(
        record["job_id"],
        template_id="template",
        content=CONTENT,
        slide_count=5,
        offline=True,
        workspace=tmp_path,
    )
    assert get_job(record["job_id"], tmp_path)["status"] == "failed"
    monkeypatch.setattr(
        "slide_agent.cli.run_pipeline", lambda **kwargs: {"status": "failed"}
    )
    assert main(["run", "--template", "x", "--content", "x", "--offline"]) == 1


def test_export_reports_missing_renderer(template, tmp_path, monkeypatch):
    monkeypatch.setattr("slide_agent.exporter.find_libreoffice", lambda: None)
    report = export_presentation(template, tmp_path / "export")
    assert report["status"] == "failed"
    assert not report["artifacts"]


@pytest.mark.skipif(
    not find_libreoffice(),
    reason="LibreOffice not found; install it or set BRANDDECK_LIBREOFFICE to soffice",
)
def test_actual_office_export(template, tmp_path):
    import fitz

    original = template.read_bytes()
    count = len(Presentation(template).slides)
    report = export_presentation(
        template, tmp_path / "export", expected_slide_count=count
    )
    assert report["status"] == "passed", report["issues"]
    assert len(report["previews"]) == count
    with fitz.open(report["artifacts"]["pdf"]) as pdf:
        assert len(pdf) == count
    assert "<svg" in Path(report["artifacts"]["html"]).read_text(encoding="utf-8")
    assert template.read_bytes() == original


def test_qa_annotations_have_deterministic_ids_and_bounds(template):
    prs = Presentation(template)
    shape = prs.slides[0].shapes.title
    shape.text = "TODO"
    prs.save(template)
    plan = {"slides": [{"bullets": []} for _ in prs.slides]}
    base = inspect_presentation(template)
    first = enrich_audit(template, copy.deepcopy(base), plan)
    second = enrich_audit(template, copy.deepcopy(base), plan)
    assert first == second
    issue = next(i for i in first["issues"] if i["code"] == "placeholder_text")
    assert issue["repairable"] and issue["bounds"] and issue["id"]


def test_duplicate_layout_names_resolve_by_part(template, tmp_path):
    prs = Presentation(template)
    prs.slide_layouts[0].name = "Duplicate"
    prs.slide_layouts[1].name = "Duplicate"
    prs.slides.add_slide(prs.slide_layouts[1])
    prs.save(template)
    directory = analyze_template(template, workspace=tmp_path / "w")
    context = read_json(directory / "context.json")
    assert context["slides"][-1]["layout_index"] == 1


def test_theme_text_color_survives_replacement():
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.util import Inches

    from slide_agent.composer import _set_text_frame

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
    shape.text = "Source"
    shape.text_frame.paragraphs[0].runs[
        0
    ].font.color.theme_color = MSO_THEME_COLOR.LIGHT_1
    _set_text_frame(shape, ["Replacement", "Second paragraph"])
    for paragraph in shape.text_frame.paragraphs:
        assert paragraph.runs[0].font.color.theme_color == MSO_THEME_COLOR.LIGHT_1


def test_large_template_context_is_bounded():
    import json

    from slide_agent.model_context import compact_design, compact_patterns

    catalog = {
        "patterns": [
            {
                "id": f"p{i}",
                "roles": ["content"],
                "capacity": {"body_chars": 500},
                "placeholders": ["x" * 10000],
            }
            for i in range(100)
        ]
    }
    compact = compact_patterns(catalog)
    assert len(compact) == 48
    assert compact[0]["id"] == "p0" and compact[-1]["id"] == "p99"
    assert len(json.dumps(compact)) < 16000
    design = compact_design(
        {
            "assets": ["x" * 100000],
            "grid": {"samples": [1] * 10000},
            "canvas": {"width_inches": 10},
        }
    )
    assert len(json.dumps(design)) < 1000


def test_web_without_build_keeps_api_available(tmp_path, monkeypatch):
    import backend.main as backend_main

    monkeypatch.setattr(backend_main, "WEB_ROOT", tmp_path / "missing-dist")
    monkeypatch.setenv("BRANDDECK_WORKSPACE", str(tmp_path / "workspace"))
    with TestClient(app) as client:
        for route in ("/", "/studio"):
            response = client.get(route)
            assert response.status_code == 503
            assert "npm run build" in response.json()["detail"]
        assert client.get("/health").status_code == 200


def test_web_and_variant_job(template, tmp_path, monkeypatch):
    monkeypatch.setenv("BRANDDECK_WORKSPACE", str(tmp_path / "w"))
    with TestClient(app) as client:
        homepage = client.get("/")
        assert homepage.status_code == 200
        assert '<div id="root"></div>' in homepage.text
        entry = re.search(r'src="(/assets/[^" ]+\.js)"', homepage.text)
        assert entry is not None, "Build React first: npm run build --prefix frontend"
        assert client.get(entry.group(1)).status_code == 200
        assert client.get("/studio").text == homepage.text
        assert client.get("/assets/src/main.tsx").status_code == 404
        assert (
            client.post(
                "/api/auth/register",
                json={
                    "username": "designer",
                    "password": "test-password-123",
                    "position": "Дизайнер",
                },
            ).status_code
            == 201
        )
        with template.open("rb") as file:
            response = client.post(
                "/v1/templates/analyze",
                files={"file": ("brand.pptx", file)},
                data={"offline": "true"},
            )
        assert response.status_code == 200
        response = client.post(
            "/v1/presentations/jobs",
            data={
                "template_id": response.json()["template_id"],
                "content": CONTENT,
                "slide_count": 5,
                "offline": "true",
                "variants": "true",
            },
        )
        assert response.status_code == 202
        job = client.get(f"/v1/jobs/{response.json()['job_id']}").json()
        assert len(job["batch"]["variants"]) == 3
        for deck in job["batch"]["variants"]:
            pid = deck["presentation_id"]
            assert client.get(f"/v1/presentations/{pid}").status_code == 200
            assert client.get(f"/v1/presentations/{pid}/download").content.startswith(
                b"PK"
            )
            assert (
                client.get(f"/v1/presentations/{pid}/download?format=exe").status_code
                == 422
            )
            assert client.get(f"/v1/presentations/{pid}/preview/0").status_code == 404
