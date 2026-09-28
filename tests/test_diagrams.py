from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from slide_agent.coverage import coverage_report
from slide_agent.diagrams import (
    DIAGRAM_TYPES,
    contrast,
    diagram_style,
    normalize_item,
    render_diagram,
    render_illustration,
    tint_series,
)
from slide_agent.qa import inspect_presentation
from slide_agent.visual_planning import sanitize_diagram

DESIGN = {
    "colors": {
        "theme": [
            {"role": "dk1", "hex": "000000"},
            {"role": "lt1", "hex": "FAFCFF"},
            {"role": "accent1", "hex": "0077FF"},
            {"role": "accent2", "hex": "FFFFFF"},
            {"role": "accent3", "hex": "FF3885"},
        ]
    },
    "brand": {"content_background": "FAFCFF", "content_body": "000000"},
    "typography": {
        "primary_font": "Arial",
        "body_size_pt": 16,
        "observed_sizes_pt": [40, 24, 18, 16, 14, 12],
    },
}
STEPS = [
    "В первую неделю команда собирает инструкции.",
    "Во вторую неделю проверяет поиск.",
    "В третью неделю проводит пилот.",
    "В четвертую неделю анализирует обратную связь.",
]
VISUALS = {
    "process": {"type": "process", "items": STEPS},
    "cycle": {
        "type": "cycle",
        "items": ["Планирование", "Разработка", "Проверка качества", "Выпуск"],
    },
    "hierarchy": {
        "type": "hierarchy",
        "root": "Сервис поддержки",
        "items": [
            {"label": "Форма", "detail": "Единая точка входа"},
            {"label": "Поиск", "detail": "База знаний"},
        ],
    },
    "pyramid": {
        "type": "pyramid",
        "items": ["Стратегия: цели компании", "Тактика: планы", "Операции: задачи"],
    },
    "funnel": {
        "type": "funnel",
        "items": [
            {"label": "Обращения", "value": "1200"},
            {"label": "Автоответ", "value": "800"},
        ],
    },
    "matrix": {
        "type": "matrix",
        "axes": {"x": "Сложность", "y": "Ценность"},
        "items": [
            "Сильные: форма",
            "Слабые: инструкции",
            "Возможности: масштаб",
            "Угрозы: доверие",
        ],
    },
    "icon_grid": {
        "type": "icon_grid",
        "items": [
            "Сотруднику нужен понятный ответ со ссылкой на инструкцию.",
            "Специалисту поддержки нужна история обращения.",
            "Руководителю нужна сводка результатов пилота.",
        ],
    },
}


def _deck(tmp_path: Path, width: float, background: str = "FAFCFF"):
    prs = Presentation()
    prs.slide_width = Inches(width)
    prs.slide_height = Inches(width * 9 / 16)
    style = diagram_style(DESIGN, background)
    records = []
    for kind in DIAGRAM_TYPES:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(background)
        height = width * 9 / 16
        zone = (0.5, 1.2, width - 1.0, height - 1.7)
        records.append(
            render_diagram(slide, VISUALS[kind], zone, style, context="Пилот")
        )
    output = tmp_path / f"diagrams-{width}-{background}.pptx"
    prs.save(output)
    return output, records


@pytest.mark.parametrize("width", [13.333, 10.0])
@pytest.mark.parametrize("background", ["FAFCFF", "000000", "0077FF"])
def test_every_diagram_is_native_and_passes_layout_qa(
    tmp_path: Path, width, background
):
    output, records = _deck(tmp_path, width, background)
    report = inspect_presentation(output, expected_slide_count=len(DIAGRAM_TYPES))
    blocking = [
        issue
        for issue in report["issues"]
        if issue["code"]
        in {"text_overflow_risk", "text_overlap", "out_of_bounds", "empty_slide"}
    ]
    assert blocking == []
    prs = Presentation(output)
    for slide, record, kind in zip(prs.slides, records, DIAGRAM_TYPES):
        groups = [
            shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.GROUP
        ]
        assert [group.name for group in groups] == [f"BrandDeck Diagram {kind}"]
        children = list(groups[0].shapes)
        assert children and not any(
            child.shape_type == MSO_SHAPE_TYPE.PICTURE for child in children
        )
        assert record["type"] == kind and record["native"]
        # Group extents enclose the children (python-pptx recalculates them).
        assert groups[0].width > 0 and groups[0].height > 0


def test_diagram_text_uses_template_type_scale(tmp_path: Path):
    output, _ = _deck(tmp_path, 13.333)
    sizes = set()
    for slide in Presentation(output).slides:
        for group in slide.shapes:
            for shape in group.shapes:
                if shape.has_text_frame:
                    sizes.update(
                        run.font.size.pt
                        for paragraph in shape.text_frame.paragraphs
                        for run in paragraph.runs
                        if run.font.size
                    )
    assert sizes and sizes <= {40, 24, 18, 16, 14, 12}


def test_icon_grid_records_matched_pictograms(tmp_path: Path):
    _, records = _deck(tmp_path, 13.333)
    grid = next(record for record in records if record["type"] == "icon_grid")
    assert grid["icon_matches"] == 3
    assert len(set(grid["icons"])) == 3


def test_independent_quantities_do_not_become_pyramid_levels_or_cycle():
    items = [
        {"label": "2 подразделения", "detail": "Охват пилота"},
        {"label": "150 сотрудников", "detail": "Участники пилота"},
        {"label": "6 недель", "detail": "Срок пилота"},
    ]
    for kind in ("pyramid", "cycle"):
        visual = sanitize_diagram({"type": kind, "items": items})
        assert visual == {"type": "icon_grid", "items": items}
    assert sanitize_diagram({"type": "pyramid", "items": items[:2]}) == {
        "type": "icon_grid",
        "items": items[:2],
    }
    detail_values = [
        {"label": "Срок", "detail": "6 недель"},
        {"label": "Сотрудники", "detail": "150 человек"},
        {"label": "Подразделения", "detail": "2 отдела"},
    ]
    assert sanitize_diagram({"type": "pyramid", "items": detail_values}) == {
        "type": "icon_grid",
        "items": detail_values,
    }
    levels = [{"label": f"{number} уровень"} for number in (1, 2, 3)]
    assert sanitize_diagram({"type": "pyramid", "items": levels})["type"] == "pyramid"
    ordered_stages = [
        {"label": "2 недели", "detail": "Проектирование"},
        {"label": "3 месяца", "detail": "Разработка"},
        {"label": "4 дня", "detail": "Запуск"},
    ]
    assert sanitize_diagram({"type": "process", "items": ordered_stages})["type"] == "process"
    ambiguous = [
        {"label": "2 отдела", "detail": "150 сотрудников"},
        {"label": "6 недель", "detail": "Срок пилота"},
        {"label": "3 человека", "detail": "Команда"},
    ]
    assert sanitize_diagram({"type": "cycle", "items": ambiguous})["type"] == "cycle"


def test_narrow_metric_cards_fill_column_and_pyramid_apex_has_number():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    zone = (5.08, 0.84, 4.22, 3.84)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    record = render_diagram(
        slide,
        {
            "type": "icon_grid",
            "items": ["2 подразделения", "150 сотрудников", "6 недель"],
            "layout_variant": "balanced",
        },
        zone,
        {**diagram_style(DESIGN), "canvas_area": 56.25},
    )
    assert record["columns"] == 1 and record["narrow_stack"] is True
    cards = [
        shape
        for shape in slide.shapes[0].shapes
        if shape.name.startswith("BrandDeck Diagram Card ")
    ]
    assert len(cards) == 3
    assert len({card.left for card in cards}) == 1
    assert sum(card.width * card.height for card in cards) / (prs.slide_width * prs.slide_height) > 0.25

    pyramid_slide = prs.slides.add_slide(prs.slide_layouts[6])
    render_diagram(
        pyramid_slide,
        {"type": "pyramid", "items": ["Стратегия", "Тактика", "Операции"]},
        zone,
        diagram_style(DESIGN),
    )
    apex = next(
        shape
        for shape in pyramid_slide.shapes[0].shapes
        if shape.name == "BrandDeck Diagram Level 1"
    )
    assert apex.text == "01"


@pytest.mark.parametrize("zone", [(0.47, 1.94, 11.24, 4.14), (5.08, 0.84, 4.22, 3.84)])
def test_two_item_focus_grid_differs_from_balanced(zone):
    items = [
        {"label": "Поиск ответа", "detail": "До 30 минут"},
        {"label": "Поддержка", "detail": "Повторяющиеся вопросы"},
    ]
    positions = {}
    for variant in ("balanced", "focus"):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        record = render_diagram(
            slide,
            {"type": "icon_grid", "items": items, "layout_variant": variant},
            zone,
            {**diagram_style(DESIGN), "canvas_area": 100},
        )
        cards = [
            shape
            for shape in slide.shapes[0].shapes
            if shape.name.startswith("BrandDeck Diagram Card ")
        ]
        assert len(cards) == 2
        positions[variant] = [(card.left, card.top, card.width, card.height) for card in cards]
        if variant == "focus":
            assert record["emphasis"] == "first"
        if zone[2] >= 7.5:
            assert all(card.height <= Inches(3.01) for card in cards)
            assert all(card.top > Inches(zone[1]) for card in cards)
    assert positions["balanced"] != positions["focus"]


def test_icon_grid_variants_keep_claims_and_change_card_geometry(tmp_path: Path):
    claims = [
        "Сотрудник получает ответ со ссылкой на инструкцию",
        "Специалист поддержки видит историю обращения",
        "Руководитель получает сводку результатов пилота",
        "Команда обновляет регламент после проверки",
    ]
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    positions = {}
    for variant in ("balanced", "columns", "focus"):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        record = render_diagram(
            slide,
            {"type": "icon_grid", "items": claims, "layout_variant": variant},
            (0.5, 1.7, 12.3, 5.1),
            diagram_style(DESIGN),
        )
        assert record["layout_variant"] == variant
        group = slide.shapes[0]
        cards = {
            shape.name: shape
            for shape in group.shapes
            if shape.name.startswith("BrandDeck Diagram Card ")
        }
        texts = {
            shape.name: shape.text
            for shape in group.shapes
            if shape.name.startswith("BrandDeck Diagram Text ")
        }
        assert [texts[f"BrandDeck Diagram Text {i}"] for i in range(1, 5)] == claims
        positions[variant] = [
            (
                cards[f"BrandDeck Diagram Card {i}"].left / 914400,
                cards[f"BrandDeck Diagram Card {i}"].top / 914400,
                cards[f"BrandDeck Diagram Card {i}"].width / 914400,
                cards[f"BrandDeck Diagram Card {i}"].height / 914400,
            )
            for i in range(1, 5)
        ]

    balanced, columns, focus = (positions[name] for name in positions)
    assert balanced[0][1] == balanced[1][1]
    assert balanced[2][1] > balanced[0][1]
    assert len({round(card[0], 3) for card in columns}) == 1
    assert [card[1] for card in columns] == sorted(card[1] for card in columns)
    assert focus[0][3] > focus[1][3] * 2
    assert focus[1][0] > focus[0][0] + focus[0][2]
    output = tmp_path / "icon-grid-variants.pptx"
    prs.save(output)
    assert not any(
        issue["code"] in {"text_overflow_risk", "text_overlap", "out_of_bounds"}
        for issue in inspect_presentation(output, expected_slide_count=3)["issues"]
    )


def test_focus_grid_limits_card_coverage_on_compact_canvas():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    style = {**diagram_style(DESIGN), "canvas_area": 10.833333 * 7.5}
    record = render_diagram(
        slide,
        {"type": "icon_grid", "items": ["Первое", "Второе", "Третье", "Четвёртое"],
         "layout_variant": "focus"},
        (0.25, 0.85, 10.34, 5.82),
        style,
    )
    group = slide.shapes[0]
    assert record["layout_variant"] == "focus"
    assert group.width / 914400 * group.height / 914400 <= style["canvas_area"] * 0.701


@pytest.mark.parametrize(
    "background", ["FFFFFF", "FAFCFF", "000000", "0077FF", "EBF3F9", "7F56D9"]
)
def test_diagram_colours_keep_text_contrast(background):
    style = diagram_style(DESIGN, background)
    assert contrast(style["accent_fill"], style["accent_text"]) >= 4.5
    assert contrast(style["surface"], style["surface_text"]) >= 4.5
    assert contrast(style["background"], style["text"]) >= 4.5
    for fill, text in tint_series(style, 5):
        assert contrast(fill, text) >= 4.5


def test_inline_labels_preserve_the_exact_source_sentence(tmp_path: Path):
    item = normalize_item("Безопасность: шифрование и контроль доступа")
    assert (
        item["inline"]
        and item["head"] + item["tail"] == "Безопасность: шифрование и контроль доступа"
    )
    output, _ = _deck(tmp_path, 13.333)
    source = "# Пилот\n## План работ\n" + "\n".join(f"- {step}" for step in STEPS)
    plan = {
        "slides": [
            {"title": "Пилот"},
            {"title": "План работ", "visual": VISUALS["process"]},
        ]
    }
    report = coverage_report(source, plan, output)
    step_rows = [row for row in report["units"] if row["text"] in STEPS]
    assert len(step_rows) == 4 and all(row["in_pptx"] for row in step_rows)


def test_illustration_is_native_and_topic_aware(tmp_path: Path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    record = render_illustration(
        slide, (5, 1, 4, 4), diagram_style(DESIGN), text="Безопасность данных"
    )
    assert record == {
        "type": "illustration",
        "icon": "shield-check",
        "keyword_match": True,
        "native": True,
    }
    group = slide.shapes[0]
    assert group.shape_type == MSO_SHAPE_TYPE.GROUP
    assert not any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in group.shapes)


def test_too_few_items_is_rejected():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    with pytest.raises(ValueError):
        render_diagram(
            slide,
            {"type": "cycle", "items": ["A", "B"]},
            (0, 0, 8, 4),
            diagram_style(DESIGN),
        )
