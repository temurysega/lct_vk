"""Exemplar leftovers: the composer removes them and the audit reports them."""

import io
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from slide_agent.composer import (
    _add_table,
    _contain_generated_text,
    _expand_small_image_zone,
    _remove_exemplar_leftovers,
    _remove_unmatched_content_images,
)
from slide_agent.diagrams import contrast
from slide_agent.exporter import export_presentation, find_libreoffice
from slide_agent.qa import inspect_presentation
from slide_agent.render_audit import inspect_rendered_fill

DESIGN = {"canvas": {"width_inches": 13.333, "height_inches": 7.5}}


def _rect(slide, name, box, kind=MSO_SHAPE.ROUNDED_RECTANGLE, text=""):
    shape = slide.shapes.add_shape(kind, *(Inches(value) for value in box))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0x20, 0x20, 0x20)
    if text:
        shape.text_frame.text = text
    return shape


def _text(slide, name, box, text):
    shape = slide.shapes.add_textbox(*(Inches(value) for value in box))
    shape.name = name
    shape.text_frame.text = text
    return shape


def _deck() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def test_photo_slide_removes_repeated_raster_rule_strips():
    prs = _deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    stripe = io.BytesIO()
    Image.new("RGB", (100, 2), "#B5C1D3").save(stripe, format="PNG")
    rule = slide.shapes.add_picture(
        io.BytesIO(stripe.getvalue()), Inches(4), Inches(2), Inches(4), Inches(0.008)
    )
    slide.shapes.add_picture(
        io.BytesIO(stripe.getvalue()), Inches(4), Inches(3), Inches(4), Inches(0.008)
    )
    photo = io.BytesIO()
    Image.new("RGB", (600, 400), "#345678").save(photo, format="PNG")
    large = slide.shapes.add_picture(
        io.BytesIO(photo.getvalue()), Inches(5), Inches(1.5), Inches(4), Inches(3)
    )
    slot = _remove_unmatched_content_images(
        slide, {}, DESIGN, {rule.image.sha1}, keep_slot=True
    )
    assert slot.element is large.element
    assert [shape.element for shape in slide.shapes] == [large.element]


def test_filled_picture_placeholder_does_not_leak_source_qr():
    prs = _deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    image = io.BytesIO()
    Image.new("RGB", (100, 100), "white").save(image, format="PNG")
    picture = slide.shapes.add_picture(
        io.BytesIO(image.getvalue()), Inches(0.7), Inches(0.7), Inches(1.2), Inches(1.2)
    )
    placeholder = OxmlElement("p:ph")
    placeholder.set("idx", "2")
    placeholder.set("type", "pic")
    picture._element.xpath("./p:nvPicPr/p:nvPr")[0].append(placeholder)
    picture = slide.shapes[0]
    assert picture.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER
    assert picture.element.tag == qn("p:pic")

    _remove_unmatched_content_images(slide, {}, DESIGN, set())
    assert len(slide.shapes) == 0


def test_repeated_exemplar_fragments_inside_content_are_removed():
    prs = _deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bitmap = io.BytesIO()
    Image.new("RGB", (20, 20), "#287AC5").save(bitmap, format="PNG")
    content = slide.shapes.add_picture(
        io.BytesIO(bitmap.getvalue()), Inches(3), Inches(3), Inches(0.08), Inches(0.08)
    )
    logo = slide.shapes.add_picture(
        io.BytesIO(bitmap.getvalue()), Inches(0.1), Inches(6.8), Inches(0.4), Inches(0.4)
    )
    _remove_unmatched_content_images(slide, {}, DESIGN, {content.image.sha1})
    assert [shape.element for shape in slide.shapes] == [logo.element]


def test_dark_brand_table_keeps_text_readable():
    prs = _deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    design = {
        "typography": {"primary_font": "Arial", "body_size_pt": 16},
        "brand": {"background": "000000", "heading": "0077FF", "body": "FFFFFF", "accent": "FFFFFF"},
        "colors": {"theme": [{"role": "lt2", "hex": "FFFFFF"}]},
    }
    _add_table(slide, {"headers": ["Этап"], "rows": [["Ответ"]]}, (1, 1, 4, 2), design)
    table = slide.shapes[0].table
    for row in range(2):
        cell = table.cell(row, 0)
        fill = str(cell.fill.fore_color.rgb)
        text = str(cell.text_frame.paragraphs[0].runs[0].font.color.rgb)
        assert contrast(fill, text) >= 4.5


def test_generated_text_stays_on_canvas():
    prs = _deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    body = _text(slide, "BrandDeck Body", (4.1, 6.75, 2.15, 1), "Текст")
    _contain_generated_text(slide, prs.slide_width, prs.slide_height)
    assert body.top + body.height <= prs.slide_height


def test_small_photo_zone_expands_only_into_free_space():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _text(slide, "BrandDeck Body", (0.3, 1.2, 2.5, 3.0), "Текст")
    design = {"canvas": {"width_inches": 10, "height_inches": 5.625}}
    small = (3.1, 1.2, 3.3, 1.8)
    expanded = _expand_small_image_zone(slide, small, design)
    assert expanded[2] >= 6 and expanded[3] >= 3.5
    _text(slide, "BrandDeck Note", (7.0, 2.0, 1.0, 0.5), "Заметка")
    assert _expand_small_image_zone(slide, small, design) == small


def test_photo_zone_of_an_underfilled_slide_grows_even_when_it_is_not_small():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _text(slide, "BrandDeck Body", (6.4, 1.9, 2.5, 3.0), "Текст")
    # 3.2 x 5.1 inches is a sixth of the slide: not small, so it stays...
    photo = (9.1, 1.8, 3.2, 5.1)
    assert _expand_small_image_zone(slide, photo, DESIGN) == photo
    # ...until the render found the slide below a quarter filled.
    grown = _expand_small_image_zone(slide, photo, DESIGN, underfilled=True)
    assert grown[:2] == photo[:2] and grown[2] > photo[2]
    assert grown[0] + grown[2] <= 13.333 - 0.3


def test_narrow_photo_zone_can_expand_width_without_growing_height():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _text(slide, "BrandDeck Body", (5.89, 2.96, 2.71, 3.89), "Текст")
    design = {"canvas": {"width_inches": 13.333, "height_inches": 7.5}}
    narrow = (8.87, 2.96, 3.48, 3.89)
    expanded = _expand_small_image_zone(slide, narrow, design)
    assert expanded[:2] == narrow[:2]
    assert expanded[2] > 4.0
    assert expanded[3] == narrow[3]
    _text(slide, "BrandDeck Note", (12.45, 3.0, 0.4, 0.5), "Не закрывать")
    assert _expand_small_image_zone(slide, narrow, design) == narrow


def test_cleanup_removes_scaffolding_and_keeps_design():
    prs = _deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, "Band", (0, 5.8, 13.333, 1.7))
    label = _rect(slide, "Sample label", (1.2, 1.5, 2.1, 0.5), text="Текст")
    icon = slide.shapes.add_group_shape()
    icon.name = "Icon"
    icon.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.5), Inches(1.5), Inches(0.5), Inches(0.5))
    card = _rect(slide, "Card", (6.0, 1.0, 6.4, 1.9), text="Заголовок")
    inside = _rect(slide, "Card icon", (6.1, 1.1, 0.5, 0.5), kind=MSO_SHAPE.OVAL)
    under = _rect(slide, "Under text", (0.5, 3.4, 2.8, 1.4))
    host = _rect(slide, "Host panel", (6.0, 3.2, 6.4, 2.3))
    avatar = _rect(slide, "Avatar", (0.5, 6.3, 0.9, 0.9), kind=MSO_SHAPE.OVAL)
    column = _text(slide, "Column", (0.5, 0.1, 5.0, 1.0), "Образец")
    _rect(slide, "Highlight", (0.6, 0.4, 4.5, 0.3))  # a bar in the text flow
    arrow = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(0.7), Inches(3.5), Inches(3.0), Inches(4.5)
    )
    arrow.name = "Arrow"
    divider = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(5.2), Inches(3.0), Inches(5.2), Inches(5.0)
    )
    divider.name = "Divider"
    original = {shape.element for shape in slide.shapes}
    had_text = {label.element, card.element, column.element}

    # The composer erased the sample label, filled the card and added text.
    label.text_frame.text = ""
    card.text_frame.text = "Пилот охватывает два подразделения"
    column.text_frame.text = "Новый заголовок блока"
    body = _text(slide, "BrandDeck Body", (0.5, 3.0, 4.0, 2.0), "Новый текст слайда")
    _text(slide, "BrandDeck Text 9", (6.2, 3.4, 6.0, 1.8), "Текст на панели")
    caption = _text(slide, "BrandDeck Caption", (1.6, 6.3, 3.0, 0.9), "Вопросы?")
    original |= {caption.element}  # a template caption the composer filled

    removed = {
        item["shape"]: item["reason"]
        for item in _remove_exemplar_leftovers(slide, DESIGN, original, had_text)
    }
    assert removed == {
        "Sample label": "erased_sample",
        "Icon": "attached_to_erased",
        "Under text": "under_content",
        "Avatar": "empty_photo_frame",
        "Highlight": "under_content",
        "Arrow": "under_content",
    }
    kept = {shape.name for shape in slide.shapes}
    assert {"Band", "Card", "Card icon", "Host panel", "Divider", body.name} <= kept
    assert inside.name in kept and host.name in kept and under.name not in kept
    assert icon.name not in kept and avatar.name not in kept


def test_empty_icon_tile_goes_and_a_tile_with_its_icon_stays():
    prs = _deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    empty = _rect(slide, "Tile", (0.3, 1.45, 0.62, 0.62))
    _text(slide, "BrandDeck Text 1", (1.1, 1.6, 2.2, 0.3), "Документы внутри")
    full = _rect(slide, "Tile 2", (0.3, 3.0, 0.62, 0.62))
    icon = _rect(slide, "Icon", (0.45, 3.15, 0.32, 0.32), kind=MSO_SHAPE.OVAL)
    _text(slide, "BrandDeck Text 2", (1.1, 3.15, 2.2, 0.3), "Данные не уходят")
    original = {empty.element, full.element, icon.element}
    removed = _remove_exemplar_leftovers(slide, DESIGN, original, set())
    assert removed == [{"shape": "Tile", "reason": "empty_photo_frame"}]


def _template(path: Path) -> None:
    prs = _deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, "Sample card", (0.6, 1.4, 3.0, 1.6), text="Текст")
    _rect(slide, "Accent", (5.0, 2.0, 2.0, 1.5))
    _rect(slide, "Avatar", (0.6, 5.2, 0.9, 0.9), kind=MSO_SHAPE.OVAL)
    _text(slide, "Speaker", (1.7, 5.2, 3.0, 0.9), "Имя спикера")
    _rect(slide, "Highlight", (1.8, 5.5, 2.5, 0.2))
    _rect(slide, "Card", (8.0, 4.0, 3.0, 1.5), text="Заголовок")
    _rect(slide, "Card icon", (8.1, 4.1, 0.4, 0.4), kind=MSO_SHAPE.OVAL)
    prs.save(path)


def test_audit_reports_leftovers_against_the_template(tmp_path: Path):
    template = tmp_path / "template.pptx"
    _template(template)
    prs = Presentation(template)
    slide = prs.slides[0]
    for shape in slide.shapes:
        if shape.name == "Sample card":
            shape.text_frame.text = ""
        if shape.name in {"Speaker", "Card"}:
            shape.text_frame.text = "Вопросы?"
            shape.name = f"BrandDeck Text {shape.shape_id}"
    _text(slide, "BrandDeck Body", (4.4, 1.5, 3.2, 1.2), "Текст поверх акцента")
    output = tmp_path / "output.pptx"
    prs.save(output)

    report = inspect_presentation(output, template_path=template)
    found = {
        (issue["code"], issue["shape"].split(" / ")[0]) for issue in report["issues"]
    }
    assert found == {
        ("emptied_template_block", "Sample card"),
        ("template_overlap", "Accent"),
        ("template_overlap", "Highlight"),
        ("empty_photo_frame", "Avatar"),
    }
    assert report["status"] == "warning" and report["score"] < 100


def test_audit_accepts_clean_slide(tmp_path: Path):
    template = tmp_path / "template.pptx"
    _template(template)
    prs = Presentation(template)
    slide = prs.slides[0]
    for shape in list(slide.shapes):
        if shape.name in {"Accent", "Avatar", "Highlight"}:
            shape.element.getparent().remove(shape.element)
        elif shape.has_text_frame:
            shape.text_frame.text = "Новый текст"
            shape.name = f"BrandDeck Text {shape.shape_id}"
    output = tmp_path / "output.pptx"
    prs.save(output)
    codes = {issue["code"] for issue in inspect_presentation(output, template_path=template)["issues"]}
    assert not codes & {"template_overlap", "emptied_template_block", "empty_photo_frame"}


@pytest.mark.skipif(not find_libreoffice(), reason="LibreOffice is unavailable")
def test_rendered_fill_flags_sparse_content_slides(tmp_path: Path):
    prs = _deck()
    for _ in range(3):
        prs.slides.add_slide(prs.slide_layouts[6])
    cover, sparse, full = prs.slides
    _text(cover, "BrandDeck Title", (1, 3, 8, 1), "Обложка")
    line = _text(sparse, "BrandDeck Body", (0.8, 1.0, 11.0, 5.5), "Один короткий тезис.")
    line.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    picture = tmp_path / "picture.png"
    Image.new("RGB", (800, 600), (40, 90, 160)).save(picture)
    _text(full, "BrandDeck Body", (0.8, 1.0, 4.5, 1.0), "Текст рядом с изображением.")
    image = full.shapes.add_picture(str(picture), Inches(5.8), Inches(0.9), Inches(6.8), Inches(5.1))
    image.name = "BrandDeck Image"
    output = tmp_path / "output.pptx"
    prs.save(output)

    assert export_presentation(output, tmp_path / "exports", expected_slide_count=3)["status"] == "passed"
    report = inspect_rendered_fill(
        tmp_path / "exports" / "output.pdf", output, ["cover", "content", "content"]
    )
    assert [item["slide"] for item in report["slides"]] == [2, 3]
    assert [(i["slide"], i["code"]) for i in report["issues"]] == [(2, "slide_underfilled")]
    assert report["issues"][0]["repairable"] is True


@pytest.mark.skipif(not find_libreoffice(), reason="LibreOffice is unavailable")
def test_rendered_fill_uses_cards_inside_generated_group(tmp_path: Path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10.833333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _text(slide, "BrandDeck Title", (0.25, 0.03, 10.34, 0.64), "Заголовок")
    diagram = slide.shapes.add_group_shape()
    diagram.name = "BrandDeck Diagram icon_grid"
    for index, (x, y) in enumerate(((0.25, 0.9), (5.52, 0.9), (0.25, 4.25), (5.52, 4.25)), 1):
        card = diagram.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x), Inches(y), Inches(5.07), Inches(2.74),
        )
        card.name = f"BrandDeck Diagram Card {index}"
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(230, 235, 240)
        label = diagram.shapes.add_textbox(
            Inches(x + 0.2), Inches(y + 0.3), Inches(4.6), Inches(1.0)
        )
        label.name = f"BrandDeck Diagram Text {index}"
        label.text_frame.text = f"Карточка {index}"
    output = tmp_path / "output.pptx"
    prs.save(output)
    assert export_presentation(output, tmp_path / "exports", expected_slide_count=1)["status"] == "passed"

    report = inspect_rendered_fill(tmp_path / "exports/output.pdf", output, ["content"])
    assert report["status"] == "passed"
    assert 0.68 <= report["slides"][0]["coverage"] <= 0.75


def test_diagram_slide_drops_the_exemplar_chart_beside_it():
    prs = _deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chart = slide.shapes.add_group_shape()
    chart.name = "Shape chart"
    bar = chart.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(3), Inches(3), Inches(5), Inches(0.3)
    )
    label = chart.shapes.add_textbox(Inches(0.6), Inches(2.8), Inches(1.2), Inches(0.6))
    label.text_frame.text = "Этап"
    bar.fill.solid()
    loose_bar = _rect(slide, "Loose bar", (9.5, 5.5, 3.0, 0.3))
    accent = _rect(slide, "Title accent", (11.5, 0.3, 1.0, 0.4))
    original = {shape.element for shape in slide.shapes}
    had_text = {chart.element}

    label.text_frame.text = ""
    diagram = slide.shapes.add_group_shape()
    diagram.name = "BrandDeck Diagram icon_grid"
    diagram.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.8), Inches(6.0), Inches(4.9)
    ).text_frame.text = "Новый тезис"

    removed = {
        item["shape"]: item["reason"]
        for item in _remove_exemplar_leftovers(
            slide, DESIGN, original, had_text, diagram_only=True
        )
    }
    assert removed == {"Shape chart": "erased_sample", "Loose bar": "beside_diagram"}
    kept = {shape.name for shape in slide.shapes}
    assert accent.name in kept and loose_bar.name not in kept


def test_audit_reports_an_emptied_chart_under_new_content(tmp_path: Path):
    template = tmp_path / "template.pptx"
    prs = _deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chart = slide.shapes.add_group_shape()
    chart.name = "Shape chart"
    chart.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(2.0), Inches(12), Inches(4.4)
    ).fill.solid()
    chart.shapes.add_textbox(Inches(1), Inches(2.2), Inches(2), Inches(0.4)).text_frame.text = "Этап"
    prs.save(template)

    output = Presentation(template)
    group = output.slides[0].shapes[0]
    for child in group.shapes:
        if child.has_text_frame:
            child.text_frame.text = ""
    diagram = output.slides[0].shapes.add_group_shape()
    diagram.name = "BrandDeck Diagram icon_grid"
    diagram.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.8), Inches(6.0), Inches(4.9)
    ).text_frame.text = "Новый тезис"
    output.save(tmp_path / "output.pptx")

    report = inspect_presentation(tmp_path / "output.pptx", template_path=template)
    assert ("template_overlap", "Shape chart") in {
        (issue["code"], issue["shape"].split(" / ")[0]) for issue in report["issues"]
    }


def test_an_emptied_text_slot_of_the_sample_goes():
    prs = _deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slot = _rect(slide, "Sample card", (0.5, 1.8, 3.5, 2.0), text="Образец")
    filled = _rect(slide, "Sample card 2", (4.3, 1.8, 3.5, 2.0), text="Образец")
    original = had_text = {slot.element, filled.element}
    # The composer renamed both slots; the plan had text for one of them.
    slot.text_frame.text, slot.name = "", "BrandDeck Text 5"
    filled.text_frame.text, filled.name = "Новый текст", "BrandDeck Text 6"
    removed = _remove_exemplar_leftovers(slide, DESIGN, original, had_text)
    assert removed == [{"shape": "BrandDeck Text 5", "reason": "erased_sample"}]


def test_audit_reports_an_emptied_renamed_slot(tmp_path: Path):
    template = tmp_path / "template.pptx"
    _template(template)
    prs = Presentation(template)
    card = next(shape for shape in prs.slides[0].shapes if shape.name == "Card")
    card.text_frame.text, card.name = "", "BrandDeck Text 7"
    output = tmp_path / "output.pptx"
    prs.save(output)
    found = {
        (issue["code"], issue["shape"])
        for issue in inspect_presentation(output, template_path=template)["issues"]
    }
    assert ("emptied_template_block", "BrandDeck Text 7") in found
