"""Grouped source text remains an editable slot when an exemplar is reused."""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from slide_agent.composer import _clear_group_text, _fill_slide, _infer_text_groups


def _text(group, name: str, value: str, x: float, y: float, w: float, h: float, size: float):
    shape = group.shapes.add_textbox(
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.name = name
    shape.text = value
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    return shape


def _grouped_slide():
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    navigation = slide.shapes.add_group_shape()
    page = _text(navigation, "page", "02", 11.2, 0.4, 1.2, 1.8, 90)

    heading_group = slide.shapes.add_group_shape()
    heading = _text(
        heading_group,
        "heading",
        "A descriptive source heading",
        0.5,
        0.9,
        6.8,
        0.6,
        27,
    )

    body_group = slide.shapes.add_group_shape()
    year = _text(body_group, "year", "1950", 0.5, 1.8, 3.1, 1.9, 90)
    label = _text(body_group, "label", "THE CLAIM", 0.5, 2.8, 1.3, 0.25, 12)
    body = _text(
        body_group,
        "paragraph",
        "The source paragraph spans several lines of the slide.",
        0.5,
        3.2,
        4.0,
        1.4,
        15,
    )

    footer_group = slide.shapes.add_group_shape()
    footer = _text(
        footer_group,
        "footer",
        "FIELD GUIDE · SOURCE DECK",
        0.5,
        7.18,
        2.6,
        0.2,
        9,
    )
    return presentation, slide, heading, body, year, page, footer, label


def test_grouped_exemplar_selects_heading_and_prose_instead_of_numbers():
    _, slide, heading, body, year, page, footer, label = _grouped_slide()

    titles, subtitles, bodies = _infer_text_groups(slide, [], [], [], 13.333, 7.5)

    assert titles == [heading]
    assert subtitles == [body]
    assert all(shape not in subtitles + bodies for shape in (year, page, footer, label))


def test_grouped_exemplar_replaces_selected_text_and_clears_samples(tmp_path: Path):
    presentation, slide, heading, body, year, page, footer, label = _grouped_slide()
    design = {
        "canvas": {"width_inches": 13.333, "height_inches": 7.5},
        "spacing": {"typical_left_margin_inches": 0.5, "typical_right_margin_inches": 0.5},
        "typography": {"primary_font": "Arial", "body_size_pt": 15},
    }

    _fill_slide(
        slide,
        {
            "role": "content",
            "title": "New supported heading",
            "body": "New supported paragraph.",
            "bullets": [],
        },
        design,
        set(),
    )
    output = tmp_path / "grouped-output.pptx"
    presentation.save(output)
    restored = Presentation(output)

    assert heading.text == "New supported heading"
    # A lone source body slot establishes its position/style; composition
    # replaces it with an editable body text box in that area.
    assert body.text == ""
    assert any(
        shape.name == "BrandDeck Body" and shape.text == "New supported paragraph."
        for shape in restored.slides[0].shapes
    )
    assert year.text == page.text == footer.text == label.text == ""
    assert any(
        shape.name.startswith("BrandDeck") and shape.text == "New supported heading"
        for group in restored.slides[0].shapes
        if group.shape_type == 6
        for shape in group.shapes
        if shape.has_text_frame
    )


def test_group_clear_keeps_assigned_nested_child():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    outer = slide.shapes.add_group_shape()
    inner = outer.shapes.add_group_shape()
    kept = _text(inner, "keep", "Use this", 1, 1, 2, 0.5, 20)
    erased = _text(inner, "erase", "Erase this", 1, 2, 2, 0.5, 20)

    _clear_group_text(outer, {kept.element})

    assert kept.text == "Use this"
    assert erased.text == ""


def test_translated_group_is_not_treated_as_a_slide_coordinate_slot():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    _text(group, "offset-title", "Heading inside offset group", 0.4, 0.5, 5, 0.7, 27)
    group.element.grpSpPr.xfrm.off.x += Inches(3)

    titles, subtitles, bodies = _infer_text_groups(slide, [], [], [], 13.333, 7.5)

    assert not titles and not subtitles and not bodies
