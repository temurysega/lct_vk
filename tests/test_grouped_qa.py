"""QA sees text filled into a source group's editable exemplar frame."""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from slide_agent.qa import inspect_presentation


def test_text_inside_source_group_counts_as_slide_content(tmp_path: Path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.name = "Source heading group"
    nested = group.shapes.add_group_shape()
    nested.name = "Source text bundle"
    title = nested.shapes.add_textbox(Inches(0.6), Inches(0.7), Inches(8), Inches(1))
    title.name = "BrandDeck Text 1"
    title.text = "Новый заголовок в группе шаблона"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    output = tmp_path / "grouped-title.pptx"
    presentation.save(output)

    report = inspect_presentation(output, expected_slide_count=1)

    assert report["slides"][0]["text_shape_count"] == 1
    assert all(issue["code"] != "empty_slide" for issue in report["issues"])


def test_nested_source_group_text_still_receives_overflow_check(tmp_path: Path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.name = "Source content group"
    text = group.shapes.add_textbox(Inches(0.6), Inches(1.0), Inches(1.5), Inches(0.3))
    text.name = "BrandDeck Text 2"
    text.text = "Длинная фраза, которая никак не помещается в эту узкую рамку"
    text.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    output = tmp_path / "grouped-overflow.pptx"
    presentation.save(output)

    report = inspect_presentation(output, expected_slide_count=1)

    assert any(issue["code"] == "text_overflow_risk" for issue in report["issues"])
