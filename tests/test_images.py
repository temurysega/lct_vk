import base64
import io
import json
import threading
import zipfile
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path
from typing import ClassVar

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Inches, Pt

from examples.create_demo_assets import create_template
from slide_agent.analyzer import analyze_template
from slide_agent.composer import _fill_slide, compose_presentation
from slide_agent.imagegen import ImageSettings, generate_images
from slide_agent.images import (
    ImageLibrary,
    attach_images,
    collect_content_images,
    crop_to_fill,
    extract_markdown_images,
    label_from_filename,
    normalize_image,
)
from slide_agent.qa import inspect_presentation
from slide_agent.service import generate_deck, generate_variants
from slide_agent.utils import read_json

CONTENT = """# Пилот сервиса поддержки
## Команда проекта
- Специалисты поддержки отвечают сотрудникам.
- Аналитик готовит сводку.
## Проблема
- Сотрудники ищут инструкции в нескольких источниках.
- Ответы на повторяющиеся вопросы отвлекают специалистов.
## Итоги
- Время ответа сократилось.
- Сотрудники довольны сервисом.
"""


def _png(path: Path, size=(800, 500), color="#0077FF") -> Path:
    Image.new("RGB", size, color).save(path)
    return path


def _bytes(size=(640, 480), fmt="PNG", color="#FF3885") -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", size, color).save(buffer, format=fmt)
    return buffer.getvalue()


def test_normalize_image_bounds_and_rejects_bad_input():
    _, suffix, size = normalize_image(_bytes((4000, 1000), "JPEG"))
    assert suffix == "jpg" and max(size) == 2400
    assert normalize_image(_bytes((300, 300)))[1] == "png"
    for data in (b"", b"not an image", _bytes((40, 40))):
        with pytest.raises(ValueError):
            normalize_image(data)
    bomb = io.BytesIO()
    Image.new("1", (8000, 6000)).save(bomb, format="PNG")
    with pytest.raises(ValueError, match="too many pixels"):
        normalize_image(bomb.getvalue())


@pytest.mark.parametrize(
    "image,box", [((1600, 900), (4, 4)), ((500, 1500), (6, 3)), ((1000, 1000), (3, 2))]
)
def test_crop_to_fill_matches_box_aspect(image, box):
    left, top, right, bottom = crop_to_fill(image, box)
    visible = image[0] * (1 - left - right) / (image[1] * (1 - top - bottom))
    assert visible == pytest.approx(box[0] / box[1])


def test_markdown_images_stay_inside_the_content_folder(tmp_path: Path):
    folder = tmp_path / "content"
    folder.mkdir()
    _png(folder / "team.png")
    _png(tmp_path / "secret.png")
    text = '# Команда\nНаша команда.\n\n![Команда проекта](team.png)\n![x](../secret.png)\n![y](https://example.com/a.png)\n<img src="team.png" alt="Команда">'
    cleaned, references = extract_markdown_images(text, folder)
    assert "![" not in cleaned and "<img" not in cleaned
    local = [reference for reference in references if reference["path"]]
    assert [reference["label"] for reference in local] == ["Команда проекта", "Команда"]
    assert all(reference["path"].parent == folder.resolve() for reference in local)
    assert "Команда" in local[0]["context"]


def test_pictures_are_extracted_from_pptx_and_docx_content(tmp_path: Path):
    source = Presentation()
    logo = _png(tmp_path / "logo.png", (200, 200), "#000000")
    photo = _png(tmp_path / "photo.png", (1200, 800), "#123456")
    for index in range(4):
        slide = source.slides.add_slide(source.slide_layouts[5])
        slide.shapes.title.text = f"Слайд {index}"
        slide.shapes.add_picture(
            str(logo), Inches(9), Inches(0.2), Inches(0.6), Inches(0.6)
        )
    picture = source.slides[1].shapes.add_picture(
        str(photo), Inches(1), Inches(2), Inches(5), Inches(3.3)
    )
    picture._element.xpath("./p:nvPicPr/p:cNvPr")[0].set("descr", "Команда поддержки")
    content = tmp_path / "material.pptx"
    source.save(content)
    library = ImageLibrary(tmp_path / "store")
    collect_content_images(content, "text", library)
    assert [asset.label for asset in library.assets.values()] == ["Команда поддержки"]
    assert "Слайд 1" in next(iter(library.assets.values())).context

    docx = tmp_path / "material.docx"
    document = (
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><w:body>'
        "<w:p><w:r><w:t>Рост выручки по кварталам</w:t></w:r></w:p>"
        '<w:p><w:r><w:drawing><wp:inline><wp:docPr id="1" name="p" descr="График роста"/>'
        '<a:graphic><a:graphicData><a:blip r:embed="rId5"/></a:graphicData></a:graphic>'
        "</wp:inline></w:drawing></w:r></w:p></w:body></w:document>"
    )
    rels = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId5" Type="image" Target="media/image1.png"/></Relationships>'
    )
    with zipfile.ZipFile(docx, "w") as archive:
        archive.writestr("word/document.xml", document)
        archive.writestr("word/_rels/document.xml.rels", rels)
        archive.writestr("word/media/image1.png", _bytes((900, 600)))
    library = ImageLibrary(tmp_path / "store2")
    collect_content_images(docx, "text", library)
    asset = next(iter(library.assets.values()))
    assert asset.label == "График роста" and "Рост выручки" in asset.context


def test_attach_images_matches_by_text_and_places_unlabeled_uploads(tmp_path: Path):
    library = ImageLibrary(tmp_path / "store")
    labeled = library.add(
        _bytes(color="#111111"), source="upload", label="команда проекта поддержки"
    )
    unlabeled = library.add(
        _bytes(color="#222222"),
        source="upload",
        label=label_from_filename("IMG_2041.jpg"),
    )
    plan = {
        "slides": [
            {"title": "Пилот", "role": "cover", "bullets": [], "visual": None},
            {
                "title": "Проблема",
                "role": "content",
                "bullets": ["Долгий поиск инструкций."],
                "visual": None,
            },
            {
                "title": "Команда проекта",
                "role": "content",
                "bullets": ["Поддержка и аналитика."],
                "visual": None,
            },
            {"title": "Спасибо", "role": "closing", "bullets": [], "visual": None},
        ]
    }
    summary = attach_images(plan, library)
    assert plan["slides"][2]["visual"]["asset_id"] == labeled.asset_id
    assert plan["slides"][2]["visual"]["match"] == "lexical"
    assert plan["slides"][1]["visual"] == {
        "type": "image",
        "asset_id": unlabeled.asset_id,
        "match": "order",
        "match_score": 0.0,
    }
    assert plan["slides"][0]["visual"] is None and plan["slides"][3]["visual"] is None
    assert summary["placed"] == 2 and set(plan["assets"]) == {
        labeled.asset_id,
        unlabeled.asset_id,
    }


def test_planned_brief_image_is_reported_with_its_slide(tmp_path: Path):
    library = ImageLibrary(tmp_path / "store")
    asset = library.add(_bytes(), source="upload", label="server room")
    plan = {
        "slides": [
            {"title": "Data stays inside", "role": "content", "visual": {"type": "image", "asset_id": asset.asset_id}}
        ]
    }
    summary = attach_images(plan, library)
    assert summary["placed"] == 1
    assert summary["placements"] == [
        {"asset_id": asset.asset_id, "slide": 1, "match": "planned", "score": None}
    ]


def test_generated_deck_places_images_without_distortion(tmp_path: Path):
    template = tmp_path / "clean-blue.pptx"
    create_template(template, "clean-blue")
    workspace = tmp_path / "workspace"
    team = _png(tmp_path / "команда проекта.png", (1600, 900))
    result = generate_deck(
        template=template,
        content=CONTENT,
        workspace=workspace,
        slide_count=5,
        offline=True,
        images=[team],
    )
    assert result["status"] == "completed"
    assert result["images"]["placements"][0]["match"] == "lexical"
    assert result["visuals"]["images"] == 1
    prs = Presentation(result["output"])
    pictures = [
        shape
        for slide in prs.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        and shape.name == "BrandDeck Image"
    ]
    assert len(pictures) == 1
    assert not [
        issue for issue in result["qa"]["issues"] if issue["code"] == "image_distorted"
    ]
    plan = read_json(Path(result["presentation_dir"]) / "deck_plan.final.json")
    asset_id = result["images"]["placements"][0]["asset_id"]
    assert plan["assets"][asset_id]["source"] == "upload"


def test_variants_share_one_image_placement(tmp_path: Path):
    template = tmp_path / "tech-dark.pptx"
    create_template(template, "tech-dark")
    batch = generate_variants(
        template=template,
        content=CONTENT,
        workspace=tmp_path / "workspace",
        slide_count=5,
        offline=True,
        export_formats=(),
        images=[_png(tmp_path / "команда проекта.png")],
    )
    placements = {
        json.dumps(deck["images"]["placements"]) for deck in batch["variants"]
    }
    assert len(placements) == 1 and batch["status"] == "completed"


def test_qa_flags_a_stretched_picture(tmp_path: Path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    image = _png(tmp_path / "wide.png", (1600, 400))
    picture = slide.shapes.add_picture(
        str(image), Inches(1), Inches(1), Inches(3), Inches(3)
    )
    picture.name = "BrandDeck Image"
    output = tmp_path / "stretched.pptx"
    prs.save(output)
    issues = inspect_presentation(output)["issues"]
    assert any(
        issue["code"] == "image_distorted" and issue["severity"] == "warning"
        for issue in issues
    )


def _analyzed(tmp_path: Path):
    template = tmp_path / "clean-blue.pptx"
    create_template(template, "clean-blue")
    analyzed = analyze_template(template, workspace=tmp_path / "workspace")
    return analyzed, read_json(analyzed / "design_system.json")


def test_picture_placeholder_receives_the_image_or_an_illustration(tmp_path: Path):
    analyzed, design = _analyzed(tmp_path)
    library = ImageLibrary(tmp_path / "store")
    asset = library.add(_bytes((1200, 900)), source="upload", label="офис")
    base = {
        "subtitle": "",
        "body": "",
        "speaker_notes": "",
        "pattern_id": "layout-8",
        "layout_index": 8,
        "master_index": 0,
    }
    plan = {
        "title": "t",
        "assets": library.plan_assets(),
        "slides": [
            {
                **base,
                "title": "Офис",
                "role": "content",
                "bullets": ["Новый офис."],
                "visual": {"type": "image", "asset_id": asset.asset_id},
            },
            {
                **base,
                "title": "Безопасность данных",
                "role": "content",
                "bullets": ["Защита."],
                "visual": None,
            },
        ],
    }
    output = tmp_path / "placeholders.pptx"
    result = compose_presentation(
        template_dir=analyzed, plan=plan, output_path=output, design_system=design
    )
    placements = [record.get("placement") for record in result["visuals"]["slides"]]
    assert placements == ["template_placeholder", "template_slot"]
    assert result["visuals"]["slides"][1]["type"] == "illustration"
    prs = Presentation(output)
    first = [
        shape
        for shape in prs.slides[0].placeholders
        if shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE
    ]
    assert first and first[0].name == "BrandDeck Image"
    assert not [
        s
        for s in prs.slides[1].placeholders
        if s.placeholder_format.type == PP_PLACEHOLDER.PICTURE
    ]
    report = inspect_presentation(output, design_system=design, expected_slide_count=2)
    assert not [
        issue for issue in report["issues"] if issue["code"] == "image_distorted"
    ]


def _exemplar(tmp_path: Path, text_box):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    source_photo = _png(tmp_path / "source.png", (900, 900), "#FF0000")
    slide.shapes.add_picture(
        str(source_photo), Inches(6), Inches(1.5), Inches(3), Inches(3)
    )
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(5), Inches(0.8))
    title.text = "Исходный заголовок"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    box = slide.shapes.add_textbox(*(Inches(value) for value in text_box))
    box.text = "Исходный текст примера"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
    return slide


def test_template_photo_slot_is_reused_beside_text(tmp_path: Path):
    slide = _exemplar(tmp_path, (0.5, 1.5, 7, 3.5))
    library = ImageLibrary(tmp_path / "store")
    asset = library.add(_bytes((1600, 900)), source="upload")
    design = {
        "canvas": {"width_inches": 10, "height_inches": 7.5},
        "spacing": {},
        "typography": {"primary_font": "Arial", "body_size_pt": 16},
        "colors": {"theme": []},
    }
    context = {
        "assets": library.plan_assets(),
        "records": [],
        "slide": 1,
        "title": "T",
        "text": "",
    }
    spec = {
        "title": "T",
        "role": "content",
        "subtitle": "",
        "body": "",
        "bullets": ["Новый текст."],
        "visual": {"type": "image", "asset_id": asset.asset_id},
    }
    _fill_slide(slide, spec, design, set(), context)
    picture = next(
        shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    assert picture.name == "BrandDeck Image" and picture.left == Inches(6)
    assert context["records"][0]["placement"] == "template_slot"
    text = next(
        shape
        for shape in slide.shapes
        if shape.has_text_frame and "Новый текст" in shape.text
    )
    assert text.left + text.width <= picture.left


def test_photo_under_centred_text_is_not_used_as_a_slot(tmp_path: Path):
    slide = _exemplar(tmp_path, (6.2, 2, 2.6, 2))
    library = ImageLibrary(tmp_path / "store")
    asset = library.add(_bytes((1600, 900)), source="upload")
    design = {
        "canvas": {"width_inches": 10, "height_inches": 7.5},
        "spacing": {},
        "typography": {"primary_font": "Arial", "body_size_pt": 16},
        "colors": {"theme": []},
    }
    context = {
        "assets": library.plan_assets(),
        "records": [],
        "slide": 1,
        "title": "T",
        "text": "",
    }
    spec = {
        "title": "T",
        "role": "content",
        "subtitle": "",
        "body": "",
        "bullets": ["Новый текст."],
        "visual": {"type": "image", "asset_id": asset.asset_id},
    }
    _fill_slide(slide, spec, design, set(), context)
    assert context["records"][0]["placement"] == "zone"
    sources = [
        shape
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        and shape.name != "BrandDeck Image"
    ]
    assert sources == []


class _ImageAPI(BaseHTTPRequestHandler):
    requests: ClassVar[list[dict]] = []

    def do_POST(self):
        body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        type(self).requests.append(
            {"path": self.path, "auth": self.headers.get("Authorization"), **body}
        )
        payload = {
            "data": [{"b64_json": base64.b64encode(_bytes((1024, 768))).decode()}]
        }
        encoded = json.dumps(payload).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(encoded)))
        self.end_headers()
        self.wfile.write(encoded)

    def log_message(self, *args):
        pass


def test_image_generation_fills_requests_through_openai_compatible_api(tmp_path: Path):
    server = HTTPServer(("127.0.0.1", 0), _ImageAPI)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        settings = ImageSettings(
            base_url=f"http://127.0.0.1:{server.server_port}/v1",
            api_key="secret",
            model="black-forest-labs/FLUX.1-schnell",
        )
        plan = {
            "slides": [
                {"title": "Пилот", "role": "cover", "bullets": [], "visual": None},
                {
                    "title": "Офис будущего",
                    "role": "content",
                    "bullets": ["Гибкие рабочие места"],
                    "visual": {"type": "image", "request": "светлый офис"},
                },
            ]
        }
        library = ImageLibrary(tmp_path / "store")
        offline = generate_images(plan, {}, library, offline=True, settings=settings)
        assert offline["generated"] == [] and not _ImageAPI.requests
        report = generate_images(
            plan,
            {"colors": {"theme": [{"role": "accent1", "hex": "0077FF"}]}},
            library,
            offline=False,
            settings=settings,
        )
    finally:
        server.shutdown()
    request = _ImageAPI.requests[0]
    assert (
        request["path"] == "/v1/images/generations"
        and request["auth"] == "Bearer secret"
    )
    assert request["model"] == "black-forest-labs/FLUX.1-schnell"
    assert "Офис будущего" in request["prompt"] and "#0077FF" in request["prompt"]
    assert report["generated"][0]["slide"] == 2
    asset_id = plan["slides"][1]["visual"]["asset_id"]
    assert plan["assets"][asset_id]["source"] == "generated"


def test_unavailable_image_generation_falls_back_to_illustration(tmp_path: Path):
    analyzed, design = _analyzed(tmp_path)
    plan = {
        "title": "t",
        "assets": {},
        "slides": [
            {
                "title": "Безопасность данных",
                "role": "content",
                "subtitle": "",
                "body": "",
                "bullets": ["Защита."],
                "speaker_notes": "",
                "visual": {"type": "image", "request": "щит"},
                "pattern_id": "layout-1",
                "layout_index": 1,
                "master_index": 0,
            }
        ],
    }
    settings = ImageSettings(
        base_url="http://127.0.0.1:9/v1", api_key="", model="m", timeout_seconds=2
    )
    report = generate_images(
        plan, design, ImageLibrary(tmp_path / "store"), offline=False, settings=settings
    )
    assert report["errors"] and plan["slides"][0]["visual"].get("asset_id") is None
    result = compose_presentation(
        template_dir=analyzed,
        plan=plan,
        output_path=tmp_path / "o.pptx",
        design_system=design,
    )
    assert result["visuals"]["slides"][0]["type"] == "illustration"


def test_api_accepts_image_uploads_and_sanitises_names(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("BRANDDECK_WORKSPACE", str(tmp_path / "workspace"))
    template = tmp_path / "brand.pptx"
    create_template(template, "clean-blue")
    from backend.main import app

    with TestClient(app) as client:
        assert (
            client.post(
                "/api/auth/register",
                json={
                    "username": "designer",
                    "password": "test-password-123",
                    "position": "Дизайнер",
                },
            ).status_code
            == 201
        )
        with template.open("rb") as source:
            template_id = client.post(
                "/v1/templates/analyze",
                files={"file": ("brand.pptx", source)},
                data={"offline": "true"},
            ).json()["template_id"]
        rejected = client.post(
            "/v1/presentations/jobs",
            data={"template_id": template_id, "content": CONTENT, "offline": "true"},
            files=[("image_files", ("evil.svg", b"<svg/>", "image/svg+xml"))],
        )
        assert rejected.status_code == 400
        response = client.post(
            "/v1/presentations/jobs",
            data={
                "template_id": template_id,
                "content": CONTENT,
                "slide_count": "5",
                "offline": "true",
            },
            files=[("image_files", ("команда:проекта.png", _bytes(), "image/png"))],
        )
        assert response.status_code == 202
        job = client.get(f"/v1/jobs/{response.json()['job_id']}").json()
        assert job["status"] == "completed"
        manifest = client.get(f"/v1/presentations/{job['presentation_id']}").json()
        assert manifest["images"]["placements"][0]["match"] == "lexical"
        saved = list((tmp_path / "workspace" / "users").rglob("01-*.png"))
        assert [path.name for path in saved] == ["01-команда_проекта.png"]
