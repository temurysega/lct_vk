from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from slide_agent.analyzer import analyze_template
from slide_agent.service import generate_deck
from slide_agent.utils import read_json

CONTENT = """# Native rebuild
## Problem
- Converted slides contain disconnected fragments.
- Replacing text breaks the composition.
## Approach
- Extract the brand rules.
- Rebuild every composition with native objects.
## Result
- Editable slides.
- Predictable spacing and typography.
"""


def _fragmented_template(path: Path) -> None:
    prs = Presentation()
    logo = path.with_suffix(".png")
    Image.new("RGB", (80, 30), "#00D3E3").save(logo)
    for slide_index in range(4):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(7, 12, 24)
        slide.shapes.add_picture(
            str(logo), Inches(11.7), Inches(6.8), Inches(0.8), Inches(0.3)
        )
        title = slide.shapes.add_textbox(
            Inches(0.7), Inches(0.6), Inches(8.5), Inches(0.8)
        )
        title.name = f"PDF text title {slide_index}"
        title.text = f"Source slide {slide_index + 1}"
        title.text_frame.paragraphs[0].runs[0].font.name = "VK Sans Display"
        title.text_frame.paragraphs[0].runs[0].font.size = Pt(34)
        for fragment_index in range(5):
            fragment = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(0.7 + fragment_index * 1.3),
                Inches(2.0 + (fragment_index % 2) * 1.1),
                Inches(1.0),
                Inches(0.7),
            )
            fragment.name = f"PDF vector {slide_index}-{fragment_index}"
            fragment.fill.solid()
            fragment.fill.fore_color.rgb = RGBColor(19, 31, 52)
    prs.save(path)


def test_fragmented_pdf_source_is_rebuilt_as_native_grid(tmp_path: Path):
    template = tmp_path / "fragmented.pptx"
    workspace = tmp_path / "workspace"
    _fragmented_template(template)

    analyzed = analyze_template(template, workspace=workspace)
    design = read_json(analyzed / "design_system.json")
    catalog = read_json(analyzed / "pattern_catalog.json")
    assert design["source_model"]["fragmented"] is True
    assert design["source_model"]["composition_mode"] == "native_grid"
    assert design["typography"]["heading_font"] == "VK Sans Display"
    assert any(
        pattern.get("source_kind") == "native_grid" for pattern in catalog["patterns"]
    )

    result = generate_deck(
        template=analyzed,
        content=CONTENT,
        workspace=workspace,
        slide_count=5,
        offline=True,
    )
    assert result["composition_mode"] == "native_grid"
    assert result["qa"]["status"] in {"passed", "warning"}
    generated = Presentation(result["output"])
    assert all(
        not shape.name.lower().startswith("pdf ")
        for slide in generated.slides
        for shape in slide.shapes
    )
    assert all(
        any(shape.name.startswith("BrandDeck Native") for shape in slide.shapes)
        for slide in generated.slides
    )
