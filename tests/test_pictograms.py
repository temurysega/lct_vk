import json
import math
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

from slide_agent.icon_builder import parse_icon_module
from slide_agent.pictograms import (
    ICON_ROOT,
    add_icon,
    assign_icons,
    choose_icon,
    icon_bundle,
    icon_names,
    keyword_config,
)
from slide_agent.vector import arc_to_cubics, element_commands, parse_path


def test_path_parser_handles_relative_commands_and_compressed_arc_flags():
    commands = parse_path("M2 3h4v2l-1 1s1 1 2 2a2 2 0 012 2z")
    assert commands[0] == ("M", 2.0, 3.0)
    assert commands[1] == ("L", 6.0, 3.0)
    assert commands[2] == ("L", 6.0, 5.0)
    assert commands[3] == ("L", 5.0, 6.0)
    # "s" reflects nothing after a line: first control point is the pen.
    assert commands[4][:3] == ("C", 5.0, 6.0)
    arc = [command for command in commands if command[0] == "C"][-1]
    assert arc[-2:] == pytest.approx((9.0, 10.0))
    assert commands[-1] == ("Z",)


def test_arc_to_cubics_stays_on_the_circle():
    segments = arc_to_cubics(0, 10, 10, 10, 0, False, True, 20, 10)
    assert segments[-1][-2:] == (20, 10)
    for command in segments:
        x, y = command[-2], command[-1]
        assert math.hypot(x - 10, y - 10) == pytest.approx(10, abs=1e-6)
    assert len(segments) == 2  # a half circle is split into quarter arcs


def test_every_bundled_icon_becomes_native_custom_geometry(tmp_path: Path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    names = icon_names()
    assert len(names) >= 120
    for index, name in enumerate(names):
        add_icon(
            slide.shapes,
            name,
            (0.2 + index % 20 * 0.45, 0.2 + index // 20 * 0.45, 0.4),
            "0077FF",
        )
    output = tmp_path / "icons.pptx"
    prs.save(output)
    shapes = list(Presentation(output).slides[0].shapes)
    assert len(shapes) == len(names)
    for shape in shapes:
        assert shape.shape_type != MSO_SHAPE_TYPE.PICTURE
        geometry = shape._element.spPr.find(qn("a:custGeom"))
        assert geometry is not None
        paths = geometry.findall(f"{qn('a:pathLst')}/{qn('a:path')}")
        assert paths and all(path.get("fill") == "none" for path in paths)
        line = shape._element.spPr.find(qn("a:ln"))
        assert line.get("cap") == "rnd"
        assert shape._element.find(qn("p:style")) is None


def test_shapes_other_than_paths_convert():
    circle = element_commands("circle", {"cx": "12", "cy": "12", "r": "4"})
    assert circle[0] == ("M", 16.0, 12.0) and circle[-1] == ("Z",)
    rect = element_commands(
        "rect", {"x": "2", "y": "2", "width": "20", "height": "8", "rx": "2"}
    )
    assert sum(1 for command in rect if command[0] == "C") == 4
    polyline = element_commands("polyline", {"points": "1 2 3 4 5 6"})
    assert [command[0] for command in polyline] == ["M", "L", "L"]


def test_keyword_config_matches_bundle_and_builder_format():
    config = keyword_config()
    bundle = icon_bundle()
    assert set(config["icons"]) == set(bundle["icons"])
    assert config["fallback"] in bundle["icons"]
    assert bundle["license"] == "ISC"
    assert (
        (ICON_ROOT / "LICENSE-lucide.txt").read_text(encoding="utf-8").startswith("ISC")
    )
    module = 'createLucideIcon("X", [["path", { d: "M1 1h2", key: "a" }], ["circle", { cx: "1", cy: "2", r: "3", key: "b" }]]);'
    assert parse_icon_module(module) == [
        ["path", {"d": "M1 1h2"}],
        ["circle", {"cx": "1", "cy": "2", "r": "3"}],
    ]


@pytest.mark.parametrize(
    ("text", "icon"),
    [
        ("Безопасность данных сотрудников", "shield-check"),
        ("Команда проекта", "users"),
        ("Рост выручки на 20%", "trending-up"),
        ("Риски внедрения", "triangle-alert"),
        ("Launch the pilot in Q3", "rocket"),
        ("Обучение специалистов", "book-open"),
    ],
)
def test_icon_choice_is_deterministic_for_russian_and_english(text, icon):
    assert choose_icon(text)[0] == icon
    assert choose_icon(text)[0] == icon


def test_unrelated_text_has_no_confident_icon_and_lists_avoid_repeats():
    assert choose_icon("Лорем ипсум долор")[0] is None
    icons = assign_icons(["Команда поддержки", "Команда аналитиков"])
    assert icons[0] != icons[1]
    assert assign_icons(["Команда", "абвгд"], require_all=True) == [None, None]


def test_keywords_config_is_valid_json_with_lowercase_entries():
    raw = json.loads((ICON_ROOT / "keywords.json").read_text(encoding="utf-8"))
    for keywords in raw["icons"].values():
        assert all(keyword == keyword.lower() for keyword in keywords)


def test_icon_box_geometry_is_square(tmp_path: Path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = add_icon(slide.shapes, "rocket", (1, 1, 0.5), "112233")
    assert shape.width == shape.height == Inches(0.5)
    assert shape.name == "BrandDeck Icon rocket"
