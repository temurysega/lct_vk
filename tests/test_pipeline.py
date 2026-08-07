from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from examples.create_demo_assets import create_template
from slide_agent.analyzer import analyze_template
from slide_agent.service import generate_deck
from slide_agent.utils import read_json

CONTENT = """# Service Alpha
## Problem
- Reports take 10 hours every week.
- Five systems must be checked manually.
## Solution
- One automated data pipeline.
- A prioritized action dashboard.
## Results
- Preparation time fell from 10 hours to 30 minutes.
- Coverage increased from 50% to 90%.
## Plan
- Connect two more systems.
- Roll out to three teams.
"""


def test_offline_pipeline_preserves_template_examples(tmp_path: Path):
    template = tmp_path / "tech-dark.pptx"
    workspace = tmp_path / "workspace"
    create_template(template, "tech-dark")
    logo = tmp_path / "logo.png"
    Image.new("RGB", (80, 80), "#7F56D9").save(logo)
    source_prs = Presentation(template)
    for source_slide in source_prs.slides:
        source_slide.shapes.add_picture(
            str(logo), Inches(11.5), Inches(0.5), Inches(0.55), Inches(0.55)
        )
    source_prs.save(template)

    analyzed = analyze_template(template, workspace=workspace)
    manifest = read_json(analyzed / "manifest.json")
    catalog = read_json(analyzed / "pattern_catalog.json")
    assert manifest["layout_count"] >= 3
    assert catalog["pattern_count"] >= 3
    assert any(item["example_slide_indices"] for item in catalog["patterns"])
    assert any(
        item.get("source_kind") == "slide_exemplar" for item in catalog["patterns"]
    )
    exemplar = next(
        item
        for item in catalog["patterns"]
        if item.get("source_kind") == "slide_exemplar"
    )
    assert exemplar["capacity"]["body_chars"] > 0
    assert exemplar["capacity"]["rhetorical_pattern"]

    result = generate_deck(
        template=analyzed,
        content=CONTENT,
        workspace=workspace,
        slide_count=6,
        offline=True,
    )
    assert result["qa"]["status"] in {"passed", "warning"}
    assert result["slide_count"] == 6
    plan = read_json(Path(result["presentation_dir"]) / "deck_plan.json")
    assert all(slide.get("pattern_selection") for slide in plan["slides"])
    output = Path(result["output"])
    assert output.exists() and output.stat().st_size > 10_000
    generated = Presentation(output)
    assert len(generated.slides) == 6
    assert generated.slides[0].shapes.title.text == "Service Alpha"
    assert any(
        shape.name == "Brand decoration tech-dark"
        for shape in generated.slides[0].shapes
    )
    assert any(
        shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        for shape in generated.slides[0].shapes
    )


def test_analysis_is_content_addressed(tmp_path: Path):
    template = tmp_path / "clean-blue.pptx"
    create_template(template, "clean-blue")
    first = analyze_template(template, workspace=tmp_path / "workspace")
    second = analyze_template(template, workspace=tmp_path / "workspace")
    assert first == second
