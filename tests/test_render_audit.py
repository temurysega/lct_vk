from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from slide_agent.exporter import export_presentation, find_libreoffice
from slide_agent.render_audit import (
    contrast_ratio,
    inspect_rendered_contrast,
    repair_rendered_contrast,
)


def test_contrast_ratio() -> None:
    assert contrast_ratio((0, 0, 0), (255, 255, 255)) == pytest.approx(21)
    assert contrast_ratio((0, 119, 255), (32, 32, 32)) < 4.5


@pytest.mark.skipif(not find_libreoffice(), reason="LibreOffice is unavailable")
def test_low_contrast_is_found_and_repaired_on_office_render(tmp_path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(32, 32, 32)
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1))
    shape.name = "BrandDeck Body"
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.text = "Текст на темном фоне должен читаться"
    paragraph.runs[0].font.size = Pt(24)
    paragraph.runs[0].font.color.rgb = RGBColor(0, 119, 255)
    output = tmp_path / "output.pptx"
    prs.save(output)

    result = export_presentation(output, tmp_path / "exports", expected_slide_count=1)
    assert result["status"] == "passed"
    rendered = tmp_path / "exports" / "output.pdf"
    before = inspect_rendered_contrast(rendered, output)
    assert any(issue["code"] == "rendered_low_contrast" for issue in before["issues"])
    assert repair_rendered_contrast(output, before["issues"])

    result = export_presentation(output, tmp_path / "exports", expected_slide_count=1)
    assert result["status"] == "passed"
    after = inspect_rendered_contrast(rendered, output)
    assert after["issues"] == []
