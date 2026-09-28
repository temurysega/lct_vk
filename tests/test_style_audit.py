"""Appendix 1 template-conformance checks and the composer's compliance."""

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from slide_agent.composer import _TYPE_SCALE, _add_table, _snap_size, _theme_colors
from slide_agent.qa import inspect_presentation

DESIGN = {
    "canvas": {"width_inches": 13.333, "height_inches": 7.5},
    "typography": {"primary_font": "Arial", "body_size_pt": 16, "observed_sizes_pt": [12, 14, 16, 24, 32]},
    "colors": {"theme": [{"role": "dk1", "hex": "000000"}, {"role": "accent1", "hex": "0077FF"}]},
    "brand": {"background": "FFFFFF", "accent": "000000", "body": "000000"},
}


def _deck() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def _text(slide, name, box, text, size, color="000000"):
    shape = slide.shapes.add_textbox(*(Inches(value) for value in box))
    shape.name = name
    shape.text_frame.text = text
    run = shape.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    return shape


def _codes(path: Path, template: Path | None = None) -> set[tuple[int, str]]:
    report = inspect_presentation(path, design_system=DESIGN, template_path=template)
    return {(issue.get("slide"), issue["code"]) for issue in report["issues"]}


def test_style_checks_flag_size_scale_colour_margin_and_chart(tmp_path: Path):
    prs = _deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _text(slide, "BrandDeck Body", (1, 1, 6, 1), "Мелкий текст", 8)
    _text(slide, "BrandDeck Text 7", (1, 2.5, 6, 1), "Кегль вне шкалы", 19)
    _text(slide, "BrandDeck Caption", (1, 4, 6, 1), "Чужой цвет", 16, color="00B050")
    _text(slide, "BrandDeck Edge", (0.05, 5.5, 6, 1), "У самого края", 16)
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("Ряд", (1, 2))
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7.5), Inches(1), Inches(5), Inches(4), data
    )
    frame.name = "BrandDeck Chart"
    clean = prs.slides.add_slide(prs.slide_layouts[6])
    _text(clean, "BrandDeck Body", (1, 1, 6, 1), "Нормальный текст", 16, color="0077FF")
    output = tmp_path / "output.pptx"
    prs.save(output)

    codes = _codes(output)
    assert {
        (1, "text_too_small"),
        (1, "font_off_scale"),
        (1, "off_palette_color"),
        (1, "margin_intrusion"),
        (1, "chart_unlabeled"),
    } <= codes
    assert not {code for slide, code in codes if slide == 2} & {
        "text_too_small", "font_off_scale", "off_palette_color", "margin_intrusion"
    }


def test_pie_chart_flags_cramped_legend_and_duplicated_numbers(tmp_path: Path):
    prs = _deck()
    data = CategoryChartData()
    data.categories = ["Продажи", "Поддержка"]
    data.add_series("Значения", [60, 90])
    cramped = prs.slides.add_slide(prs.slide_layouts[6])
    frame = cramped.shapes.add_chart(
        XL_CHART_TYPE.PIE, Inches(3.5), Inches(2), Inches(3.2), Inches(4.5), data
    )
    frame.name = "BrandDeck Chart"
    frame.chart.has_legend = True
    frame.chart.legend.position = XL_LEGEND_POSITION.RIGHT
    frame.chart.plots[0].has_data_labels = True
    labels = frame.chart.plots[0].data_labels
    labels.show_value = True
    labels.show_percentage = True

    roomy = prs.slides.add_slide(prs.slide_layouts[6])
    second = roomy.shapes.add_chart(
        XL_CHART_TYPE.PIE, Inches(3), Inches(2), Inches(5), Inches(4), data
    )
    second.name = "BrandDeck Chart"
    second.chart.has_legend = True
    second.chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    second.chart.plots[0].has_data_labels = True
    second.chart.plots[0].data_labels.show_value = False
    second.chart.plots[0].data_labels.show_percentage = True
    output = tmp_path / "charts.pptx"
    prs.save(output)

    codes = _codes(output)
    assert (1, "chart_legend_too_narrow") in codes
    assert (1, "chart_redundant_labels") in codes
    assert not {code for slide, code in codes if slide == 2} & {
        "chart_legend_too_narrow", "chart_redundant_labels", "chart_too_small"
    }


def test_tints_of_the_palette_and_template_layouts_pass(tmp_path: Path):
    template = tmp_path / "template.pptx"
    _deck().save(template)
    prs = Presentation(template)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(2))
    card.name = "BrandDeck Diagram Card 1"
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0x80, 0xBB, 0xFF)  # accent over white
    output = tmp_path / "output.pptx"
    prs.save(output)
    codes = {code for _, code in _codes(output, template)}
    assert "off_palette_color" not in codes and "foreign_layout" not in codes


def test_cyrillic_font_substitution_is_reported_only_for_cjk_template(tmp_path: Path):
    prs = _deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = _text(slide, "BrandDeck Body", (1, 1, 8, 1), "Русский текст", 16)
    shape.text_frame.paragraphs[0].runs[0].font.name = "Arial"
    output = tmp_path / "cyrillic.pptx"
    prs.save(output)

    cjk_design = {
        **DESIGN,
        "typography": {
            **DESIGN["typography"],
            "primary_font": "游ゴシック",
            "families": ["游ゴシック"],
        },
    }
    with_fallback = inspect_presentation(output, design_system=cjk_design)
    assert "foreign_fonts" not in {issue["code"] for issue in with_fallback["issues"]}
    assert with_fallback["font_substitutions"][0]["to"] == "Arial"

    latin_design = {
        **DESIGN,
        "typography": {**DESIGN["typography"], "primary_font": "Play", "families": ["Play"]},
    }
    without_fallback = inspect_presentation(output, design_system=latin_design)
    assert "foreign_fonts" in {issue["code"] for issue in without_fallback["issues"]}


def test_composer_snaps_sizes_to_the_template_scale():
    token = _TYPE_SCALE.set((12.0, 14.0, 16.0, 24.0))
    try:
        assert _snap_size(15.5, 9) == 14.0
        assert _snap_size(27.0, 9) == 24.0
        assert _snap_size(10.0, 9) == 12.0  # nearest larger step within reach
    finally:
        _TYPE_SCALE.reset(token)
    assert _snap_size(13.3, 9) == 13.3  # without a scale nothing changes


def test_grey_brand_accent_does_not_replace_the_theme_accent():
    assert _theme_colors(DESIGN)["accent1"] == "0077FF"
    prs = _deck()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_table(slide, {"headers": ["Метрика", "Цель"], "rows": [["Время", "до 30 минут"]]},
               (1, 1, 6, 4), DESIGN)
    table = slide.shapes[0].table
    header = str(table.cell(0, 0).fill.fore_color.rgb)
    assert header != "000000" and header.startswith("00")
    # A two-row table would leave the slide nearly empty: it fills its zone.
    assert sum(row.height for row in table.rows) == Inches(4)

    dense = prs.slides.add_slide(prs.slide_layouts[6])
    rows = [[f"Показатель {index}", "Значение с пояснением"] for index in range(6)]
    _add_table(dense, {"headers": ["Метрика", "Цель"], "rows": rows}, (1, 1, 11, 6), DESIGN)
    assert sum(row.height for row in dense.shapes[0].table.rows) < Inches(6)  # rows follow text
