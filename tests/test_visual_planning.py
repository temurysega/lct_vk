from itertools import pairwise
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from examples.create_demo_assets import create_template
from slide_agent.planner import _fallback_plan, _slide_requirements, normalize_plan
from slide_agent.service import generate_deck
from slide_agent.utils import read_json
from slide_agent.visual_planning import add_offline_visuals, infer_diagram

PLAN_TEXT = """# Пилот
## План работ
- В первую неделю команда собирает инструкции.
- Во вторую неделю проверяет поиск.
- В третью неделю проводит пилот.
- В четвертую неделю анализирует обратную связь.
## Пользователи
- Сотруднику нужен понятный ответ со ссылкой на инструкцию.
- Специалисту поддержки нужна история обращения.
- Руководителю нужна сводка результатов пилота.
## Детали
- Лорем ипсум долор сит амет.
- Консектетур адиписцинг элит.
"""


def test_offline_rules_choose_diagrams_and_keep_source_wording():
    plan = _fallback_plan(PLAN_TEXT, 5)
    slides = {slide["title"]: slide for slide in plan["slides"]}
    process = slides["План работ"]["visual"]
    assert process["type"] == "process" and process["origin"] == "offline_rules"
    assert process["items"][0] == "В первую неделю команда собирает инструкции."
    assert slides["План работ"]["bullets"] == []
    assert slides["Пользователи"]["visual"]["type"] == "icon_grid"
    # Text without confident pictograms stays a plain list.
    assert slides["Детали"]["visual"] is None and len(slides["Детали"]["bullets"]) == 2


def test_sequence_markers_and_headings_select_types():
    steps = [
        "Сначала собираем данные.",
        "Затем обучаем модель.",
        "После этого проверяем качество.",
    ]
    assert infer_diagram("Подход", steps)["type"] == "process"
    assert infer_diagram("SWOT-анализ", ["a", "b", "c", "d"])["type"] == "matrix"
    assert infer_diagram("Воронка продаж", ["Лиды", "Сделки"])["type"] == "funnel"
    hierarchy = infer_diagram("Структура команды", ["Аналитики", "Разработчики"])
    assert hierarchy["type"] == "hierarchy" and hierarchy["root"] == "Структура команды"
    assert (
        infer_diagram("Итоги", ["Первое наблюдение слишком длинное " * 5, "b", "c"])
        is None
    )


def test_pictogram_grids_are_limited_and_not_adjacent():
    slides = [
        {
            "title": f"Команда {index}",
            "role": "content",
            "visual": None,
            "bullets": ["Команда поддержки", "Сотрудники офиса"],
        }
        for index in range(6)
    ]
    add_offline_visuals(slides)
    chosen = [index for index, slide in enumerate(slides) if slide.get("visual")]
    assert len(chosen) == 2
    assert all(b - a > 1 for a, b in pairwise(chosen))


def test_model_visuals_are_allow_listed_and_bounded():
    raw = {
        "slides": [
            {"title": "A", "role": "cover"},
            {
                "title": "B",
                "role": "content",
                "visual": {
                    "type": "process",
                    "items": [{"label": "x" * 200, "detail": "y"}] * 9,
                },
            },
            {
                "title": "C",
                "role": "content",
                "visual": {"type": "matrix", "items": ["a", "b", "c"]},
            },
            {
                "title": "D",
                "role": "content",
                "visual": {"type": "image", "asset_id": "unknown"},
            },
            {
                "title": "E",
                "role": "content",
                "visual": {"type": "image", "asset_id": "abc"},
            },
            {
                "title": "F",
                "role": "content",
                "visual": {"type": "exec_code", "items": ["rm"]},
            },
            {"title": "G", "role": "closing"},
        ]
    }
    plan = normalize_plan(raw, source_text="x", slide_count=None, asset_ids={"abc"})
    visuals = [slide["visual"] for slide in plan["slides"]]
    assert len(visuals[1]["items"]) == 6 and len(visuals[1]["items"][0]["label"]) == 80
    assert visuals[2]["type"] == "icon_grid"
    assert visuals[3] is None and visuals[4] == {"type": "image", "asset_id": "abc"}
    assert visuals[5] is None


def test_layout_adapter_contract_is_unchanged():
    plain = {"title": "T", "role": "content", "bullets": ["a", "b"], "visual": None}
    diagram = {
        "title": "T",
        "role": "content",
        "bullets": [],
        "visual": {"type": "process", "items": ["один", "два", "три"]},
    }
    image = {
        "title": "T",
        "role": "content",
        "bullets": ["a"],
        "visual": {"type": "image", "asset_id": "x"},
    }
    keys = {
        "role",
        "title_chars",
        "body_chars",
        "bullet_count",
        "needs_data",
        "needs_image",
        "rhetorical_pattern",
        "has_visual",
    }
    for slide in (plain, diagram, image):
        assert set(_slide_requirements(slide)) == keys
    requirements = _slide_requirements(diagram)
    assert requirements["rhetorical_pattern"] == "bullet_list"
    assert requirements["bullet_count"] == 3 and requirements["body_chars"] == len(
        "одиндватри"
    )
    assert _slide_requirements(image)["needs_image"] is True


def test_offline_deck_contains_native_diagrams_and_literal_coverage(tmp_path: Path):
    template = tmp_path / "editorial.pptx"
    create_template(template, "editorial")
    result = generate_deck(
        template=template,
        content=PLAN_TEXT,
        workspace=tmp_path / "ws",
        slide_count=5,
        offline=True,
    )
    assert result["status"] == "completed"
    assert result["visuals"]["diagrams"] >= 2 and result["visuals"]["pictograms"] >= 3
    coverage = read_json(Path(result["presentation_dir"]) / "coverage_report.json")
    assert coverage["missing_from_pptx"] == []
    groups = [
        shape
        for slide in Presentation(result["output"]).slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP
        and shape.name.startswith("BrandDeck Diagram")
    ]
    assert len(groups) == result["visuals"]["diagrams"]
    assert result["workflow"]["pictograms"]["license"] == "ISC"
    assert result["workflow"]["files"]["planner"] == "planner-v3.txt"
