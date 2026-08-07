from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from examples.create_demo_assets import create_template
from slide_agent.analyzer import analyze_template
from slide_agent.composer import _remove_unmatched_content_images, compose_presentation
from slide_agent.qa import inspect_presentation
from slide_agent.utils import read_json


def _slide(title, visual=None, role="data", layout=1):
    return {
        "title": title,
        "subtitle": "",
        "body": "",
        "bullets": [],
        "visual": visual,
        "speaker_notes": "",
        "role": role,
        "pattern_id": f"layout-{layout}",
        "layout_index": layout,
        "master_index": 0,
    }


def test_all_supported_visuals(tmp_path: Path):
    template = tmp_path / "template.pptx"
    create_template(template, "clean-blue")
    analyzed = analyze_template(template, workspace=tmp_path / "workspace")
    design = read_json(analyzed / "design_system.json")
    plan = {
        "title": "Visual test",
        "slides": [
            _slide("Visual test", role="cover", layout=0),
            _slide(
                "Metrics",
                {"type": "metric_cards", "items": [{"label": "Speed", "value": "4x"}]},
            ),
            _slide(
                "Bars",
                {
                    "type": "bar_chart",
                    "categories": ["A", "B"],
                    "series": [{"name": "Value", "values": [3, 7]}],
                },
            ),
            _slide(
                "Pie",
                {"type": "pie_chart", "categories": ["A", "B"], "values": [40, 60]},
            ),
            _slide(
                "Table",
                {
                    "type": "table",
                    "headers": ["Name", "Value"],
                    "rows": [["A", "3"], ["B", "7"]],
                },
            ),
            _slide(
                "Timeline",
                {
                    "type": "timeline",
                    "items": [
                        {"label": "Now", "detail": "Prototype"},
                        {"label": "Next", "detail": "Pilot"},
                    ],
                },
            ),
            _slide("Thank you", role="closing", layout=10),
        ],
    }
    output = tmp_path / "visuals.pptx"
    compose_presentation(
        template_dir=analyzed,
        plan=plan,
        output_path=output,
        design_system=design,
    )
    result = inspect_presentation(output, design_system=design, expected_slide_count=7)
    assert result["status"] in {"passed", "warning"}
    prs = Presentation(output)
    assert any(shape.has_chart for shape in prs.slides[2].shapes)
    assert any(shape.has_table for shape in prs.slides[4].shapes)


def test_source_content_images_are_removed_but_brand_and_background_survive(
    tmp_path: Path,
):
    source = Presentation()
    source.slide_width = Inches(13.333)
    source.slide_height = Inches(7.5)
    slide = source.slides.add_slide(source.slide_layouts[6])
    background = tmp_path / "background.png"
    brand = tmp_path / "brand.png"
    stale_icon = tmp_path / "stale-icon.png"
    Image.new("RGB", (1200, 675), "#101828").save(background)
    Image.new("RGB", (120, 40), "#00B7FF").save(brand)
    Image.new("RGB", (80, 80), "#FF0000").save(stale_icon)
    slide.shapes.add_picture(
        str(background), 0, 0, source.slide_width, source.slide_height
    )
    brand_shape = slide.shapes.add_picture(
        str(brand), Inches(0.4), Inches(6.8), Inches(1.2), Inches(0.4)
    )
    slide.shapes.add_picture(
        str(stale_icon), Inches(0.5), Inches(0.5), Inches(0.5), Inches(0.5)
    )
    brand_hash = str(brand_shape.image.sha1)

    _remove_unmatched_content_images(
        slide,
        {"role": "content", "visual": None},
        {"canvas": {"width_inches": 13.333, "height_inches": 7.5}},
        {brand_hash},
    )

    pictures = [
        shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert len(pictures) == 2
    assert brand_hash in {str(shape.image.sha1) for shape in pictures}


def test_qa_checks_autofit_overflow_and_generated_text_overlap(tmp_path: Path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    first = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(0.35))
    first.name = "BrandDeck Body"
    first.text = "A very long sentence that cannot fit inside this tiny text box. " * 3
    first.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    first.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    second = slide.shapes.add_textbox(
        Inches(1.5), Inches(1.05), Inches(2), Inches(0.35)
    )
    second.text = "Overlapping text"
    output = tmp_path / "qa-overlap.pptx"
    prs.save(output)

    result = inspect_presentation(output)
    codes = {issue["code"] for issue in result["issues"]}
    assert "text_overflow_risk" in codes
    assert "text_overlap" in codes
