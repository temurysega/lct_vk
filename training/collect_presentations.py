"""Download the pinned, attributed presentation inventory; never run source code."""

import argparse
from concurrent.futures import ThreadPoolExecutor
import hashlib
import json
from pathlib import Path
from urllib.parse import quote
from urllib.request import Request, urlopen

from pptx import Presentation


def download(url: str, destination: Path) -> None:
    if destination.exists():
        return
    destination.parent.mkdir(parents=True, exist_ok=True)
    with urlopen(Request(url, headers={"User-Agent": "BrandDeck-dataset"}), timeout=120) as response:
        data = response.read()
    temporary = destination.with_suffix(destination.suffix + ".part")
    temporary.write_bytes(data)
    temporary.replace(destination)


def collect(source: dict, root: Path) -> list[dict]:
    base = f"https://raw.githubusercontent.com/{source['repo']}/{source['revision']}/"
    for key, name in (("license_file", "LICENSE.source.txt"), ("readme", "README.source.md")):
        download(base + quote(source[key]), root / source["family"] / name)
    rows = []
    for relative in source["files"]:
        filename = source["family"] + "__" + Path(relative).name
        path = root / "originals" / filename
        url = base + quote(relative)
        download(url, path)
        presentation = Presentation(path)
        texts = [" | ".join(s.text for s in slide.shapes if s.has_text_frame)
                 for slide in presentation.slides]
        rows.append({
            "filename": filename, "repo": source["repo"], "revision": source["revision"],
            "source_path": relative, "url": url, "family": source["family"],
            "license": source["license"], "author": source.get("author", source["repo"].split("/")[0]),
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "bytes": path.stat().st_size,
            "slides": len(texts), "text_slides": sum(bool(t.strip()) for t in texts),
            "text_chars": sum(map(len, texts)), "sample": texts[:3],
        })
        print(f"{filename}: {len(texts)} slides", flush=True)
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, required=True)
    args = parser.parse_args()
    sources = json.loads((args.root / "sources.json").read_text(encoding="utf-8"))
    with ThreadPoolExecutor(max_workers=4) as pool:
        batches = list(pool.map(lambda s: collect(s, args.root), sources))
    (args.root / "inventory.json").write_text(
        json.dumps([r for batch in batches for r in batch], ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
