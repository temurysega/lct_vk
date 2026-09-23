"""Build the Russian-only layout-selection bundle (v3) from the reviewed v2 sources.

Every slide of every source receives a recorded decision in ``data/review.json``.
Rules come from a visual review of all v2 examples:

* only Russian text (English and Chinese-English sources are dropped);
* hidden slides, code, URL/contact-only and picture- or drawing-driven slides
  are dropped: the model sees text and layout features, never the pictures,
  so their observed layout is noise for it;
* successive animation builds keep only the most complete slide;
* a long paragraph is not a title when a short line sits above it;
* the slide role matches the application (content for statements, section
  for a lone heading, cover/closing as observed).

Usage (from the repository root)::

    python -m training.prepare_ru_v3 --previous ../обучение --output ../обучение_ru_v3 --renders <pdf dir>
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import sys
from collections import Counter, defaultdict
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from slide_agent.layout_selector import candidate_features
from slide_agent.planner import _score_pattern, _slide_requirements
from slide_agent.prompt_config import load_prompt
from slide_agent.service import resolve_template
from training.curation import (
    clean_text,
    element_text,
    font_size,
    is_page_marker,
    repeated_marginal_text,
)
from training.package_ru_v3 import package
from training.prepare_drive import make_record

EXPERIMENT = "ru_only_v3"
PERMUTATIONS = {"train": 3, "validation": 1, "test": 1}
_CODE = re.compile(
    r"^\s*(SELECT|FROM|WHERE|GROUP BY|ORDER BY|CREATE TABLE|INSERT INTO|LIMIT)\b"
    r"|#include|std::|\breturn\b.*;|\binline\b|\bstruct\b|\bdef\s+\w+\(|=>|[{};]\s*$"
    r"|\w+\([^)]*\)\s*(;|\{)|^\s*[{}]\s*$"
    # Console output, YAML keys and byte dumps are code too.
    r"|READY\s+STATUS|\bpod/|^[a-z_][\w.-]*:\s*$|\b0x[0-9A-Fa-f]{2}\b"
)
_CYRILLIC = re.compile(r"[\u0400-\u04ff]")
_URL = re.compile(
    r"^(https?://\S+|www\.\S+|\S+@\S+\.\w+|@\w+|[\w.-]+\.(ru|com|io|org|net|info)(/\S*)?)$",
    re.IGNORECASE,
)


def is_code(text: str) -> bool:
    return bool(_CODE.search(text))


def is_url_like(text: str) -> bool:
    return bool(_URL.match(clean_text(text)))


def cyrillic_share(text: str) -> float:
    cyrillic = len(_CYRILLIC.findall(text))
    latin = len(re.findall(r"[A-Za-z]", text))
    return cyrillic / max(1, cyrillic + latin)


def role_for(observed_roles: list[str], bullets: list[str]) -> str:
    first = (observed_roles or ["content"])[0]
    if first in {"cover", "closing"}:
        return first
    return "content" if bullets else "section"


def _usable_elements(slide: dict, repeated: set[str]) -> list[dict]:
    elements = []
    for element in slide.get("text_elements", []):
        text = element_text(element)
        placeholder = element.get("placeholder_type", "").upper()
        if (
            not re.search(r"[^\W\d_]", text, re.UNICODE)
            or is_page_marker(text)
            or text.casefold() in repeated
            or any(kind in placeholder for kind in ("SLIDE_NUMBER", "FOOTER", "DATE"))
            or text.casefold()
            in {
                "фото",
                "photo",
                "click to add title",
                "нажмите, чтобы добавить заголовок",
            }
        ):
            continue
        elements.append(element)
    return elements


def on_canvas(element: dict, width: float, height: float) -> bool:
    """Text parked entirely off the canvas is never shown in the slide show."""
    left, top = element.get("left", 0), element.get("top", 0)
    return (
        left < width
        and top < height
        and left + element.get("width", 0) > 0
        and top + element.get("height", 0) > 0
    )


def running_labels(context: dict, repeated: set[str]) -> set[str]:
    """Short non-title texts repeated in the top or bottom band of 3+ slides.

    Running headers and kickers ("BI Tableau over ClickHouse" above every
    section title) are decoration, not statements. Titles are excluded, so a
    template's placeholder headings survive.
    """
    height = context["presentation"]["slide_height_inches"]
    counts: Counter[str] = Counter()
    for slide in context["slides"]:
        title, _, _ = pick_content(slide, height, repeated)
        counts.update(
            {
                element_text(e).casefold()
                for e in _usable_elements(slide, repeated)
                if len(element_text(e)) <= 40
                and element_text(e) != title
                and (e.get("top", 0) < height * 0.3 or e.get("top", 0) > height * 0.84)
            }
        )
    return {text for text, count in counts.items() if count >= 3}


def pick_content(
    slide: dict,
    height: float,
    repeated: set[str],
    labels: frozenset[str] | set[str] = frozenset(),
) -> tuple[str, list[str], bool]:
    """Title and bullets; returns whether a paragraph-as-title was corrected."""
    elements = _usable_elements(slide, repeated)
    if not elements:
        return "", [], False
    upper = [e for e in elements if e.get("top", 0) < height * 0.7]
    pool = upper or elements

    def score(element: dict) -> tuple:
        text = element_text(element)
        placeholder = element.get("placeholder_type", "").upper()
        explicit = "TITLE" in placeholder and "SUBTITLE" not in placeholder
        return (
            int(explicit),
            font_size(element) - max(0, len(text) - 180) / 45,
            -element.get("top", 0),
            -element.get("left", 0),
        )

    title = max(pool, key=score)
    corrected = False
    if len(element_text(title)) > 90:
        # Headings are short and sit above body text.
        above = [
            e
            for e in elements
            if e is not title
            and 2 <= len(element_text(e)) <= 60
            and e.get("top", 0) < title.get("top", 0)
        ]
        if above:
            title = min(above, key=lambda e: e.get("top", 0))
            corrected = True
    bullets: list[str] = []
    for element in sorted(
        elements, key=lambda e: (round(e.get("top", 0), 1), e.get("left", 0))
    ):
        if element is title:
            continue
        for paragraph in element.get("paragraphs", []):
            text = clean_text(paragraph.get("text", ""))
            if (
                text
                and not is_page_marker(text)
                and re.search(r"\w", text)
                and text not in bullets
                and element_text(element).casefold() not in labels
            ):
                bullets.append(text)
    return element_text(title), bullets, corrected


def recurring_images(prs) -> set[str]:
    """Picture hashes on at least 40% of slides: logos and brand decoration."""
    counts: Counter[str] = Counter()

    def visit(shapes, seen: set[str]) -> None:
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                visit(shape.shapes, seen)
            elif hasattr(shape, "image"):
                try:
                    seen.add(shape.image.sha1)
                except (AttributeError, KeyError, ValueError):
                    pass

    for slide in prs.slides:
        seen: set[str] = set()
        visit(slide.shapes, seen)
        counts.update(seen)
    limit = max(3, len(prs.slides) * 0.4)
    return {digest for digest, count in counts.items() if count >= limit}


def media_profile(
    slide, width: int, height: int, brand: frozenset[str] | set[str] = frozenset()
) -> dict:
    """Share of the slide covered by content pictures/charts and drawings.

    Full-slide backgrounds and recurring brand pictures are not content.
    """
    area = max(1, width * height)
    picture = chart = 0
    drawings = 0

    def visit(shapes) -> None:
        nonlocal picture, chart, drawings
        for shape in shapes:
            size = max(0, int(shape.width or 0)) * max(0, int(shape.height or 0))
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                visit(shape.shapes)
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
                try:
                    digest = shape.image.sha1
                except (AttributeError, KeyError, ValueError):
                    digest = ""
                if size < 0.8 * area and digest not in brand:
                    picture += size
            elif getattr(shape, "has_chart", False):
                chart += size
            elif shape.shape_type in {
                MSO_SHAPE_TYPE.AUTO_SHAPE,
                MSO_SHAPE_TYPE.FREEFORM,
                MSO_SHAPE_TYPE.LINE,
            } or shape._element.tag == qn("p:cxnSp"):
                has_text = (
                    getattr(shape, "has_text_frame", False) and shape.text.strip()
                )
                blip = shape._element.find(f"{qn('p:spPr')}/{qn('a:blipFill')}")
                if blip is not None:
                    if size < 0.8 * area:
                        picture += size
                elif not has_text:
                    drawings += 1

    visit(slide.shapes)
    return {
        "hidden": slide._element.get("show") == "0",
        "picture_ratio": round(min(1.0, picture / area), 3),
        "chart_ratio": round(min(1.0, chart / area), 3),
        "drawings": drawings,
    }


def is_russian(title: str, bullets: list[str], role: str, brand: bool = False) -> bool:
    """Names, e-mails and product terms are Latin; the heading decides."""
    text = " ".join([title, *bullets])
    title_letters = len(_CYRILLIC.findall(title))
    if role in {"cover", "closing"}:
        # An organizer template may open with its brand name ("VK Tech").
        return brand or title_letters >= 3
    if title_letters >= 4 and cyrillic_share(title) >= 0.6:
        return True
    return len(_CYRILLIC.findall(text)) >= 10 and (
        cyrillic_share(title) >= 0.4 or cyrillic_share(text) >= 0.5
    )


def rejection(
    title: str, bullets: list[str], role: str, media: dict, brand: bool = False
) -> str | None:
    if media["hidden"]:
        return "hidden_slide"
    if not is_russian(title, bullets, role, brand):
        return "not_russian"
    if is_url_like(title) or (
        bullets and sum(map(is_url_like, bullets)) * 2 >= len(bullets)
    ):
        return "url_or_contacts_only"
    if is_code(title) or any(is_code(b) for b in bullets):
        return "code"
    visual = media["picture_ratio"] + media["chart_ratio"]
    captions = bool(bullets) and sum(map(len, bullets)) / len(bullets) <= 30
    if (
        visual >= 0.5
        or (visual >= 0.25 and len(bullets) <= 2)
        # Short captions under screenshots or logos are not statements.
        or (visual >= 0.25 and captions)
    ):
        return "picture_or_chart_driven"
    if not bullets and media["drawings"] >= 6 and role not in {"cover", "closing"}:
        return "drawing_driven"
    return None


def build_duplicates(rows: list[dict]) -> set[int]:
    """Indices of animation builds whose bullets are contained in a later build."""
    by_title: dict[str, list[int]] = defaultdict(list)
    for index, row in enumerate(rows):
        by_title[row["slide"]["title"].casefold()].append(index)
    drop = set()
    for indices in by_title.values():
        for a in indices:
            for b in indices:
                # Builds are neighbouring slides; equal placeholder text on
                # distant layouts (templates) is not an animation.
                adjacent = (
                    0 < abs(rows[a]["source_slide"] - rows[b]["source_slide"]) <= 2
                )
                if a == b or b in drop or not adjacent:
                    continue
                bullets_a, bullets_b = (
                    rows[a]["slide"]["bullets"],
                    rows[b]["slide"]["bullets"],
                )
                if set(bullets_a) < set(bullets_b) or (
                    bullets_a == bullets_b and a < b
                ):
                    drop.add(a)
                    break
    return drop


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write(path: Path, value) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def extract(template: Path, cache: Path, source: dict) -> tuple[list[dict], list[dict]]:
    directory = resolve_template(template, cache, client=None)
    context = json.loads((directory / "context.json").read_text(encoding="utf-8"))
    catalog = json.loads(
        (directory / "pattern_catalog.json").read_text(encoding="utf-8")
    )["patterns"]
    by_id = {p["id"]: p for p in catalog}
    width_in = context["presentation"]["slide_width_inches"]
    height_in = context["presentation"]["slide_height_inches"]
    for slide in context["slides"]:
        slide["text_elements"] = [
            e
            for e in slide.get("text_elements", [])
            if on_canvas(e, width_in, height_in)
        ]
    repeated = repeated_marginal_text(context)
    labels = running_labels(context, repeated)
    prs = Presentation(template)
    brand = recurring_images(prs)
    rows, review = [], []
    for slide in context["slides"]:
        number = slide["index"]
        original = (
            source["original_slide_map"][number - 1]
            if source.get("original_slide_map")
            else number
        )
        entry = {
            "template": template.name,
            "source_slide": number,
            "original_slide": original,
        }
        title, bullets, corrected = pick_content(slide, height_in, repeated, labels)
        positive = by_id.get(f"slide-{number}")
        media = media_profile(
            prs.slides[number - 1], prs.slide_width, prs.slide_height, brand
        )
        entry.update(
            title=title[:160],
            bullets=len(bullets),
            media=media,
            title_corrected=corrected,
        )
        role = role_for(positive.get("roles", []) if positive else [], bullets)
        reason = None
        if not title:
            reason = "no_editable_text"
        elif positive is None:
            reason = "missing_observed_layout"
        else:
            reason = rejection(
                title, bullets, role, media, brand=bool(source.get("organizer"))
            )
        # Training and the application both show five bullets to the model;
        # only extreme slides (tag clouds, dense tables) are dropped.
        if reason is None and (
            len(title) > 160 or len(bullets) > 12 or any(len(b) > 400 for b in bullets)
        ):
            reason = "text_too_dense"
        if reason is None and any(
            e.get("left", 0) < -0.02
            or e.get("top", 0) < -0.02
            or e.get("left", 0) + e.get("width", 0) > width_in + 0.02
            or e.get("top", 0) + e.get("height", 0) > height_in + 0.02
            for e in _usable_elements(slide, repeated)
        ):
            reason = "text_outside_slide"
        if reason is None and positive["capacity"].get("out_of_canvas_body_zones", 0):
            reason = "unresolved_body_geometry"
        alternatives = []
        if reason is None:
            sample = {
                "title": title,
                "role": role,
                "body": "",
                "bullets": bullets,
                "visual": None,
            }
            requirements = _slide_requirements(sample)
            seen = {json.dumps(candidate_features(positive), sort_keys=True)}
            for pattern in sorted(
                catalog,
                key=lambda p: (
                    -_score_pattern(p, requirements, reuse_count=0, avoided=False)[0]
                ),
            ):
                signature = json.dumps(candidate_features(pattern), sort_keys=True)
                if signature not in seen:
                    seen.add(signature)
                    alternatives.append(pattern)
            if len(alternatives) < 2:
                reason = "too_few_distinct_candidates"
        entry["decision"] = "kept" if reason is None else "removed"
        entry["reason"] = reason
        review.append(entry)
        if reason is None:
            rows.append(
                {
                    "slide": sample,
                    "requirements": requirements,
                    "positive": positive,
                    "alternatives": alternatives[:12],
                    "group": "pending",
                    "source_slide": number,
                    "template": template.name,
                    "template_sha256": sha(template),
                    "design_family": source["design_family"],
                    "source_url": source.get("url") or source.get("source_url", ""),
                    "original_slide": original,
                    "review_index": len(review) - 1,
                }
            )
    # Templates show each layout once with placeholder text: no builds there.
    builds = set() if source.get("organizer") else build_duplicates(rows)
    for index in sorted(builds, reverse=True):
        review[rows[index]["review_index"]].update(
            decision="removed", reason="animation_build_duplicate"
        )
        rows.pop(index)
    return rows, review


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--previous", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument(
        "--cache", type=Path, default=Path("slide-workspace/ru-v3-cache")
    )
    parser.add_argument(
        "--renders", type=Path, help="PDF renders of the sources for previews/"
    )
    args = parser.parse_args()
    previous, output = args.previous.resolve(), args.output.resolve()
    if (output / "weights").exists() or (output / "runs").exists():
        raise ValueError("Do not replace a bundle that already has training runs")
    sys.path.insert(0, str(previous / "scripts"))
    from data_utils import check_split, content_group

    inventory = json.loads(
        (previous / "data/source_inventory.json").read_text(encoding="utf-8")
    )
    organizer_cases = json.loads(
        (previous / "data/organizer_test.json").read_text(encoding="utf-8")
    )["cases"]
    if output.exists():
        shutil.rmtree(output)
    for folder in ("data", "templates", "test_templates", "licenses", "scripts", "qa"):
        (output / folder).mkdir(parents=True, exist_ok=True)

    parts: dict[str, list[dict]] = {"train": [], "validation": [], "test": []}
    review: list[dict] = []
    kept_sources = []
    # Examples come from the render-checked copies in templates/ (VK included);
    # the organizers' untouched originals with media go to test_templates/.
    for case in organizer_cases:
        shutil.copy2(previous / case["template"], output / case["template"])
    for source in inventory:
        name = source["filename"]
        split = source.get("split")
        source["organizer"] = source.get("design_family") == "vk-brand"
        template = previous / "templates" / name
        if sha(template) != source["sha256"]:
            raise ValueError("Changed reviewed source: " + name)
        if source.get("language", "ru") != "ru":
            review.append(
                {
                    "template": name,
                    "decision": "removed",
                    "reason": "source_not_russian",
                }
            )
            continue
        rows, decisions = extract(template, args.cache.resolve(), source)
        review.extend(decisions)
        for row in rows:
            for permutation in range(PERMUTATIONS[split]):
                record = make_record(row, permutation)
                record["provenance"].update(
                    split=split,
                    language="ru",
                    original_slide=row["original_slide"],
                    license=source.get(
                        "license",
                        "user-provided-for-training-not-for-public-redistribution",
                    ),
                    quality="reviewed_v3",
                    **(
                        {"evaluation_suite": "organizer"}
                        if source.get("organizer")
                        else {}
                    ),
                )
                record["provenance"]["group"] = content_group(record)
                parts[split].append(record)
        shutil.copy2(template, output / "templates" / name)
        kept = sum(d["decision"] == "kept" for d in decisions)
        kept_sources.append(
            {
                **{k: v for k, v in source.items() if k != "organizer"},
                "kept_examples": kept,
                "reviewed_slides": len(decisions),
            }
        )
        print(
            f"{name[:60]:60} {split:10} kept {kept:3} of {len(decisions):3}", flush=True
        )

    # One content group may live in one split only; keep one row per permutation.
    splits_of = defaultdict(set)
    for split, records in parts.items():
        for record in records:
            splits_of[record["provenance"]["group"]].add(split)
    dropped = Counter()
    for split, records in parts.items():
        seen, kept_rows = set(), []
        for record in records:
            group = record["provenance"]["group"]
            if len(splits_of[group]) > 1:
                dropped["cross_split_content_rows"] += 1
                continue
            key = (group, record.get("permutation", 0))
            if key in seen:
                dropped["within_split_duplicate_rows"] += 1
                continue
            seen.add(key)
            kept_rows.append(record)
        parts[split] = kept_rows
        (output / "data" / f"{split}.jsonl").write_text(
            "".join(json.dumps(r, ensure_ascii=False) + "\n" for r in kept_rows),
            encoding="utf-8",
        )
    check_split(parts["train"], parts["validation"], parts["test"])
    assert all(
        r["provenance"]["design_family"] != "vk-brand"
        for s in ("train", "validation")
        for r in parts[s]
    )

    cases = []
    for case in organizer_cases:
        path = output / case["template"]
        with ZipFile(path) as archive:
            media = {
                n: hashlib.sha256(archive.read(n)).hexdigest()
                for n in archive.namelist()
                if n.startswith("ppt/media/")
            }
        cases.append(
            {
                **case,
                "sha256": sha(path),
                "media": media,
                "layout_examples": sum(
                    r["provenance"]["template"] == Path(case["template"]).name
                    for r in parts["test"]
                ),
            }
        )
    assert all(case["layout_examples"] > 0 for case in cases)
    write(
        output / "data/organizer_test.json",
        {
            "cases": cases,
            "note": "Original organizer PPTX with all media. Test is evaluated after the validation decision and never selects the adapter.",
        },
    )
    write(output / "data/source_inventory.json", kept_sources)
    write(
        output / "data/review.json",
        {
            "rules": __doc__.split("Usage")[0].strip(),
            "dropped_rows": dict(dropped),
            "slides": review,
        },
    )

    stats = {}
    for split, records in parts.items():
        unique = {r["provenance"]["group"]: r for r in records}
        stats[split] = {
            "rows": len(records),
            "unique_groups": len(unique),
            "templates": len({r["provenance"]["template"] for r in records}),
            "families": sorted({r["provenance"]["design_family"] for r in records}),
            "unique_by_language": dict(
                Counter(r["provenance"]["language"] for r in unique.values())
            ),
            "roles": dict(
                Counter(
                    json.loads(r["messages"][1]["content"])["requirements"]["role"]
                    for r in unique.values()
                )
            ),
        }
    reasons = Counter(d["reason"] for d in review if d["decision"] == "removed")
    previous_manifest = json.loads(
        (previous / "data/manifest.json").read_text(encoding="utf-8")
    )
    write(
        output / "data/manifest.json",
        {
            "schema_version": 5,
            "profile": "cpu",
            "task": "layout_candidate_selection",
            "label_kind": "weak_observed_exemplar_not_human_preference",
            "split_unit": "design_family",
            "experiment": EXPERIMENT,
            "organizer_templates_test_only": True,
            "train_rows": len(parts["train"]),
            "validation_rows": len(parts["validation"]),
            "test_rows": len(parts["test"]),
            "train_unique_groups": stats["train"]["unique_groups"],
            "split_stats": stats,
            "review": {
                "slides_reviewed": sum("source_slide" in d for d in review),
                "removed_by_reason": dict(reasons),
            },
            "quality_gate": {
                "split_leakage": False,
                "organizer_originals_preserved": True,
                "final_render_failures": 0,
                "previous": previous_manifest.get("quality_gate"),
            },
            "layout_prompt_sha256": hashlib.sha256(
                load_prompt("layout").encode()
            ).hexdigest(),
            "files": {},
        },
    )
    packaged = package(previous, output, args.renders)
    print(
        json.dumps(
            {
                "packaged": packaged,
                "splits": stats,
                "removed_by_reason": dict(reasons),
                "dropped_rows": dict(dropped),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
