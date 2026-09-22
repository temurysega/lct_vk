"""Conservative, colour-independent PPTX family audit; no slide-level splitting."""

from __future__ import annotations

import hashlib
import json
from itertools import combinations
from pathlib import Path

from pptx import Presentation


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fingerprint(path: Path) -> dict:
    presentation = Presentation(path)
    width, height = presentation.slide_width, presentation.slide_height

    def geometry(shapes):
        # Ignore text, palette, font, IDs and ZIP timestamps to catch recolours.
        return sorted(
            (
                int(s.shape_type or -1),
                round((s.left or 0) / width, 2),
                round((s.top or 0) / height, 2),
                round((s.width or 0) / width, 2),
                round((s.height or 0) / height, 2),
            )
            for s in shapes
        )

    def signature(value):
        return hashlib.sha256(json.dumps(value, sort_keys=True).encode()).hexdigest()

    slide_geometry = [geometry(s.shapes) for s in presentation.slides]
    masters = [geometry(m.shapes) for m in presentation.slide_masters]
    # Blank/default masters alone do not establish a design family.
    distinctive_masters = [signature(m) for m in masters if len(m) >= 7]
    return {
        "sha256": sha256(path),
        "slides": len(presentation.slides),
        "text_shapes": sum(
            s.has_text_frame and bool(s.text.strip())
            for sl in presentation.slides
            for s in sl.shapes
        ),
        "geometry_sha256": signature(sorted(signature(s) for s in slide_geometry)),
        "slide_geometry": sorted({signature(s) for s in slide_geometry if len(s) >= 3}),
        "distinctive_masters": sorted(set(distinctive_masters)),
    }


def cluster_sources(sources: list[dict]) -> list[dict]:
    """Union declared families, duplicates and conservative structural matches."""
    parent = list(range(len(sources)))

    def root(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    edges = []
    for (i, a), (j, b) in combinations(enumerate(sources), 2):
        fa, fb = a["fingerprint"], b["fingerprint"]
        sa, sb = set(fa["slide_geometry"]), set(fb["slide_geometry"])
        shared = len(sa & sb)
        similarity = shared / max(1, min(len(sa), len(sb)))
        reason = None
        if a["design_family"] == b["design_family"]:
            reason = "declared_common_design_system"
        elif fa["sha256"] == fb["sha256"]:
            reason = "byte_duplicate"
        elif fa["geometry_sha256"] == fb["geometry_sha256"]:
            reason = "same_geometry_ignoring_text_and_colour"
        elif set(fa["distinctive_masters"]) & set(fb["distinctive_masters"]):
            reason = "shared_distinctive_master_geometry"
        elif shared >= 3 and similarity >= 0.65:
            reason = "shared_slide_geometry"
        if reason:
            parent[root(j)] = root(i)
            edges.append(
                {
                    "left": a["name"],
                    "right": b["name"],
                    "reason": reason,
                    "shared_layouts": shared,
                    "containment": similarity,
                }
            )
    for i, source in enumerate(sources):
        members = [
            s["design_family"] for j, s in enumerate(sources) if root(i) == root(j)
        ]
        source["design_family"] = min(members)
    return edges
